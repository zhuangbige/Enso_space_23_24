from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/daiyue/Documents/Maryland")
OUT = ROOT / "outputs/research_2026/papers/ENSO暖量到达_去向与海气耦合命运_中文小论文_投稿措辞修订版_20260808.docx"
FIG = ROOT / "outputs/research_2026/three_questions_figure_package_20260805/01_main_figures"

FONT_CN = "Songti SC"
FONT_HEAD = "PingFang SC"
FONT_LATIN = "Times New Roman"
FONT_MATH = "Cambria Math"

INK = "111111"
MUTED = "555555"
BLUE = "1F4E79"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F8"
GRID = "AAB4BE"


def set_run_font(run, size=11, bold=False, italic=False, color=INK,
                 latin=FONT_LATIN, east_asia=FONT_CN):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=INK)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.45)
    sec.bottom_margin = Cm(2.35)
    sec.left_margin = Cm(2.55)
    sec.right_margin = Cm(2.55)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.45

    for style_name, size, before, after in (
        ("Heading 1", 14, 14, 7),
        ("Heading 2", 12, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_HEAD
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15

    caption = doc.styles["Caption"]
    caption.font.name = FONT_LATIN
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.15

    footer_p = sec.footer.paragraphs[0]
    add_page_number(footer_p)

    doc.core_properties.title = "厄尔尼诺暖量到达、去向与海气耦合命运"
    doc.core_properties.subject = "基于SODA3的20°C阈值暖量预算与事件对比"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "ENSO; SODA3; 暖水体积; 暖量输送; 海气耦合"


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run("厄尔尼诺暖量到达、去向与海气耦合命运"),
                 size=18, bold=True, east_asia=FONT_HEAD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("——基于 SODA3 的 20°C 阈值暖量预算及 1997、2015、2023 事件对比"),
                 size=12, bold=True, color=INK, east_asia=FONT_HEAD)


def add_abstract(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run("摘要："), size=10.5, bold=True, east_asia=FONT_HEAD)
    text = (
        "暖水体积（WWV）及其充放电过程是厄尔尼诺—南方涛动（ENSO）海洋记忆的核心表征，"
        "但标准化体积输送不能直接回答暖量是否真正进入东太平洋、到达后从何处离开，以及相似的最东部太平洋（FEP）先暖为何产生不同海气耦合命运。"
        "本文基于SODA3.15.2构建20°C阈值以上相对热含量H₂₀和暖水体热能输送Qwarm,20，"
        "在160°E—80°W、10°S—5°N范围建立西—中—东三箱预算，并结合OISST、ERA5、GPCP和CPC OLR，"
        "比较冷舌型（CT）与中央太平洋型（CP）合成以及1997/98、2015/16和2023/24三个事件。"
        "结果表明：CT与CP均能把暖量输送到东部，但CT发展期跨140°W输入更强；1997输入最强，而2015和2023的发展期输入分别为3.30和3.91 ZJ，说明2023后续出口偏弱不能归因于暖量未到达。"
        "到达后的主要分配通道为5°N北向输出、海表热损失、上层海洋储量变化及阈值预算剩余项，10°S不是稳健主出口。"
        "次年1—5月，1997、2015和2023的5°N总输出依次为3.81、0.88和0.10 ZJ，形成强排放—部分排放—近闭合阶梯。"
        "诊断分解显示，2015与2023的南向地转项几乎相同，但2015的Ekman暖量输送更强，使合成后的总诊断出口保持北向。"
        "13事件统计进一步表明，中太平洋OHC与随后西风响应、以及西风响应与次年Ekman出口分别显著正相关；"
        "离线换风试验可在固定海洋压力场下大幅改变2015与2023的诊断出口。"
        "因此，暖量到达东部只是必要条件，事件命运取决于盆地尺度次表层预置能否转化为深对流—风场接通，并通过增强的风驱Ekman暖量输送提高5°N总诊断北向输出。"
    )
    set_run_font(p.add_run(text), size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("关键词："), size=10.5, bold=True, east_asia=FONT_HEAD)
    set_run_font(p.add_run("ENSO；暖水体积；阈值暖量；暖量预算；Ekman输送；地转输送；海气耦合"), size=10.5)


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc, text, *, no_indent=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0 if no_indent else 22)
    p.paragraph_format.keep_together = keep
    set_run_font(p.add_run(text), size=11)
    return p


def add_lead(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(label), size=11, bold=True, color=INK, east_asia=FONT_HEAD)
    set_run_font(p.add_run(text), size=11)
    return p


def add_equation(doc, equation, number):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.7), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    run = p.add_run(equation)
    set_run_font(run, size=11.5, latin=FONT_MATH, east_asia=FONT_CN)
    p.add_run("\t")
    set_run_font(p.add_run(f"({number})"), size=10.5)
    return p


def add_table(doc, caption, headers, rows, widths_dxa, font_size=9.3, header_fill=LIGHT_BLUE):
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = True
    set_run_font(cap.add_run(caption), size=9.5, bold=True, east_asia=FONT_HEAD)

    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    mark_repeat_header(table.rows[0])

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        set_run_font(p.add_run(str(header)), size=font_size, bold=True, east_asia=FONT_HEAD)

    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(str(value)) > 16 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            set_run_font(p.add_run(str(value)), size=font_size)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i])

    after = doc.add_paragraph()
    after.paragraph_format.first_line_indent = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, filename, caption, width_in=6.25, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / filename), width=Inches(width_in))
    cap = doc.add_paragraph(style="Caption")
    set_run_font(cap.add_run(caption), size=9.5, east_asia=FONT_CN)
    return p


def add_reference(doc, index, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(f"[{index}] {text}"), size=9.5)
    return p


def build():
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_abstract(doc)

    add_heading(doc, "1 引言", 1)
    add_body(doc, (
        "ENSO是热带海洋—大气系统中最强的年际变率，其发展和衰减依赖赤道太平洋上层海洋的热量重新分配。"
        "Jin（1997）的充放电振子把纬向平均温跃层深度与经向暖水输送失衡视为ENSO的慢记忆；"
        "Meinen和McPhaden（2000）的观测进一步表明，赤道太平洋20°C等温面以上WWV异常通常领先Niño-3.4海表温度数月。"
        "这一框架说明暖水的积累和释放对事件相位转换至关重要，但WWV本身是体积量，不能直接表示穿过固定断面的能量通量。"
    ))
    add_body(doc, (
        "围绕ENSO多样性，Kug等（2009）和Capotondi等（2015）指出冷舌型或东太平洋型事件与中央太平洋型事件在海表温度结构、次表层预兆和生命周期上存在系统差异。"
        "Chakravorty等（2021）利用SODA复合、surface center-of-heat index、D20 EOF和170°E、5°N、10°S暖水体积输送表明，CT事件具有较清晰的赤道充水、向东输送和成熟后向极放水，而CP事件的充放电循环较弱。"
        "然而，其输送量被标准化，尚不能回答多少暖量跨140°W进入东部、多少经南北边界离开、多少由海表散失或跨20°C阈值转换。"
    ))
    add_body(doc, (
        "2023/24事件进一步暴露了这一缺口。观测和模式研究显示，2023/24具有显著海洋增暖，但热带太平洋风压、降水和遥相关响应弱于历史强事件；"
        "可能的上游调制包括外海盆增暖、多年La Niña遗产、持续东风和NPMM—MJO条件（Hong等，2024；Kim等，2025；Lin等，2025；Peng等，2025；Zhang等，2025；Fan等，2026）。"
        "这些研究解释了大气为何可能偏弱，却仍未把弱耦合定量落实到‘到达后的暖量从哪里离开’这一海洋收支问题。"
    ))
    add_body(doc, (
        "本文利用SODA3温度、盐度、流速、海表高度、风应力和净加热场同网格的优势，把经典WWV诊断推进为20°C阈值以上相对热含量H₂₀、暖水体热能输送Qwarm,20、三箱控制体预算以及5°N地转—Ekman分解。"
        "研究依次回答三个问题：（1）暖量是否以及通过什么路径到达东太平洋；（2）到达后的暖量在北向输出、南向输出、海表散热和海洋储存之间如何分配；"
        "（3）为什么相似的FEP先暖和东部输入会产生不同海气耦合与暖量命运。"
    ))
    add_lead(doc, "核心假设：", "暖量到达东部只是事件发展的必要条件；其成熟后命运与盆地尺度次表层预置、大气对流—风场接通以及5°N处Ekman和地转诊断分量的相对贡献有关。")

    add_heading(doc, "2 数据与方法", 1)
    add_heading(doc, "2.1 数据资料与研究设计", 2)
    add_body(doc, (
        "海洋主诊断使用SODA3.15.2。该产品在同一网格上提供温度、盐度、三维流速、SSH、风应力和net_heating，"
        "因此可以在一致的断面和控制体内计算储量、平流、表面通量及动力分量。OISST v2.1用于独立定义FEP和Niño-3.4表面异常；"
        "ERA5用于约束u850、SLP、风应力和表面热通量；GPCP与CPC OLR用于诊断深对流位置。"
    ))
    add_table(doc, "表1  数据资料及其在研究中的作用",
              ["资料", "主要变量", "时段/分辨率", "用途"],
              [
                  ["SODA3.15.2", "T、S、u/v、SSH、τx、net_heating", "月平均；1980—2022，补充2024年1—5月", "H₂₀、Qwarm,20、三箱预算、OHC/D20和诊断分解"],
                  ["OISST v2.1", "SST", "月平均；1982—2024", "FEP、Niño-3.4和表面触发"],
                  ["ERA5", "u850、SLP、τx、短长波、潜热和感热", "月平均；与事件期共同覆盖", "大气闭合、独立Ekman和海表热损失"],
                  ["GPCP / CPC OLR", "降水、OLR", "月平均", "对流中心东移与深对流独立约束"],
              ], [1500, 2600, 2200, 3060], font_size=9.0)
    add_body(doc, (
        "研究采用两层设计。类型合成选取与Chakravorty等（2021）重叠且落在SODA3共同期内的4个CT事件（1982、1986、1997、2006）和4个CP事件（1991、1994、2002、2004），"
        "发展期定义为Jun(0)—Nov(0)，成熟—衰减期定义为Dec(0)—May(+1)。过程对照选取1997/98、2015/16和2023/24，发展年使用Jan—Jun，次年输出使用Jan—May。"
        "CT/CP合成回答群体差异，三事件对照用于识别强协同、风驱增强和出口近闭合三种诊断状态。"
    ))

    add_heading(doc, "2.2 异常、趋势与统计推断", 2)
    add_body(doc, (
        "所有异常按日历月分别去除气候态和线性趋势。CT/CP合成以1982—2022为共同基准；三事件诊断分解以1980—2014为主基准，并以1982—2022检验敏感性。"
        "CT/CP标量序列采用中心3个月平均，三事件断面分解使用未平滑月值。类型差异以整事件为抽样单位进行20,000次bootstrap，报告均值、95%区间和效应量。"
        "FEP逐日历月回归使用1982—2024年的43个年度样本，每月OLS自由度为41。13个暖事件的机制关系同时报告Pearson相关、事件bootstrap区间、事件置换p值和逐一删事件范围；"
        "1997、2015和2023仅用于过程一致性对照，不对三个个例实施显著性检验。"
    ))

    add_heading(doc, "2.3 20°C阈值暖量及其输送", 2)
    add_body(doc, (
        "为把WWV从体积量推进到能量量纲，定义20°C阈值超温和指示函数。20°C的选择来自D20/WWV的动力学意义，"
        "不被解释为海水热力学的绝对参考温度。"
    ))
    add_equation(doc, "χ₂₀(T)=max(T−Tᵣ,0),   I₂₀(T)=1(T>Tᵣ),   Tᵣ=20 °C", 1)
    add_equation(doc, "H₂₀(t)=ρ₀cₚ ∭Ω χ₂₀[T(x,y,d,t)] dV", 2)
    add_equation(doc, "WWV=∭Ω I₂₀ dV,   ΔT̄₂₀=(∭Ωχ₂₀dV)/(∭ΩI₂₀dV),   H₂₀=ρ₀cₚ·WWV·ΔT̄₂₀", 3)
    add_equation(doc, "Qwarm,20,Γ(t)=ρ₀cₚ ∬Γ vₙ(t)χ₂₀[T(t)] dA", 4)
    add_table(doc, "表2  式（1）—（4）的全部符号、物理意义与单位",
              ["符号", "物理意义", "单位/约定"],
              [
                  ["T(x,y,d,t)", "经度x、纬度y、向下深度d和时间t处的海水温度", "°C；温差等价用K"],
                  ["Tᵣ", "阈值参考温度", "20°C"],
                  ["χ₂₀(T)", "高于20°C的超温；低于阈值时取0", "K"],
                  ["I₂₀(T)", "阈值指示函数：T>Tᵣ取1，否则取0", "无量纲"],
                  ["max、1(·)", "最大值算子与真假条件指示算子", "数学算子"],
                  ["ρ₀", "参考海水密度", "1025 kg m⁻³"],
                  ["cₚ", "定压比热容", "3990 J kg⁻¹ K⁻¹"],
                  ["Ω", "指定海洋控制体；不是地球自转角速度", "空间区域"],
                  ["dV、dA", "体积微元和断面面积微元", "m³、m²"],
                  ["H₂₀(t)", "控制体内20°C阈值以上相对热含量", "J，文中积分后用ZJ"],
                  ["WWV", "20°C等温面以上暖水体积", "m³"],
                  ["ΔT̄₂₀", "暖水体相对20°C的体积平均超温", "K"],
                  ["Γ", "计算Qwarm,20的固定边界断面", "如140°W或5°N"],
                  ["vₙ", "垂直于断面Γ的速度分量", "m s⁻¹；正向需随断面声明"],
                  ["Qwarm,20,Γ", "穿过Γ的20°C阈值暖水体热能输送；不是传统总海洋热输送", "W/PW；时间积分后为J/ZJ"],
                  ["x、y、d、t", "经度、纬度、向下深度和时间坐标", "度、度、m、月"],
              ], [1500, 5800, 2060], font_size=8.8)
    add_body(doc, (
        "式（3）说明H₂₀与WWV之间并非任意换算：H₂₀等于暖水体积、平均超温、密度与比热的乘积。"
        "式（4）则把单位体积的阈值热能与法向流速相乘并对断面积分，量纲为功率。"
        "本文统一称其为‘20°C阈值暖水体热能输送’，避免与ρ₀cₚ∬vₙT dA形式的传统总海洋热输送混淆。"
    ))

    add_heading(doc, "2.4 固定控制体预算、符号与阈值非线性", 2)
    add_body(doc, (
        "控制体范围为10°S—5°N、0—300 m，并按经度划分为西箱160°E—180°、中箱180°—140°W和东箱140°W—80°W。"
        "为避免纬向向东为正、经向向北为正与箱体收支符号混用，所有预算边界统一采用‘入箱为正’。"
    ))
    add_equation(doc, "FΓ=−ρ₀cₚ ∬Γ (v·nout)χ₂₀(T)dA,   FΓ>0 表示进入箱体", 5)
    add_equation(doc, "dH₂₀/dt=Cₓ+F₁₀S+F₅N+Qsurface,warm+R₂₀", 6)
    add_equation(doc, "Qsurface,warm=∬A I₂₀(Ts)Qnet dA,   Qnet>0 表示海洋吸热", 7)
    add_equation(doc, "Dχ₂₀/Dt=I₂₀(T)·DT/Dt   (T≠Tᵣ)", 8)
    add_table(doc, "表3  式（5）—（8）的全部符号、物理意义与符号约定",
              ["符号", "物理意义", "单位/约定"],
              [
                  ["FΓ", "边界Γ对箱体H₂₀预算的暖量输入项", "W/PW；FΓ>0入箱，FΓ<0出箱"],
                  ["v", "SODA3三维海流速度向量", "m s⁻¹"],
                  ["nout", "边界Γ的箱体外法向单位向量", "无量纲"],
                  ["v·nout", "海流在外法向上的速度；正值表示流出", "m s⁻¹"],
                  ["负号−", "把外向速度转换为预算中的入箱符号", "符号约定"],
                  ["Cₓ", "东、西边界暖量输入项之和，即纬向辐合", "W/PW"],
                  ["F₁₀S、F₅N", "10°S南界和5°N北界的入箱暖量项", "W/PW；展示输出时取相反号"],
                  ["dH₂₀/dt", "箱体阈值相对热含量的时间变化率", "W/PW"],
                  ["A", "箱体海表面积", "m²"],
                  ["Ts", "SODA3最上层温度，用作海表阈值掩膜", "°C"],
                  ["Qnet", "进入海洋为正的净表面热通量", "W m⁻²"],
                  ["Qsurface,warm", "作用于Ts>20°C海表区域的净热输入", "W/PW；负值表示热损失"],
                  ["R₂₀", "由储量倾向减去全部显式项得到的诊断剩余", "W/PW"],
                  ["D/Dt", "随流体运动的物质导数", "s⁻¹乘以被导变量单位"],
                  ["T≠Tᵣ", "链式法则在非阈值点的适用条件", "阈值面需单独处理"],
              ], [1500, 5800, 2060], font_size=8.8)
    add_body(doc, (
        "式（6）不是把χ₂₀假定为线性保守示踪剂后直接得到的无剩余守恒律，而是以R₂₀定义完成的离散诊断恒等式。"
        "式（8）表明，在T不等于阈值且不可压缩的绝热平流中，阈值算子遵循链式法则；"
        "但扩散、混合、夹卷和外部加热使水体跨越20°C等温面时会产生非线性水团转换。"
        "月平均资料无法把这些过程逐项解析，故R₂₀共同包含跨阈值转换、垂向和斜压交换、混合扩散、涡动协方差、同化增量以及时间—空间离散误差，不能预先命名为某一种洋流。"
    ))
    add_body(doc, (
        "式（7）采用表层温度I₂₀(Ts)掩膜，是月平均有限层厚资料下的表面阈值近似。"
        "若改用I(D20>0)，会把作用在低于20°C表层的热通量直接计入H₂₀，因此只作为暖层足迹敏感性，不作为主定义。"
    ))

    add_heading(doc, "2.5 FEP残差与海气接通时序", 2)
    add_body(doc, (
        "FEP定义为5°S—5°N、90°W—80°W平均SST异常。为了分离最东部局地增暖与同期盆地尺度Niño-3.4共同变化，"
        "对每个日历月分别使用1982—2024年的43个年度样本进行OLS回归。"
    ))
    add_equation(doc, "FEP′(m,y)=aₘ+bₘN₃.₄′(m,y)+ε(m,y)", 9)
    add_equation(doc, "FEPres(m,y)=ε(m,y)=FEP′(m,y)−[aₘ+bₘN₃.₄′(m,y)]", 10)
    add_table(doc, "表4  式（9）—（10）的全部符号与统计含义",
              ["符号", "物理/统计含义", "单位/约定"],
              [
                  ["FEP′(m,y)", "年份y、日历月m的去趋势FEP SST异常", "°C；撇号′表示异常"],
                  ["N₃.₄′(m,y)", "同月去趋势Niño-3.4 SST异常", "°C"],
                  ["m、y", "日历月索引和年份索引", "m=1…12，y=1982…2024"],
                  ["aₘ", "日历月m独立回归的截距", "°C"],
                  ["bₘ", "FEP对Niño-3.4的同月回归斜率", "无量纲"],
                  ["ε(m,y)", "回归未解释部分", "°C"],
                  ["FEPres", "控制Niño-3.4同期幅度后的FEP统计残差", "°C；不解释为独立物理强迫"],
                  ["[ ]", "回归期望响应aₘ+bₘN₃.₄′", "°C"],
              ], [1700, 5700, 1960], font_size=8.9)
    add_body(doc, (
        "操作性海洋转变月定义为东部减西部的0—300 m OHC异常和D20异常连续两个月同时为正的首月；"
        "操作性大气闭合月定义为中西太平洋u850西风异常为正且OLR异常为负、并连续维持两个月的首月。"
        "这些判据只用于三个事件的过程对齐，不作为普适ENSO阈值。"
    ))

    add_heading(doc, "2.6 5°N地转—Ekman诊断分解", 2)
    add_body(doc, (
        "在5°N、140°W—80°W、0—300 m断面上，规定北向为正。地转经向速度由海表高度梯度和斜压密度梯度共同重建。"
        "坐标x向东，深度d=−z向下为正；去除水平一致的参考压力后，有："
    ))
    add_equation(doc, "fvg=(1/ρ₀)∂p/∂x", 11)
    add_equation(doc, "p(x,d)=ρ₀gη(x)+g∫₀ᵈρ′(x,d′)dd′", 12)
    add_equation(doc, "vg(d)=(g/f)∂η/∂x+[g/(ρ₀f)]∫₀ᵈ∂ρ′(x,d′)/∂x dd′", 13)
    add_equation(doc, "My=−τx/(ρ₀f)", 14)
    add_equation(doc, "Qwarm,20,Ekman=ρ₀cₚ∫x My(x)·χ̄₂₀,0–50m(x)dx", 15)
    add_equation(doc, "Qwarm,20,total=Qwarm,20,geo+Qwarm,20,Ekman+Qwarm,20,res", 16)
    add_table(doc, "表5  式（11）—（16）的全部符号、物理意义和正方向",
              ["符号", "物理意义", "单位/约定"],
              [
                  ["f", "5°N处科氏参数", "s⁻¹；北半球为正"],
                  ["vg", "由压力梯度重建的地转经向速度", "m s⁻¹；北向为正"],
                  ["p", "去除水平一致参考项后的海水压力", "Pa"],
                  ["∂/∂x", "沿向东坐标x的偏导数", "每米变化率"],
                  ["g", "重力加速度", "9.81 m s⁻²"],
                  ["η", "相对于基准的海表高度", "m"],
                  ["d、d′", "向下为正的目标深度和积分深度变量", "m"],
                  ["ρ′", "密度异常ρ−ρref(d)；只有水平密度差进入斜压压力梯度", "kg m⁻³"],
                  ["ρref(d)", "仅随深度变化的参考密度剖面", "kg m⁻³"],
                  ["My", "Ekman层垂向积分的经向体积输送（每单位 zonal 宽度）", "m² s⁻¹；北向为正"],
                  ["τx", "纬向风应力", "N m⁻²；向东为正"],
                  ["负号−", "北半球Ekman输送位于风应力右侧；东向τx对应南向My", "符号约定"],
                  ["χ̄₂₀,0–50m", "0—50 m层内χ₂₀的深度平均", "K"],
                  ["dx", "沿5°N断面的 zonal 长度微元", "m"],
                  ["Qwarm,20,geo", "以vg代入式（4）得到的地转暖量输送", "W/PW；北向为正"],
                  ["Qwarm,20,Ekman", "由My和上层平均超温得到的一阶Ekman暖量输送", "W/PW；北向为正"],
                  ["Qwarm,20,total", "以SODA3总经向速度计算的暖量输送", "W/PW；北向为正"],
                  ["Qwarm,20,res", "总量减地转和Ekman后的诊断残差", "W/PW；非独立动力模态"],
              ], [1700, 5660, 2000], font_size=8.55)
    add_body(doc, (
        "密度项写为ρ′而非ρ，是为了明确水平一致的背景密度不产生斜压压力梯度；由于ρref仅随深度变化，实际数值上∂ρ′/∂x等于∂ρ/∂x。"
        "式（16）是诊断恒等式，不是完整动力模态分解；Qwarm,20,res包含年龄转流、瞬变涡动、惯性过程以及Ekman层近似误差。"
        "在转换为东部箱体预算时，5°N北向输出为正，对应F₅N=−Qwarm,20,total。"
    ))

    add_heading(doc, "2.7 离线换风敏感性试验", 2)
    add_body(doc, (
        "为定量估计局地风驱对诊断出口的贡献，在不重新积分海洋或大气的条件下，固定某一事件的温度、地转项和残差，仅替换5°N局地风应力："
    ))
    add_equation(doc, "Qrec(i←j)=Qgeo(i)+Qres(i)+QEkman[T(i),τx(j)]", 17)
    add_table(doc, "表6  式（17）的全部符号与试验边界",
              ["符号", "含义", "约定/边界"],
              [
                  ["Qrec(i←j)", "事件i海洋状态接受事件j风应力后的重构总出口", "ZJ；箭头表示j的风放入i"],
                  ["i", "被固定的海洋事件索引", "i∈{1997,2015,2023}"],
                  ["j", "提供替换风应力的事件索引", "j∈{1997,2015,2023}"],
                  ["Qgeo(i)", "事件i原有的地转暖量输送", "固定不响应新风"],
                  ["Qres(i)", "事件i原有的诊断残差", "固定不响应新风"],
                  ["T(i)", "事件i的上层温度及χ̄₂₀权重", "固定海洋热状态"],
                  ["τx(j)", "事件j的5°N局地纬向风应力", "唯一被替换的外部输入"],
                  ["QEkman[T(i),τx(j)]", "以i的暖层结构和j的风计算的Ekman项", "离线一阶重构"],
              ], [1800, 5300, 2260], font_size=8.9)
    add_body(doc, (
        "该试验不允许温度、SSH、密度、压力场和残差对新风场作动力调整，因此只能称为离线反事实敏感性试验，不能称为完整因果检验。"
    ))

    add_heading(doc, "3 结果", 1)
    add_heading(doc, "3.1 暖量确实进入东太平洋，但事件间到达强度不同", 2)
    add_body(doc, (
        "赤道0—300 m OHC经度—时间演变显示，1997、2015和2023均存在向东传播或东移的次表层暖异常，但盆地尺度支撑明显不同（图1）。"
        "1997的暖异常最强且从西—中太平洋连续向东扩展；2015仍具有清晰的中东太平洋暖层调整；2023则较早出现最东部表面增暖，但中太平洋深层暖量支撑较弱。"
    ))
    add_figure(doc, "Fig01_eastward_OHC_arrival_hovmoller.png",
               "图1  1997/98、2015/16和2023/24赤道太平洋0—300 m OHC异常的经度—时间演变。图由本研究基于海洋再分析资料绘制。",
               width_in=6.2)
    add_body(doc, (
        "CT/CP合成进一步表明，两类事件都能把暖量送入东部。Jun(0)—Nov(0)期间，CT跨140°W东向Qwarm,20累计为6.25 ZJ，"
        "95%区间为[2.91, 9.60] ZJ；CP为2.97 ZJ，区间为[1.91, 4.03] ZJ。同期东部H₂₀分别增加7.81和3.71 ZJ。"
        "因此类型差异首先表现为到达的强度、持续性和垂向结构，而不是简单的‘到达/未到达’二分。"
    ))
    add_figure(doc, "Fig02_CT_CP_three_box_arrival_and_fate.png",
               "图2  CT与CP事件的三箱暖量到达和成熟后分配。上排为主要断面Qwarm,20，下排为累计储量、边界、海表和剩余项。",
               width_in=6.2)
    add_table(doc, "表7  三事件发展年1—6月的东部暖量到达指标",
              ["事件", "140°W东向Qwarm,20 (ZJ)", "东部H₂₀变化 (ZJ)", "东部SST (°C)", "东部D20 (m)"],
              [
                  ["1997", "+12.13", "+9.29", "+0.35", "+17.27"],
                  ["2015", "+3.30", "+3.53", "+0.98", "+15.09"],
                  ["2023", "+3.91", "+3.88", "+0.93", "+12.03"],
              ], [1100, 2600, 2100, 1700, 1860], font_size=9.0)
    add_body(doc, (
        "1997跨140°W输入和东部储热约为2015/2023的三倍。更关键的是，2015和2023的发展期输入十分接近："
        "二者跨140°W输入分别为3.30和3.91 ZJ，东部H₂₀增加3.53和3.88 ZJ。"
        "由此可以排除‘2023只是最东部表面局地变暖、暖量并未进入东部’这一解释。"
    ))
    add_lead(doc, "问题一的答案：", "暖量确实通过赤道上层海洋和140°W通道进入东太平洋；1997输入最强，2015与2023的输入近似。后续差异不能用‘暖量来没来’解释。")

    add_heading(doc, "3.2 到达后的主要分配通道是5°N、海表热损失和储量/剩余过程", 2)
    add_body(doc, "为便于讨论成熟—衰减期放电，结果部分统一将5°N北向输出定义为正；该展示符号与控制体预算中的入箱项F₅N相反，即5°N北向输出=−F₅N。")
    add_figure(doc, "Fig03_event_arrival_export_surface_loss.png",
               "图3  1997、2015和2023的发展期输入、东部储热、次年5°N输出及海表热损失。三事件形成强排放—部分排放—近闭合阶梯。",
               width_in=6.15)
    add_table(doc, "表8  CT/CP成熟—衰减期东部箱体暖量分配（ZJ）",
              ["类型", "H₂₀变化", "5°N北向输出", "10°S南向输出", "海表热损失", "预算剩余"],
              [
                  ["CT", "−7.86", "+2.74", "+0.68", "+3.69", "+2.09"],
                  ["CP", "−3.03", "−0.96", "−0.48", "+1.76", "+0.82"],
              ], [1050, 1500, 1850, 1750, 1650, 1560], font_size=8.8)
    add_body(doc, (
        "Dec(0)—May(+1)期间，CT东部H₂₀减少7.86 ZJ，其中5°N北向输出2.74 ZJ、海表热损失3.69 ZJ；"
        "CP东部H₂₀减少3.03 ZJ，但5°N项为−0.96 ZJ，即合成意义上仍向东部箱体输入。"
        "CT−CP的5°N北向输出差为3.70 ZJ，bootstrap 95%区间为[1.65, 5.65] ZJ。"
        "10°S项的组内区间跨零，说明南界不是稳健主出口。ERA5独立估算同样显示CT海表热损失显著强于CP。"
    ))
    add_table(doc, "表9  三事件次年1—5月的东部暖量命运（ZJ）",
              ["事件", "5°N总输出", "海表热损失", "东部H₂₀变化", "4.75°N单边估计"],
              [
                  ["1997", "+3.81", "+4.37", "−9.73", "+3.85"],
                  ["2015", "+0.88", "+1.16", "−11.10", "+0.92"],
                  ["2023", "+0.10", "+0.40", "−9.14", "−1.24"],
              ], [1100, 1900, 1900, 2050, 2410], font_size=8.9)
    add_body(doc, (
        "1997同时具有强北向输出和强海表热损失；2015两者均减弱但仍有净排放；2023的居中通量面5°N输出仅0.10 ZJ，海表热损失为0.40 ZJ。"
        "三事件的东部H₂₀都明显下降，但下降量并不按出口强度单调排列，说明储量减少不能等同于某一条具体去向。"
        "2023的5°N净方向对通量面温度插值敏感：采用4.75°N单边温度时为−1.24 ZJ，采用4.75°N与5.25°N居中温度时为+0.10 ZJ。"
        "因此最稳健的表述是‘出口显著受抑或近闭合’，而不是断言其必然净向南。"
    ))
    add_lead(doc, "问题二的答案：", "到达东部的暖量主要经5°N北向动力出口、海表热损失、上层海洋储量变化和R₂₀所代表的未显式过程重新分配；10°S不是稳健主出口。")

    add_heading(doc, "3.3 相似FEP先暖为何产生不同暖量命运", 2)
    add_heading(doc, "3.3.1 表面起点相似，不代表次表层海洋状态相同", 3)
    add_figure(doc, "Fig05_FEP_subsurface_preconditioning.png",
               "图4  1997、2015和2023的FEP、0—300 m OHC与D20过程。2023最东部表面先暖，但中太平洋OHC和温跃层东移支撑最弱。",
               width_in=6.2)
    add_body(doc, (
        "2023年3—4月FEP异常最强，但其3月中太平洋OHC仅约0.39 GJ m⁻²，1997为1.57 GJ m⁻²；"
        "赤道带纬向积分正OHC异常在1997和2023分别为10.94和3.42 PJ m⁻¹。"
        "同时，1997东—西D20差已经转为+5.03 m，而2023仍为−10.94 m。"
        "在13个暖事件中，给定FEP暖度后的2023中太平洋OHC残差为−0.61 GJ m⁻²，排名最低。"
        "因此2023最突出的异常不是‘表面不暖’，而是FEP暖度与盆地尺度次表层预置严重不匹配。"
    ))

    add_heading(doc, "3.3.2 决定性分叉发生在表面增暖之后、深对流东移之前", 3)
    add_figure(doc, "Fig06_FEP_atmospheric_closure.png",
               "图5  FEP—OHC—D20—u850—OLR的三事件过程闭合。1997和2015较早建立海洋—大气接通，2023到5月才弱接通。",
               width_in=6.3)
    add_body(doc, (
        "1997在3月已经出现明显中西太平洋西风和负OLR异常，表面增暖迅速连接深对流和风场；2015的地下海洋条件在1月或更早预置，大气在3月闭合。"
        "2023的FEP残差从2月持续为正，但3—4月中太平洋OLR仍为正，地下海洋和大气到5月才同时换符号，且幅度和持续性均弱。"
        "5—6月平均u850在1997、2015和2023中分别为5.25、3.37和0.29 m s⁻¹，相应的−OLR为29.90、20.49和7.63 W m⁻²。"
        "这表明FEP先暖只是表面触发，只有当暖异常进入能够组织中东太平洋深对流和气压梯度的背景，才会形成持续风驱。"
    ))

    add_heading(doc, "3.3.3 大气接通差异通过Ekman与地转诊断分量的合成落实为出口差异", 3)
    add_figure(doc, "Fig04_event_5N_dynamical_decomposition.png",
               "图6  1997、2015和2023次年1—5月5°N总Qwarm,20的地转、Ekman和残差诊断分解。",
               width_in=6.15)
    add_table(doc, "表10  三事件次年1—5月5°N Qwarm,20诊断分解（ZJ）",
              ["事件", "总量", "Ekman", "总地转", "SSH地转", "密度地转", "残差"],
              [
                  ["1997", "+3.81", "+3.25", "+0.59", "+0.36", "+0.23", "−0.03"],
                  ["2015", "+0.88", "+2.44", "−1.28", "−1.43", "+0.15", "−0.28"],
                  ["2023", "+0.10", "+1.56", "−1.27", "−1.32", "+0.05", "−0.20"],
              ], [950, 1200, 1300, 1400, 1400, 1600, 1510], font_size=8.6)
    add_body(doc, (
        "1997的Ekman和地转项均为北向；2015的总诊断出口由+2.44 ZJ的北向Ekman项、−1.28 ZJ的南向地转项与−0.28 ZJ残差合成；"
        "2023的南向地转项为−1.27 ZJ，与2015几乎相同，而Ekman项只有1.56 ZJ，合成后的总诊断出口接近零。"
        "2015相对2023多出的净输出为0.79 ZJ，Ekman差为0.88 ZJ，相当于净差的112%，其余分量部分抵消该差异。"
        "因此2015与2023的直接差异不是地转压力场强弱，而是Ekman暖量输送强度不同，使总诊断输出在相似南向地转项背景下显著分离。"
    ))

    add_figure(doc, "Fig07_1997_2023_coupled_mechanism_synthesis.png",
               "图7  1997与2023从SST—风场、暖量释放、OHC/D20预置到深对流、降水和SLP梯度的全过程对照。",
               width_in=5.5, page_break=True)

    add_heading(doc, "3.3.4 十三事件统计与换风试验提供定量过程证据", 3)
    add_figure(doc, "Fig08_third_question_coupled_attribution.png",
               "图8  十三个暖事件的FEP—OHC—u850—次年5°N Ekman联系，以及1997、2015和2023的离线风应力替换试验。",
               width_in=6.15)
    add_body(doc, (
        "13个暖事件中，3—4月中太平洋OHC与5—6月中西太平洋u850的相关为r=0.89，事件bootstrap 95%区间为[0.39, 0.97]，"
        "逐一删事件后的r范围为[0.79, 0.93]。5—6月u850与次年1—5月5°N Ekman Qwarm,20的相关为r=0.75，"
        "95%区间为[0.18, 0.92]，逐一删事件范围为[0.57, 0.80]。"
        "这些结果支持‘次表层预置—西风接通—Ekman出口’的跨事件过程联系，但不把相关解释为单向因果。"
    ))
    add_table(doc, "表11  固定海洋状态、替换5°N风应力后的重构总Qwarm,20（ZJ）",
              ["固定海洋状态", "1997风", "2015风", "2023风"],
              [
                  ["1997", "+3.81", "+2.70", "+1.86"],
                  ["2015", "+2.14", "+0.88", "+0.08"],
                  ["2023", "+2.10", "+0.88", "+0.10"],
              ], [2600, 2250, 2250, 2260], font_size=9.1)
    add_body(doc, (
        "保持2023温度、地转项和残差不变，使用2023自身风时总出口为0.10 ZJ；换入2015风后增至0.88 ZJ，换入1997风后增至2.10 ZJ。"
        "反向把2023风放入2015海洋状态，出口由0.88降至0.08 ZJ。"
        "在2015与2023两种固定海洋状态下，风应力替换均能显著改变重构出口值，说明局地风驱是诊断出口差异的重要一阶贡献；"
        "但由于海洋状态不随新风调整，这一结果仍是敏感性证据，而不是完整耦合因果试验。"
    ))
    add_lead(doc, "问题三的答案：", "FEP先暖只是表面起点。盆地尺度次表层暖量库存和D20/SSH结构决定可用的海洋预置，大气背景决定深对流与风场能否接通；局地Ekman暖量输送的强弱最终对应5°N总诊断北向输出的大小。")

    add_heading(doc, "4 讨论", 1)
    add_heading(doc, "4.1 与既有充放电及2023事件研究的关系", 2)
    add_body(doc, (
        "Jin（1997）的充放电振子以盆地平均温跃层或WWV表征ENSO慢记忆，本文不替代这一理论，而是在其基础上增加固定控制体的能量收支："
        "把盆地尺度蓄热进一步拆成暖量何时跨140°W到达、成熟后经何种边界和海表过程离开。"
        "Chakravorty等（2021）回答了CT与CP充放电过程是否不同，本文则把标准化暖水体积输送推进到具有能量量纲的H₂₀与Qwarm,20，并显式保留阈值非线性剩余项。"
        "新增贡献不在于重新定义WWV，而在于建立WWV→H₂₀→Qwarm,20→三箱去向→地转/Ekman诊断分量→离线换风的连续诊断链。"
        "这一框架区分了三个常被混淆的层次：跨断面输入不等于箱体储存，箱体储量下降不等于单一出口，表面先暖也不等于海气耦合已经接通。"
    ))
    add_body(doc, (
        "Peng等（2025）强调海洋动力对2023/24事件生成的重要作用；本文结果与此并不矛盾，因为2023的发展期暖量确实跨140°W进入东部，且输入与2015接近。"
        "本文补充的是到达之后的阶段：相似输入并未产生相似的次表层预置—深对流—风场接通和5°N诊断出口。"
        "因此，‘事件如何生成’与‘到达后的暖量如何分配’是相互衔接但不能彼此替代的两个问题。"
    ))

    add_heading(doc, "4.2 ‘暖量去哪了’的现实意义", 2)
    add_body(doc, (
        "暖量去向决定厄尔尼诺如何衰减、热量如何从赤道输出以及海表异常能维持多久。"
        "如果暖量通过5°N和海表热损失快速释放，事件更容易形成明确的放电和相位转换；若出口近闭合，热量可能更多留在上层海洋、通过垂向混合或跨阈值转换重新分配。"
        "这一差异会影响后续季节的SST持续性、热带对流加热位置以及遥相关强度。"
        "因此，5°N Qwarm,20、Ekman与地转诊断分量的相对贡献和海表热损失不仅是事后诊断，也可作为海洋再分析和耦合模式评估的过程指标。"
    ))

    add_heading(doc, "4.3 结论的稳健部分与不能过度声称的部分", 2)
    add_body(doc, (
        "两种距平基准下，三事件总出口均保持1997>2015>2023的排序，Ekman和地转符号不变；线性状态方程也不改变地转方向。"
        "最稳健的结论是到达、分配和诊断分量合成结果的相对排序，而不是单一绝对值。"
        "同时，本文存在五项边界：第一，欧拉控制体预算不能识别同一水团的连续拉格朗日轨迹；第二，R₂₀尚未拆分混合、夹卷、跨等温面转换和同化增量；"
        "第三，2024年5°N净方向对通量面温度插值敏感；第四，13事件样本仍小；第五，换风试验不允许压力场和残差响应新风。"
    ))
    add_body(doc, (
        "现有证据把2023失配位置锁定在FEP表面增暖之后、深对流东移之前，但不能把外海盆增暖、多年LaNiña遗产、冷NPMM、降水偶极或WWB不足中的某一个指定为本文已经独立证明的唯一上游根因。"
        "下一步最有价值的工作不是继续堆叠相关图，而是利用五日三维热收支拆分R₂₀，并在海洋模式或耦合模式中关闭或替换5°N异常风应力，检验第三问的因果充分性。"
    ))

    add_heading(doc, "5 结论", 1)
    add_body(doc, (
        "本文基于SODA3建立20°C阈值相对热含量和暖水体热能输送诊断，并用独立海气资料回答了暖量到达、去向与耦合命运三个递进问题。主要结论如下。"
    ))
    conclusions = [
        "暖量确实跨140°W进入东太平洋。CT与CP均可到达，但CT输入和储存更强；1997输入最强，2015与2023的发展期输入近似。",
        "到达后的主要分配通道为5°N北向动力出口、海表热损失、储量变化以及预算剩余过程；10°S不是稳健主出口。1997、2015和2023分别表现为强排放、部分排放和出口近闭合。",
        "2015与2023具有近似的南向地转项，但2015的风驱Ekman暖量输送更强，使合成后的总诊断出口保持北向；2023的总诊断出口则接近零。",
        "FEP先暖不是充分条件。盆地尺度次表层预置决定大气能否有效接通，深对流—风场响应决定Ekman出口强度；离线换风试验定量支持局地风驱的一阶作用，但不等同于完整因果检验。",
    ]
    for idx, item in enumerate(conclusions, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"（{idx}）"), size=11, bold=True, east_asia=FONT_HEAD)
        set_run_font(p.add_run(item), size=11)

    add_heading(doc, "数据与可重复性说明", 1)
    add_body(doc, (
        "SODA3、OISST、ERA5、GPCP和CPC OLR均为公开资料。本文图表由本地分析脚本生成；"
        "直接计算模块包括三箱暖量预算、CT/CP海气耦合、三事件5°N诊断分解及第三问事件级归因。"
        "正式投稿前应将派生数据、事件表、绘图脚本和环境文件整理为带版本号的公开存档。"
    ))

    add_heading(doc, "参考文献", 1)
    refs = [
        "Jin, F.-F. (1997). An equatorial ocean recharge paradigm for ENSO. Part I: Conceptual model. Journal of the Atmospheric Sciences, 54, 811–829. https://doi.org/10.1175/1520-0469(1997)054<0811:AEORPF>2.0.CO;2",
        "Meinen, C. S., & McPhaden, M. J. (2000). Observations of warm water volume changes in the equatorial Pacific and their relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559. https://doi.org/10.1175/1520-0442(2000)013<3551:OOWWVC>2.0.CO;2",
        "Sun, D.-Z. (2000). The heat sources and sinks of the 1986–87 El Niño. Journal of Climate, 13, 3533–3550. https://doi.org/10.1175/1520-0442(2000)013<3533:THSASO>2.0.CO;2",
        "Bosc, C., & Delcroix, T. (2008). Observed equatorial Rossby waves and ENSO-related warm water volume changes in the equatorial Pacific Ocean. Journal of Geophysical Research: Oceans, 113, C06003. https://doi.org/10.1029/2007JC004613",
        "Kug, J.-S., Jin, F.-F., & An, S.-I. (2009). Two types of El Niño events: Cold Tongue El Niño and Warm Pool El Niño. Journal of Climate, 22, 1499–1515. https://doi.org/10.1175/2008JCLI2624.1",
        "Capotondi, A., et al. (2015). Understanding ENSO diversity. Bulletin of the American Meteorological Society, 96, 921–938. https://doi.org/10.1175/BAMS-D-13-00117.1",
        "Carton, J. A., Chepurin, G. A., & Chen, L. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983. https://doi.org/10.1175/JCLI-D-17-0149.1",
        "Adler, R. F., et al. (2018). The Global Precipitation Climatology Project monthly analysis: An update. Atmosphere, 9, 138. https://doi.org/10.3390/atmos9020138",
        "Hersbach, H., et al. (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146, 1999–2049. https://doi.org/10.1002/qj.3803",
        "Huguenin, M. F., Holmes, R. M., & England, M. H. (2020). Key role of diabatic processes in regulating warm water volume variability over ENSO events. Journal of Climate, 33, 9945–9964. https://doi.org/10.1175/JCLI-D-20-0198.1",
        "Chakravorty, S., Perez, R. C., Gnanaseelan, C., & Anderson, B. T. (2021). Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 126, e2020JC017075. https://doi.org/10.1029/2020JC017075",
        "Huang, B., et al. (2021). Improvements of the Daily Optimum Interpolation Sea Surface Temperature (DOISST) version 2.1. Journal of Climate, 34, 2923–2939. https://doi.org/10.1175/JCLI-D-20-0166.1",
        "Li, X., Hu, Z.-Z., Huang, B., & Jin, F.-F. (2022). Oceanic meridional transports and their roles in warm water volume variability and ENSO in the tropical Pacific. Climate Dynamics, 59, 245–261. https://doi.org/10.1007/s00382-021-06124-w",
        "Verdy, A., Mazloff, M. R., Cornuelle, B. D., & Kim, S. B. (2023). Balancing volume, temperature, and salinity budgets during 2014–2018 in the Tropical Pacific Ocean State Estimate. Journal of Geophysical Research: Oceans, 128, e2022JC019576. https://doi.org/10.1029/2022JC019576",
        "Hong, C.-C., Sullivan, A., & Chang, C.-C. (2024). Impact of North Atlantic Tripole and Extratropical North Pacific extreme SSTs on the 2023/24 El Niño. Geophysical Research Letters, 51, e2024GL110639. https://doi.org/10.1029/2024GL110639",
        "Kim, J.-W., Tian, B., & Yu, J.-Y. (2025). Unusual and persistent easterlies restrained the 2023/24 El Niño development after a triple-dip La Niña. npj Climate and Atmospheric Science, 8, 34. https://doi.org/10.1038/s41612-024-00890-0",
        "Lin, Y.-F., et al. (2025). Diverse NPMM conditions deviate the 2023/24 El Niño from the 1997/1998 and 2015/2016 extreme El Niño events. npj Climate and Atmospheric Science, 8, 133. https://doi.org/10.1038/s41612-025-01013-z",
        "Peng, Q., et al. (2025). Strong 2023–2024 El Niño generated by ocean dynamics. Nature Geoscience, 18, 471–478. https://doi.org/10.1038/s41561-025-01700-9",
        "Zhang, L., et al. (2025). The 2023/24 El Niño event exhibited unusually weak extratropical teleconnections. Communications Earth & Environment, 6, 595. https://doi.org/10.1038/s43247-025-02584-8",
        "Fan, H., Lin, S., Zhang, L., et al. (2026). Multi-year La Niña dampens the Southern Oscillation during subsequent El Niño through tropical Pacific-rooted dynamics. The Innovation Geoscience, 4, 100224. https://doi.org/10.59717/j.xinn-geo.2026.100224",
    ]
    for i, ref in enumerate(refs, 1):
        add_reference(doc, i, ref)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
