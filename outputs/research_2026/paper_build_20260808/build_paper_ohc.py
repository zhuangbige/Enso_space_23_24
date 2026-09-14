#!/usr/bin/env python3
"""Build the OHC-only Chinese ENSO short paper as a new DOCX."""

from __future__ import annotations

from pathlib import Path
import os
import sys

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(HERE))
import build_paper as base  # noqa: E402

# Monochrome single-family override. ``OHC_PAPER_FONT`` is only used for
# renderer compatibility checks; the deliverable defaults to Arial Unicode MS.
PAPER_FONT = os.environ.get("OHC_PAPER_FONT", "Arial Unicode MS")
base.FONT_CN = PAPER_FONT
base.FONT_HEAD = PAPER_FONT
base.FONT_LATIN = PAPER_FONT
_base_set_run_font = base.set_run_font


def _set_run_font_unified(
    run,
    size=11,
    bold=False,
    italic=False,
    color=base.INK,
    latin=None,
    east_asia=None,
):
    return _base_set_run_font(
        run,
        size=size,
        bold=bold,
        italic=italic,
        color=color,
        latin=latin or PAPER_FONT,
        east_asia=east_asia or PAPER_FONT,
    )


base.set_run_font = _set_run_font_unified

_base_add_figure = base.add_figure


def _add_figure_with_alt(doc, filename, caption, width_in=6.25, page_break=False):
    """Add a figure and use its scientific caption as accessible alt text."""
    paragraph = _base_add_figure(
        doc,
        filename,
        caption,
        width_in=width_in,
        page_break=page_break,
    )
    doc_properties = paragraph._p.xpath(".//wp:docPr")
    if doc_properties:
        doc_properties[0].set("title", caption.split("。", 1)[0])
        doc_properties[0].set("descr", caption)
    return paragraph


base.add_figure = _add_figure_with_alt


OUT = Path(os.environ.get(
    "OHC_PAPER_OUT",
    ROOT / "outputs/research_2026/papers/ENSO东太平洋OHC到达_释放与海气耦合命运_中文小论文_OHC统一版_20260808.docx",
))
PHASE_CSV = ROOT / "data/derived/research_2026/soda3_ohc_only_phase_summary_13events.csv"
TYPE_CSV = ROOT / "data/derived/research_2026/soda3_ohc_only_ct_cp_summary.csv"
CORR_CSV = ROOT / "data/derived/research_2026/soda3_ohc_only_coupling_correlations.csv"
CHAIN_CSV = ROOT / "data/derived/research_2026/soda3_ohc_only_three_event_coupling_chain.csv"
SWAP_CSV = ROOT / "data/derived/research_2026/soda3_ohc_2015_2023_wind_swap_summary.csv"

FIG_HOV = ROOT / "outputs/research_2026/three_questions_figure_package_20260805/01_main_figures/Fig01_eastward_OHC_arrival_hovmoller.png"
FIG_EVENT = ROOT / "outputs/research_2026/ohc_only_three_questions/Fig02_OHC_event_storage_release_transport.png"
FIG_TYPE = ROOT / "outputs/research_2026/ohc_only_three_questions/Fig03_OHC_CT_CP_comparison.png"
FIG_CLOSURE = ROOT / "outputs/research_2026/ohc_only_three_questions/Fig04_OHC_atmospheric_closure.png"
FIG_COUPLING = ROOT / "outputs/research_2026/soda3_ohc_minimal/soda3_ohc_coupling_minimal.png"
FIG_SWAP = ROOT / "outputs/research_2026/ohc_only_three_questions/soda3_ohc_2015_2023_wind_swap.png"

EVENTS = (1997, 2015, 2023)
REFERENCES = (0, 10, 20)


def configure_document(doc: Document) -> None:
    """Narrative-proposal preset with a monochrome Chinese academic override."""
    base.configure_document(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.45)
    section.bottom_margin = Cm(2.35)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.55)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = base.FONT_LATIN
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(base.INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), base.FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT_CN)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 12, 6),
        "Heading 3": (12, 8, 4),
    }
    for name, (size, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = base.FONT_HEAD
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(base.INK)
        style._element.rPr.rFonts.set(qn("w:ascii"), base.FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT_HEAD)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True

    doc.core_properties.title = "厄尔尼诺东太平洋上层海洋热含量的到达、释放与海气耦合命运"
    doc.core_properties.subject = "基于SODA3的0—300 m OHC与标准海洋热输送诊断"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "ENSO; SODA3; OHC; ocean heat transport; Ekman; air-sea coupling"


def add_title_block(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    base.set_run_font(
        paragraph.add_run("厄尔尼诺东太平洋上层海洋热含量的到达、释放与海气耦合命运"),
        size=18,
        bold=True,
        color=base.INK,
        east_asia=base.FONT_HEAD,
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(18)
    base.set_run_font(
        paragraph.add_run("——基于 SODA3 的 0—300 m OHC、标准热输送及 1997、2015、2023 事件对比"),
        size=11.5,
        bold=True,
        color=base.INK,
        east_asia=base.FONT_HEAD,
    )


def fmt(value: float) -> str:
    return f"{float(value):+.2f}"


def range_text(row: pd.Series, pattern: str) -> str:
    values = [float(row[pattern.format(reference=reference)]) for reference in REFERENCES]
    return f"{min(values):+.2f}—{max(values):+.2f}"


def type_row(types: pd.DataFrame, event_class: str, phase: str, metric: str) -> pd.Series:
    return types.loc[
        types.event_class.eq(event_class)
        & types.phase.eq(phase)
        & types.metric.eq(metric)
    ].iloc[0]


def add_abstract(doc: Document, development: pd.DataFrame, decay: pd.DataFrame) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(7)
    base.set_run_font(paragraph.add_run("摘要："), size=10.5, bold=True, east_asia=base.FONT_HEAD)
    text = (
        "热带太平洋上层海洋热含量是厄尔尼诺—南方涛动海洋记忆的重要表征，但盆地尺度暖水体积和标准化输送不能直接区分东太平洋的热储存、边界热输送与海表失热。"
        "本文基于SODA3.15.2，以固定0—300 m海层OHC作为唯一储量定义，在5°S—5°N、140°W—80°W建立东太平洋控制体；"
        "边界采用传统ρ₀cₚ∫v(T−Tref)dA热输送，并以0、10和20°C三个参考温度检验开放断面的质量输送敏感性。"
        "结合OISST、ERA5、GPCP和CPC OLR，比较冷舌型（CT）与中央太平洋型（CP）合成以及1997/98、2015/16和2023/24事件。"
        f"发展年1—6月，1997、2015和2023东部OHC分别增加{development.loc[1997, 'east_ohc300_change_zj']:.2f}、"
        f"{development.loc[2015, 'east_ohc300_change_zj']:.2f}和{development.loc[2023, 'east_ohc300_change_zj']:.2f} ZJ；"
        "2015与2023仅相差约5%，说明2023后续弱响应不能归因于东部没有获得热量。"
        f"次年1—5月三事件东部OHC分别减少{abs(decay.loc[1997, 'east_ohc300_change_zj']):.2f}、"
        f"{abs(decay.loc[2015, 'east_ohc300_change_zj']):.2f}和{abs(decay.loc[2023, 'east_ohc300_change_zj']):.2f} ZJ。"
        "1997的5°N总热输送在三个参考温度下均为北向；2015和2023的总输送可随参考温度换号，但Ekman贡献始终为正且保持1997>2015>2023。"
        "离线换风试验进一步显示，把2023风应力异常施加于固定2015海洋状态会使北向热输送减少0.80—2.85 ZJ，反向替换则增加0.77—2.81 ZJ。"
        "结果表明，东部OHC增加是事件发展的重要海洋条件，但并不决定后续命运；次表层热量分布、深对流—风场接通和风驱Ekman贡献共同控制成熟后的热量释放。"
        "同时，开放断面总热输送的绝对正负号必须在体积收支或参考温度敏感性约束下解释。"
    )
    base.set_run_font(paragraph.add_run(text), size=10.5)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    base.set_run_font(paragraph.add_run("关键词："), size=10.5, bold=True, east_asia=base.FONT_HEAD)
    base.set_run_font(
        paragraph.add_run("ENSO；上层海洋热含量；海洋热输送；Ekman输送；东太平洋；海气耦合"),
        size=10.5,
    )


def build() -> None:
    phase = pd.read_csv(PHASE_CSV)
    primary = phase.loc[phase.baseline.eq("1980-2014") & phase.event_year.isin(EVENTS)]
    development = primary.loc[primary.phase.eq("development_Jan0_Jun0")].set_index("event_year")
    decay = primary.loc[primary.phase.eq("decay_Jan1_May1")].set_index("event_year")
    types = pd.read_csv(TYPE_CSV)
    correlations = pd.read_csv(CORR_CSV)
    chain = pd.read_csv(CHAIN_CSV).set_index("year")
    swap = pd.read_csv(SWAP_CSV)
    swap["donor_wind_anomaly"] = swap.donor_wind_anomaly.astype(str)
    swap = swap.set_index(["target_ocean_event", "donor_wind_anomaly", "reference_temperature_c"])

    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_abstract(doc, development, decay)

    base.add_heading(doc, "1 引言", 1)
    base.add_body(doc, (
        "ENSO的发展和衰减依赖赤道太平洋上层海洋热量的跨盆地重新分配。Jin（1997）的充放电振子把纬向平均温跃层变化和经向热量输送失衡视为ENSO慢记忆；"
        "Meinen和McPhaden（2000）的观测进一步表明，赤道太平洋暖水体积通常领先Niño-3.4海表温度。"
        "这些工作奠定了海洋预置—东传—放电的理论框架，但体积指标和标准化输送不能直接给出固定区域内储存了多少能量，也不能保证跨开放断面的热输送具有唯一参考。"
    ))
    base.add_body(doc, (
        "Chakravorty等（2021）利用SODA2.1.6比较CT与CP事件，发现CT具有较清晰的西太平洋蓄热、赤道东传和成熟后极向放电，而CP的充放电循环较弱。"
        "本文由此把问题向东推进：东太平洋上层海洋是否真正获得热量；成熟后这些热量如何通过储量变化、边界输送和海表通量释放；"
        "以及为什么2015和2023具有相近东部储热却产生不同海气响应。"
    ))
    base.add_body(doc, (
        "2023/24事件为这一问题提供了关键对照。已有研究指出该事件海洋增暖显著，但中西太平洋降水、风压响应和遥相关弱于历史强事件，可能受到持续东风、多年La Niña遗产、外海盆增暖和北太平洋背景态调制。"
        "这些研究主要解释大气响应为何偏弱；本文进一步检验弱耦合是否对应不同的东部OHC释放和5°N风驱热输送。"
    ))
    base.add_lead(doc, "研究问题：", "（1）东太平洋0—300 m OHC是否增加；（2）成熟后OHC如何释放和重新分配；（3）相似东部储热为何没有形成相同海气耦合与风驱输出。")

    base.add_heading(doc, "2 数据与方法", 1)
    base.add_heading(doc, "2.1 数据、区域与事件窗口", 2)
    base.add_body(doc, (
        "海洋主资料为SODA3.15.2月平均温度、三维流速、风应力、垂向速度和净海表加热。东部控制体固定为5°S—5°N、140°W—80°W、0—300 m。"
        "OISST v2.1用于定义FEP和Niño-3.4 SST异常；ERA5提供850 hPa纬向风、海平面气压和表面通量；GPCP与CPC OLR用于约束深对流位置。"
    ))
    base.add_table(
        doc,
        "表1  数据资料及其作用",
        ["资料", "变量", "主要时段", "本文作用"],
        [
            ["SODA3.15.2", "T、u/v/w、τx、net_heating", "1980—2024月平均", "0—300 m OHC、边界热输送、Ekman和箱体收支"],
            ["OISST v2.1", "SST", "1982—2024月平均", "FEP、Niño-3.4与表面触发"],
            ["ERA5", "u850、SLP、风应力、热通量", "事件共同期", "大气接通和表面失热"],
            ["GPCP / CPC OLR", "降水、OLR", "事件共同期", "深对流独立约束"],
        ],
        [1500, 2450, 1900, 3510],
        font_size=8.9,
        header_fill="F4F6F9",
    )
    base.add_body(doc, (
        "类型合成采用4个CT事件（1982、1986、1997、2006）和4个CP事件（1991、1994、2002、2004）。为保证OHC、边界输送和海气指标使用完全相同的资料窗口，"
        "发展期统一为Jan(0)—Jun(0)，成熟后释放期为Jan(+1)—May(+1)。过程对照选取1997/98、2015/16和2023/24。"
        "所有结果均以事件发展年标记，例如2015表示2015/16事件。"
    ))

    base.add_heading(doc, "2.2 固定深度OHC与区域储量", 2)
    base.add_body(doc, (
        "本文只使用标准固定深度OHC，不引入温度阈值或水团掩膜。首先计算每个网格0—300 m的垂向热含量，再对东部控制体面积分。"
    ))
    base.add_equation(doc, "OHC₀₋₃₀₀(x,y,t)=ρ₀cₚ ∫₀³⁰⁰ [T(x,y,z,t)−Tref] dz", 1)
    base.add_equation(doc, "OHCΩ(t)=∬Ω OHC₀₋₃₀₀(x,y,t) dA", 2)
    base.add_body(doc, (
        "主结果取Tref=0°C。对固定体积而言，改变常数参考温度只给原始储量增加时间不变的常数，因此去除逐月气候态和趋势后的OHC异常及其阶段变化不受Tref影响。"
        "区域积分储量以ZJ表示，1 ZJ=10²¹ J。"
    ))

    base.add_heading(doc, "2.3 标准边界热输送及参考温度敏感性", 2)
    base.add_equation(doc, "QΓ(Tref,t)=ρ₀cₚ ∬Γ vₙ(x,z,t)[T(x,z,t)−Tref] dA", 3)
    base.add_equation(doc, "VΓ(t)=∬Γ vₙ dA", 4)
    base.add_equation(doc, "QΓ(T₂)−QΓ(T₁)=−ρ₀cₚ(T₂−T₁)VΓ", 5)
    base.add_body(doc, (
        "式（5）说明开放断面热输送不同于固定体积储量：只要断面净体积输送不严格为零，热输送就依赖参考温度。"
        "因此本文以Tref=0°C作为传统主计算，并同时报告10和20°C敏感性。20°C在这里只是代数参考温度，所有水体仍参与积分，不构成阈值诊断。"
    ))
    base.add_equation(doc, "EΓ=Σₖ Q′Γ(tₖ)Δtₖ,   Δtₖ=days_in_monthₖ×86400 s", 6)
    base.add_body(doc, (
        "非线性输送项遵循‘先计算逐月v(T−Tref)，再对完整乘积按日历月去除1980—2014气候态和线性趋势’的顺序；不把速度异常和温度异常分别相乘。"
        "PW到ZJ的转换使用每月实际秒数加权。1982—2022基准作为敏感性检验。"
    ))

    base.add_heading(doc, "2.4 东部控制体热收支", 2)
    base.add_equation(doc, "dOHCΩ/dt=Q₁₄₀W−Q₈₀W+Q₅S−Q₅N+Qsurface+Q₃₀₀m+R", 7)
    base.add_body(doc, (
        "Q140W和Q80W以东向为正，Q5S和Q5N以北向为正；因此式（7）中140°W与5°S为入箱项，80°W与5°N为出箱项。"
        "Qsurface进入海洋为正，结果部分为便于解释改报正的海表热损失。Q300m为300 m底界向上平流热输入，R由储量变化减去全部显式项得到，"
        "包含未解析混合、扩散、涡动协方差、同化增量和离散误差。R的存在意味着预算闭合是诊断恒等式，而不是独立物理验证。"
    ))

    base.add_heading(doc, "2.5 Ekman热输送与离线换风", 2)
    base.add_equation(doc, "My(x,t)=−τx(x,t)/(ρ₀f)", 8)
    base.add_equation(doc, "QEkman(Tref)=ρ₀cₚ ∫ My(x)[T̄₀₋₅₀m(x)−Tref] dx", 9)
    base.add_equation(doc, "QnonEkman=Qtotal−QEkman", 10)
    base.add_body(doc, (
        "在5°N，西向风应力对应北向Ekman体积输送。式（9）以0—50 m平均温度作为输送温度权重，是一阶近似而非完整垂向动力分解。"
        "本文不再把非Ekman项进一步命名为地转项；它只是总热输送减去一阶Ekman估计后的诊断剩余。"
    ))
    base.add_equation(doc, "Qrec(i←j)=QnonEkman(i)+QEkman[T(i),τx(j)]", 11)
    base.add_body(doc, (
        "离线换风固定目标事件i的温度和非Ekman输送，只把5°N去趋势风应力异常替换为事件j。该试验检验局地风驱的一阶贡献，不允许温度、压力场、环流或残差对新风场作出响应，故不属于耦合模式因果试验。"
    ))

    base.add_heading(doc, "2.6 FEP残差、海气链与统计", 2)
    base.add_equation(doc, "FEPres(m,y)=FEP′(m,y)−[aₘ+bₘN₃.₄′(m,y)]", 12)
    base.add_body(doc, (
        "FEP为5°S—5°N、90°W—80°W平均SST异常。式（12）逐日历月回归并移除同期Niño-3.4线性分量，用于识别相对于盆地尺度强度的额外最东部增暖。"
        "海气链依次检查SST—降水、降水—SLP和SLP—风应力的条件残差。CT/CP均以整事件进行20,000次bootstrap；13个暖事件关系报告Pearson相关和双侧p值，"
        "三事件只用于过程对照，不进行三例显著性检验。"
    ))
    base.add_table(
        doc,
        "表2  主要符号、物理意义与单位",
        ["符号", "物理意义", "单位/符号约定"],
        [
            ["ρ₀、cₚ", "参考海水密度与定压比热", "1025 kg m⁻³；3990 J kg⁻¹ K⁻¹"],
            ["T、Tref", "海温与常数参考温度", "°C；主值0°C，敏感性10/20°C"],
            ["Ω、Γ", "固定东部控制体与其边界断面", "5°S—5°N、140°W—80°W、0—300 m"],
            ["OHCΩ", "区域积分0—300 m热含量", "J；阶段变化用ZJ"],
            ["vₙ、VΓ", "断面法向速度与体积输送", "m s⁻¹；Sv"],
            ["QΓ、EΓ", "瞬时热输送与阶段累计热量", "W/PW；J/ZJ"],
            ["τx、f、My", "纬向风应力、科氏参数、经向Ekman体积输送", "N m⁻²；s⁻¹；m² s⁻¹"],
            ["R", "由储量变化减去显式项得到的预算剩余", "W/PW；不预先解释为单一过程"],
            ["m、y、aₘ、bₘ", "日历月、年份、月回归截距与斜率", "统计量"],
        ],
        [1450, 5070, 2840],
        font_size=8.7,
        header_fill="F4F6F9",
    )

    base.add_heading(doc, "3 结果", 1)
    base.add_heading(doc, "3.1 OHC确实向东重分配，但事件幅度不同", 2)
    base.add_figure(
        doc,
        str(FIG_HOV),
        "图1  1997/98、2015/16和2023/24赤道太平洋0—300 m OHC经度—时间演变。异常基于1980—2024逐日历月气候态并去除同月线性趋势；虚线标出180°、140°W和90°W。",
        width_in=6.3,
    )
    base.add_body(doc, (
        "1997在发展年春季已出现从中太平洋向东扩展的连续正OHC异常，夏秋季在140°W以东形成强而深厚的热储存。2015同样建立了中东太平洋正异常，但幅度明显较小。"
        "2023则首先表现为西太平洋和最东部暖异常并存，中太平洋桥接较迟且较弱；尽管如此，140°W以东在发展年仍形成持续正OHC。"
    ))
    base.add_table(
        doc,
        "表3  三事件发展期东部OHC与水平热辐合",
        ["事件", "东部OHC变化（ZJ）", "净水平热辐合范围（ZJ）", "140°W热输送范围（ZJ）"],
        [
            [str(year), fmt(development.loc[year, "east_ohc300_change_zj"]), range_text(development.loc[year], "horizontal_convergence_tref{reference:02d}_zj"), range_text(development.loc[year], "qvt_140w_east_tref{reference:02d}_zj")]
            for year in EVENTS
        ],
        [1200, 2450, 2900, 2810],
        font_size=8.8,
        header_fill="F4F6F9",
    )
    base.add_body(doc, (
        "1997东部OHC增加14.79 ZJ，约为2015和2023的2.4—2.6倍。2015与2023分别增加5.77和6.08 ZJ，仅相差0.31 ZJ。"
        "三事件的全边界净水平热辐合在三个参考温度下均为正，因此东部储热增加与控制体净热输入相互一致。"
        "相反，2015和2023的单一140°W热输送可随Tref换号，说明‘东部获得热量’应由储量和闭合边界辐合共同判断，而不能只由一个开放断面给出唯一ZJ数。"
    ))
    base.add_lead(doc, "问题一的答案：", "1997、2015和2023的发展期东部0—300 m OHC均增加；2015与2023增量近似。2023后续弱响应不能解释为东太平洋没有获得热量。")

    base.add_heading(doc, "3.2 CT具有稳健东部增热，CP在同一窗口内较弱", 2)
    base.add_figure(
        doc,
        str(FIG_TYPE),
        "图2  OHC统一框架下的CT—CP对比。左图为Jan(0)—Jun(0)东部OHC变化及事件bootstrap 95%区间；右图为次年Jan—May 5°N总热输送和Ekman贡献的Tref=0、10、20°C范围。",
        width_in=6.25,
    )
    ct_dev = type_row(types, "CT", "development_Jan0_Jun0", "east_ohc300_change_zj")
    cp_dev = type_row(types, "CP", "development_Jan0_Jun0", "east_ohc300_change_zj")
    base.add_table(
        doc,
        "表4  CT与CP的OHC储量和成熟后热输送",
        ["类型", "发展期东部OHC变化", "衰减期东部OHC变化", "5°N总输送范围", "5°N Ekman范围"],
        [
            ["CT", f"{ct_dev['mean']:+.2f} [{ct_dev['ci95_low']:+.2f}, {ct_dev['ci95_high']:+.2f}]", f"{type_row(types, 'CT', 'decay_Jan1_May1', 'east_ohc300_change_zj')['mean']:+.2f}", "+3.75—+6.44", "+1.86—+5.07"],
            ["CP", f"{cp_dev['mean']:+.2f} [{cp_dev['ci95_low']:+.2f}, {cp_dev['ci95_high']:+.2f}]", f"{type_row(types, 'CP', 'decay_Jan1_May1', 'east_ohc300_change_zj')['mean']:+.2f}", "−5.49—−0.15", "+0.06—+0.37"],
        ],
        [850, 2630, 2200, 1870, 1810],
        font_size=8.6,
        header_fill="F4F6F9",
    )
    base.add_body(doc, (
        "CT发展期东部OHC平均增加8.30 ZJ，95%区间不跨零；CP平均变化为−0.37 ZJ，区间跨零。"
        "在统一Jan—Jun窗口内，OHC结果支持CT具有清楚的东部增热，却不支持CP也具有同等稳健的到达量。"
        "成熟后CT的5°N总输送在三个参考温度下均为北向，CP均为南向或接近零；CT的Ekman贡献也显著更强。"
        "因此本文承接Chakravorty等（2021）的群体差异，但不再声称CT和CP在所有阶段都把相同性质的热量送入东部。"
    ))

    base.add_heading(doc, "3.3 成熟后东部OHC显著释放，但开放断面去向并非总能唯一确定", 2)
    base.add_figure(
        doc,
        str(FIG_EVENT),
        "图3  三事件东部OHC到达、成熟后储量释放、5°N总热输送和Ekman贡献。彩色竖线表示Tref=0、10、20°C范围，黑色菱形为10°C结果。",
        width_in=6.25,
    )
    base.add_table(
        doc,
        "表5  三事件次年Jan—May的东部热量释放与5°N输送",
        ["事件", "东部OHC变化", "海表热损失", "5°N总输送范围", "5°N Ekman范围", "净水平辐合范围"],
        [
            [str(year), fmt(decay.loc[year, "east_ohc300_change_zj"]), fmt(decay.loc[year, "east_surface_loss_zj"]), range_text(decay.loc[year], "qvt_5n_north_tref{reference:02d}_zj"), range_text(decay.loc[year], "qekman_5n_north_tref{reference:02d}_zj"), range_text(decay.loc[year], "horizontal_convergence_tref{reference:02d}_zj")]
            for year in EVENTS
        ],
        [850, 1510, 1470, 1860, 1800, 1870],
        font_size=8.4,
        header_fill="F4F6F9",
    )
    base.add_body(doc, (
        "三事件东部OHC均大幅减少，其中2015减少18.51 ZJ，略大于1997和2023。海表异常热损失却依次为4.37、1.17和0.40 ZJ，"
        "说明储量下降不能简单等同于向大气散热。三事件的净水平热辐合在所有参考温度下均为负，表明成熟后控制体总体处于水平热量输出状态。"
    ))
    base.add_body(doc, (
        "1997的5°N总热输送为+7.04—+8.67 ZJ，参考温度不改变北向结论。2015和2023的范围分别为−3.22—+3.73和−9.12—+1.58 ZJ，绝对方向不唯一。"
        "然而Ekman项在三个参考温度下始终为正，且1997>2015>2023。由此可以稳健地判断2015比2023具有更强北向风驱贡献，却不能仅凭标准开放断面热输送断言2015总出口一定向北或2023一定向南。"
    ))
    base.add_lead(doc, "问题二的答案：", "东部OHC成熟后明显释放，1997存在稳健5°N北向输出；2015与2023的总出口方向受参考温度影响，但2015的Ekman北向贡献和海表失热均强于2023。")

    base.add_heading(doc, "3.4 相似东部储热为何产生不同海气耦合命运", 2)
    base.add_heading(doc, "3.4.1 FEP先暖不等于盆地尺度海洋—大气接通", 3)
    base.add_figure(
        doc,
        str(FIG_CLOSURE),
        "图4  1997、2015和2023从FEP表面增暖、0—300 m OHC重新分布到u850与OLR响应的经度—时间演变。紫虚线表示本文操作性定义的大气接通月份。",
        width_in=6.3,
    )
    base.add_body(doc, (
        "1997的中太平洋OHC正异常在春季已建立，并迅速对应西风和负OLR异常；2015的OHC预置较弱但仍在3月前后形成大气接通。"
        "2023虽具有持续正FEP残差，中太平洋OHC向东扩展却明显偏迟，3—4月深对流和西风响应仍弱，到5月才出现有限接通。"
        "在三事件条件残差中，2023最弱环节为SLP→风应力，残差为−0.16个标准差；2015对应值为+1.40。"
    ))

    base.add_heading(doc, "3.4.2 事件统计把OHC预置连接到风场和Ekman输送", 3)
    base.add_figure(
        doc,
        str(FIG_COUPLING),
        "图5  中太平洋0—300 m OHC分布及其与随后西风响应的13事件关系。阴影区为180°—140°W候选海洋桥接区；散点回归用于过程联系而非因果阈值。",
        width_in=6.25,
    )
    def corr(x: str, y: str) -> pd.Series:
        return correlations.loc[correlations.x.eq(x) & correlations.y.eq(y)].iloc[0]

    r_ohc_nino = corr("east_ohc300_change_zj", "mature_djf_nino34_c")
    r_ohc_precip = corr("east_ohc300_change_zj", "mar_jun_sst_to_precip_residual_z")
    r_ek_intensity = corr("qekman_5n_north_tref00_zj", "may_jun_intensity_coupling_residual_z")
    r_ek_precip = corr("qekman_5n_north_tref00_zj", "mar_jun_sst_to_precip_residual_z")
    base.add_body(doc, (
        f"13个暖事件中，发展期东部OHC增加与成熟期Niño-3.4相关为r={r_ohc_nino.pearson_r:.2f}（p={r_ohc_nino.p_value:.3f}），"
        f"与SST→降水耦合残差相关为r={r_ohc_precip.pearson_r:.2f}（p={r_ohc_precip.p_value:.3f}）。"
        f"次年5°N Ekman热输送与发展期强度耦合残差相关为r={r_ek_intensity.pearson_r:.2f}（p={r_ek_intensity.p_value:.3f}），"
        f"与SST→降水残差相关为r={r_ek_precip.pearson_r:.2f}（p={r_ek_precip.p_value:.3f}）。"
        "这些关系支持‘海洋预置—对流响应—风驱再分配’的事件级联系，但不能证明单向因果。"
    ))

    base.add_heading(doc, "3.4.3 离线换风隔离局地风应力的一阶贡献", 3)
    base.add_figure(
        doc,
        str(FIG_SWAP),
        "图6  固定海洋温度和非Ekman输送、替换5°N风应力异常后的OHC热输送重构。左图为Tref=0°C三事件矩阵，右图显示2015—2023结论对参考温度的敏感性。",
        width_in=6.25,
    )
    base.add_table(
        doc,
        "表6  2015—2023换风对重构总5°N热输送的影响（ZJ）",
        ["换风情景", "Tref=0°C", "Tref=10°C", "Tref=20°C"],
        [
            ["2015海洋换入2023风", fmt(swap.loc[(2015, "2023", 0.0), "change_from_control_zj"]), fmt(swap.loc[(2015, "2023", 10.0), "change_from_control_zj"]), fmt(swap.loc[(2015, "2023", 20.0), "change_from_control_zj"])],
            ["2023海洋换入2015风", fmt(swap.loc[(2023, "2015", 0.0), "change_from_control_zj"]), fmt(swap.loc[(2023, "2015", 10.0), "change_from_control_zj"]), fmt(swap.loc[(2023, "2015", 20.0), "change_from_control_zj"])],
        ],
        [3300, 2020, 2020, 2020],
        font_size=8.9,
        header_fill="F4F6F9",
    )
    base.add_body(doc, (
        "把2023风应力异常施加于固定2015海洋状态，总热输送相对控制减少0.80—2.85 ZJ；把2015风施加于固定2023海洋状态则增加0.77—2.81 ZJ。"
        "三个参考温度下方向完全一致，且两种反向试验的幅度近似对称。"
        "这说明局地风应力是2015—2023 Ekman差异的重要一阶贡献者；但重构总值本身仍随Tref改变，故不能声称风应力单独使两个事件的总出口状态完全互换。"
    ))
    base.add_lead(doc, "问题三的答案：", "2015与2023虽有近似东部OHC增量，但2023次表层桥接和SLP—风应力响应较弱；较弱风应力对应较弱Ekman北向热输送，构成两事件成熟后命运差异的稳健诊断贡献。")

    base.add_heading(doc, "4 讨论", 1)
    base.add_heading(doc, "4.1 从充放电框架到固定区域OHC收支", 2)
    base.add_body(doc, (
        "本文不替代Jin（1997）的充放电振子，也不否认WWV对ENSO海洋记忆的价值，而是使用文献中更通行的固定深度OHC回答区域能量问题。"
        "这一转换把‘盆地是否充水’与‘140°W以东储存了多少热量’区分开，也把‘储量下降’与‘某一开放断面输出’区分开。"
        "相对于Chakravorty等（2021），最稳健的新结果不是重新证明暖水可以东传，而是量化东部OHC增量、控制体净水平热辐合及后续OHC释放。"
    ))

    base.add_heading(doc, "4.2 OHC框架的优势与代价", 2)
    base.add_body(doc, (
        "OHC的优势是定义标准、单位明确、固定体积储量不受参考温度选择影响，便于与海洋再分析、观测和模式热收支直接比较。"
        "代价是传统热输送不是独立于质量输送的唯一量。2015和2023的140°W或5°N总热输送可随Tref换号，正是式（5）中参考温度乘净体积输送项的结果。"
        "因此，完全采用OHC后，论文能够稳健回答‘东部是否增热’和‘储量是否释放’，却不能在未完成质量闭合时把每条开放边界解释成唯一的热量去向。"
    ))

    base.add_heading(doc, "4.3 物理意义与预测启示", 2)
    base.add_body(doc, (
        "2015与2023具有近似东部OHC增量却产生不同风场和Ekman贡献，说明初始海洋热量并非事件命运的充分条件。"
        "热量在经度方向上的分布、温跃层结构以及深对流能否向中太平洋扩展，决定大气是否形成持续西风响应；"
        "风应力进一步调节成熟后的经向热量再分配和海表热释放。对业务监测而言，东部OHC增量应与中太平洋OHC桥接、OLR和风应力效率联合使用，而不能作为单一预报指标。"
    ))

    base.add_heading(doc, "4.4 局限性", 2)
    base.add_body(doc, (
        "本文仍有六项限制。第一，SODA3月平均资料不能分离涡动协方差和高频风爆发；第二，0—50 m温度权重只是一阶Ekman近似；"
        "第三，东部箱体预算剩余尚未拆分混合、扩散、夹卷和同化增量；第四，开放断面总热输送仍存在参考温度和质量闭合问题；"
        "第五，CT/CP每组只有4个事件，13事件相关样本也有限；第六，离线换风不允许海洋和大气对新风场动态调整。"
        "因此本文使用‘对应、贡献和敏感性’描述机制，不使用‘决定、必要条件或完整因果证明’等强措辞。"
    ))

    base.add_heading(doc, "5 结论", 1)
    base.add_lead(doc, "第一，", "1997、2015和2023发展期东部0—300 m OHC均增加，2015与2023增量仅相差约5%。2023后续弱响应不能归因于东部没有获得热量。")
    base.add_lead(doc, "第二，", "成熟后三事件东部OHC均显著减少；1997具有稳健的5°N北向总热输送，2015与2023的总输送方向受参考温度影响，但2015的Ekman北向贡献和海表热损失均强于2023。")
    base.add_lead(doc, "第三，", "CT在统一Jan—Jun窗口内具有稳健东部增热，CP信号较弱；成熟后CT的5°N总输送和Ekman贡献也明显强于CP，延续了不同ENSO类型充放电差异。")
    base.add_lead(doc, "第四，", "相似东部储热并不保证相似海气命运。2023的中太平洋次表层桥接和SLP—风应力响应较弱；离线换风在三个参考温度下均支持局地风驱对2015—2023 Ekman差异的重要贡献。")
    base.add_body(doc, (
        "总体而言，完全基于OHC的结果把研究主题从单一暖水路径追踪收束为东太平洋上层海洋热量的到达、释放和风驱再分配。"
        "最稳健的闭环是：东部OHC增加并不决定事件强度；次表层空间预置影响对流和风场接通，风应力进一步调节Ekman热输送和成熟后OHC释放。"
    ))

    base.add_heading(doc, "数据与可重复性说明", 1)
    base.add_body(doc, (
        "SODA3、OISST、ERA5、GPCP和CPC OLR均为公开资料。本文OHC储量、边界热输送、参考温度敏感性、CT/CP合成、三事件换风和图件均由项目脚本直接生成；"
        "所有PW到ZJ转换使用每月实际秒数。正式投稿前应公开派生事件表、计算脚本、环境文件和质量闭合检查。"
    ))

    base.add_heading(doc, "参考文献", 1)
    references = [
        "Jin, F.-F. (1997). An equatorial ocean recharge paradigm for ENSO. Part I: Conceptual model. Journal of the Atmospheric Sciences, 54, 811–829. https://doi.org/10.1175/1520-0469(1997)054<0811:AEORPF>2.0.CO;2",
        "Meinen, C. S., & McPhaden, M. J. (2000). Observations of warm water volume changes in the equatorial Pacific and their relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559. https://doi.org/10.1175/1520-0442(2000)013<3551:OOWWVC>2.0.CO;2",
        "Sun, D.-Z. (2000). The heat sources and sinks of the 1986–87 El Niño. Journal of Climate, 13, 3533–3550. https://doi.org/10.1175/1520-0442(2000)013<3533:THSASO>2.0.CO;2",
        "Kug, J.-S., Jin, F.-F., & An, S.-I. (2009). Two types of El Niño events: Cold Tongue El Niño and Warm Pool El Niño. Journal of Climate, 22, 1499–1515. https://doi.org/10.1175/2008JCLI2624.1",
        "Capotondi, A., et al. (2015). Understanding ENSO diversity. Bulletin of the American Meteorological Society, 96, 921–938. https://doi.org/10.1175/BAMS-D-13-00117.1",
        "Carton, J. A., Chepurin, G. A., & Chen, L. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983. https://doi.org/10.1175/JCLI-D-17-0149.1",
        "Hersbach, H., et al. (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146, 1999–2049. https://doi.org/10.1002/qj.3803",
        "Huguenin, M. F., Holmes, R. M., & England, M. H. (2020). Key role of diabatic processes in regulating warm water volume variability over ENSO events. Journal of Climate, 33, 9945–9964. https://doi.org/10.1175/JCLI-D-20-0198.1",
        "Chakravorty, S., Perez, R. C., Gnanaseelan, C., & Anderson, B. T. (2021). Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 126, e2020JC017075. https://doi.org/10.1029/2020JC017075",
        "Huang, B., et al. (2021). Improvements of the Daily Optimum Interpolation Sea Surface Temperature version 2.1. Journal of Climate, 34, 2923–2939. https://doi.org/10.1175/JCLI-D-20-0166.1",
        "Li, X., Hu, Z.-Z., Huang, B., & Jin, F.-F. (2022). Oceanic meridional transports and their roles in warm water volume variability and ENSO in the tropical Pacific. Climate Dynamics, 59, 245–261. https://doi.org/10.1007/s00382-021-06124-w",
        "Verdy, A., Mazloff, M. R., Cornuelle, B. D., & Kim, S. B. (2023). Balancing volume, temperature, and salinity budgets during 2014–2018 in the Tropical Pacific Ocean State Estimate. Journal of Geophysical Research: Oceans, 128, e2022JC019576. https://doi.org/10.1029/2022JC019576",
        "Hong, C.-C., Sullivan, A., & Chang, C.-C. (2024). Impact of North Atlantic Tripole and Extratropical North Pacific extreme SSTs on the 2023/24 El Niño. Geophysical Research Letters, 51, e2024GL110639. https://doi.org/10.1029/2024GL110639",
        "Kim, J.-W., Tian, B., & Yu, J.-Y. (2025). Unusual and persistent easterlies restrained the 2023/24 El Niño development after a triple-dip La Niña. npj Climate and Atmospheric Science, 8, 34. https://doi.org/10.1038/s41612-024-00890-0",
        "Lin, Y.-F., et al. (2025). Diverse NPMM conditions deviate the 2023/24 El Niño from the 1997/1998 and 2015/2016 extreme El Niño events. npj Climate and Atmospheric Science, 8, 133. https://doi.org/10.1038/s41612-025-01013-z",
        "Peng, Q., et al. (2025). Strong 2023–2024 El Niño generated by ocean dynamics. Nature Geoscience, 18, 471–478. https://doi.org/10.1038/s41561-025-01700-9",
        "Zhang, L., et al. (2025). The 2023/24 El Niño event exhibited unusually weak extratropical teleconnections. Communications Earth & Environment, 6, 595. https://doi.org/10.1038/s43247-025-02584-8",
    ]
    for index, reference in enumerate(references, 1):
        base.add_reference(doc, index, reference)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
