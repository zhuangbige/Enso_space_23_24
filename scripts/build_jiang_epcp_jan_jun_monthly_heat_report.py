#!/usr/bin/env python3
"""Build a Chinese DOCX report of monthly Jan--Jun ENSO heat allocation.

The report follows the month-by-month time layout of the June
event-comparison report.  It does not average JFM, MJJ, or January--June,
and it does not create a new event index or classification threshold.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data/derived/research_2026"
REPORTS = ROOT / "outputs/research_2026/reports"

BUDGET_FILE = DERIVED / "soda3_three_box_ohc_anomaly_budget_monthly.csv"
MANIFEST_FILE = DERIVED / "jiang2026_epcp_fep_event_manifest_1963_2023.csv"
REFERENCE_DOCX = REPORTS / "enso_event_comparison_unified_1980_2014_jan_jun_2026.docx"
OUTPUT_FILE = REPORTS / "Jiang_EP_CP_FEP_事件逐月热量分配_发展年1至6月_CN_20260814.docx"

EVENTS = (1982, 1986, 1987, 1991, 1994, 1997, 2002, 2004, 2006, 2009, 2015, 2018, 2019, 2023)
EP_EVENTS = (1982, 1986, 1997, 2015)
CP_EVENTS = (1987, 1991, 1994, 2002, 2004, 2006, 2009, 2018, 2019)
MONTHS = tuple(range(1, 7))
MONTH_NAMES = {1: "1月", 2: "2月", 3: "3月", 4: "4月", 5: "5月", 6: "6月"}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5D6875"
GRID = "D6DEE8"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F6F8FA"
WHITE = "FFFFFF"

PAGE_WIDTH_DXA = 15840
PAGE_HEIGHT_DXA = 12240
MARGIN_DXA = 864
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - 2 * MARGIN_DXA
TABLE_INDENT_DXA = 120

TABLE_WIDTHS = [
    650,
    1662,
    1680,
    1680,
    1680,
    1680,
    1680,
    1680,
    1720,
]
CLASS_WIDTHS = [1300, 1100, 1100, 1650, 8962]


def set_run_font(run, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    rfonts.set(qn("w:cs"), "Times New Roman")


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    settings = {
        "Heading 1": (15, BLUE, 14, 7),
        "Heading 2": (12.5, BLUE, 10, 5),
        "Heading 3": (11.5, DARK_BLUE, 7, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = False
        style.paragraph_format.keep_together = False
        style.paragraph_format.page_break_before = False
        style.paragraph_format.widow_control = False


def set_cell_margins(cell, top: int = 60, start: int = 90, bottom: int = 60, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths sum to {sum(widths)}, expected {CONTENT_WIDTH_DXA}")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    set_table_borders(table)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def format_cell(cell, *, size: float = 8.2, bold: bool = False, color: str = "000000", align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 8.5, color=MUTED)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("ENSO发展年逐月热量分配    SODA3 0–300 m")
    set_run_font(run, 8.5, bold=True, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_plain_paragraph(doc: Document, text: str, *, bold_lead: str | None = None, after: float = 5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 10.5, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, 10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.5)


def add_table_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.widow_control = False
    run = paragraph.add_run(text)
    set_run_font(run, 9.5, bold=True, color=INK)


def short_evidence(text: str) -> str:
    if text.startswith("Explicit"):
        return "Jiang正文明确给出"
    if text.startswith("Jiang-consistent"):
        return "与Jiang正文总体计数一致的工作重建；本地缺少S6逐事件标签"
    return "Jiang列出22次事件，但正文成熟类型计数只覆盖21次；暂不判定"


def load_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    manifest = pd.read_csv(MANIFEST_FILE)
    columns = ["time", "year", "month"]
    for box in ("central", "east"):
        columns.extend(
            [
                f"{box}_ohc300_anom_zj",
                f"{box}_storage_tendency_anom_pw",
                f"{box}_west_theta_heat_in_corr_anom_pw",
                f"{box}_east_theta_heat_in_corr_anom_pw",
                f"{box}_south_theta_heat_in_corr_anom_pw",
                f"{box}_north_theta_heat_in_corr_anom_pw",
                f"{box}_surface_anom_pw",
                f"{box}_residual_anom_pw",
            ]
        )
    budget = pd.read_csv(BUDGET_FILE, usecols=columns, parse_dates=["time"])
    budget = budget.loc[budget.year.isin(EVENTS) & budget.month.isin(MONTHS)].copy()
    counts = budget.groupby("year")["month"].nunique()
    if set(counts.index.astype(int)) != set(EVENTS) or not counts.eq(6).all():
        raise ValueError(f"Incomplete Jan--Jun event coverage: {counts.to_dict()}")
    if budget[columns[3:]].isna().any().any():
        missing = budget.columns[budget.isna().any()].tolist()
        raise ValueError(f"Missing monthly heat values: {missing}")
    return manifest, budget.sort_values(["year", "month"]).reset_index(drop=True)


def transformed_box(frame: pd.DataFrame, box: str) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "月份": frame["month"].map(MONTH_NAMES),
            "OHC′\n(ZJ)": frame[f"{box}_ohc300_anom_zj"],
            "储热倾向\n(PW)": frame[f"{box}_storage_tendency_anom_pw"],
            "西侧输入\n(PW)": frame[f"{box}_west_theta_heat_in_corr_anom_pw"],
            "向东输出\n(PW)": -frame[f"{box}_east_theta_heat_in_corr_anom_pw"],
            "向南输出\n(PW)": -frame[f"{box}_south_theta_heat_in_corr_anom_pw"],
            "向北输出\n(PW)": -frame[f"{box}_north_theta_heat_in_corr_anom_pw"],
            "海表损失\n(PW)": -frame[f"{box}_surface_anom_pw"],
            "未解析储热\n(PW)": frame[f"{box}_residual_anom_pw"],
        }
    )


def add_data_table(doc: Document, data: pd.DataFrame) -> None:
    table = doc.add_table(rows=1, cols=len(data.columns))
    header = table.rows[0]
    repeat_header(header)
    for cell, label in zip(header.cells, data.columns):
        cell.text = str(label)
        shade_cell(cell, HEADER_FILL)
        format_cell(cell, size=8.0, bold=True, color=INK)

    for row_index, row in data.iterrows():
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index == 0:
                text = str(value)
            else:
                number = float(value)
                text = f"{number:+.3f}"
            cells[index].text = text
            if row_index % 2 == 1:
                shade_cell(cells[index], LIGHT_FILL)
            format_cell(cells[index], size=8.2)
    set_table_geometry(table, TABLE_WIDTHS)


def add_classification_table(doc: Document, manifest: pd.DataFrame) -> None:
    rows = []
    for row in manifest.itertuples(index=False):
        rows.append(
            [
                row.event_label,
                {"EP": "EP", "CP": "CP", "unresolved": "未判定"}[row.mature_type_working],
                "是" if bool(row.fep_dominant) else "否",
                "有，发展年1–6月" if int(row.event_year) in EVENTS else "无，早于本热收支样本",
                short_evidence(row.classification_evidence),
            ]
        )
    headers = ["事件", "成熟类型", "FEP主导", "本报告热量表", "分类依据"]
    table = doc.add_table(rows=1, cols=len(headers))
    repeat_header(table.rows[0])
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        shade_cell(cell, HEADER_FILL)
        format_cell(cell, size=8.4, bold=True, color=INK)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            if row_index % 2 == 1:
                shade_cell(cells[index], LIGHT_FILL)
            align = WD_ALIGN_PARAGRAPH.LEFT if index == 4 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(cells[index], size=8.2, align=align)
    set_table_geometry(table, CLASS_WIDTHS)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=30, start=90, bottom=30, end=90)


def classification_label(year: int, manifest: pd.DataFrame) -> tuple[str, str, str]:
    row = manifest.loc[manifest.event_year.eq(year)].iloc[0]
    mature = str(row.mature_type_working)
    if year == 2023:
        group = "FEP主导的EP事件（单独延伸）"
    elif mature == "EP":
        group = "历史EP事件"
    else:
        group = "历史CP事件"
    fep = "是" if bool(row.fep_dominant) else "否"
    return group, fep, short_evidence(str(row.classification_evidence))


def event_summary(frame: pd.DataFrame) -> str:
    central_east = -frame["central_east_theta_heat_in_corr_anom_pw"]
    east_west = frame["east_west_theta_heat_in_corr_anom_pw"]
    east_ohc = frame["east_ohc300_anom_zj"]
    peak_index = int(central_east.idxmax())
    peak_month = MONTH_NAMES[int(frame.loc[peak_index, "month"])]
    return (
        f"逐月变化：中箱向东输出由1月{central_east.iloc[0]:+.3f} PW变为6月{central_east.iloc[-1]:+.3f} PW，"
        f"东箱西侧输入由{east_west.iloc[0]:+.3f} PW变为{east_west.iloc[-1]:+.3f} PW，"
        f"东箱OHC′由{east_ohc.iloc[0]:+.3f} ZJ变为{east_ohc.iloc[-1]:+.3f} ZJ。"
        f"1—6月内中箱向东输出最大值出现在{peak_month}，为{central_east.loc[peak_index]:+.3f} PW。"
    )


def add_event_page(doc: Document, year: int, number: int, manifest: pd.DataFrame, budget: pd.DataFrame) -> None:
    doc.add_page_break()
    if number == 1:
        doc.add_heading("3  发展年1—6月逐事件热量分配", level=1)
        add_plain_paragraph(
            doc,
            "以下每次事件分别给出中箱和东箱的六个月月值。表中没有季节平均、事件平均或新指数。",
        )
    label = f"{year}/{str(year + 1)[-2:]}"
    group, fep, evidence = classification_label(year, manifest)
    heading = doc.add_heading(f"3.{number}  {label}：{group}", level=2)
    heading.paragraph_format.keep_with_next = False
    heading.paragraph_format.keep_together = False
    heading.paragraph_format.page_break_before = False
    heading.paragraph_format.widow_control = False
    add_plain_paragraph(
        doc,
        f"分类：{group}；FEP主导：{fep}；依据：{evidence}。时间为发展年{year}年1月至6月，逐月列示。",
    )
    frame = budget.loc[budget.year.eq(year)].copy().reset_index(drop=True)
    add_plain_paragraph(doc, event_summary(frame))
    add_table_title(doc, "中箱：180°—140°W，10°S—5°N，0—300 m")
    add_data_table(doc, transformed_box(frame, "central"))
    add_table_title(doc, "东箱：140°W—80°W，10°S—5°N，0—300 m")
    add_data_table(doc, transformed_box(frame, "east"))


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Jiang EP/CP与FEP事件的逐月热量分配")
    set_run_font(run, 22, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("SODA3 0–300 m两箱收支｜发展年1—6月｜逐月呈现，不作季节平均")
    set_run_font(run, 12, color=DARK_BLUE)

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(12)
    run = metadata.add_run("时间排布参照：enso_event_comparison截至6月版本    完成日期：2026年8月14日")
    set_run_font(run, 9.5, color=MUTED)


def build_report() -> Path:
    if not REFERENCE_DOCX.exists():
        raise FileNotFoundError(REFERENCE_DOCX)
    manifest, budget = load_data()
    REPORTS.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = "Jiang EP/CP与FEP事件的逐月热量分配"
    doc.core_properties.subject = "发展年1至6月逐月SODA3两箱热收支"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "ENSO; EP; CP; FEP; SODA3; OHC; heat transport"

    add_title_block(doc)

    doc.add_heading("结论", level=1)
    add_plain_paragraph(
        doc,
        "本报告只完成热量部分。所有事件按发展年1月至6月逐月排列，不使用JFM、MJJ或1—6月平均。历史EP、历史CP和2023分别列出，不用组平均代替单个事件。2026只提供时间排布参照，不进入Jiang历史事件分类或SODA3热收支表。",
    )
    add_plain_paragraph(
        doc,
        "2023年中箱向东输出从1月+0.072 PW持续增强到6月+1.096 PW，东箱西侧输入从+0.115 PW增强到+1.410 PW，东箱OHC′从−2.885 ZJ升至+8.230 ZJ。同期中箱OHC′在4月达到+3.149 ZJ，5—6月约为+2.1 ZJ。逐月结果直接表明，暖量在春季跨过140°W并主要进入东箱，而不是持续留在中箱。",
    )

    doc.add_heading("1  时间、区域与符号", level=1)
    add_plain_paragraph(
        doc,
        "时间使用每个事件发展年的1月、2月、3月、4月、5月和6月六个独立月值。事件名1982/83对应1982年1—6月，2023/24对应2023年1—6月。没有跨月平均，也没有把相邻月份合成季节。",
    )
    add_plain_paragraph(
        doc,
        "中箱为180°—140°W、10°S—5°N，东箱为140°W—80°W、10°S—5°N，垂向范围均为0—300 m。OHC′单位为ZJ；储热倾向、边界输送、海表热通量和未解析项单位为PW。异常沿用现有SODA3逐日历月线性期望，基准期为1980—2022年。",
    )
    add_plain_paragraph(
        doc,
        "方向列采用统一去向口径。西侧输入为正表示热量从西侧进入箱体；向东、向南和向北输出为正表示热量离开箱体；海表损失为正表示海洋向大气放热。任何负值均表示相反方向。未解析储热为正表示未解析过程增加箱内储热，为负表示未解析过程减少储热；该项不能命名为某一种具体物理过程。",
    )

    doc.add_heading("2  Jiang事件分类与FEP说明", level=1)
    add_plain_paragraph(
        doc,
        "EP和CP描述事件成熟期的海温空间类型。FEP是far eastern Pacific，即远东太平洋5°S—5°N、90°W—80°W的早期增暖区域；FEP主导描述发展期起点，不等同于成熟期EP。FEP区域位于本报告东箱内部，但远小于东箱，因此FEP海温不能替代东箱总OHC或边界热输送。",
    )
    add_plain_paragraph(
        doc,
        "Jiang正文明确的FEP主导事件为1965/66、1969/70、1972/73、1976/77和2023/24。1969/70成熟为CP，其余四次成熟为EP，这直接说明FEP主导与EP/CP是两个不同分类维度。1980年后的SODA3样本中，2023/24是唯一的FEP主导事件，因此本报告把它作为FEP-EP单独延伸，不并入历史EP平均。",
    )
    add_plain_paragraph(
        doc,
        "本地Jiang正文没有Supporting Figure S6的非FEP逐事件标签。1982年后的EP/CP标签因此保留为与正文总体计数一致的工作重建：历史EP为1982/83、1986/87、1997/98和2015/16；历史CP为1987/88、1991/92、1994/95、2002/03、2004/05、2006/07、2009/10、2018/19和2019/20。物理量本身不依赖该标签。",
    )
    add_table_title(doc, "表1  Jiang事件分类、FEP属性与本报告热收支覆盖范围")
    add_classification_table(doc, manifest)

    ordered_events = (*EP_EVENTS, *CP_EVENTS, 2023)
    for number, year in enumerate(ordered_events, start=1):
        add_event_page(doc, year, number, manifest, budget)

    doc.add_page_break()
    doc.add_heading("4  逐月结果的直接归纳", level=1)
    add_plain_paragraph(
        doc,
        "1997年6月中箱向东输出为+1.978 PW，东箱西侧输入为+2.110 PW，东箱OHC′为+16.268 ZJ；2015年6月对应为+0.728 PW、+0.569 PW和+13.168 ZJ；2023年6月为+1.096 PW、+1.410 PW和+8.230 ZJ。2023的跨140°W输送并不弱，但东箱累计暖量规模低于两个强EP事件。",
    )
    add_plain_paragraph(
        doc,
        "2023年1—3月东箱海表损失均为负，4月和5月仍为负，表示东箱主要从大气获得异常热量；到6月海表损失才转为+0.093 PW。与此同时，东箱向北输出在不同月份改变符号，说明不能把整个发展期概括成持续北向释放。逐月表比季节平均更清楚地保留了这些转折。",
    )
    add_plain_paragraph(
        doc,
        "EP/CP分类不能由某一个月、某一条边界或东箱OHC′单独决定。月度表显示，同一类型内部也存在明显差异，部分CP事件也会在个别月份出现较强东向输送。分类用于组织比较，热量去向仍以每次事件、每个月的直接物理量为准。",
    )

    doc.add_heading("5  数据与分类来源", level=1)
    add_plain_paragraph(
        doc,
        "热量数据来自现有SODA3.15.2固定0—300 m、质量守恒两箱OHC异常收支文件soda3_three_box_ohc_anomaly_budget_monthly.csv。分类来自jiang2026_epcp_fep_event_manifest_1963_2023.csv，其中正文明确标签与工作重建标签分别保留。",
    )
    add_plain_paragraph(
        doc,
        "时间排布参照enso_event_comparison_unified_1980_2014_jan_jun_2026.docx：发展年1月至6月逐月观察。该文档的2026诊断不进入本报告的历史热事件统计。Jiang文献为Jiang, Chen, Zhang and Tseng (2026), Early establishment of far eastern Pacific warming shapes distinctive air-sea coupling in the 2023—2024 El Niño, Geophysical Research Letters, e2026GL123227。",
    )

    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_report())
