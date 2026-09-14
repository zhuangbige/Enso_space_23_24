#!/usr/bin/env python3
"""Replace the correlation branch of the ENSO report with 2026 profile evidence."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization_"
    "Strength_Profile_2026_Revision.docx"
)
PROFILE_FIGURE = (
    ROOT
    / "outputs/research_2026/enso_strength_vertical_profiles_2026/"
    "fig_jfm_amj_strong_moderate_difference_2026_vertical_temperature.png"
)


def first_text_run(paragraph: Paragraph):
    for run in paragraph.runs:
        if run.text:
            return run
    return paragraph.runs[0] if paragraph.runs else None


def replace_text(paragraph: Paragraph, text: str, run_template=None) -> Paragraph:
    if run_template is None:
        run_template = first_text_run(paragraph)
    rpr = (
        deepcopy(run_template._r.rPr)
        if run_template is not None and run_template._r.rPr is not None
        else None
    )
    paragraph.clear()
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return paragraph


def clone_after(
    reference: Paragraph, template: Paragraph, text: str, run_template=None
) -> Paragraph:
    element = deepcopy(template._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    return replace_text(paragraph, text, run_template or first_text_run(template))


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix!r}")


def paragraph_index(paragraphs: list[Paragraph], target: Paragraph) -> int:
    for index, paragraph in enumerate(paragraphs):
        if paragraph._p is target._p:
            return index
    raise ValueError("Target paragraph is not in the document paragraph list")


def set_page_break_before(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.page_break_before = True


def replace_eleventh_image(document: Document) -> None:
    if len(document.inline_shapes) != 11:
        raise ValueError(f"Expected 11 figures, found {len(document.inline_shapes)}")
    if not PROFILE_FIGURE.exists():
        raise FileNotFoundError(PROFILE_FIGURE)

    # Match the manuscript style by removing the raster title and source strip;
    # the scientific axes, panel labels, colorbar, and discriminator box remain.
    with Image.open(PROFILE_FIGURE) as image:
        if image.size != (5165, 2351):
            raise ValueError(f"Unexpected profile-figure size: {image.size}")
        cropped = image.crop((0, 130, 5165, 2271))
        buffer = BytesIO()
        cropped.save(buffer, format="PNG", optimize=True)
        blob = buffer.getvalue()
        aspect_ratio = cropped.height / cropped.width

    shape = document.inline_shapes[10]
    blip = shape._inline.xpath(".//a:blip")[0]
    relationship_id = blip.get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = blob

    # The original Fig. 8 part carried a source-rectangle crop. It must not be
    # inherited by the wider four-column replacement figure.
    for source_rectangle in shape._inline.xpath(".//a:srcRect"):
        parent = source_rectangle.getparent()
        if parent is not None:
            parent.remove(source_rectangle)

    shape.width = Inches(6.3)
    shape.height = int(shape.width * aspect_ratio)
    shape._inline.docPr.set("name", "Figure 8 vertical profile comparison")
    shape._inline.docPr.set(
        "descr",
        "JFM and AMJ equatorial-Pacific longitude-depth temperature anomalies "
        "for strong and moderate El Nino composites, their difference, and 2026",
    )


def revise_methods(document: Document) -> None:
    title = document.paragraphs[0]
    replace_text(title, "Data and Vertical-Profile Diagnostics Used in This Report")
    replace_text(
        document.paragraphs[1],
        "This report compares the January-June evolution of sea surface temperature, "
        "upper-ocean heat storage, and wind fields in 1997, 2015, and 2023, then uses "
        "strength-conditioned longitude-depth profiles to diagnose whether the "
        "January-June 2026 oceanic development pathway more closely resembles historical "
        "strong or moderate El Niño events.",
    )
    replace_text(
        document.paragraphs[5],
        "Wind fields and the real-time 2026 ocean state. The case-study maps are "
        "overlaid with ERA5 850-hPa wind anomalies, defined using same-calendar-month "
        "linear detrending over 1982-2024. Because the local SODA3 archive ends in 2024, "
        "the 2026 vertical profiles use NOAA/NCEP GODAS potential temperature relative "
        "to its 1991-2020 monthly climatology.",
    )
    replace_text(
        document.paragraphs[6],
        "Events and time windows. The case studies are 1997, 2015, and 2023. The working "
        "strength classification, based on peak amplitude during the first development "
        "winter, defines strong events as 1982, 1991, 1997, 2009, 2015, and 2023, and "
        "moderate events as 1986, 1994, and 2002. The extended 1986-1988 episode is assigned "
        "to the moderate class by its first winter and should be retained as a sensitivity "
        "case. JFM represents the initial reservoir and AMJ the subsequent transfer stage.",
    )

    old_heading = find_paragraph(
        document, "Cross-Event Spreading Metric and JJA Coupling Residual"
    )
    next_heading = find_paragraph(
        document, "1997: Full-Depth Warming Leads, Followed by a Surface Surge"
    )
    paragraphs = document.paragraphs
    start = paragraph_index(paragraphs, old_heading)
    stop = paragraph_index(paragraphs, next_heading)
    heading_template = deepcopy(old_heading)
    subheading_template = deepcopy(paragraphs[start + 1])
    body_template = deepcopy(paragraphs[start + 2])
    compact_body_template = deepcopy(paragraphs[start + 3])
    heading_run = first_text_run(old_heading)
    subheading_run = first_text_run(paragraphs[start + 1])
    body_run = first_text_run(paragraphs[start + 2])
    compact_body_run = first_text_run(paragraphs[start + 3])

    for paragraph in paragraphs[start + 1 : stop]:
        remove_paragraph(paragraph)

    replace_text(
        old_heading,
        "Historical Strength Classes and 2026 Vertical-Profile Diagnostics",
        heading_run,
    )
    set_page_break_before(old_heading)
    ref = clone_after(
        old_heading,
        subheading_template,
        "Comparison Design",
        subheading_run,
    )
    ref = clone_after(
        ref,
        body_template,
        "Historical monthly SODA3 temperature anomalies are cosine-latitude-weighted over "
        "5°S-5°N and interpolated to a common 121°E-81°W, 0-300 m grid. Event fields are "
        "averaged within the strong and moderate classes, and the strong-minus-moderate "
        "difference isolates the structure associated with event strength. The 2026 GODAS "
        "column is processed on the same common grid and summarized for JFM and AMJ.",
        body_run,
    )
    ref = clone_after(
        ref, subheading_template, "Process Metrics", subheading_run
    )
    ref = clone_after(
        ref,
        compact_body_template,
        "Heat-reservoir depth is the positive-temperature-anomaly centroid depth within "
        "the west-central Pacific (121°E-150°W, 0-300 m) during JFM. It distinguishes a "
        "deep preconditioned reservoir from a shallower anomaly that is already closer to "
        "the surface.",
        compact_body_run,
    )
    ref = clone_after(
        ref,
        compact_body_template,
        "Eastward-transfer rate is represented by the longitude displacement of the "
        "basinwide positive-anomaly centroid from JFM to AMJ, divided by the three months "
        "between the seasonal centers. This is a pathway-speed proxy, not a Lagrangian water-"
        "parcel velocity.",
        compact_body_run,
    )
    ref = clone_after(
        ref,
        compact_body_template,
        "Shoaling and upper-ocean release are evaluated with two complementary quantities: "
        "the fraction of positive eastern-Pacific heat in the upper 100 m during AMJ, and "
        "the JFM-to-AMJ gain of eastern-Pacific 0-100 m heat scaled by the initial west-"
        "central 100-300 m reservoir. The latter is an efficiency proxy; values above one "
        "indicate that continuing recharge, local surface forcing, or product differences "
        "also contribute and therefore must not be interpreted as literal conversion "
        "efficiency.",
        compact_body_run,
    )
    ref = clone_after(
        ref, subheading_template, "Interpretive Limits", subheading_run
    )
    clone_after(
        ref,
        body_template,
        "The historical composites and the 2026 target come from different ocean products "
        "and anomaly baselines. Overlapping GODAS-SODA3 event checks reproduce the main "
        "longitudinal-vertical evolution, but direct amplitude differences remain less "
        "secure than the sign, location, eastward progression, and shoaling of the anomaly. "
        "Accordingly, the 2026 diagnosis is a process-pathway classification as of June "
        "2026, not a forecast of the final 2026/27 Niño amplitude.",
        body_run,
    )


def revise_results(document: Document) -> None:
    heading_1997 = find_paragraph(document, "1997: Full-Depth Warming")
    heading_2023 = find_paragraph(document, "2023: Far-Eastern Surface Warming")
    replace_text(
        heading_2023,
        "2023: Far-Eastern Surface Warming Leads, with Full-Depth Warming Catching Up "
        "About One Month Later",
        first_text_run(heading_1997),
    )
    heading_2023.paragraph_format.left_indent = Inches(0)
    heading_2023.paragraph_format.first_line_indent = Inches(0)

    body_2015 = find_paragraph(document, "In January 2015")
    body_2023 = find_paragraph(document, "In January 2023")
    body_2023_text = body_2023.text
    replace_text(body_2023, body_2023_text, first_text_run(body_2015))
    body_2023.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body_2023.paragraph_format.left_indent = Inches(0)
    body_2023.paragraph_format.first_line_indent = Inches(0)

    replace_text(
        find_paragraph(document, "2015: Persistent Subsurface Warming"),
        "2015 Event — Persistent Subsurface Warming, Brief Surface Cooling",
    )
    transition = find_paragraph(document, "Thus, the distinctive feature of 2023")
    replace_text(
        transition,
        "Thus, the distinctive feature of 2023 was not merely a southward displacement of "
        "heat. Rather, heat was not concentrated near the equatorial warm core but spread "
        "to both sides of the equator, especially southward, producing a broader and deeper "
        "heat distribution. This comparison also shows why a single 0-300 m OHC value is "
        "insufficient for diagnosing a developing event: the same integrated storage can "
        "represent a deep west-central reservoir or heat that has already propagated eastward "
        "and shoaled into the eastern 50-100 m layer.",
    )

    section_heading = find_paragraph(
        document, "From Case Studies to Cross-Event Evidence"
    )
    section_index = paragraph_index(document.paragraphs, section_heading)
    result_intro = document.paragraphs[section_index + 1]
    figure_setup = document.paragraphs[section_index + 2]
    image_paragraph = document.paragraphs[section_index + 3]
    caption = document.paragraphs[section_index + 4]
    result_paragraph = document.paragraphs[section_index + 5]

    replace_text(
        section_heading,
        "Strength-Conditioned Vertical Structure and the 2026 Development Path",
    )
    set_page_break_before(section_heading)
    replace_text(
        result_intro,
        "The case-study sections identify several forms of SSTA-OHC desynchronization, "
        "but they do not by themselves establish which historical intensity pathway best "
        "describes 2026. Figure 8 therefore conditions the vertical structure on event "
        "strength and places the January-June 2026 evolution beside the strong and moderate "
        "composites.",
    )
    replace_text(
        figure_setup,
        "The comparison is organized as strong composite, moderate composite, strong minus "
        "moderate, and 2026, with JFM and AMJ in the two rows. This layout directly tests "
        "three process questions: whether the initial warm reservoir is deep, how quickly its "
        "centroid shifts eastward, and how efficiently the anomaly enters the eastern-Pacific "
        "upper ocean. The dashed box highlights 150°W-81°W and 50-100 m, the layer that most "
        "clearly separates the historical strength classes in AMJ.",
    )
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    replace_eleventh_image(document)
    replace_text(
        caption,
        "Fig. 8: JFM and AMJ equatorial-Pacific (5°S-5°N) longitude-depth potential-"
        "temperature anomalies for the historical strong composite, moderate composite, "
        "strong-minus-moderate difference, and January-June 2026. Historical fields use "
        "SODA3.15.2 same-calendar-month detrended anomalies over 1980-2024; 2026 uses NOAA/"
        "NCEP GODAS anomalies relative to the 1991-2020 GODAS monthly climatology. Strong "
        "events are 1982, 1991, 1997, 2009, 2015, and 2023; moderate events are 1986, 1994, "
        "and 2002. The dashed box marks the eastern-Pacific 50-100 m diagnostic layer. "
        "Spatial progression is more directly comparable than absolute amplitude across "
        "the two products.",
    )
    replace_text(
        result_paragraph,
        "Initial reservoir depth is not the feature that makes 2026 strong-like. In JFM, "
        "the west-central warm-reservoir centroid is 137.2 m in 2026, essentially identical "
        "to the moderate composite (137.4 m) and shallower than the strong composite "
        "(155.6 m). The January centroid is likewise near 140 m. Thus, 2026 did not begin "
        "from the canonical deep-reservoir configuration of the historical strong composite.",
    )
    extra_1 = clone_after(
        result_paragraph,
        result_intro,
        "The subsequent eastward transfer, however, is unambiguously strong-like. From JFM "
        "to AMJ, the positive-anomaly centroid shifts 56.8° eastward in 2026, equivalent to "
        "a seasonal-center displacement rate of about 18.9° longitude per month. The strong "
        "and moderate composites shift 42.4° (14.1° per month) and 7.2° (2.4° per month), "
        "respectively. The monthly trajectory gives the same result: from January to June, "
        "the 2026 centroid moves about 74.6° eastward and shoals from 140 to 95 m, compared "
        "with 60.6° and 156 to 90 m for the strong composite, and only 11.4° with little "
        "shoaling for the moderate composite.",
    )
    extra_2 = clone_after(
        extra_1,
        result_intro,
        "The decisive separation appears in the eastern-Pacific 50-100 m layer during AMJ. "
        "Its mean anomaly reaches +2.78°C in 2026, compared with +1.58°C in the strong "
        "composite and approximately 0.00°C in the moderate composite. The eastern-Pacific "
        "upper-100 m positive fraction is 0.625 in 2026, between the strong (0.699) and "
        "moderate (0.553) composites, while the upper-ocean release proxy is 1.13, above "
        "the strong (0.89) and moderate (0.64) values. Because that proxy exceeds one and "
        "the 2026 column uses GODAS, it indicates unusually rapid delivery plus continuing "
        "heat input rather than a literal efficiency greater than 100%. The combined evidence "
        "therefore classifies 2026 as a noncanonical strong-pathway case: a moderate-like "
        "initial depth followed by exceptionally rapid eastward transfer and shoaling.",
    )
    extra_2.paragraph_format.keep_together = True


def revise_conclusions(document: Document) -> None:
    conclusion_intro = find_paragraph(document, "The analysis supports three conclusions")
    index = paragraph_index(document.paragraphs, conclusion_intro)
    heading = document.paragraphs[index - 1]
    conclusion_1 = document.paragraphs[index + 1]
    conclusion_2 = document.paragraphs[index + 2]
    conclusion_3 = document.paragraphs[index + 3]

    synthesis_run = first_text_run(find_paragraph(document, "Interpretive Limits"))
    replace_text(heading, "Synthesis", synthesis_run)
    replace_text(conclusion_intro, "The revised analysis supports three conclusions:")
    replace_text(
        conclusion_1,
        "Conclusion 1. The ocean did not lack heat in 2023, but during the early development "
        "phase the heat was located mainly in the western basin, below the surface, and south "
        "of the equator. It did not overlap promptly with surface warming in the eastern "
        "equatorial Pacific, indicating slower establishment of development-phase air-sea "
        "coupling rather than persistent decoupling during the mature phase.",
    )
    replace_text(
        conclusion_2,
        "Conclusion 2. Compared with 1997 and 2015, warm anomalies in 2023 were broader and "
        "deeper. This event-level result motivates diagnosing not only total OHC but also the "
        "longitude, depth, and meridional allocation of the warm reservoir.",
    )
    replace_text(
        conclusion_3,
        "Conclusion 3. As of June 2026, the initial heat-reservoir depth is moderate-like, "
        "but the subsequent eastward displacement, shoaling, and warming of the eastern "
        "50-100 m layer are strong-like and in several measures exceed the historical strong "
        "composite. The present evidence therefore supports a shallow-initial-reservoir, "
        "fast-transfer route into a strong-event pathway, while the cross-product design and "
        "incomplete 2026/27 life cycle preclude treating this as a final intensity forecast.",
    )


def audit(document: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    forbidden = (
        "partial r",
        "permutation p",
        "JJA Pattern-Adjusted Composite Coupling Residual",
        "Cross-Event Spreading Metric",
    )
    for phrase in forbidden:
        if phrase in text:
            raise ValueError(f"Removed correlation branch remains: {phrase}")
    required = (
        "Strong-Conditioned Vertical Structure".replace("Strong-", "Strength-"),
        "+2.78°C",
        "noncanonical strong-pathway case",
    )
    for phrase in required:
        if phrase not in text:
            raise ValueError(f"Required revised content is missing: {phrase}")
    if len(document.inline_shapes) != 11:
        raise ValueError("The report must retain ten original figures plus revised Fig. 8")
    if len(document.element.xpath(".//m:oMathPara")) != 2:
        raise ValueError("Only the two retained display equations should remain")


def strip_embedded_fonts(path: Path) -> None:
    """Remove embedded-font payloads that trigger a LibreOffice import crash."""
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(path) as source, ZipFile(
            temporary, "w", compression=ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                if info.filename.startswith("word/fonts/"):
                    continue
                if info.filename == "word/_rels/fontTable.xml.rels":
                    continue
                data = source.read(info.filename)
                if info.filename == "word/fontTable.xml":
                    root = etree.fromstring(data)
                    for element in root.xpath(
                        ".//*[starts-with(local-name(), 'embed')]",
                        namespaces=namespaces,
                    ):
                        parent = element.getparent()
                        if parent is not None:
                            parent.remove(element)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone="yes",
                    )
                target.writestr(info, data)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    revise_methods(document)
    revise_results(document)
    revise_conclusions(document)
    document.core_properties.title = (
        "ENSO Development-Phase SSTA-OHC Desynchronization and the 2026 Vertical Profile"
    )
    document.core_properties.subject = (
        "Revised manuscript replacing the event-correlation branch with strength-conditioned "
        "vertical-profile evidence"
    )
    audit(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    strip_embedded_fonts(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
