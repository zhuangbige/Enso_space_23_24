#!/usr/bin/env python3
"""Replace OHC maps and vertical sections in the 2026-08-16 ENSO report."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "outputs/research_2026/reports"
SOURCE = REPORT_DIR / "ENSO发展期_SSTA-OHC不同步_阶段性汇报_暖核移动补充版_20260816.docx"
OUTPUT = REPORT_DIR / "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_20260818.docx"

OHC_MAPS = {
    year: ROOT
    / "outputs/research_2026/soda3_ohc300_monthly_wind_panels_30s30n_1980_2024_detrended"
    / f"fig_{year}_jan_jun_soda3_ohc300_30s30n_era5_850_windA_1980_2024.png"
    for year in (1997, 2015, 2023)
}

SSTA_MAPS = {
    year: ROOT
    / f"outputs/research_2026/event_{year}_monthly_panels_1980_2024_detrended"
    / f"fig_{year}_jan_jun_oisst_ssta_era5_850_windA_1980_2024.png"
    for year in (1997, 2015, 2023)
}

FULL_PACIFIC_SECTIONS = {
    year: ROOT
    / "outputs/research_2026/soda3_full_pacific_meridional_sections_120e80w_1980_2024_detrended"
    / (
        f"fig_{year}_jan_jun_soda3_full_pacific_meridional_section_"
        "120e80w_30s30n_1980_2024_detrended.png"
    )
    for year in (1997, 2015, 2023)
}

FONT = "Noto Sans CJK SC"
FIGURE_WIDTH = Inches(6.20)
EVENT_FIGURE_WIDTH = Inches(5.95)


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def clear_pagination_properties(paragraph: Paragraph) -> None:
    """Remove Word properties that render as black square marks in Show/Hide mode."""
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        for tag in ("keepNext", "keepLines", "pageBreakBefore", "suppressLineNumbers"):
            element = p_pr.find(qn(f"w:{tag}"))
            if element is not None:
                p_pr.remove(element)
    for element in paragraph._p.xpath('.//*[local-name()="lastRenderedPageBreak"]'):
        element.getparent().remove(element)


def apply_run_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{slot}"), FONT)


def replace_text(paragraph: Paragraph, text: str) -> Paragraph:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)
    apply_run_font(run)
    return paragraph


def paragraph_by_prefix(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def picture_paragraph_before(caption: Paragraph) -> Paragraph:
    previous = caption._p.getprevious()
    if previous is None:
        raise ValueError(f"No picture paragraph before caption {caption.text!r}")
    paragraph = Paragraph(previous, caption._parent)
    if not paragraph._p.xpath('.//*[local-name()="drawing"]'):
        raise ValueError(f"Paragraph before caption is not a picture: {caption.text!r}")
    return paragraph


def set_picture(
    paragraph: Paragraph,
    path: Path,
    alt_text: str,
    *,
    width=FIGURE_WIDTH,
) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run()
    apply_run_font(run)
    run.add_picture(str(path), width=width)
    doc_properties = paragraph._p.xpath('.//*[local-name()="docPr"]')
    if doc_properties:
        doc_properties[0].set("descr", alt_text)


def remove_standalone_page_break_before(paragraph: Paragraph) -> None:
    previous = paragraph._p.getprevious()
    if previous is None:
        return
    previous_paragraph = Paragraph(previous, paragraph._parent)
    has_page_break = bool(
        previous_paragraph._p.xpath(
            './/*[local-name()="br" and @*[local-name()="type"]="page"]'
        )
    )
    if not previous_paragraph.text.strip() and has_page_break:
        previous_paragraph._p.getparent().remove(previous_paragraph._p)


def insert_page_break_before(paragraph: Paragraph) -> Paragraph:
    previous = paragraph._p.getprevious()
    if previous is not None:
        page_break = Paragraph(previous, paragraph._parent)
        if not page_break.text and not page_break._p.xpath('.//*[local-name()="drawing"]'):
            clear_paragraph(page_break)
        else:
            page_break_xml = OxmlElement("w:p")
            paragraph._p.addprevious(page_break_xml)
            page_break = Paragraph(page_break_xml, paragraph._parent)
    else:
        page_break_xml = OxmlElement("w:p")
        paragraph._p.addprevious(page_break_xml)
        page_break = Paragraph(page_break_xml, paragraph._parent)
    page_break.style = "Normal"
    clear_pagination_properties(page_break)
    page_break.add_run().add_break(WD_BREAK.PAGE)
    return page_break


def add_page_break(document: Document) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    clear_pagination_properties(paragraph)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def remove_content_after(paragraph: Paragraph) -> None:
    body = paragraph._p.getparent()
    current = paragraph._p.getnext()
    while current is not None and current.tag != qn("w:sectPr"):
        following = current.getnext()
        body.remove(current)
        current = following


def add_body(document: Document, text: str, *, after: float = 6.0) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.widow_control = False
    run = paragraph.add_run(text)
    apply_run_font(run, size=10.5)
    return paragraph


def format_plain_heading(paragraph: Paragraph, level: int) -> Paragraph:
    sizes = {1: 21.0, 2: 15.0, 3: 12.5}
    paragraph.style = "Normal"
    clear_pagination_properties(paragraph)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.first_line_indent = None
    for run in paragraph.runs:
        apply_run_font(run, size=sizes[level], bold=True)
    return paragraph


def add_heading(
    document: Document,
    text: str,
    level: int,
) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    run = paragraph.add_run(text)
    apply_run_font(run, bold=True)
    return format_plain_heading(paragraph, level)


def add_picture(
    document: Document,
    path: Path,
    alt_text: str,
    *,
    width=FIGURE_WIDTH,
) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    set_picture(paragraph, path, alt_text, width=width)
    return paragraph


def add_caption(document: Document, text: str) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.widow_control = False
    run = paragraph.add_run(text)
    apply_run_font(run, size=8.2)
    return paragraph


def replace_ohc_maps(document: Document) -> None:
    ssta_captions = {
        1997: paragraph_by_prefix(document, "图1（a）"),
        2015: paragraph_by_prefix(document, "图2（a）"),
        2023: paragraph_by_prefix(document, "图3（a）"),
    }
    captions = {
        1997: paragraph_by_prefix(document, "图1（b）"),
        2015: paragraph_by_prefix(document, "图2（b）"),
        2023: paragraph_by_prefix(document, "图3（b）"),
    }
    for year, caption in ssta_captions.items():
        set_picture(
            picture_paragraph_before(caption),
            SSTA_MAPS[year],
            f"{year}年一月至六月OISST海温异常与ERA5八百五十百帕风异常",
            width=EVENT_FIGURE_WIDTH,
        )
    for year, caption in captions.items():
        set_picture(
            picture_paragraph_before(caption),
            OHC_MAPS[year],
            f"{year}年一月至六月30度南北纬SODA3上300米OHC异常与ERA5八百五十百帕风异常",
            width=EVENT_FIGURE_WIDTH,
        )
    replace_text(
        captions[1997],
        "图1（b） 1997年1—6月SODA3 0—300 m OHC异常及ERA5 850 hPa风异常。范围扩展为30°S—30°N、120°E—80°W；OHC和风均采用1980—2024年同月线性去趋势基准，三次事件使用相同±3 GJ m⁻²色标。",
    )
    replace_text(
        captions[2015],
        "图2（b） 2015年1—6月SODA3 0—300 m OHC异常及ERA5 850 hPa风异常。空间范围、距平定义和色标同图1（b）。",
    )
    replace_text(
        captions[2023],
        "图3（b） 2023年1—6月SODA3 0—300 m OHC异常及ERA5 850 hPa风异常。空间范围、距平定义和色标同图1（b）。",
    )

    event_paragraphs = {
        1997: (
            paragraph_by_prefix(document, "1997："),
            paragraph_by_prefix(document, "FEP中OHC"),
            ssta_captions[1997],
            captions[1997],
            paragraph_by_prefix(document, "图1表明"),
        ),
        2015: (
            paragraph_by_prefix(document, "2015："),
            paragraph_by_prefix(document, "2015年1月表层"),
            ssta_captions[2015],
            captions[2015],
            paragraph_by_prefix(document, "图2表明"),
        ),
        2023: (
            paragraph_by_prefix(document, "2023："),
            paragraph_by_prefix(document, "2023年1月，"),
            ssta_captions[2023],
            captions[2023],
            paragraph_by_prefix(document, "图3表明"),
        ),
    }
    for heading, introduction, ssta_caption, ohc_caption, conclusion in event_paragraphs.values():
        format_plain_heading(heading, 1)
        heading.paragraph_format.space_after = Pt(3)
        for run in heading.runs:
            apply_run_font(run, size=21)
        introduction.paragraph_format.line_spacing = 1.05
        introduction.paragraph_format.space_after = Pt(3)
        conclusion.paragraph_format.line_spacing = 1.05
        conclusion.paragraph_format.space_before = Pt(1)
        conclusion.paragraph_format.space_after = Pt(1)
        for paragraph in (introduction, conclusion):
            for run in paragraph.runs:
                apply_run_font(run, size=9.8)
        for caption in (ssta_caption, ohc_caption):
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.line_spacing = 1.0
            caption.paragraph_format.space_before = Pt(1)
            caption.paragraph_format.space_after = Pt(2)
            for run in caption.runs:
                apply_run_font(run, size=7.8)


def rebuild_section_part(document: Document) -> None:
    heading = paragraph_by_prefix(document, "连续垂向结构：")
    old_caption = paragraph_by_prefix(document, "图4")
    old_equatorial_picture = deepcopy(picture_paragraph_before(old_caption)._p)
    remove_standalone_page_break_before(heading)
    remove_content_after(heading)
    replace_text(heading, "连续剖面结构：赤道暖异常东移与经向扩展")
    format_plain_heading(heading, 1)
    insert_page_break_before(heading)

    add_body(
        document,
        "经度—深度剖面保留原来的5°S—5°N赤道带，用于追踪暖舌由西向东、由深向浅的演变；随后对固定120°E—80°W范围作完整经度平均，用纬度—深度剖面比较三次事件的南北结构。",
    )
    body = document.element.body
    body.insert(len(body) - 1, old_equatorial_picture)
    equatorial_picture = Paragraph(old_equatorial_picture, document._body)
    equatorial_picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_pagination_properties(equatorial_picture)
    add_caption(
        document,
        "图4  1997、2015和2023年1—6月赤道太平洋热含量密度异常的经度—深度剖面。5°S—5°N余弦纬度加权；SODA3原生上300 m深度层；固定经度范围120°E—80°W。",
    )
    add_heading(document, "赤道带剖面解释", 2)
    add_body(document, "1997年初中太平洋约100—200 m存在显著暖核，春季向东上翻并形成深厚的东部暖异常，支持整层领先于表层的判断。")
    add_body(document, "2015年次表层暖舌从年初持续连接中、东太平洋；即使海表短暂偏冷，深层暖异常仍然延续，显示出较强的海洋记忆。")
    add_body(document, "2023年初暖量更偏西、偏深，远东海表先变暖，暖舌随后向东上翻；到初夏，深层支撑仍弱于1997和2015。")

    add_page_break(document)
    add_heading(document, "固定完整经度平均的纬度—深度剖面", 2)
    add_body(
        document,
        "以下各图每个月都对同一120°E—80°W范围取平均。该诊断用于比较盆地平均的南北与垂直结构；由于东西向信号被平均，暖核东移过程仍需结合图4解释。",
    )

    meridional_text = {
        1997: (
            "图5  1997年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。范围30°S—30°N、0—300 m；1980—2024年同月线性去趋势，统一±6 MJ m⁻³色标。",
            "1997年暖异常主要集中在赤道附近100—200 m，并在春季明显增强和上抬；南侧约5°—15°S存在较弱的深层正异常，而北侧5°—15°N同期出现较强冷异常。",
        ),
        2015: (
            "图6  2015年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。定义同图5。",
            "2015年暖异常从年初开始偏浅并集中在赤道附近，5°—15°N冷异常随春季增强；南半球次表层暖异常不连续，进一步说明该事件的经向结构相对偏北。",
        ),
        2023: (
            "图7  2023年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。定义同图5。",
            "2023年1—2月约150—250 m暖异常向南扩展至10°—15°S最明显，随后暖异常逐步向赤道集中并上抬。该结果支持2023年早期存在较宽的南侧深层暖异常，但不能单独表征其经向输送机制。",
        ),
    }
    for index, year in enumerate((1997, 2015, 2023)):
        if index > 0:
            add_page_break(document)
        add_heading(document, f"{year}年完整经度平均", 3)
        add_picture(
            document,
            FULL_PACIFIC_SECTIONS[year],
            f"{year}年一月至六月固定完整太平洋经度平均的纬度深度热含量密度异常剖面",
            width=Inches(6.20),
        )
        caption, interpretation = meridional_text[year]
        add_caption(document, caption)
        add_body(document, interpretation)

    add_heading(document, "剖面结论", 2)
    add_body(
        document,
        "赤道带剖面表明，三次事件都具有暖异常向东上翻的共同过程；完整经度平均的纬度—深度剖面进一步显示经向结构并不相同。1997以赤道主暖舌为主并伴随较弱南侧分支；2015的赤道以北结构更突出；2023发展早期的南侧次表层暖异常最清楚。",
    )
    add_body(
        document,
        "上述剖面描述的是温度异常换算的热含量密度异常及其空间配置，不能仅凭图形区分水平平流、波动调整、垂向交换或表面热通量。关于“向南扩散”的机制表述，后续仍应结合经向热输送或热收支诊断进行定量验证。",
    )


def set_document_font_and_remove_numbering(document: Document) -> None:
    for style in document.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = FONT
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{slot}"), FONT)

    containers = [document.paragraphs]
    for section in document.sections:
        containers.extend([section.header.paragraphs, section.footer.paragraphs])
    for paragraphs in containers:
        for paragraph in paragraphs:
            p_pr = paragraph._p.get_or_add_pPr()
            num_pr = p_pr.find(qn("w:numPr"))
            if num_pr is not None:
                p_pr.remove(num_pr)
            clear_pagination_properties(paragraph)
            for run in paragraph.runs:
                apply_run_font(run)


def audit_no_bullets(document: Document) -> None:
    forbidden = ("•", "●", "▪", "◦", "‣")
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph._p.xpath('./*[local-name()="pPr"]/*[local-name()="numPr"]'):
            raise ValueError(f"Paragraph {index} still has numbering")
        if paragraph._p.xpath(
            './*[local-name()="pPr"]/*[local-name()="keepNext" or '
            'local-name()="keepLines" or local-name()="pageBreakBefore" or '
            'local-name()="suppressLineNumbers"]'
        ):
            raise ValueError(f"Paragraph {index} still has black-square pagination properties")
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            raise ValueError(f"Paragraph {index} still uses a Heading style")
        if paragraph.text.lstrip().startswith(forbidden):
            raise ValueError(f"Paragraph {index} starts with a bullet character")


def main() -> int:
    required = [
        SOURCE,
        *SSTA_MAPS.values(),
        *OHC_MAPS.values(),
        *FULL_PACIFIC_SECTIONS.values(),
    ]
    for path in required:
        if not path.exists():
            raise FileNotFoundError(path)

    document = Document(SOURCE)
    replace_ohc_maps(document)
    rebuild_section_part(document)
    set_document_font_and_remove_numbering(document)
    audit_no_bullets(document)

    document.core_properties.title = "ENSO发展期SSTA—OHC不同步：30度图、赤道带与纬向剖面更新"
    document.core_properties.subject = "阶段性汇报：OHC—风异常与暖核经向结构"
    document.core_properties.comments = "2026-08-18：保留30度OHC图，剖面改为赤道带与固定完整经度平均"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
