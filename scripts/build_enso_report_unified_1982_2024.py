#!/usr/bin/env python3
"""Recompute the ENSO stage report with a unified 1982--2024 anomaly baseline."""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import numpy as np
import pandas as pd
import xarray as xr
from PIL import Image

CODEX_DOCUMENT_SITE = Path(
    "/Users/daiyue/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/python/lib/python3.12/site-packages"
)
if CODEX_DOCUMENT_SITE.exists():
    sys.path.append(str(CODEX_DOCUMENT_SITE))

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import analyze_3d_warm_heat_spreading_decoupling as cross  # noqa: E402
import analyze_soda3_all_warm_event_core_trajectories as core_centroid  # noqa: E402
import build_soda3_full_pacific_meridional_sections as meridional_builder  # noqa: E402
import build_soda3_tropical_ohc300 as tropical_ohc  # noqa: E402
import plot_enso_year_jan_jun_ssta_winda as sst_wind  # noqa: E402
import plot_soda3_full_pacific_meridional_sections as meridional_plot  # noqa: E402
import plot_soda3_monthly_ohc_vertical_sections as vertical  # noqa: E402
import plot_soda3_monthly_ohc_wind_events as ohc_wind  # noqa: E402


BASELINE = (1982, 2024)
EVENTS = (1997, 2015, 2023)
WORK = ROOT / "outputs/research_2026/enso_report_unified_1982_2024"
FIGURES = WORK / "figures"
DERIVED = ROOT / "data/derived/research_2026/enso_report_unified_1982_2024"
SOURCE_DOCX = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "论文图注无粗体终版_20260818.docx"
)
OUTPUT_DOCX = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "统一1982_2024基准_论文图注无粗体版_20260818.docx"
)

OHC_SOURCE = (
    ROOT
    / "data/interim/research_2026/soda3_ohc_tropical_30s30n/"
    "soda3_ohc300_jan_jun_30s30n_1980_2024_1deg.nc"
)
OHC_UNIFIED = DERIVED / "soda3_ohc300_jan_jun_30s30n_1982_2024_detrended.nc"
VERTICAL_UNIFIED = (
    DERIVED
    / "soda3_ohc_vertical_sections_continuous_1997_2015_2023_1982_2024_detrended.nc"
)
MERIDIONAL_UNIFIED = (
    DERIVED
    / "soda3_ohc_meridional_sections_full_pacific_120e80w_"
    "1997_2015_2023_1982_2024_detrended.nc"
)
CENTROID_UNIFIED = DERIVED / "soda3_13_warm_events_monthly_core_centroids_1982_2024.csv"
EVENT_METRICS_CSV = DERIVED / "enso_3d_warm_heat_spreading_event_metrics_1982_2024.csv"
RELATIONSHIPS_CSV = DERIVED / "enso_3d_warm_heat_spreading_relationships_1982_2024.csv"
POINTS_CSV = DERIVED / "enso_cross_event_broadening_jja_coupling_points_1982_2024.csv"
STATS_CSV = DERIVED / "enso_cross_event_broadening_jja_coupling_statistics_1982_2024.csv"
SUMMARY_JSON = DERIVED / "unified_1982_2024_report_diagnostics.json"

# Crop out only the global raster title and bottom method strip. Panel titles,
# axes, colorbars, units, annotations and statistics remain in the figures.
CROP_BOXES = {
    1: (0, 140, 4376, 1975),
    2: (0, 140, 4138, 1875),
    3: (0, 140, 4376, 1975),
    4: (0, 140, 4138, 1875),
    5: (0, 140, 4376, 1975),
    6: (0, 140, 4138, 1875),
    7: (0, 115, 6147, 2620),
    8: (0, 150, 4156, 2168),
    9: (0, 150, 4156, 2168),
    10: (0, 150, 4156, 2168),
    11: (0, 180, 3360, 1840),
}


def build_sst_and_wind_figures() -> tuple[list[Path], Path]:
    output_dir = FIGURES / "sst_wind"
    wind_dir = DERIVED / "sst_wind"
    output_dir.mkdir(parents=True, exist_ok=True)
    wind_dir.mkdir(parents=True, exist_ok=True)
    sst_wind.SST_BASELINE = BASELINE
    sst_wind.WIND_BASELINE = BASELINE
    sst_wind.DERIVED_DIR = wind_dir

    sst = sst_wind.open_oisst()
    era5 = sst_wind.open_era5_wind()
    figures: list[Path] = []
    try:
        for event in EVENTS:
            ssta = sst_wind.same_month_detrended_target(
                sst, target_year=event, baseline=BASELINE
            )
            u850 = sst_wind.same_month_detrended_target(
                era5.u850, target_year=event, baseline=BASELINE
            )
            v850 = sst_wind.same_month_detrended_target(
                era5.v850, target_year=event, baseline=BASELINE
            )
            wind = xr.Dataset({"u850a": u850, "v850a": v850})
            sst_wind.validate(ssta, wind)
            sst_wind.save_derived(ssta, u850, v850, target_year=event)
            png, _ = sst_wind.plot(
                ssta, wind, target_year=event, output_dir=output_dir
            )
            figures.append(png)
    finally:
        close_sst = getattr(sst, "close", None)
        close_wind = getattr(era5, "close", None)
        if callable(close_sst):
            close_sst()
        if callable(close_wind):
            close_wind()
    return figures, wind_dir


def build_unified_ohc_file() -> Path:
    with xr.open_dataset(OHC_SOURCE, engine="netcdf4") as source:
        absolute = source.ohc_300m.load()
        source_attrs = dict(source.attrs)
    times = pd.DatetimeIndex(absolute.time.values)
    anomaly = tropical_ohc.detrended_selected_month_anomaly(
        absolute.values, times, BASELINE[0], BASELINE[1]
    )
    dataset = xr.Dataset(
        {
            "ohc_300m": absolute.astype(np.float32),
            "ohc_300m_anomaly_detrended": (
                ("time", "lat", "lon"),
                anomaly.astype(np.float32),
            ),
        },
        coords=absolute.coords,
        attrs=source_attrs,
    )
    dataset.attrs.update(
        anomaly_definition=(
            f"{BASELINE[0]}-{BASELINE[1]} calendar-month climatology and "
            "same-calendar-month linear trend removed"
        ),
        calculation_note=(
            "Absolute 1980-2024 archive retained; anomaly fit restricted to "
            f"{BASELINE[0]}-{BASELINE[1]}."
        ),
    )
    for name in dataset.data_vars:
        dataset[name].attrs.update(units="J m-2")
    dataset.ohc_300m.attrs["long_name"] = "Ocean heat content integrated over 0-300 m"
    dataset.ohc_300m_anomaly_detrended.attrs["long_name"] = (
        "Detrended calendar-month anomaly of 0-300 m ocean heat content"
    )
    OHC_UNIFIED.parent.mkdir(parents=True, exist_ok=True)
    encoding = {
        name: {
            "zlib": True,
            "complevel": 4,
            "shuffle": True,
            "dtype": "float32",
            "_FillValue": np.float32(-1.0e20),
        }
        for name in dataset.data_vars
    }
    dataset.to_netcdf(OHC_UNIFIED, engine="netcdf4", encoding=encoding)
    return OHC_UNIFIED


def build_ohc_wind_figures(wind_dir: Path) -> list[Path]:
    output_dir = FIGURES / "ohc_wind"
    ohc_wind.OHC_BASELINE = "1982-2024"
    ohc_wind.WIND_BASELINE = "1982-2024"
    ohc_wind.WIND_DIR = wind_dir
    ohc_wind.setup_style()
    with ohc_wind.open_dataset(OHC_UNIFIED) as source:
        ohc = ohc_wind.validate_ohc(source).load()
    figures = []
    for event in EVENTS:
        png, _ = ohc_wind.plot_event(ohc, event, output_dir, 3.0)
        figures.append(png)
    return figures


def build_equatorial_section() -> Path:
    output_dir = FIGURES / "equatorial_section"
    vertical.BASELINE = BASELINE
    vertical.setup_style()
    with vertical.open_dataset(vertical.TEMP_FILE) as source:
        vertical.validate_input(source)
        equatorial = vertical.equatorial_mean_temperature(source).load()
    anomaly = vertical.detrended_same_month_anomaly(equatorial)
    diagnostics = vertical.build_diagnostics(anomaly, list(EVENTS))
    vertical.save_diagnostics(diagnostics, VERTICAL_UNIFIED)
    png, _ = vertical.plot_combined(diagnostics, list(EVENTS), output_dir, 16.0)
    return png


def build_meridional_sections() -> list[Path]:
    output_dir = FIGURES / "meridional_sections"
    meridional_builder.BASELINE = BASELINE
    meridional_builder.OUT_FILE = MERIDIONAL_UNIFIED
    meridional_builder.assemble_output()
    meridional_plot.BASELINE = BASELINE
    meridional_plot.setup_style()
    with xr.open_dataset(MERIDIONAL_UNIFIED, engine="netcdf4") as source:
        meridional_plot.validate(source)
        dataset = source.load()
    figures = []
    for event in EVENTS:
        png, _ = meridional_plot.plot_event(dataset, event, output_dir, 6.0)
        figures.append(png)
    return figures


def build_cross_event() -> tuple[Path, pd.Series, pd.DataFrame]:
    output_dir = FIGURES / "cross_event"
    output_dir.mkdir(parents=True, exist_ok=True)

    # Recompute the moving warm-core centroid with the same 1982-2024 field.
    vertical.BASELINE = BASELINE
    manifest = core_centroid.load_event_manifest()
    centroid_fields = core_centroid.build_heat_density(core_centroid.EVENTS)
    centroids = core_centroid.build_centroid_table(centroid_fields, manifest)
    centroids.to_csv(CENTROID_UNIFIED, index=False, float_format="%.6f")

    cross.BASELINE = BASELINE
    cross.CENTROID_CSV = CENTROID_UNIFIED
    cross.CROSS_EVENT_PNG = output_dir / "enso_cross_event_broadening_jja_coupling_en.png"
    cross.CROSS_EVENT_PDF = output_dir / "enso_cross_event_broadening_jja_coupling_en.pdf"
    fields, depth, longitude = cross.load_anomaly_fields()
    monthly = cross.build_monthly_table(fields, depth, longitude)
    events = cross.build_event_table(monthly)
    relationships = cross.build_relationships(events)

    total = events.total_full_0_300_gjm2_jfm.to_numpy(dtype=float)
    off_equatorial = events.offeq_full_0_300_gjm2_jfm.to_numpy(dtype=float)
    coupling = events.coupling_pattern_residual_z_jja.to_numpy(dtype=float)
    points = pd.DataFrame(
        {
            "year": events.year.to_numpy(dtype=int),
            "group": events.group.astype(str),
            "jfm_total_three_band_ohc_gjm2": total,
            "jfm_offequatorial_ohc_gjm2": off_equatorial,
            "jfm_offequatorial_ohc_residual_gjm2": cross.residualize(
                off_equatorial, total[:, None]
            ),
            "jja_pattern_adjusted_coupling_residual_z_raw": coupling,
            "jja_coupling_residual_after_total_ohc_control_z": cross.residualize(
                coupling, total[:, None]
            ),
        }
    )
    primary = relationships.loc[
        relationships.test_id.eq("component_broadening")
    ].iloc[0]
    events.to_csv(EVENT_METRICS_CSV, index=False)
    relationships.to_csv(RELATIONSHIPS_CSV, index=False)
    points.to_csv(POINTS_CSV, index=False)
    primary.to_frame().T.to_csv(STATS_CSV, index=False)
    cross.plot_cross_event_support(events, relationships)
    return cross.CROSS_EVENT_PNG, primary, points


def crop_figures(paths: list[Path]) -> tuple[dict[str, bytes], list[Path]]:
    if len(paths) != 11:
        raise ValueError(f"Expected 11 figures, got {len(paths)}")
    cropped_dir = FIGURES / "cropped_for_docx"
    cropped_dir.mkdir(parents=True, exist_ok=True)
    replacements: dict[str, bytes] = {}
    cropped_paths: list[Path] = []
    for index, path in enumerate(paths, start=1):
        box = CROP_BOXES[index]
        with Image.open(path) as image:
            # Matplotlib's tight bounding box can differ by one pixel across
            # backends. Clamp that harmless edge difference while retaining
            # the scientifically meaningful crop coordinates.
            right = min(box[2], image.width)
            bottom = min(box[3], image.height)
            if image.width < box[2] - 2 or image.height < box[3] - 2:
                raise ValueError(
                    f"Figure {index} is too small for crop {box}: {image.size} ({path})"
                )
            cropped = image.crop((box[0], box[1], right, bottom))
            buffer = BytesIO()
            cropped.save(buffer, format="PNG", optimize=True)
            replacements[f"word/media/image{index}.png"] = buffer.getvalue()
            cropped_path = cropped_dir / f"figure_{index:02d}.png"
            cropped.save(cropped_path, format="PNG", optimize=True)
            cropped_paths.append(cropped_path)
    return replacements, cropped_paths


def unicode_number(value: float, digits: int) -> str:
    return f"{value:.{digits}f}".replace("-", "−")


def set_paragraph_text(paragraph, text: str) -> None:
    size = None
    color = None
    font_name = "PingFang SC"
    if paragraph.runs:
        size = paragraph.runs[0].font.size
        color = paragraph.runs[0].font.color.rgb
        font_name = paragraph.runs[0].font.name or font_name
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = size
    run.font.bold = None
    if color is not None:
        run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def update_document_text(document: Document, primary: pd.Series) -> None:
    # Change only anomaly-baseline statements; the 1982-2014 cross-validation
    # training window is intentionally unchanged.
    for paragraph in document.paragraphs:
        if "1980" not in paragraph.text:
            continue
        updated = (
            paragraph.text.replace("1980—2024", "1982—2024")
            .replace("1980–2024", "1982–2024")
            .replace("1980-2024", "1982-2024")
        )
        set_paragraph_text(paragraph, updated)

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("风场与大气耦合量。"):
            set_paragraph_text(
                paragraph,
                "风场与大气耦合量。空间图叠加ERA5 850 hPa风异常，风异常采用"
                "1982—2024年同历月线性去趋势基准。跨事件耦合响应使用GPCP降水、"
                "ERA5海平面气压和ERA5纬向风应力；历史预期关系以1982—2014年为训练期，"
                "并按同历月逐年留一拟合。",
            )

    r_text = unicode_number(float(primary.partial_r), 2)
    p_value = float(primary.permutation_p)
    q_value = float(primary.permutation_q_fdr)
    p_text = f"{p_value:.3f}"
    q_text = f"{q_value:.3f}"
    loeo_low = unicode_number(float(primary.loeo_r_min), 2)
    loeo_high = unicode_number(float(primary.loeo_r_max), 2)

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("在13次暖事件中，JFM离赤道暖量越多"):
            set_paragraph_text(
                paragraph,
                "在13次暖事件中，JFM离赤道暖量越多，随后JJA型态调整后的海气耦合响应越弱"
                f"（偏相关r={r_text}，置换检验p={p_text}，FDR q={q_text}）。"
                f"逐事件剔除后，相关仍保持在{loeo_low}至{loeo_high}之间，说明该关系并非由"
                "1997、2015或2023中的某一个事件单独决定。结果表明，除三纬带平均暖含量之外，"
                "暖量是否集中在赤道附近与随后海气耦合响应的相对强弱有关，并与耦合转化效率降低的解释一致。",
            )
        elif paragraph.text.startswith("结论三。13次暖事件中"):
            set_paragraph_text(
                paragraph,
                "结论三。13次暖事件中，在控制三纬带面积加权平均OHC后，JFM离赤道暖量越多，"
                f"随后JJA海气耦合响应越弱（r={r_text}，p={p_text}）。这一结果支持暖量空间集中度"
                "影响发展期耦合响应的解释，但尚不能证明经向扩散是直接原因。",
            )


def write_docx_media(
    intermediate: Path, output: Path, replacements: dict[str, bytes]
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(intermediate) as source, ZipFile(
            temporary, "w", compression=ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                data = replacements.get(info.filename, source.read(info.filename))
                copied = ZipInfo(info.filename, date_time=info.date_time)
                copied.compress_type = ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.internal_attr = info.internal_attr
                copied.flag_bits = info.flag_bits
                copied.create_system = info.create_system
                target.writestr(copied, data)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()


def build_docx(replacements: dict[str, bytes], primary: pd.Series) -> Path:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)
    document = Document(SOURCE_DOCX)
    update_document_text(document, primary)
    with NamedTemporaryFile(
        dir=OUTPUT_DOCX.parent, suffix=".docx", delete=False
    ) as handle:
        intermediate = Path(handle.name)
    try:
        document.save(intermediate)
        write_docx_media(intermediate, OUTPUT_DOCX, replacements)
    finally:
        if intermediate.exists():
            intermediate.unlink()
    return OUTPUT_DOCX


def audit_docx(path: Path) -> dict[str, int]:
    document = Document(path)
    bold = len(document.element.xpath(".//w:b | .//w:bCs"))
    numbered = len(document.element.xpath(".//w:numPr"))
    equations = len(document.element.xpath(".//m:oMath"))
    captions = sum(paragraph.text.startswith("Fig. ") for paragraph in document.paragraphs)
    baseline_1980 = sum("1980" in paragraph.text for paragraph in document.paragraphs)
    if len(document.inline_shapes) != 11:
        raise ValueError("Final document does not contain 11 figures")
    if equations != 8:
        raise ValueError(f"Expected 8 native equations, found {equations}")
    if captions != 11:
        raise ValueError(f"Expected 11 external Fig. captions, found {captions}")
    if bold:
        raise ValueError(f"Bold formatting remains: {bold}")
    if numbered:
        raise ValueError(f"List numbering remains: {numbered}")
    if baseline_1980:
        raise ValueError("An editable 1980 baseline statement remains")
    return {
        "paragraphs": len(document.paragraphs),
        "figures": len(document.inline_shapes),
        "equations": equations,
        "captions": captions,
        "bold_tags": bold,
        "numbering_tags": numbered,
    }


def main() -> int:
    FIGURES.mkdir(parents=True, exist_ok=True)
    DERIVED.mkdir(parents=True, exist_ok=True)

    sst_figures, wind_dir = build_sst_and_wind_figures()
    build_unified_ohc_file()
    ohc_figures = build_ohc_wind_figures(wind_dir)
    equatorial_figure = build_equatorial_section()
    meridional_figures = build_meridional_sections()
    cross_figure, primary, points = build_cross_event()

    ordered_figures = [
        sst_figures[0],
        ohc_figures[0],
        sst_figures[1],
        ohc_figures[1],
        sst_figures[2],
        ohc_figures[2],
        equatorial_figure,
        *meridional_figures,
        cross_figure,
    ]
    replacements, cropped_paths = crop_figures(ordered_figures)
    build_docx(replacements, primary)
    audit = audit_docx(OUTPUT_DOCX)

    point_2023 = points.loc[points.year.eq(2023)].iloc[0]
    diagnostics = {
        "anomaly_baseline": list(BASELINE),
        "coupling_training_window_unchanged": [1982, 2014],
        "n_events": int(primary.n_events),
        "partial_r": float(primary.partial_r),
        "residual_rank_r": float(primary.partial_spearman_r),
        "permutation_p": float(primary.permutation_p),
        "fdr_q": float(primary.permutation_q_fdr),
        "bootstrap_95pct": [
            float(primary.bootstrap_ci_low),
            float(primary.bootstrap_ci_high),
        ],
        "leave_one_event_out_r": [
            float(primary.loeo_r_min),
            float(primary.loeo_r_max),
        ],
        "year_2023_x_gjm2": float(
            point_2023.jfm_offequatorial_ohc_residual_gjm2
        ),
        "year_2023_y_z": float(
            point_2023.jja_coupling_residual_after_total_ohc_control_z
        ),
        "docx_audit": audit,
        "cropped_figures": [str(path) for path in cropped_paths],
        "output_docx": str(OUTPUT_DOCX),
    }
    SUMMARY_JSON.write_text(
        json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(diagnostics, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
