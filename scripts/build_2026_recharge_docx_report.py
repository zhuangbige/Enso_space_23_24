#!/usr/bin/env python3
"""Build a Word report for the published-method 2026 ENSO recharge analysis."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs/research_2026/enso_recharge_published_methods_2026"
FIG1 = OUT_DIR / "fig01_published_wwv_and_transport_jan_jun_2026.png"
FIG2 = OUT_DIR / "fig02_published_d20_profiles_jfm_amj_2026.png"
FIG3 = OUT_DIR / "fig03_published_vertical_temperature_sections_2026.png"
FIG4 = OUT_DIR / "fig04_published_d20_hovmoller_jan_jun_2026.png"
OUTPUT = OUT_DIR / "2026_ENSO_Recharge_Published_Methods_Report_CN.docx"

LATIN_FONT = "Times New Roman"
EAST_ASIA_FONT = "STSong"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(85, 85, 85)


def set_run_font(run, size: float | None = None, color=BLACK, italic=False) -> None:
    # LibreOffice may apply w:hAnsi to an entire mixed Chinese/Latin run and
    # then fail to fall back from Times New Roman for CJK glyphs.  Use a
    # CJK-capable font for every font slot whenever the run contains Chinese;
    # pure Latin runs (including the reference list) remain Times New Roman.
    has_cjk = any("\u3400" <= ch <= "\u9fff" for ch in (run.text or ""))
    run_font = EAST_ASIA_FONT if has_cjk else LATIN_FONT
    run.font.name = run_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), run_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), run_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), run_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), run_font)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = False
    run.italic = italic


def configure_style(style, size: float, before: float, after: float, line: float, align=None) -> None:
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    style.font.size = Pt(size)
    style.font.bold = False
    style.font.color.rgb = BLACK
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        fmt.alignment = align


def add_text(doc: Document, text: str, style: str = "Normal", first_indent=True):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.29)
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 15 if level == 1 else 12.5)
    return p


def add_hyperlink(paragraph, text: str, url: str, size: float = 9.5):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    r_pr.append(color)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_el)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_reference(doc: Document, leading_text: str, link_text: str | None = None, url: str | None = None):
    p = doc.add_paragraph(style="Reference")
    r = p.add_run(leading_text)
    set_run_font(r, 7.5)
    if link_text and url:
        add_hyperlink(p, link_text, url, 7.5)
    return p


def add_equation(doc: Document, body_xml: str) -> None:
    p = doc.add_paragraph(style="Equation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(f'<m:oMathPara {nsdecls("m")}><m:oMath>{body_xml}</m:oMath></m:oMathPara>'))


def mr(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def sub(base: str, subscript: str) -> str:
    return f"<m:sSub><m:e>{mr(base)}</m:e><m:sub>{mr(subscript)}</m:sub></m:sSub>"


def integral(lower: str, upper: str, expression: str) -> str:
    if upper:
        upper_pr = '<m:supHide m:val="0"/>'
        upper_xml = f"<m:sup>{mr(upper)}</m:sup>"
    else:
        upper_pr = '<m:supHide m:val="1"/>'
        upper_xml = ""
    return (
        '<m:nary><m:naryPr><m:chr m:val="∫"/><m:limLoc m:val="subSup"/>'
        f'<m:subHide m:val="0"/>{upper_pr}</m:naryPr>'
        f"<m:sub>{mr(lower)}</m:sub>{upper_xml}<m:e>{expression}</m:e></m:nary>"
    )


def add_picture(doc: Document, path: Path, width=6.45) -> None:
    p = doc.add_paragraph(style="Figure")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width))


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, 9)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, 9, GRAY)


def build() -> None:
    if not all(path.exists() for path in (FIG1, FIG2, FIG3, FIG4)):
        raise FileNotFoundError("Required figures have not been generated")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    configure_style(doc.styles["Normal"], 11, 0, 6, 1.25, WD_ALIGN_PARAGRAPH.JUSTIFY)
    configure_style(doc.styles["Heading 1"], 15, 16, 8, 1.15)
    configure_style(doc.styles["Heading 2"], 12.5, 11, 5, 1.15)
    configure_style(doc.styles["Caption"], 9, 3, 7, 1.1, WD_ALIGN_PARAGRAPH.JUSTIFY)

    for name in ("Reference", "Equation", "Figure"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    configure_style(doc.styles["Reference"], 7.5, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.LEFT)
    doc.styles["Reference"].paragraph_format.left_indent = Inches(0.3)
    doc.styles["Reference"].paragraph_format.first_line_indent = Inches(-0.3)
    configure_style(doc.styles["Equation"], 11, 4, 6, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    configure_style(doc.styles["Figure"], 10, 6, 2, 1.0, WD_ALIGN_PARAGRAPH.CENTER)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("2026 ENSO recharge diagnosis")
    set_run_font(hr, 8.5, GRAY)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    tr = title.add_run("2026年ENSO早期充电过程")
    set_run_font(tr, 22)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("暖水体积、断面输送与温跃层演变的文献方法诊断")
    set_run_font(sr, 13, GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(20)
    mrn = meta.add_run("研究汇报  |  2026年8月24日  |  Yue Dai")
    set_run_font(mrn, 10, GRAY)

    add_heading(doc, "摘要", 1)
    add_text(
        doc,
        "本报告使用现有文献和官方机构定义的诊断量，分析2026年1—6月赤道太平洋暖水充电过程。暖水体积采用NOAA/PMEL标准WWV，温跃层采用20°C等温面深度D20，暖水输送采用已发表研究中的170°E、5°N和10°S固定断面。结果显示，2026年JFM至4月发生快速充电，主要由170°E东向暖水输送和5°N北侧向赤道输送支持，而10°S南侧没有同步贡献。AMJ期间WWV继续维持高值，D20加深信号由西—中太平洋向中—东太平洋扩展。该结构接近历史very-strong事件常见的温跃层演变，但早期WWV不能单独用于判断2026年最终ENSO强度。",
    )

    add_heading(doc, "1 研究目的与理论基础", 1)
    add_text(
        doc,
        "ENSO充放电理论将赤道太平洋上层暖水体积视为海气耦合系统的重要记忆量。El Niño发展前，离赤道质量输送辐合可以增加赤道暖水储量；随后暖水向中—东太平洋重新分配并影响温跃层深度和海表增暖（Jin, 1997; Meinen & McPhaden, 2000）。本报告不构造新的相似度指数，而是依次回答三个问题：2026年暖水体积是否异常偏多、暖水通过哪些已发表断面输入、温跃层加深信号如何随月份和经度演变。",
    )

    add_heading(doc, "2 数据与方法", 1)
    add_heading(doc, "2.1 标准暖水体积", 2)
    add_text(
        doc,
        "WWV采用Meinen and McPhaden（2000）和NOAA/PMEL的标准定义，即计算5°S—5°N、120°E—80°W范围内20°C等温面以上的暖水体积。主时间序列直接采用PMEL官方WWV异常，单位为10¹⁴ m³。",
    )
    add_equation(doc, mr("WWV′ = ∫ₐ D′₂₀ dA"))

    add_heading(doc, "2.2 D20与暖水输送", 2)
    add_text(
        doc,
        "D20为20°C等温面深度，用于表示热带太平洋温跃层结构；正异常表示温跃层加深和暖水层增厚（Li et al., 2022）。暖水输送采用Chakravorty et al.（2021）的固定断面和海表至逐格、逐月D20的积分范围。东向和向赤道输送统一规定为正值，单位为Sv。",
    )
    inner_z = integral("0", "D20", mr("u dz"))
    add_equation(doc, sub("M", "170E") + mr(" = ") + integral("10°S", "5°N", inner_z + mr(" dy")))
    inner_n = integral("0", "D20", mr("v dz"))
    add_equation(doc, sub("M", "5N") + mr(" = −") + integral("140°E", "80°W", inner_n + mr(" dx")))
    inner_s = integral("0", "D20", mr("v dz"))
    add_equation(doc, sub("M", "10S") + mr(" = ") + integral("140°E", "80°W", inner_s + mr(" dx")))

    add_heading(doc, "2.3 资料、气候态与历史分组", 2)
    add_text(
        doc,
        "断面输送和D20使用NCEP Global Ocean Data Assimilation System（GODAS）月平均位温及流速资料，异常相对于1991—2020同历月气候态。GODAS自算WWV与PMEL官方WWV在所选10个年份1—6月共60个月上的相关为0.997，说明D20识别和面积积分具有良好一致性。历史事件依据NOAA/CPC Relative Oceanic Niño Index（RONI）分为moderate、strong和very strong。moderate组包括1994、2002和2023年，strong组包括1986和2009年，very-strong组包括1982、1991、1997和2015年。2026年尚未成熟，因此不预先分类。",
    )

    add_heading(doc, "3 WWV与暖水输送异常", 1)
    add_picture(doc, FIG1)
    add_caption(
        doc,
        "Fig. 1. 2026年1—6月赤道太平洋暖水体积及暖水输送异常与历史ENSO强度组的比较。（a）NOAA/PMEL标准WWV异常；（b）170°E、10°S—5°N断面的东向暖水输送异常；（c）5°N、140°E—80°W断面的向赤道暖水输送异常；（d）10°S、140°E—80°W断面的向赤道暖水输送异常。蓝、橙和红线分别为moderate、strong和very-strong事件组均值，阴影为组内最小—最大范围，黑线为2026年。WWV取自PMEL官方序列；输送异常相对于GODAS 1991—2020同历月气候态。",
    )

    add_heading(doc, "3.1 暖水体积变化", 2)
    add_text(
        doc,
        "2026年WWV异常从1月的+1.02 × 10¹⁴ m³增加至4月的+2.84 × 10¹⁴ m³，5月略降至+2.70 × 10¹⁴ m³，6月重新升至+3.00 × 10¹⁴ m³。JFM和AMJ平均分别为+1.46和+2.85 × 10¹⁴ m³，均高于very-strong事件的同期合成均值。黑线在4—6月仍大致位于very-strong事件范围内，因此更准确的结论是2026年处于历史高WWV区间，而不是前所未有。",
    )

    add_heading(doc, "3.2 170°E东向输送", 2)
    add_text(
        doc,
        "170°E断面在1—4月保持+17.3至+25.9 Sv的东向输送异常，与WWV快速上升同期。5月输送降至约+1.0 Sv，6月恢复至+19.1 Sv。这表明西太平洋暖水向东输入是早期充电的重要组成部分，并存在5月短暂停顿和6月恢复。",
    )

    add_heading(doc, "3.3 南北侧输送的不对称性", 2)
    add_text(
        doc,
        "5°N向赤道输送在1月达到+41.5 Sv，2—3月仍约为+27 Sv，随后持续减弱，5月转为−3.9 Sv。相比之下，10°S断面除1月外在2—5月均为负值。JFM平均的5°N和10°S向赤道输送分别为+32.1和−3.4 Sv。这一对比说明2026年早期充电主要由北侧而非南北两侧同步输入支持，具有清楚的半球不对称性。",
    )

    add_heading(doc, "3.4 阶段划分", 2)
    add_text(
        doc,
        "综合四个面板，1—4月可定义为快速充电阶段，170°E东向输送与5°N向赤道输送同时增强，WWV快速增加。5月是输入暂停阶段，两个主要输送量明显减弱，WWV小幅回落。6月170°E东向输送恢复，WWV重新达到高值，但南北侧输入仍然较弱，因此更接近高WWV背景下的维持和再分配。",
    )

    doc.add_page_break()
    add_heading(doc, "4 D20季节平均经度剖面", 1)
    add_picture(doc, FIG2)
    add_caption(
        doc,
        "Fig. 2. 2026年JFM和AMJ赤道太平洋D20异常经度剖面与历史ENSO强度组的比较。D20为20°C等温面深度，正异常表示温跃层加深。各剖面为5°S—5°N余弦纬度加权平均，范围为120°E—80°W。实线为组均值或2026年，阴影为各历史组事件范围；所有异常均相对于GODAS 1991—2020同历月气候态。",
    )

    add_heading(doc, "4.1 JFM：西部储热更突出", 2)
    add_text(
        doc,
        "JFM期间，2026年D20在西太平洋至日期变更线附近总体为正，西部局地加深接近20 m，随后向中—东太平洋逐渐减弱。这说明早期高WWV首先表现为西—中太平洋暖水层增厚，而不是东太平洋温跃层已经全面加深。与历史组比较，2026西部正异常高于多数合成均值，但中—东太平洋尚未形成very-strong组中较连续的加深结构。",
    )

    add_heading(doc, "4.2 AMJ：加深中心明显东移", 2)
    add_text(
        doc,
        "AMJ期间，2026年D20从日期变更线以东持续增加，中—东太平洋大部达到约15—25 m正异常，最东端局地超过40 m。黑线与very-strong组在中—东太平洋的正异常形态最接近，并在最东端达到其历史范围上部。与此同时，西太平洋局地减弱或转为弱负异常，说明暖水并非在全盆地同步增加，而是从西部储存状态转向中—东太平洋的重新配置。",
    )

    add_heading(doc, "4.3 剖面图给出的阶段转换", 2)
    add_text(
        doc,
        "JFM与AMJ的对照把充电和再分配两个阶段分开：JFM以西部暖水层增厚为主，AMJ则以中—东太平洋温跃层加深为主。该变化与Fig. 1中1—4月WWV快速上升、5月输入减弱但暖水储量仍保持高值的时间关系一致。",
    )

    doc.add_page_break()
    add_heading(doc, "5 垂向位温异常结构", 1)
    add_picture(doc, FIG3)
    add_caption(
        doc,
        "Fig. 3. 历史ENSO强度组和2026年JFM、AMJ赤道太平洋垂向位温异常。位温在5°S—5°N进行余弦纬度加权平均，显示0—300 m、120°E—80°W范围；所有面板使用同一色标，异常相对于GODAS 1991—2020同历月气候态。红色表示暖异常，蓝色表示冷异常。",
    )

    add_heading(doc, "5.1 JFM：暖异常主要储存在次表层", 2)
    add_text(
        doc,
        "2026年JFM的暖异常从西太平洋延伸至中太平洋，主体位于约50—200 m，暖核并未普遍贴近海表；中—东太平洋浅层仍夹有冷异常。相比moderate和strong合成，2026的次表层暖异常更宽、更连续，但其东伸程度仍弱于AMJ。这一结构表明早期增暖首先体现为海洋内部储热，尚不能等同于已经形成成熟的东太平洋海表增暖。",
    )

    add_heading(doc, "5.2 AMJ：暖核沿温跃层向东延伸", 2)
    add_text(
        doc,
        "到AMJ，2026年正位温异常沿倾斜的温跃层从中太平洋向东太平洋连续延伸，东部暖异常逐渐变浅并接近上层海洋。其空间形态与very-strong组最相似，尤其是中—东太平洋约50—150 m的连续暖带；但2026最东端暖异常更集中，仍需避免把形态相似直接解释为最终强度相同。",
    )

    add_heading(doc, "5.3 与D20的相互验证", 2)
    add_text(
        doc,
        "D20只描述20°C等温面的垂向位移，而位温剖面同时展示异常的深度、厚度和倾斜结构。Fig. 2中的中—东太平洋D20加深，在Fig. 3中对应为沿温跃层向东、向上的暖异常带；两种独立表达得到一致结果，因此东移信号并非单一等温面指标造成。",
    )

    doc.add_page_break()
    add_heading(doc, "6 D20月—经度演变", 1)
    add_picture(doc, FIG4)
    add_caption(
        doc,
        "Fig. 4. 历史ENSO强度组和2026年1—6月赤道太平洋D20异常的月—经度演变。D20在5°S—5°N进行余弦纬度加权平均；红色表示温跃层加深，蓝色表示温跃层变浅，黑色等值线为零异常。四个面板使用相同色标，异常相对于GODAS 1991—2020同历月气候态。时间由下至上推进。",
    )

    add_heading(doc, "6.1 历史强度组的结构差异", 2)
    add_text(
        doc,
        "Moderate组D20异常整体较弱，正负异常较为破碎，缺少持续的中—东太平洋加深信号。Strong组在西—中太平洋存在一定正异常，但东太平洋仍以负异常为主。Very-strong组呈现最有组织的经度演变：西太平洋温跃层偏浅，而中—东太平洋正D20异常从春季开始增强并向东扩展，反映暖水从西部储存区向中—东太平洋重新配置。",
    )

    add_heading(doc, "6.2 2026年的逐月演变", 2)
    add_text(
        doc,
        "2026年1—2月大部分赤道太平洋已有弱正D20异常，但中—东太平洋仍夹杂局地负异常。3—4月中太平洋至东太平洋正异常明显增强。5—6月正D20异常覆盖中—东太平洋，最东端加深尤其明显，同时西—中太平洋局部出现弱负异常或正异常减弱。因此，2026年的变化不是全盆地同步加深，而是包含西—中太平洋暖水向中—东太平洋重新配置。",
    )

    add_heading(doc, "6.3 与very-strong事件的比较", 2)
    add_text(
        doc,
        "从形态上看，2026年4—6月比moderate和strong合成更接近very-strong组的中—东太平洋温跃层加深结构。但2026早期西太平洋没有very-strong合成中同样明显的负D20异常，而且最东太平洋在5—6月的加深局部达到或超过very-strong均值。因而只能说2026具备very-strong事件常见的部分海洋结构特征，不能将其解释为对最终事件强度的确定预测。",
    )

    add_heading(doc, "7 四幅图的综合物理解释", 1)
    add_text(
        doc,
        "四幅图共同支持一条连续的物理过程：1—4月170°E东向输入和5°N北侧向赤道输入共同增加赤道太平洋WWV；JFM新增暖量主要储存在西—中太平洋次表层。随后暖水沿温跃层向东重新分配，使中—东太平洋D20逐渐加深，并在AMJ形成连续的次表层暖带。5月断面输入减弱并不意味着此前积累的暖水立即消失，而是标志着过程从快速输入转向储存、再分配和维持；6月东向输送恢复，东太平洋温跃层加深信号继续增强。",
    )

    add_heading(doc, "8 结论与后续建议", 1)
    add_text(
        doc,
        "第一，2026年上半年已经完成一次强度较大的早期暖水充电，JFM和AMJ WWV均处于历史事件的高值范围。第二，早期充电路径具有明显南北不对称性，主要由170°E东向输送和5°N北侧向赤道输送支持，10°S没有同步贡献。第三，D20演变显示暖水从西—中太平洋向中—东太平洋重新配置，AMJ结构与very-strong事件合成较接近。第四，这些诊断描述的是海洋预条件和充电路径，不足以单独预测2026年最终ENSO强度。",
    )
    add_text(
        doc,
        "下一步应继续沿用已发表框架。待后续月份资料完整后，可按照Chakravorty et al.（2021）施加3个月平滑，检查断面输送结论是否稳定；同时应逐事件比较1997、2015、2023与2026的WWV增长时间、南北输送和D20东移过程，避免小样本组平均掩盖个例差异。SODA3可用于复核三维温度和D20结构，而实际体积输送仍应优先使用包含流速资料的GODAS或其他海洋再分析。",
    )

    add_heading(doc, "9 方法限制", 1)
    add_text(
        doc,
        "三条输送曲线不能直接相加为闭合WWV收支。170°E是盆地内部断面，5°N和10°S采用文献中的诊断边界，当前计算没有包括东边界输送、跨20°C等温面的水团转化以及完整表面收支。WWV与GODAS输送使用的异常基准期也不同，因此只能比较其时间演变与物理一致性，不能进行逐项算术闭合。历史分组事件数有限，尤其strong组只有两个事件，组间差异仅作描述性解释。",
    )

    add_heading(doc, "参考文献", 1)
    add_reference(
        doc,
        "Behringer, D. W., Ji, M., & Leetmaa, A. (1998). An improved coupled model for ENSO prediction and implications for ocean initialization. Part I: The ocean data assimilation system. Monthly Weather Review, 126, 1013–1021. ",
        "https://doi.org/10.1175/1520-0493(1998)126<1013:AICMFE>2.0.CO;2",
        "https://doi.org/10.1175/1520-0493(1998)126%3C1013:AICMFE%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "Chakravorty, S., Perez, R. C., Gnanaseelan, C., & Anderson, B. T. (2021). Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 126, e2020JC017075. ",
        "https://doi.org/10.1029/2020JC017075",
        "https://doi.org/10.1029/2020JC017075",
    )
    add_reference(
        doc,
        "Jin, F.-F. (1997). An equatorial ocean recharge paradigm for ENSO. Part I: Conceptual model. Journal of the Atmospheric Sciences, 54, 811–829. ",
        "https://doi.org/10.1175/1520-0469(1997)054<0811:AEORPF>2.0.CO;2",
        "https://doi.org/10.1175/1520-0469(1997)054%3C0811:AEORPF%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "Li, X., Hu, Z.-Z., Zhao, S., Ding, R., & Zhang, B. (2022). On the asymmetry of the tropical Pacific thermocline fluctuation associated with ENSO recharge and discharge. Geophysical Research Letters, 49, e2022GL099242. ",
        "https://doi.org/10.1029/2022GL099242",
        "https://doi.org/10.1029/2022GL099242",
    )
    add_reference(
        doc,
        "Meinen, C. S., & McPhaden, M. J. (2000). Observations of warm water volume changes in the equatorial Pacific and their relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559. ",
        "https://doi.org/10.1175/1520-0442(2000)013<3551:OOWWVC>2.0.CO;2",
        "https://doi.org/10.1175/1520-0442(2000)013%3C3551:OOWWVC%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "NOAA Climate Prediction Center. (2026). Relative Oceanic Niño Index and ENSO strength classification. Accessed August 24, 2026. ",
        "https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/enso/roni/",
        "https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/enso/roni/",
    )
    add_reference(
        doc,
        "NOAA Pacific Marine Environmental Laboratory. (2026). Equatorial Pacific warm water volume data. Accessed August 24, 2026. ",
        "https://www.pmel.noaa.gov/tao/wwv/",
        "https://www.pmel.noaa.gov/tao/wwv/",
    )
    add_reference(
        doc,
        "NOAA Physical Sciences Laboratory. (2026). NCEP Global Ocean Data Assimilation System (GODAS). Accessed August 24, 2026. ",
        "https://psl.noaa.gov/data/gridded/data.godas.html",
        "https://psl.noaa.gov/data/gridded/data.godas.html",
    )

    core = doc.core_properties
    core.title = "2026 ENSO Early Recharge Process"
    core.subject = "Published-method WWV, transport, and D20 diagnostics"
    core.author = "Yue Dai"
    core.keywords = "ENSO, recharge, WWV, D20, GODAS, PMEL"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
