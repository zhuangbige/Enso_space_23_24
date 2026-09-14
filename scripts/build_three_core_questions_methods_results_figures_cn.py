#!/usr/bin/env python3
"""Build a focused Chinese Methods-and-Results DOCX around three fixed questions."""

from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_warm_heat_fate_complete_report_cn as style  # noqa: E402


OUT_DIR = ROOT / "outputs/research_2026/reports"
OUTPUT = OUT_DIR / "三个核心问题_方法与结果_公式严谨性修订中间版_20260807.docx"
FIGURE_PACKAGE = ROOT / "outputs/research_2026/three_questions_figure_package_20260805"

THREE_BOX = (
    ROOT
    / "data/derived/research_2026/soda3_three_box_warm_heat_fate_cumulative_summary.csv"
)
COUPLING = (
    ROOT / "data/derived/research_2026/soda3_air_sea_coupled_warm_heat_summary.csv"
)
CASE_PHASE = (
    ROOT
    / "data/derived/research_2026/enso_1997_2015_2023_warm_heat_fate_phase_summary.csv"
)
CASE_GEO = (
    ROOT
    / "data/derived/research_2026/enso_1997_2015_2023_5n_geostrophic_fate_events.csv"
)
CLOSURE = (
    ROOT
    / "data/derived/research_2026/jiang_fep_soda3_atmospheric_closure_1997_2015_2023.csv"
)
COMPARE = (
    ROOT
    / "outputs/research_2026/compare_1997_2023_coupling/1997_2023_monthly_comparison.csv"
)
ATTR_EVENT = (
    ROOT
    / "data/derived/research_2026/third_question_coupled_attribution_event_table.csv"
)
ATTR_REL = (
    ROOT
    / "data/derived/research_2026/third_question_coupled_attribution_relationships.csv"
)
ATTR_CF = (
    ROOT
    / "data/derived/research_2026/third_question_coupled_attribution_counterfactual.csv"
)

style.FIGURES.update(
    {
        "arrival_hov": FIGURE_PACKAGE
        / "01_main_figures/Fig01_eastward_OHC_arrival_hovmoller.png",
        "three_box": FIGURE_PACKAGE
        / "01_main_figures/Fig02_CT_CP_three_box_arrival_and_fate.png",
        "case_bridge": FIGURE_PACKAGE
        / "01_main_figures/Fig03_event_arrival_export_surface_loss.png",
        "case_geofate": FIGURE_PACKAGE
        / "01_main_figures/Fig04_event_5N_dynamical_decomposition.png",
        "fep_ocean": FIGURE_PACKAGE
        / "01_main_figures/Fig05_FEP_subsurface_preconditioning.png",
        "fep_atmos": FIGURE_PACKAGE
        / "01_main_figures/Fig06_FEP_atmospheric_closure.png",
        "case_coupling": FIGURE_PACKAGE
        / "01_main_figures/Fig07_1997_2023_coupled_mechanism_synthesis.png",
        "third_attribution": FIGURE_PACKAGE
        / "01_main_figures/Fig08_third_question_coupled_attribution.png",
    }
)


def one(frame: pd.DataFrame, **filters) -> pd.Series:
    return style.summary_value(frame, **filters)


def fmt(value: float, digits: int = 2) -> str:
    return f"{float(value):+.{digits}f}"


def set_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(82)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(14)
    style.set_run_font(
        kicker.add_run("论文主体草稿  |  方法与结果"),
        font=style.HEAD_FONT,
        size=11,
        bold=True,
        color=style.GOLD,
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    style.set_run_font(
        title.add_run("厄尔尼诺暖量的到达、去向与海气耦合命运"),
        font=style.HEAD_FONT,
        size=27,
        bold=True,
        color=style.NAVY,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_after = Pt(8)
    style.set_run_font(
        subtitle.add_run("围绕三个固定科学问题的完整方法与结果部分"),
        font=style.HEAD_FONT,
        size=15,
        bold=True,
        color=style.DARK_BLUE,
    )

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.first_line_indent = Pt(0)
    desc.paragraph_format.space_after = Pt(62)
    style.set_run_font(
        desc.add_run("SODA3.15.2 · OISST · ERA5 · GPCP/OLR"),
        size=11,
        italic=True,
        color=style.MUTED,
    )

    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.first_line_indent = Pt(0)
    status.paragraph_format.space_after = Pt(5)
    style.set_run_font(
        status.add_run("分析状态：前两问主体闭合；第三问形成过程机制答案"),
        font=style.HEAD_FONT,
        size=11,
        bold=True,
        color=style.GREEN,
    )
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Pt(0)
    style.set_run_font(date.add_run("2026年8月7日 · 公式严谨性修订版"), size=10.5, color=style.MUTED)
    doc.add_page_break()


def audit_inputs(
    three: pd.DataFrame,
    phase: pd.DataFrame,
    geo: pd.DataFrame,
    closure: pd.DataFrame,
) -> None:
    required_events = {1997, 2015, 2023}
    if set(phase.event_year.unique()) != required_events:
        raise ValueError("Unexpected case-event set")
    if set(geo.event_year.unique()) != required_events:
        raise ValueError("Unexpected geostrophic-event set")
    if set(closure.event_year.unique()) != required_events:
        raise ValueError("Unexpected closure-event set")
    if not {
        "development_Jun0_to_Nov0",
        "mature_decay_Dec0_to_May1",
    }.issubset(set(three.phase.unique())):
        raise ValueError("Missing CT/CP phases")
    main_geo = geo.loc[geo.baseline.eq("1980-2014")].set_index("event_year")
    if not (
        main_geo.loc[1997, "total_zj"]
        > main_geo.loc[2015, "total_zj"]
        > main_geo.loc[2023, "total_zj"]
    ):
        raise ValueError("Unexpected event outlet ordering")


def build_report() -> Path:
    three = pd.read_csv(THREE_BOX)
    coupling = pd.read_csv(COUPLING)
    phase = pd.read_csv(CASE_PHASE)
    geo = pd.read_csv(CASE_GEO)
    closure = pd.read_csv(CLOSURE)
    compare = pd.read_csv(COMPARE)
    attr_event = pd.read_csv(ATTR_EVENT).set_index("year")
    attr_rel = pd.read_csv(ATTR_REL).set_index("link_id")
    attr_cf = pd.read_csv(ATTR_CF).set_index(["target_ocean_event", "donor_wind"])
    audit_inputs(three, phase, geo, closure)

    doc = Document()
    bullet_id, number_id = style.configure_document(doc)
    props = doc.core_properties
    props.title = "厄尔尼诺暖量的到达、去向与海气耦合命运：方法与结果"
    props.subject = "围绕三个固定科学问题的SODA3暖量预算与海气耦合诊断"
    props.keywords = "ENSO; FEP; H20; Qwarm,20; SODA3; Ekman; geostrophic; warm-heat fate"
    props.comments = "Formula-rigour revision generated from verified project outputs on 2026-08-07."
    set_cover(doc)

    style.add_section_heading(doc, "1 研究问题与总体设计")
    style.add_para(
        doc,
        "本方法与结果部分只围绕三个固定问题组织，不再以“发生—分型—放大”等更早框架替代：第一，暖量是否、以及通过什么路径到达东太平洋；第二，到达后的暖量在北向输出、南向输出、海表散热和海洋储存之间如何分配；第三，为什么相似的FEP先暖和东部暖量输入，会通过不同的海气耦合产生完全不同的暖量命运。",
    )
    for text in (
        "暖量是否、以及通过什么路径到达东太平洋？",
        "到达后的暖量在北向输出、南向输出、海表散热和海洋储存之间如何分配？",
        "为什么相似的FEP先暖和东部暖量输入，会通过不同的海气耦合产生完全不同的暖量命运？",
    ):
        style.add_list_item(doc, text, number_id)
    style.add_callout(
        doc,
        "总体策略",
        "采用两层互补设计：4个CT与4个CP事件合成用于检验群体差异；1997、2015和2023三个过程个例用于分辨“强协同—风驱克服—出口近闭合”的动力层级。合成结论和个例数值分别解释，不直接混合绝对量。",
        fill=style.LIGHT_BLUE,
    )

    style.add_table(
        doc,
        "表1  两层研究设计与三个问题的对应关系",
        ["设计层", "事件与阶段", "主要资料", "回答的问题"],
        [
            [
                "CT/CP合成",
                "CT: 1982/1986/1997/2006；CP: 1991/1994/2002/2004；Jun(0)—Nov(0)与Dec(0)—May(+1)",
                "SODA3、OISST、GPCP、ERA5",
                "暖量到达与群体平均去向",
            ],
            [
                "三事件过程对照",
                "1997/98、2015/16、2023/24；发展年Jan—Jun与次年Jan—May",
                "扩展SODA3、OISST、ERA5/OLR",
                "相似起点为何产生不同耦合和出口",
            ],
        ],
        [1450, 3260, 2100, 2550],
        font_size=8.4,
    )

    style.add_section_heading(doc, "2 数据与方法", page_break=True)
    doc.add_heading("2.1 数据资料", level=2)
    style.add_table(
        doc,
        "表2  数据、变量与用途",
        ["资料", "主要变量", "时段/分辨率", "用途"],
        [
            [
                "SODA3.15.2",
                "温度、盐度、SSH、u/v、纬向风应力、net_heating",
                "月平均；主历史期1980—2022，2024官方年度文件补充Jan—May",
                "H20/Q20、三箱预算、D20/OHC与5°N动力分解",
            ],
            [
                "OISST",
                "SST",
                "月平均；1982—2024",
                "FEP、Niño-3.4及东部表面增暖",
            ],
            [
                "ERA5",
                "u850、SLP、风应力、短波/长波/潜热/感热、OLR相关量",
                "月平均；与事件共同覆盖",
                "大气接通、Ekman独立估算和海表热损失",
            ],
            [
                "GPCP/CPC OLR",
                "降水、顶层长波辐射",
                "月平均",
                "对流中心位置和深对流独立约束",
            ],
        ],
        [1450, 2730, 2700, 2480],
        font_size=8.1,
    )

    doc.add_heading("2.2 空间区域、阶段和符号", level=2)
    style.add_para(
        doc,
        "CT/CP三箱预算的纬度范围为10°S—5°N、深度为0—300 m，按经度划分为西箱160°E—180°、中箱180°—140°W和东箱140°W—80°W。三事件扩展资料的东部储热区采用5°S—5°N、140°W—80°W，以保持1980—2024共同网格；两套设计共享5°N北界。合成发展期定义为Jun(0)—Nov(0)，成熟—衰减期为Dec(0)—May(+1)；三事件过程对照受扩展资料月份约束，使用Jan(0)—Jun(0)和Jan(+1)—May(+1)。",
    )
    style.add_para(
        doc,
        "箱体预算统一采用入箱符号：对任一边界Γ，FΓ>0表示暖量进入箱体，FΓ<0表示离开箱体；正文展示“北向输出”或“海表热损失”时，仅把相应入箱项乘以−1以便直观阅读。5°N动力分解另规定北向为正，并在换算到箱体预算时显式变号。不同设计的阶段和纬度范围不同，因此绝对量只在同一设计内部比较。",
    )

    doc.add_heading("2.3 异常、趋势和统计推断", level=2)
    style.add_para(
        doc,
        "所有异常均按日历月分别估计气候态和线性趋势。CT/CP合成使用1982—2022作为统一基准；1997/2015/2023过程对照以1980—2014为主基准，并用1982—2022检验基准敏感性。FEP与Niño-3.4使用OISST 1982—2024逐月去趋势异常。CT/CP标量时间序列使用中心3个月平均；三事件动力分解使用未平滑月值，避免跨年份端点错误连接。",
    )
    style.add_para(
        doc,
        "CT/CP推断以整事件为抽样单位进行20,000次bootstrap，报告事件均值、95%区间和效应量，不把相邻月份视为独立样本。FEP逐日历月回归使用1982—2024年的43个年度样本，每月OLS自由度为41；机制链的事件级关系使用13个暖事件，并同时报告Pearson r、事件置换p、事件bootstrap区间和逐一删事件范围，不以单一p值作为结论。1997、2015和2023是过程个例，不进行三例显著性检验，其用途是检验时序和动力一致性。",
    )

    doc.add_heading("2.4 问题一：暖量到达的诊断", level=2)
    style.add_equation(
        doc,
        "χ20(T) = max(T − Tr, 0),    I20(T) = 1(T > Tr),    Tr = 20°C",
        "χ20是20°C阈值以上的超温，I20为相应指示函数。20°C由D20/WWV的动力学意义选取，不被解释为绝对热力学参考态（Meinen和McPhaden，2000；Chakravorty等，2021）。",
    )
    style.add_equation(
        doc,
        "H20(t) = ρ₀ cₚ ∭Ω χ20[T(t)] dV",
        "H20是本研究定义的20°C阈值以上相对热含量（relative heat content above 20°C），不是绝对热含量，也不等同于常规0—300 m OHC。ρ₀=1025 kg m⁻³，cₚ=3990 J kg⁻¹ K⁻¹；体积分转换为ZJ。",
    )
    style.add_equation(
        doc,
        "Qwarm,20(t) = ρ₀ cₚ ∬Γ vₙ(t) χ20[T(t)] dA",
        "Qwarm,20是本研究定义的20°C阈值暖水体热能输送，单位PW；它只计算T>20°C部分，不是传统总海洋热输送ρ₀cₚ∬vₙT dA。分别计算180°、140°W、80°W以及5°N/10°S断面。其结构受经典WWV输送诊断启发，但阈值能量加权是本研究的诊断定义。",
    )
    style.add_para(
        doc,
        "问题一通过三类相互约束的证据回答：跨140°W向东Qwarm,20衡量暖量是否进入东部；东部H20变化衡量到达后是否储存；OHC、D20和SSH的东西差异衡量表面FEP增暖是否得到盆地尺度次表层结构支撑。",
    )

    doc.add_heading("2.5 问题二：暖量去向的固定控制体预算", level=2)
    style.add_equation(
        doc,
        "FΓ = −ρ₀ cₚ ∬Γ (v · nout) χ20(T) dA,    FΓ > 0 表示入箱",
        "nout为箱体外法向。该定义用于所有东西和南北边界，从根本上避免“纬向向东为正、经向向北为正”与箱体收支符号混用。东西向辐合定义为Cx≡ΣΓxFΓ。",
    )
    style.add_equation(
        doc,
        "dH20/dt = Cₓ + F10S + F5N + Qsurface,warm + R20",
        "所有F项均以进入箱体为正。R20由储量倾向减去全部显式项得到，因此该式是离散诊断恒等式，而不是假定χ20为线性保守示踪剂后得到的无剩余守恒律。",
    )
    style.add_equation(
        doc,
        "Qsurface,warm = ∬A I20(Ts) Qnet dA,    Qnet > 0 表示海洋吸热",
        "Ts采用SODA3最上层温度。I20(Ts)来自阈值函数对表层温度的一阶响应；用I(D20>0)会把作用在低于20°C表层的热通量直接计入H20，因而只适合作为足迹敏感性而非主定义。由于使用月平均、有限层厚资料，本项仍是表层阈值近似。",
    )
    style.add_para(
        doc,
        "阈值算子需要单独说明。对T≠Tr且不可压缩的绝热平流，链式法则给出Dχ20/Dt=I20DT/Dt，平流本身不会凭空制造暖量；但扩散、混合、夹卷和外部加热使水体穿越20°C时，将出现集中在阈值面的非线性转化项（形式上与δ[T−Tr]有关）。当前月平均资料不能逐项解析这些转化，因此R20不被命名为某一种物理过程，而共同包含跨20°C阈值转化、垂向交换、混合扩散、涡动协方差、SODA同化增量及月平均/网格离散误差。ERA5净表面热通量及其分量用于独立检验SODA3海表热损失。",
    )

    doc.add_heading("2.6 问题三：FEP—海洋结构—大气接通—出口分解", level=2)
    style.add_equation(
        doc,
        "FEPres(m,y) = FEP′(m,y) − [aₘ + bₘ Niño3.4′(m,y)]",
        "FEP定义为5°S—5°N、90°W—80°W。按日历月用43个年度样本回归移除同期Niño-3.4相关部分；每月OLS自由度为41。FEPres是本研究用于控制盆地平均暖事件幅度的统计残差，不被解释为独立物理强迫。",
    )
    style.add_para(
        doc,
        "操作性海洋转变月定义为东部减西部的0—300 m OHC异常和D20异常连续两个月同时为正的首月；操作性大气闭合月定义为中西太平洋u850西风异常为正且OLR异常为负，并连续维持两个月的首月。该符号判据只用于三个事件的过程对齐，不作为普适ENSO阈值。",
    )
    style.add_equation(
        doc,
        "Qwarm,20,total = Qwarm,20,geo + Qwarm,20,Ekman + Qwarm,20,res",
        "5°N、140°W—80°W、0—300 m断面使用同一SODA3产品分解总北向暖水体热能输送。Qwarm,20,res严格按总量减地转和Ekman定义，因此这是诊断恒等式，不是完整动力模态分解；残差包含非地转流、瞬变、涡动、惯性过程和近似误差。",
    )
    style.add_equation(
        doc,
        "vg(d) = (g/f) ∂η/∂x + [g/(ρ₀f)] ∫₀ᵈ ∂ρ′(x,d′)/∂x dd′",
        "ρ′=ρ−ρref(d)，其中ρref只随深度变化，因此∂ρ′/∂x=∂ρ/∂x；改写为ρ′是为了明确只有水平密度差贡献斜压压力梯度，不改变现有数值计算。地转关系与热成风积分遵循Bosc和Delcroix（2008）及Li等（2022）的断面输送框架；EOS-80为主，线性状态方程作敏感性。",
    )
    style.add_equation(
        doc,
        "My = −τx/(ρ₀f)",
        "约定τx向东为正、My向北为正；因此在北半球5°N，东向风应力对应向南Ekman输送，西向风应力异常对应向北输送。该符号来自M=τ×k/(ρ₀f)（Bosc和Delcroix，2008；Li等，2022）。",
    )
    style.add_equation(
        doc,
        "Qwarm,20,Ekman = ρ₀ cₚ ∫x My(x) · χ̄20,0–50m(x) dx",
        "χ̄20,0–50m为0—50 m平均超温。本式是把第一阶Ekman体积输送转换为阈值暖水体热能输送的离线诊断近似。",
    )

    doc.add_heading("2.7 数值质量控制", level=2)
    for text in (
        "三箱异常预算代数闭合误差不超过1.11×10⁻¹⁶ PW；160°E、180°和140°W门与既有独立Q20序列的最小相关为1.00000000。",
        "1980—2022历史5°N断面重建总Q20与既有三箱北界逐月值的最大绝对误差为2.75×10⁻⁹ PW，相关为1.000000000000000。",
        "EOS-80地转项的SSH与密度子项逐月闭合误差为6.11×10⁻¹⁶ PW；总量—地转—Ekman—残差的事件闭合误差小于4.45×10⁻¹⁶ ZJ。",
        "2024扩展文件与历史断面在拼接前强制使用完全一致的深度和经度标签，避免微小浮点坐标差异产生重复层或重复经度。",
    ):
        style.add_list_item(doc, text, bullet_id)

    style.add_section_heading(doc, "3 结果一：暖量是否以及如何到达东太平洋")
    doc.add_heading("3.1 CT与CP都能把暖量送到东部，但CT更强", level=2)
    q_ct = one(
        three,
        contrast="CT",
        phase="development_Jun0_to_Nov0",
        box="gate",
        component="q20_220e_10s5n",
    )
    q_cp = one(
        three,
        contrast="CP",
        phase="development_Jun0_to_Nov0",
        box="gate",
        component="q20_220e_10s5n",
    )
    h_ct = one(
        three,
        contrast="CT",
        phase="development_Jun0_to_Nov0",
        box="east",
        component="storage_tendency",
    )
    h_cp = one(
        three,
        contrast="CP",
        phase="development_Jun0_to_Nov0",
        box="east",
        component="storage_tendency",
    )
    style.add_para(
        doc,
        f"Jun(0)—Nov(0)期间，CT跨140°W东向Q20累计为{fmt(q_ct.mean_zj)} ZJ，95%区间[{fmt(q_ct.ci_low_zj)},{fmt(q_ct.ci_high_zj)}]；CP为{fmt(q_cp.mean_zj)} ZJ，区间[{fmt(q_cp.ci_low_zj)},{fmt(q_cp.ci_high_zj)}]。两组区间都为正，说明CP也会把暖量送入东太平洋，不能把CT/CP差异简化为“到达/未到达”。",
    )
    style.add_para(
        doc,
        f"同期东部H20储量倾向积分在CT和CP中分别为{fmt(h_ct.mean_zj)}和{fmt(h_cp.mean_zj)} ZJ。CT的输入和储存幅度更大，表明类型差异首先表现为暖量到达的强度、持续性和垂向结构，而不是单纯的有无。80°W异常通过量接近零，说明事件暖量主要在140°W—80°W东部箱体内储存和重新分配。",
    )
    style.add_figure(
        doc,
        "arrival_hov",
        "图1  1997/98、2015/16和2023/24赤道太平洋0—300 m OHC经度—时间演变，展示暖量向东到达及盆地深层支撑差异。",
        width=6.15,
    )
    style.add_figure(
        doc,
        "three_box",
        "图2  SODA3三箱暖量输送与收支。上排显示CT/CP跨180°E、140°W和80°W的Q20演变；下排显示各箱累计收支。",
        width=6.15,
    )

    doc.add_heading("3.2 1997输入最强，2015与2023的发展期输入近似", level=2)
    phase_main = phase.loc[
        phase.baseline.eq("1980-2014")
        & phase.phase.eq("development_Jan0_Jun0")
    ].set_index("event_year")
    arrival_rows = []
    for year in (1997, 2015, 2023):
        row = phase_main.loc[year]
        arrival_rows.append(
            [
                str(year),
                fmt(row.q20_140w_east_zj),
                fmt(row.east_h20_change_zj),
                fmt(row.east_sst_mean_c),
                fmt(row.east_d20_mean_m),
            ]
        )
    style.add_table(
        doc,
        "表3  三事件发展年Jan—Jun东部暖量到达指标",
        ["事件", "140°W东向Q20 (ZJ)", "东部H20变化 (ZJ)", "东部SST (°C)", "东部D20 (m)"],
        arrival_rows,
        [1000, 2240, 2050, 1900, 2170],
        font_size=8.4,
    )
    style.add_para(
        doc,
        "1997跨140°W东向输入和东部储热均约为2015/2023的三倍。2015与2023则十分接近：2015的东向Q20为+3.30 ZJ、东部H20增加+3.53 ZJ；2023分别为+3.91和+3.88 ZJ。因此2023后续出口偏弱不能解释为“暖量没有到达东部”。",
    )
    style.add_figure(
        doc,
        "case_bridge",
        "图3  1997、2015和2023的发展期输入、东部储热、次年5°N输出及海表热损失对比。",
        width=6.2,
    )
    style.add_callout(
        doc,
        "问题一的答案",
        "暖量确实通过赤道上层海洋和140°W通道进入东太平洋；CT与CP均可到达，1997最强，2015与2023发展期输入近似。事件差异不是“暖量来没来”，而是到达强度、次表层支撑和后续分配不同。",
        fill=style.LIGHT_BLUE,
    )

    style.add_section_heading(doc, "4 结果二：到达后的暖量去了哪里")
    doc.add_heading("4.1 CT/CP合成显示关键分叉位于5°N北界", level=2)
    budget_metrics = {
        "storage": "storage_tendency",
        "north_in": "north_q20_in",
        "south_in": "south_q20_in",
        "surface": "surface_warm",
        "residual": "residual_warm",
    }
    budget_rows = []
    for group in ("CT", "CP"):
        vals = {
            name: one(
                three,
                contrast=group,
                phase="mature_decay_Dec0_to_May1",
                box="east",
                component=metric,
            ).mean_zj
            for name, metric in budget_metrics.items()
        }
        budget_rows.append(
            [
                group,
                fmt(vals["storage"]),
                fmt(-vals["north_in"]),
                fmt(-vals["south_in"]),
                fmt(-vals["surface"]),
                fmt(vals["residual"]),
            ]
        )
    style.add_table(
        doc,
        "表4  成熟—衰减期东部箱体暖量分配（ZJ）",
        ["类型", "H20变化", "5°N北向输出", "10°S南向输出", "海表热损失", "预算剩余"],
        budget_rows,
        [950, 1600, 1780, 1780, 1650, 1600],
        font_size=8.4,
    )
    north_diff = one(
        coupling,
        contrast="CT_minus_CP",
        phase="mature_decay_Dec0_to_May1",
        metric="soda_north_export_zj",
    )
    era_loss = one(
        coupling,
        contrast="CT_minus_CP",
        phase="mature_decay_Dec0_to_May1",
        metric="era5_qnet_loss_zj",
    )
    style.add_para(
        doc,
        f"Dec(0)—May(+1)期间，东部H20在CT和CP中分别减少7.86和3.03 ZJ。CT跨5°N北向输出+2.74 ZJ，而CP为−0.96 ZJ，即CP合成表现为向东部箱体的净输入。CT−CP北向输出差为{fmt(north_diff['mean'])} ZJ，95%区间[{fmt(north_diff['ci_low'])},{fmt(north_diff['ci_high'])}]。10°S项的组内区间跨零，没有证据表明主要暖量从南侧排出。",
    )
    style.add_para(
        doc,
        f"SODA3海表热损失在CT和CP中分别为3.69和1.76 ZJ；ERA5独立估算为4.35和1.30 ZJ，CT−CP的ERA5净损失差为{fmt(era_loss['mean'])} ZJ且95%区间为正。CT的海表散热主要由短波减少和潜热散失构成，但只有短波组间差异稳定越过95%门槛。",
    )

    doc.add_heading("4.2 三事件形成强排放—部分排放—出口近闭合阶梯", level=2)
    geo_main = geo.loc[geo.baseline.eq("1980-2014")].set_index("event_year")
    decay_main = phase.loc[
        phase.baseline.eq("1980-2014") & phase.phase.eq("decay_Jan1_May1")
    ].set_index("event_year")
    fate_rows = []
    for year in (1997, 2015, 2023):
        fate_rows.append(
            [
                str(year),
                fmt(geo_main.loc[year, "total_zj"]),
                fmt(decay_main.loc[year, "east_surface_loss_zj"]),
                fmt(decay_main.loc[year, "east_h20_change_zj"]),
                fmt(geo_main.loc[year, "one_sided_4p75_total_zj"]),
            ]
        )
    style.add_table(
        doc,
        "表5  三事件次年Jan—May的东部暖量命运",
        ["事件", "居中面5°N输出 (ZJ)", "海表热损失 (ZJ)", "东部H20变化 (ZJ)", "4.75°N单边输出 (ZJ)"],
        fate_rows,
        [950, 2250, 1900, 1900, 2360],
        font_size=8.2,
    )
    style.add_para(
        doc,
        "1997同时具有强北向输出（+3.81 ZJ）和强海表热损失（+4.37 ZJ）；2015北向输出降至+0.88 ZJ、热损失为+1.16 ZJ；2023居中通量面估计仅+0.10 ZJ、热损失为+0.40 ZJ。三个事件的东部H20都明显下降，但下降量并不按出口强度单调排列，说明不能把“储量减少”直接等同于某一条去向。",
    )

    doc.add_heading("4.3 2023出口方向敏感，但“出口受抑”稳健", level=2)
    style.add_para(
        doc,
        "2023/24的总5°N输送对通量面温度插值敏感：跨5°N两侧4.75°N和5.25°N温度居中后为+0.10 ZJ；此前采用赤道侧4.75°N单边温度为−1.24 ZJ。1997和2015两种口径只差约0.04 ZJ，而2023相差1.33 ZJ，反映2024年5°N强温度锋与流速的共置。有限体积意义上居中面温度作为主估计更合理；论文结论应表述为“2023出口显著受抑或近闭合”，不强断言净向南。",
    )
    style.add_callout(
        doc,
        "问题二的答案",
        "到达东部的暖量主要通过5°N北向动力出口、海表热损失和上层海洋储量变化重新分配；10°S不是稳健主出口。1997为强排放，2015为部分排放，2023出口近闭合。剩余项仍阻止我们声称已经追踪每一焦耳的最终归宿。",
        fill=style.LIGHT_BLUE,
    )
    style.add_figure(
        doc,
        "case_geofate",
        "图4  1997、2015和2023次年1—5月跨5°N总Q20的Ekman、地转和残差分解。",
        width=6.15,
    )

    style.add_section_heading(doc, "5 结果三：相似FEP先暖为何产生不同暖量命运")
    doc.add_heading("5.1 FEP先暖只描述表面起点，不代表相同海洋状态", level=2)
    march = compare.loc[compare.month.eq(3)].iloc[0]
    structure_rows = [
        ["FEP SSTA (°C)", fmt(march.fep_anom_c_1997), fmt(march.fep_anom_c_2023)],
        [
            "中太平洋OHC (GJ m⁻²)",
            fmt(march.ohc300_central_anom_jm2_1997 / 1.0e9),
            fmt(march.ohc300_central_anom_jm2_2023 / 1.0e9),
        ],
        [
            "赤道太平洋总正OHC (PJ m⁻¹)",
            fmt(march.ohc300_positive_total_pj_m_1997),
            fmt(march.ohc300_positive_total_pj_m_2023),
        ],
        ["东—西D20差 (m)", fmt(march.d20_ew_anom_m_1997), fmt(march.d20_ew_anom_m_2023)],
        ["中西太平洋u850 (m s⁻¹)", fmt(march.u850_cwp_anom_ms_1997), fmt(march.u850_cwp_anom_ms_2023)],
        ["中太平洋OLR (W m⁻²)", fmt(march.olr_cp_anom_wm2_1997), fmt(march.olr_cp_anom_wm2_2023)],
    ]
    style.add_table(
        doc,
        "表6  1997与2023年3月的表面、次表层和大气状态",
        ["变量", "1997", "2023"],
        structure_rows,
        [4700, 2330, 2330],
        font_size=8.8,
    )
    style.add_para(
        doc,
        "2023年3月FEP为+1.00°C，高于1997年的+0.57°C；但1997的赤道太平洋总正OHC约为2023的3.2倍，中太平洋OHC约为4倍。1997的东—西D20差已经为+5.03 m，2023仍为−10.94 m。这说明2023虽然暖异常位置更偏东，但总暖量库存和温跃层压力结构尚未达到1997式状态。",
    )
    style.add_figure(
        doc,
        "fep_ocean",
        "图5  1997、2015和2023的FEP指标、0—300 m OHC和D20演变，显示相似表面起点对应不同次表层海洋状态。",
        width=6.0,
    )

    doc.add_heading("5.2 决定性分叉位于表面增暖之后、深对流东移之前", level=2)
    timing_rows = []
    timing_labels = {
        1997: ("3月", "4月", "3月", "4月"),
        2015: ("4月", "1月或更早", "3月", "3月"),
        2023: ("2月", "5月", "5月", "5月（弱）"),
    }
    for year in (1997, 2015, 2023):
        timing_rows.append([str(year), *timing_labels[year]])
    style.add_table(
        doc,
        "表7  FEP、海洋和大气接通时间线",
        ["事件", "residual FEP持续转正", "地下海洋转变", "大气闭合", "海洋—大气均接通"],
        timing_rows,
        [900, 2180, 2060, 1860, 2360],
        font_size=8.5,
    )
    style.add_para(
        doc,
        "1997在3月已经出现u850=+4.15 m s⁻¹和OLR=−8.52 W m⁻²，表面暖异常迅速连接深对流和西风；2015的地下海洋条件在1月已经预置，大气在3月闭合；2023的residual FEP从2月持续为正，但3—4月OLR仍为正，地下海洋与大气到5月才同时换符号，而且接通幅度和持续性远弱于前两次事件。",
    )
    mj_1997 = closure.loc[
        closure.event_year.eq(1997) & closure.month.isin([5, 6])
    ]
    mj_2015 = closure.loc[
        closure.event_year.eq(2015) & closure.month.isin([5, 6])
    ]
    mj_2023 = closure.loc[
        closure.event_year.eq(2023) & closure.month.isin([5, 6])
    ]
    style.add_para(
        doc,
        f"5—6月平均u850在1997、2015和2023中分别为{fmt(mj_1997.u850_cwp_160e_140w_ms.mean())}、{fmt(mj_2015.u850_cwp_160e_140w_ms.mean())}和{fmt(mj_2023.u850_cwp_160e_140w_ms.mean())} m s⁻¹；相应的−OLR为{fmt(-mj_1997.olr_cp_170e_140w_wm2.mean())}、{fmt(-mj_2015.olr_cp_170e_140w_wm2.mean())}和{fmt(-mj_2023.olr_cp_170e_140w_wm2.mean())} W m⁻²。FEP是否暖并不决定对流效率；暖异常必须延伸到能够组织中东太平洋深对流和气压梯度的背景中。",
    )
    style.add_figure(
        doc,
        "fep_atmos",
        "图6  FEP、OHC、D20、u850和OLR/对流响应的三事件过程比较，显示海洋转变与大气闭合的时序差异。",
        width=6.15,
    )

    doc.add_heading("5.3 大气接通差异通过Ekman—地转竞争落实为暖量命运", level=2)
    decomposition_rows = []
    for year in (1997, 2015, 2023):
        row = geo_main.loc[year]
        decomposition_rows.append(
            [
                str(year),
                fmt(row.total_zj),
                fmt(row.ekman_zj),
                fmt(row.geo_eos80_zj),
                fmt(row.geo_ssh_eos80_zj),
                fmt(row.geo_density_eos80_zj),
                fmt(row.residual_eos80_zj),
            ]
        )
    style.add_table(
        doc,
        "表8  次年Jan—May跨5°N Q20动力分解（ZJ）",
        ["事件", "总量", "Ekman", "总地转", "SSH地转", "密度地转", "残差"],
        decomposition_rows,
        [800, 1250, 1350, 1400, 1550, 1560, 1450],
        font_size=8.2,
    )
    style.add_para(
        doc,
        "1997的Ekman和地转项分别为+3.25和+0.59 ZJ，二者同向；2015的地转阻力为−1.28 ZJ，但+2.44 ZJ的Ekman输出足以克服它；2023的地转阻力为−1.27 ZJ，与2015主基准几乎相同，而Ekman只有+1.56 ZJ，最终总出口接近零。",
    )
    diff_total = geo_main.loc[2015, "total_zj"] - geo_main.loc[2023, "total_zj"]
    diff_ek = geo_main.loc[2015, "ekman_zj"] - geo_main.loc[2023, "ekman_zj"]
    style.add_para(
        doc,
        f"2015相对2023多出的净北向输出为{fmt(diff_total)} ZJ，Ekman差值为{fmt(diff_ek)} ZJ，相当于净差的{100.0 * diff_ek / diff_total:.0f}%；地转和残差差异反而部分抵消这一差距。因此2015与2023的直接分水岭不是海洋地转阻力，而是大气能否提供足够强的风驱输送。",
    )

    doc.add_heading("5.4 第三问的机制答案及因果边界", level=2)
    style.add_callout(
        doc,
        "第三问的答案",
        "FEP先暖只描述表面异常的起点。真正决定暖量命运的是两项条件能否同时成立：一是足够大的盆地尺度次表层暖量库存及向东调整的D20/SSH结构；二是允许深对流东移并建立降水—SLP—风应力响应的大气背景。二者共同决定Ekman能否克服海洋压力场阻力，并最终打开5°N暖量出口。",
        fill=style.LIGHT_BLUE,
    )
    style.add_para(
        doc,
        "1997同时满足两项条件，形成地转—Ekman协同排暖；2015属于海洋预置且大气完整接通，Ekman克服南向地转阻力；2023主要先出现FEP表面增暖，盆地深层调整和对流东移均较迟且较弱，Ekman不足以克服SSH压力场阻力，出口近闭合。",
    )
    style.add_para(
        doc,
        "现有资料把失配位置锁定在“FEP表面增暖之后、深对流东移之前”，但尚不能把外海盆增暖、冷NPMM、多年La Niña遗产、FEP降水偶极或WWB不足中的任何一个指定为本项目已经独立证明的唯一上游原因。Isca IA×FEP正式集合尚未通过240/240文件、因子统计和本地备份的完成审计，因此这些因素只作为候选调制背景。",
    )
    style.add_figure(
        doc,
        "case_coupling",
        "图7  1997与2023从OHC/D20背景、SST和风场到对流、气压梯度、风应力及暖量输出的全过程对照。",
        width=4.8,
    )

    doc.add_heading("5.5 十三事件统计与风应力替换完成定量过程归因", level=2)
    style.add_para(
        doc,
        "为了避免只从1997、2015和2023三个个例推断机制，将相同定义扩展到1982—2023年的13个暖事件。3—4月中太平洋0—300 m OHC与5—6月中西太平洋u850的事件级相关为"
        f"r={attr_rel.loc['ocean_to_wind', 'pearson_r']:+.2f}，事件置换p={attr_rel.loc['ocean_to_wind', 'permutation_p']:.2g}，"
        f"事件bootstrap 95%区间[{attr_rel.loc['ocean_to_wind', 'bootstrap_ci_low']:+.2f}, {attr_rel.loc['ocean_to_wind', 'bootstrap_ci_high']:+.2f}]；"
        f"逐一删除任一事件后的r范围为[{attr_rel.loc['ocean_to_wind', 'loeo_r_min']:+.2f}, {attr_rel.loc['ocean_to_wind', 'loeo_r_max']:+.2f}]。这里把r及其区间作为主要效应量，置换p只作辅助证据。",
    )
    style.add_para(
        doc,
        "5—6月u850与次年1—5月5°N Ekman Qwarm,20的跨事件相关为"
        f"r={attr_rel.loc['wind_to_ekman', 'pearson_r']:+.2f}，事件bootstrap 95%区间[{attr_rel.loc['wind_to_ekman', 'bootstrap_ci_low']:+.2f}, {attr_rel.loc['wind_to_ekman', 'bootstrap_ci_high']:+.2f}]，"
        f"置换p={attr_rel.loc['wind_to_ekman', 'permutation_p']:.2g}，逐一删事件r范围为[{attr_rel.loc['wind_to_ekman', 'loeo_r_min']:+.2f}, {attr_rel.loc['wind_to_ekman', 'loeo_r_max']:+.2f}]。"
        "这使“海洋预置—西风接通—Ekman出口”不再只是三事件视觉时序，但由于两段指标跨越发展期到衰减期，仍解释为过程联系而非单一步骤的直接因果。",
    )
    attribution_rows = []
    for year in (1997, 2015, 2023):
        row = attr_event.loc[year]
        attribution_rows.append(
            [
                str(year),
                fmt(row.fep_total_mar_apr_c),
                fmt(row.ohc300_central_ma_gj_m2),
                fmt(row.ohc_given_fep_residual_gj_m2),
                fmt(row.u850_cwp_mj_ms),
                fmt(row.q5n_ekman_jan_may1_zj),
                fmt(row.q5n_total_jan_may1_zj),
            ]
        )
    style.add_table(
        doc,
        "表9  表面触发、次表层预置、大气接通与暖量命运",
        ["事件", "FEP", "中太OHC", "给定FEP的OHC残差", "u850", "Ekman Q20", "总Q20"],
        attribution_rows,
        [760, 1100, 1250, 1900, 1100, 1500, 1750],
        font_size=7.8,
    )
    style.add_para(
        doc,
        f"2023的3—4月FEP为{fmt(attr_event.loc[2023, 'fep_total_mar_apr_c'])} °C，是三次关键事件中最暖；"
        f"但给定FEP暖度后的中太平洋OHC残差为{fmt(attr_event.loc[2023, 'ohc_given_fep_residual_gj_m2'])} GJ m⁻²，"
        f"从低到高排名{int(attr_event.loc[2023, 'ohc_given_fep_residual_lower_rank'])}/13。"
        f"2015和1997的相应残差分别为{fmt(attr_event.loc[2015, 'ohc_given_fep_residual_gj_m2'])}和"
        f"{fmt(attr_event.loc[1997, 'ohc_given_fep_residual_gj_m2'])} GJ m⁻²。因此2023最明确的异常不是“表面不暖”，而是表面暖与次表层预置严重不匹配。",
    )
    cf_rows = []
    for target in (1997, 2015, 2023):
        cf_rows.append(
            [
                str(target),
                fmt(attr_cf.loc[(target, "1997"), "reconstructed_total_zj"]),
                fmt(attr_cf.loc[(target, "2015"), "reconstructed_total_zj"]),
                fmt(attr_cf.loc[(target, "2023"), "reconstructed_total_zj"]),
            ]
        )
    style.add_table(
        doc,
        "表10  固定海洋状态、替换5°N风应力后的重构总Qwarm,20（ZJ）",
        ["固定海洋状态", "1997风", "2015风", "2023风"],
        cf_rows,
        [2100, 2420, 2420, 2420],
        font_size=8.2,
    )
    style.add_equation(
        doc,
        "Qrec(i←j) = Qgeo(i) + Qres(i) + QEkman[T(i), τx(j)]",
        "i表示被固定的海洋温度、地转项和诊断残差，j表示替换进来的5°N局地风应力。该式是本研究定义的离线反事实敏感性试验，不重新积分海洋或大气，也不允许温度、压力场和残差响应新风场。",
    )
    style.add_para(
        doc,
        "保持2023温度、地转项和残差不变，使用2023自身风应力时重构总出口为"
        f"{fmt(attr_cf.loc[(2023, '2023'), 'reconstructed_total_zj'])} ZJ；换入2015风应力后增至"
        f"{fmt(attr_cf.loc[(2023, '2015'), 'reconstructed_total_zj'])} ZJ，换入1997风应力后为"
        f"{fmt(attr_cf.loc[(2023, '1997'), 'reconstructed_total_zj'])} ZJ。反向将2015海洋状态与2023风应力结合，"
        f"出口由{fmt(attr_cf.loc[(2015, '2015'), 'reconstructed_total_zj'])}降至"
        f"{fmt(attr_cf.loc[(2015, '2023'), 'reconstructed_total_zj'])} ZJ。该结果量化局地风驱项对诊断出口的敏感性，但不称为完整因果检验。",
    )
    style.add_callout(
        doc,
        "第三问的定量答案",
        "FEP表面先暖不是充分条件。2023相对于其FEP暖度具有13事件中最大的中太平洋OHC缺口，随后西风和Ekman出口均弱；风应力互换可在固定海洋压力场下近乎互换2015与2023的出口状态。因此次表层预置调制大气接通效率，而局地风驱Ekman决定其能否克服地转阻力。",
        fill=style.LIGHT_BLUE,
    )
    style.add_figure(
        doc,
        "third_attribution",
        "图8  十三个暖事件的FEP—OHC—u850—次年5°N Ekman联系，以及1997、2015、2023局地风应力替换试验。",
        width=6.15,
    )

    style.add_section_heading(doc, "6 稳健性、完成度与结果边界")
    doc.add_heading("6.1 基准和状态方程敏感性", level=2)
    geo_sens = geo.loc[geo.baseline.eq("1982-2022")].set_index("event_year")
    sens_rows = []
    for year in (1997, 2015, 2023):
        sens_rows.append(
            [
                str(year),
                fmt(geo_main.loc[year, "total_zj"]),
                fmt(geo_sens.loc[year, "total_zj"]),
                fmt(geo_main.loc[year, "geo_eos80_zj"]),
                fmt(geo_sens.loc[year, "geo_eos80_zj"]),
                fmt(geo_main.loc[year, "geo_linear_zj"]),
            ]
        )
    style.add_table(
        doc,
        "表11  距平基准和状态方程敏感性（ZJ）",
        ["事件", "总量 1980—2014", "总量 1982—2022", "地转 1980—2014", "地转 1982—2022", "线性EOS地转"],
        sens_rows,
        [850, 1800, 1800, 1800, 1800, 1310],
        font_size=8.0,
    )
    style.add_para(
        doc,
        "两种距平基准下，三事件总出口保持1997>2015>2023的排序，Ekman符号和地转符号均保持不变；线性状态方程也不改变地转项方向。2015/2023地转幅度对基准较敏感，因此主文强调符号、相对排序和动力竞争，不把单一绝对值解释为普适常数。",
    )

    doc.add_heading("6.2 三个问题的完成度", level=2)
    style.add_table(
        doc,
        "表12  三个核心问题的证据状态",
        ["问题", "已解决的部分", "仍保留的边界", "判断"],
        [
            [
                "1 暖量是否到达",
                "140°W Q20、东部H20、OHC/D20/SSH共同证明到达及路径",
                "不是同一团水的拉格朗日轨迹",
                "主体解决",
            ],
            [
                "2 暖量去了哪里",
                "量化5°N、10°S、海表热损失、储量变化和剩余项",
                "剩余项尚未拆成混合、夹卷、跨等温面和同化增量",
                "主要通道解决",
            ],
            [
                "3 为什么命运不同",
                "三事件过程链、13事件统计和局地风应力替换共同完成定量过程归因",
                "外海盆等上游背景的唯一因果贡献仍需完成的模式集合",
                "过程归因解决；上游根因待试验",
            ],
        ],
        [1700, 3350, 2920, 1390],
        font_size=8.2,
    )
    style.add_callout(
        doc,
        "可直接进入论文的主结论",
        "暖量是否到达东部并不足以决定事件命运。到达后的暖量分配取决于海洋压力场与风驱Ekman的竞争，而这种竞争受盆地尺度次表层支撑和大气对流接通效率共同控制。",
        fill="FFF4E5",
        color=style.GOLD,
    )

    style.add_section_heading(doc, "附录A 直接产物与复现入口")
    style.add_table(
        doc,
        "表A1  本方法与结果部分的直接计算来源",
        ["模块", "主要结果目录/数据", "主脚本"],
        [
            [
                "三箱暖量到达与去向",
                "outputs/research_2026/soda3_three_box_warm_heat_fate/",
                "scripts/analyze_soda3_three_box_warm_heat_fate.py",
            ],
            [
                "CT/CP海气耦合",
                "outputs/research_2026/soda3_air_sea_coupled_fate/",
                "scripts/analyze_soda3_air_sea_coupled_warm_heat_fate.py",
            ],
            [
                "1997/2015/2023暖量桥接",
                "data/derived/research_2026/enso_1997_2015_2023_warm_heat_fate_phase_summary.csv",
                "scripts/analyze_2015_2023_warm_heat_fate_bridge.py",
            ],
            [
                "FEP与大气闭合",
                "data/derived/research_2026/jiang_fep_soda3_atmospheric_closure_1997_2015_2023.csv",
                "scripts/analyze_jiang_soda3_transition.py",
            ],
            [
                "三事件5°N动力分解",
                "outputs/research_2026/enso_1997_2015_2023_5n_geostrophic_fate/",
                "scripts/analyze_1997_2015_2023_5n_geostrophic_fate.py",
            ],
            [
                "第三问事件级归因",
                "outputs/research_2026/third_question_coupled_attribution/",
                "scripts/analyze_third_question_coupled_attribution.py",
            ],
        ],
        [1900, 4070, 3390],
        font_size=7.7,
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
