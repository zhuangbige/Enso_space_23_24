#!/usr/bin/env python3
"""Build a Chinese scientific-paper-style DOCX for the ENSO project."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from scripts.build_enso_stage_complete_report import (
    add_hyperlink,
    add_page_field,
    mark_table_header_row,
    set_cell_margins,
    set_run_font,
    set_table_borders,
    set_table_geometry,
)


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "outputs/research_2026"
OUTDIR = FIG / "reports"
OUTPUT = OUTDIR / "ENSO类型形成与海气脱耦_完整研究论文_CN_审阅修订版_20260731.docx"

FIGURES = {
    "framework": FIG / "paper_cn/fig01_enso_two_gate_framework_cn.png",
    "pathways": FIG / "complete_enso_type_research/fig3_event_pathway_comparison.png",
    "checkpoint": FIG / "complete_enso_type_research/fig1_type_definition_and_transition_checkpoint.png",
    "ocean_process": FIG / "soda3_fep_transition/soda3_process_full_event_evidence.png",
    "oras": FIG / "oras5_independent_validation/soda3_oras5_independent_validation.png",
    "five_day": FIG / "soda3_5day_propagation/soda3_5day_d20_ssh_propagation.png",
    "nmme": FIG / "complete_enso_type_research/fig2_nmme_barrier_and_coupling_update.png",
    "reviewer_robustness": FIG / "paper_cn/fig11_reviewer_robustness.png",
    "factor_matrix": FIG / "perturbation_attribution/fig1_event_factor_and_response_matrix.png",
    "ne_bridge": FIG / "perturbation_attribution/fig4_fep_ne_wind_bridge.png",
    "forcing": FIG / "p0_agcm_prior/fig6_p0_agcm_forcing_components.png",
}

# LibreOffice on this macOS host does not reliably honor a separate East-Asia
# font when w:ascii/w:hAnsi point to a Latin-only font. Use installed CJK fonts
# for every script slot so the visual-QA render matches Word without tofu.
LATIN = "STSong"
CJK_BODY = "STSong"
CJK_HEAD = "Arial Unicode MS"
BLACK = "111111"
MUTED = "555555"
LIGHT = "F1F1F1"
GRID = "8F8F8F"


def set_repeat_table_header(row):
    mark_table_header_row(row)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_BODY)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)

    for name, size, before, after in (
        ("Heading 1", 14, 14, 8),
        ("Heading 2", 12, 11, 5),
        ("Heading 3", 10.5, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = LATIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_HEAD)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)

    caption = doc.styles["Caption"]
    caption.font.name = LATIN
    caption._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_BODY)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(BLACK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    props = doc.core_properties
    props.title = "热带太平洋暖量东传与大气反馈接通对ENSO类型形成的共同控制"
    props.subject = "ENSO类型形成、海气脱耦与春季预报障碍"
    props.author = "Daiyue"
    props.keywords = "ENSO; EP/CP; 海气耦合; 海洋热含量; 春季预报障碍"


def set_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_run_font(
        p.add_run("ENSO类型形成、海气脱耦与春季预报障碍"),
        latin=LATIN,
        east_asia=CJK_BODY,
        size=8.5,
        color=MUTED,
    )
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def add_para(doc: Document, text: str, *, indent: bool = True, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21 if indent else 0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    set_run_font(
        p.add_run(text),
        latin=LATIN,
        east_asia=CJK_BODY,
        size=10.5,
        bold=bold,
        color=BLACK,
    )


def add_keywords(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("关键词："), latin=LATIN, east_asia=CJK_HEAD,
                 size=10.5, bold=True, color=BLACK)
    set_run_font(p.add_run(text), latin=LATIN, east_asia=CJK_BODY,
                 size=10.5, color=BLACK)


def set_cell_text(cell, text: str, *, bold=False, size=8.5, center=False) -> None:
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(str(text)), latin=LATIN, east_asia=CJK_BODY,
                 size=size, bold=bold, color=BLACK)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float], *, size=8.3, new_page=False) -> None:
    cap = doc.add_paragraph()
    if new_page:
        cap.paragraph_format.page_break_before = True
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(5)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    set_run_font(cap.add_run(caption), latin=LATIN, east_asia=CJK_BODY,
                 size=9, bold=True, color=BLACK)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    set_table_borders(table, color=GRID, size=5)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), LIGHT)
        cell._tc.get_or_add_tcPr().append(shading)
        set_cell_text(cell, header, bold=True, size=size, center=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i], top=80, bottom=80, start=100, end=100)
            set_cell_text(cells[i], value, size=size, center=i == 0)
    set_table_geometry(table, widths)
    set_table_borders(table, color=GRID, size=5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, key: str, caption: str, *, width=6.05, new_page=False) -> None:
    path = FIGURES[key]
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    if new_page:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    pic = p.add_run().add_picture(str(path), width=Inches(width))
    pic._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph(style="Caption")
    set_run_font(cp.add_run(caption), latin=LATIN, east_asia=CJK_BODY,
                 size=9, color=BLACK)


def add_equation(doc: Document, label: str, equation: str, explanation: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(equation), latin=LATIN, east_asia=CJK_BODY,
                 size=10.5, italic=True, color=BLACK)
    add_para(doc, f"其中，{explanation}")


def add_reference(doc: Document, n: int, citation: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(f"[{n}] {citation} "), latin=LATIN,
                 east_asia=CJK_BODY, size=9, color=BLACK)
    add_hyperlink(p, "DOI/链接", url)


def add_landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = 1
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    return section


def build() -> None:
    for path in FIGURES.values():
        if not path.exists():
            raise FileNotFoundError(path)
    OUTDIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    set_header_footer(doc.sections[0])

    # Article opening: no report cover or executive-summary furniture.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(
        p.add_run("热带太平洋暖量东传与大气反馈接通\n对ENSO类型形成的共同控制"),
        latin=LATIN,
        east_asia=CJK_HEAD,
        size=18,
        bold=True,
        color=BLACK,
    )
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Pt(0)
    p2.paragraph_format.space_after = Pt(10)
    set_run_font(
        p2.add_run("——兼论2023/24厄尔尼诺的非典型发展路径与春季类型预报障碍"),
        latin=LATIN,
        east_asia=CJK_BODY,
        size=12,
        color=MUTED,
    )

    doc.add_heading("摘要", level=1)
    add_para(doc, (
        "东太平洋表层海温异常偏高并不必然导致东太平洋型厄尔尼诺。"
        "为识别从暖水蓄积到成熟空间型态形成之间的关键环节，本文综合使用"
        "SODA3.15.2、ORAS5、ERA5、ERSSTv5、OISST、GPCP、CPC OLR及NOAA NMME资料，"
        "比较1982—2023年13个暖事件，并将2023/24与1997/98和2015/16置于统一诊断框架。"
        "结果表明，ENSO类型形成至少包含蓄能、向东重分配—储存和大气反馈兑现三个阶段。"
        "3—6月春季转变窗口内，140°W附近Q20与D20描述暖量是否真正进入中东太平洋；5—6月降水、气压梯度"
        "和赤道纬向风应力相对于同期SST的残差，则描述Bjerknes反馈能否接上。"
        "逐月扫描中，四月首次出现稳定的类型预测增量：SODA3与ORAS5的RMSE技巧分别为"
        "+0.43和+0.50；它是春季窗口内的候选切片而非普适阈值。对降水偶极—综合接通关系的"
        "全部78组留二事件检验均保持正号，中位偏相关为0.73，最低值为0.33，说明关系具有符号"
        "稳健性但仍受1997和2015等强事件影响。早春NMME类型误差与5—6月耦合残差的偏相关为"
        "0.79；年代顺序扩展窗口中，加入该观测更新后的RMSE为0.32，低于持续性、Niño3.4单因子"
        "和校准NMME的0.73、0.56和0.43，但这属于春末观测更新而非纯春季预报。2023表现为强表层FEP"
        "路线、较弱的暖量东传—储存信号和不完整的赤道西风反馈，其成熟型态对分类方法敏感，"
        "不能视为1997的简单弱化版本。观测扰动进一步显示，经向降水偶极最接近对流—气压"
        "接通状态，但单独改变多年La Niña、Indo–Atlantic加热、NPMM或WWB均未稳健改变2023的"
        "事件类型。因此，当前证据支持“海洋门控—大气兑现”的条件链，而不支持单一触发阈值。"
    ))
    add_keywords(doc, "ENSO类型；东太平洋型；中太平洋型；海洋热含量；海气脱耦；春季预报障碍")

    doc.add_heading("1 引言", level=1)
    add_para(doc, (
        "经典Bjerknes反馈把赤道东太平洋增暖、Walker环流减弱、赤道西风异常和温跃层加深"
        "视为相互放大的过程。然而，近年的强暖事件表明，海洋异常与大气响应并不总是同步。"
        "Peng等[1]指出，2023/24事件的强海温主要由海洋动力过程维持，而风场和SOI响应明显弱于"
        "历史强事件。Jiang等[5]进一步提出远东太平洋残余增暖及经向降水偶极可组织异常东北风；"
        "Fan等[6]则强调多年La Niña背景留下的西部暖量和纬向梯度。"
    ))
    add_para(doc, (
        "这些研究共同说明，FEP异常、总热含量、西太平洋暖水和大气风场分别刻画不同阶段。"
        "如果将它们压缩成单一指数，容易把事件是否发生、成熟暖中心位置以及大气反馈强度混为"
        "一谈。本文因此关注两个条件问题：其一，在海洋已经具备暖事件潜力后，暖量何时完成"
        "向东重分配和储存；其二，在东部表层增暖已经形成后，什么决定降水—气压—赤道西风"
        "反馈能否接上。"
    ))
    add_figure(doc, "framework", "图1  ENSO类型形成的三阶段与两个动力关口。该图为本文诊断框架；两道关口分别对应暖量向东输送—储存以及大气Bjerknes反馈兑现。", width=6.1)

    doc.add_heading("2 数据与方法", level=1)
    doc.add_heading("2.1 数据资料", level=2)
    add_table(doc, "表1  主要资料及用途", ["资料", "主要变量", "时间范围", "用途"], [
        ["SODA3.15.2", "OHC、Q20、D20、SSH、风应力、流速", "1980—2024", "暖量蓄积、输送、热收支和5日传播"],
        ["ORAS5", "D20、风应力", "1980—2026年6月", "温跃层和东传走廊的独立核验"],
        ["ERA5", "u850/v850、u10、SLP、风应力、OLR", "1980—2026", "低层环流和大气反馈"],
        ["ERSSTv5/OISST", "SST与Niño指数", "1980—2026", "事件类型、FEP和空间异常"],
        ["GPCP/CPC OLR", "降水与向外长波辐射", "1982—2024", "深对流及经向降水偶极"],
        ["NOAA NMME", "季节SST回报", "1991—2020", "固定SON目标的春季预报障碍"],
        ["TAO 0°,140°W", "温度剖面与流速", "共同可用时段", "SODA3 Q20/D20现场观测核验"],
    ], [1.1, 1.75, 1.1, 2.05], size=8.0)
    add_para(doc, (
        "新计算的事件异常和扰动统计统一采用1982—2024逐月气候态并去除线性趋势；SODA3"
        "过程图采用1980—2024同月异常。空间平均使用余弦纬度权重。所有预测模型均在训练折内"
        "重新估计均值、趋势、回归系数和标准差，以避免未来信息泄漏。"
    ))

    doc.add_heading("2.2 事件样本与类型定义", level=2)
    add_para(doc, (
        "事件样本包括1982、1986、1991、1994、1997、2002、2004、2006、2009、2015、"
        "2018、2019和2023，共13个暖事件。成熟期为发展年12月至次年2月。主类型指标采用"
        "Niño3−Niño4；敏感性检验采用Takahashi等[2]的正交E/C逻辑，并以(E−C)/√2表示东部"
        "相对于中部的暖中心位置。只有两类指标符号一致时才宣布稳定EP或CP；符号不一致时"
        "记为混合或定义敏感事件。"
    ))
    add_para(doc, (
        "评审建议增加1972、1987、1994和2020。这里不把它们机械计为四个独立新增暖事件："
        "1994已在主样本内，1987与1986/87属于同一发展事件，2020不满足本文的暖事件筛选条件，"
        "而1972早于SODA3.15.2的1980年起点，只能在替代海洋资料的敏感性检验中使用。后续扩大"
        "样本应以统一筛选规则和共同资料覆盖期为前提，不能以重复计数或改变事件定义换取样本量。"
    ))

    doc.add_heading("2.3 海洋门控与大气接通", level=2)
    add_equation(doc, "式1", "TSI = 0.5 × [z(Q20₁₄₀W, Mar–Apr) + z(D20₁₄₀W, Mar–Apr)]",
                 "Q20表示穿过140°W的上层暖量输送，D20表示当地温跃层加深和储热。TSI仅作为暖事件条件下的分型更新量，不作为普适发生指数。")
    add_equation(doc, "式2", "Rₖ,m = [Yₖ,m − Ŷₖ,m(Niño3.4, SST type, trend)] / σ(εₖ,m)",
                 "Y分别为降水、东西向SLP驱动和中西太平洋纬向风应力；期望关系在每个留一事件或时间外推训练折内重建。综合接通量为三条标准化残差的平均，但正文同时报告各支路。")

    doc.add_heading("2.4 统计检验与预测验证", level=2)
    add_para(doc, (
        "关联分析同时控制早期SST型态、Niño3.4强度、趋势和海洋门控；多重比较采用"
        "Benjamini–Hochberg FDR。事件不确定性使用事件块Bootstrap。类型预测采用留一事件"
        "验证；大气接通核心关系另穷举13个事件中的全部78组留二事件检验。业务预报更新采用"
        "逐年扩展窗口，确保目标年份及其未来信息不参与训练，并与春季类型持续性、Niño3.4"
        "单因子模型及不含观测更新的校准NMME作同窗比较。"
    ))

    doc.add_heading("3 结果", level=1)
    doc.add_heading("3.1 1997、2015与2023代表不同的发展路径", level=2)
    add_para(doc, (
        "三个强暖事件均具有较大的海洋储备，但春季后的演变并不相同。1997的暖量重心迅速"
        "东移，随后大气接通残差显著增强；2015的海洋和大气响应均较强，但成熟型态更接近"
        "E/C混合；2023的FEP表层路径最突出，而大气接通残差接近零或为负。"
    ))
    add_figure(doc, "pathways", "图2  1997、2015和2023年1—6月的暖水储备、暖量重心、FEP残余海温和海气接通残差。阴影表示3—5月关键发展窗口。", width=6.05, new_page=True)

    doc.add_heading("3.2 类型定义与春季转变窗口", level=2)
    add_para(doc, (
        "Niño3−Niño4与成熟Niño3.4强度的相关为0.83，说明传统箱区型指标包含显著的幅度"
        "混杂。正交E−C指标把该相关降低至0.64，并显示2023在Niño3−Niño4上偏东、在E−C上"
        "偏中，属于定义敏感的混合事件。对1—6月进行事先一致的逐月留一扫描后，SODA3在四月"
        "首次获得+0.43的RMSE"
        "技巧，符号正确率92%；ORAS5为+0.50，符号正确率100%。换用E−C目标后仍保持正技巧。"
    ))
    add_para(doc, (
        "但剔除1982、1997和2015后，四月技巧的Bootstrap区间重新跨越零。因此，本文把3—6月"
        "整体定义为春季转变窗口；四月只是其中最早出现正增量的、由数据识别出的候选切片，"
        "用于生成动力假设，而不是事先指定或适用于所有暖事件的硬阈值。"
    ))
    add_figure(doc, "checkpoint", "图3  类型定义及春季转变窗口内的逐月类型预测技巧。四月为扫描中首次出现正增量的候选切片；误差线为事件Bootstrap区间。", width=6.05, new_page=True)

    doc.add_heading("3.3 暖量向东重分配比总热含量更接近类型形成", level=2)
    add_para(doc, (
        "跨13事件比较表明，总暖水体积主要反映事件强度；暖量重心、东部暖水比例、SSH东西"
        "梯度、障碍层和垂向平流与成熟Niño3−Niño4的联系更直接。在El Niño条件样本中，"
        "海洋重分配组合在3—6月对SST持续性基准提供正技巧，而无条件全年度走步检验并不稳定，"
        "进一步说明这些量属于条件式类型预测量。"
    ))
    add_figure(doc, "ocean_process", "图4  SODA3对暖水蓄积向EP增长转变的过程诊断。上排给出事件异常及与成熟型态的相关；下排比较全年度与暖事件条件下的逐月预测技巧。", width=6.05, new_page=True)

    doc.add_heading("3.4 ORAS5与五日资料的过程核验", level=2)
    add_para(doc, (
        "用ORAS5替换SODA3的D20和暖量重心后，四月技巧仍保持同号，说明关键结论并非单一"
        "海洋再分析的产物。SODA3五日风应力扰动与随后140°W暖水输送的最大相关出现在风领先"
        "约20天，r=0.17，事件块Bootstrap区间为[0.12, 0.26]。相关幅度不大，但其时间顺序与"
        "WWB激发Kelvin型调整和暖量东传的动力预期一致。"
    ))
    add_figure(doc, "oras", "图5  SODA3与ORAS5对温跃层位置和东传通道的独立对照。两套资料给出的关键月份和事件排序总体一致。", width=6.05)
    add_figure(doc, "five_day", "图6  五日风应力、D20与SSH传播诊断。峰值关系出现在西部风扰动领先140°W暖量输送约20天。", width=6.05, new_page=True)

    doc.add_heading("3.5 春季类型定位困难与耦合更新", level=2)
    add_para(doc, (
        "固定SON验证目标后，NMME对Niño3.4幅度的技巧高于对暖中心位置的技巧。按"
        "Niño3−Niño4标签分组，EP的春季惩罚为+0.30 ℃、CP为+0.01 ℃；换成正交E−C误差后，"
        "CP惩罚反而为+0.24而EP为+0.04。由此可见，EP或CP哪一类春季障碍更强取决于分型定义"
        "和样本中的极端事件，稳健结论应表述为“暖中心位置技巧落后于幅度技巧”。"
    ))
    add_para(doc, (
        "早春SON类型绝对误差与5—6月耦合残差的偏相关为0.79（p=0.020）。在严格逐年扩展窗口"
        "试验中，春季类型持续性、Niño3.4单因子、校准NMME和NMME加5—6月耦合更新的SON类型"
        "RMSE依次为0.73、0.56、0.43和0.32；更新相对前三者分别改善56%、43%和26%，相关由"
        "NMME的0.68提高到0.83。留一年区间为[+1%, +42%]，而年代顺序区间为[−10%, +50%]。"
        "因此，该量是可在5—6月使用的春末观测更新因子，不是3月即可获得的纯预报量，也还不是"
        "已确立的业务增益。"
    ))
    add_figure(doc, "nmme", "图7  NMME春季类型定位障碍及基于海气耦合状态的更新试验。面板(c)误差线表示事件块Bootstrap区间。", width=6.05, new_page=True)

    doc.add_heading("3.6 大气接通需要区分对流—气压和赤道西风两条支路", level=2)
    add_para(doc, (
        "在控制海洋门控、FEP和早期SST型态后，FEP经向降水偶极与综合大气接通的偏相关为"
        "+0.72，95%区间为[+0.47, +0.91]，FDR q=0.024。该偶极与降水残差和SLP梯度的联系"
        "分别为+0.79和+0.71，但与纬向风应力仅为+0.46且区间跨零。这意味着“对流已经响应”"
        "并不等价于“能推动暖量东传的赤道西风已经接上”。"
    ))
    add_para(doc, (
        "留二事件检验进一步区分了稳健关系与极端事件杠杆。综合接通关系的78组删除结果全部"
        "为正，中位偏相关为0.73，95%的组合高于0.50；同时删除1997和2015时降至最低0.33。"
        "降水支路的78组结果全部高于0.50，SLP支路也大多稳定，而风应力支路只有45%的组合"
        "高于0.50，个别组合可转为负值。这表明稳定证据集中在对流—气压支路，赤道西风支路"
        "仍需要更多事件和动力试验。"
    ))
    add_figure(doc, "reviewer_robustness", "图8  评审后补充的稳健性检验。(a) 全部78组留二事件偏相关分布；(b) 同一年代顺序扩展窗口下的SON类型RMSE基线比较。", width=6.05, new_page=True)
    add_figure(doc, "factor_matrix", "图9  13个暖事件的背景因子、海洋条件与大气响应矩阵。各列为去趋势并控制同期SST后的标准化量。", width=6.05)
    add_para(doc, (
        "矩阵中，海洋门控较强并不自动对应更强的大气兑现，反之亦然。强EP事件通常同时位于"
        "暖量东传和纬向风应力响应的正侧；2023则在FEP降水偶极上突出，却在赤道西风支路上"
        "偏弱。这种错位是本文所称海气脱耦的操作性定义：它不是海洋与大气完全无关，而是"
        "给定同期SST与海洋门控后，关键大气过程没有达到历史同类事件所暗示的响应幅度。"
    ))
    add_para(doc, (
        "进一步分解显示，降水加热可先建立经向气压结构，但要把该结构投影到赤道纬向风，"
        "仍取决于东北风入侵、背景Walker环流和天气尺度西风爆发。因而用OLR或降水单独判定"
        "Bjerknes反馈已经闭合会过早；至少应同时检查SLP梯度和赤道纬向风应力。WWB脉冲、"
        "WWB频次和东北风入侵采用多种定义后，与综合接通的控制偏相关仍较弱且符号不统一。"
        "因此本文不把这些过程强行压缩成一个拟合的“大气闸门指数”：小样本下该做法易过拟合，"
        "若把最终风应力响应纳入解释风应力的指数还会产生结果变量泄漏。"
    ))
    add_figure(doc, "ne_bridge", "图10  FEP残余、经向降水偶极、东北风入侵和纬向风应力之间的桥梁检验。2023个例符号完整，但跨事件路径的不确定区间仍较宽。", width=6.05, new_page=True)

    doc.add_heading("3.7 单因子替换不能稳健改变2023的成熟类型", level=2)
    add_para(doc, (
        "使用不包含2023的训练样本，将2023的多年La Niña记忆、Indo–Atlantic加热、NPMM、"
        "FEP降水偶极和WWB分别替换为中性或1997/2015水平，共七种单因子及组合情景。所有情景"
        "均未使E/C与Niño3−Niño4同时过零。单因子+1标准差扰动的Bootstrap区间也全部跨零。"
        "因此，当前观测统计否定了“某一个量单独控制所有年份”的简单解释。"
    ))
    add_table(doc, "表2  核心定量结果与解释边界", ["诊断", "结果", "可以支持的解释", "不能支持的解释"], [
        ["春季海洋门控", "四月切片：SODA3 +0.43；ORAS5 +0.50", "3—6月转变窗口内的强EP候选检查点", "普适类型阈值"],
        ["5—6月耦合更新", "RMSE 0.32；基线为0.73/0.56/0.43", "春末观测更新候选量", "纯3月预报或已确立业务增益"],
        ["降水偶极", "partial r=+0.72；留二中位0.73、最低0.33", "符号稳健的对流—气压接通状态", "独立上游触发器"],
        ["WWB—140°W输送", "风领先20天；r=+0.17", "时间顺序与动力预期一致", "WWB独立决定事件类型"],
        ["2023反事实替换", "两种类型指标过零次数=0", "条件链和交互更合理", "单因子可稳健改型"],
    ], [1.25, 1.35, 1.75, 1.65], size=7.9)

    doc.add_heading("4 讨论", level=1)
    doc.add_heading("4.1 FEP高为什么不一定形成EP", level=2)
    add_para(doc, (
        "FEP刻画远东太平洋表层增暖，既可能由局地表面热通量、沿岸过程或浅层温跃层变化产生，"
        "也可能是盆地尺度暖量东传的结果。只有后一过程进一步在中东太平洋形成D20、SSH和上层"
        "热含量异常，并获得持续赤道西风反馈，才容易发展成典型强EP。因而FEP是表层路径指标，"
        "不是完整海洋充放电状态的替代。"
    ))
    doc.add_heading("4.2 2023不是1997的弱化复制", level=2)
    add_para(doc, (
        "2023同时具有多年La Niña记忆、异常暖的热带北大西洋和印度洋背景、冷NPMM、较少WWB、"
        "较强东北风入侵以及偏弱纬向风应力。其FEP残余和经向降水偶极又明显偏强。这一组合"
        "说明FEP可以先组织局地对流和气压响应，但背景态仍可阻断其向赤道西风支路的转换。"
        "1997则表现为暖量重心快速东移并与强大气反馈同步，属于两道关口均通过的强EP分支。"
    ))
    doc.add_heading("4.3 与现有研究的关系", level=2)
    add_para(doc, (
        "本文的双关口框架与Fang和Mu[7]关于海洋温跃层状态和西太平洋风异常共同控制可预报性"
        "的结果一致，也与Li等[3]把大气Bjerknes反馈拆分为SST—降水加热—气压—风应力的过程"
        "框架相衔接。其增量不在于提出一个新的单变量指数，而在于把SODA3/ORAS5海洋位置、"
        "过程分支残差、NMME空间型回报和观测扰动放入同一套可证伪检验，并明确指出类型定义"
        "和极端事件对春季障碍结论的影响。"
    ))
    add_para(doc, (
        "需要强调，“双关口”目前是受观测支持、可由动力试验证伪的组织框架，而非已经确立的"
        "普适范式。13事件样本仍不足以稳定估计高阶交互；扩大样本又受到海洋再分析覆盖期、"
        "事件独立性和类型定义的共同限制。稳健性分析显示，对流—气压支路比风应力支路更稳定，"
        "因此后续工作的优先级不是继续合成指数，而是分别识别气压结构如何投影到赤道西风。"
    ))

    doc.add_heading("5 动力因果试验设计与当前进展", level=1)
    add_para(doc, (
        "观测回归只能说明候选条件与额外大气响应是否同现，不能证明人为移除某个海盆异常后"
        "风场一定改变。因此，第一阶段动力试验采用Indo–Atlantic（IA）开/关与residual-FEP"
        "开/关的2×2规定SST矩阵。主试验使用2023总IA异常；敏感性试验只保留移除趋势、"
        "Niño3.4和E/C型态后的独立IA残差。两套四格的“全开”均逐格重构2023观测SST，误差为零。"
    ))
    add_para(doc, (
        "规定SST的AGCM只能回答“给定海温以后，大气为何接上或接不上”，不能动态开启或关闭"
        "海洋门控。真正的因果分岔由后续耦合矩阵检验：海洋门控低/中/高与风场不利/中性/有利"
        "构成3×3试验。低、高门控通过在盆地内向西或向东平移等量上层热含量实现，尽量守恒"
        "盆地总热量；不直接把东部温跃层钳制到气候态，以免引入非平衡Kelvin调整并混淆因果。"
        "NPMM、WWB及外海盆背景只在主矩阵出现稳定类型翻转后顺序加入，避免用13事件资料同时"
        "拟合不可识别的多重高阶交互。"
    ))
    add_figure(doc, "forcing", "图11  2023年3—8月IA与residual-FEP规定SST强迫组成。上排为两个IA定义和FEP分量，下排为双关闭基态及完整观测异常。", width=6.05)
    add_table(doc, "表3  预注册动力实验及判据", ["阶段", "因子矩阵", "回答的问题", "判据"], [
        ["P0 AGCM", "IA on/off × FEP on/off", "给定SST后外海盆是否调制FEP大气响应", "降水、SLP、赤道纬向应力的主效应与交互"],
        ["P0耦合", "海洋门控低/中/高 × 风场不利/中性/有利", "暖量位置和风场能否共同导致类型翻转", "DJF E/C与Niño3−Niño4同时过零"],
        ["P1耦合", "NPMM × WWB，三个起始日期", "天气扰动是否依赖背景态", "集合均值交互显著且对起始日稳健"],
    ], [0.9, 1.55, 2.1, 1.45], size=8.0)
    add_para(doc, (
        "本机已完成Isca T42、40层、真实地形、RRTM辐射和Full Betts–Miller对流配置。干动力核、"
        "AMIP规定SST及自定义IA0×FEP0强迫均完成端到端技术试跑，降水、OLR、10米风和表面"
        "动量通量无NaN。8套SST边界已转换为Isca月变格式。24个月共同控制态正在以逐月重启方式"
        "连续积分，截至2026年7月31日22:43已通过前22个月且正在运行第23个月；已完成月段的主要变量无缺测且量级合理。上述结果仅证明实验环境可用，"
        "不构成科学因果证据；正式结论须等待共同控制态和集合积分。"
    ))

    doc.add_heading("6 结论", level=1)
    conclusions = [
        "FEP异常偏高或总热含量增加都不是EP型厄尔尼诺形成的充分条件。",
        "ENSO类型形成应区分蓄能、暖量向东重分配—储存和大气反馈兑现三个阶段。",
        "3—6月构成海洋门控的春季转变窗口；四月及140°W附近Q20/D20是强EP分支的候选切片，而非普适硬阈值。",
        "大气接通至少包含对流—气压和赤道西风—风应力两条支路；留二检验支持前者的符号稳健性，但后者仍受事件组合影响。",
        "2023代表强表层FEP路线、较弱海洋门控和不完整大气兑现的混合路径，而非1997的弱化复制。",
        "海气耦合残差可改善春末类型预报，但样本量和年代顺序检验仍不足以确立业务增益。",
        "最终因果闭环必须把规定SST AGCM的大气归因与守恒热量扰动的耦合模式矩阵分开完成。",
    ]
    for i, text in enumerate(conclusions, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(21)
        p.paragraph_format.first_line_indent = Pt(-21)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"（{i}）{text}"), latin=LATIN, east_asia=CJK_BODY,
                     size=10.5, color=BLACK)

    doc.add_heading("数据与代码可用性", level=1)
    add_para(doc, (
        "派生指数、留一及留二事件预测、NMME基线与更新结果、扰动矩阵和规定SST边界均保存在本地项目的"
        "data/derived/research_2026与experiments目录；核心绘图和分析脚本位于scripts目录，"
        "其中apply_reviewer_robustness_checks.py可重现本次评审后新增的稳健性图表。"
        "项目自身测试共70项，当前全部通过。"
    ))

    doc.add_heading("参考文献", level=1)
    refs = [
        ("Peng et al. (2025). Strong 2023–2024 El Niño generated by ocean dynamics. Nature Geoscience.", "https://www.nature.com/articles/s41561-025-01700-9"),
        ("Takahashi, K. et al. (2011). ENSO regimes: Reinterpreting the canonical and Modoki El Niño. Geophysical Research Letters, 38.", "https://doi.org/10.1029/2011GL047364"),
        ("Li et al. (2024). Process-separated atmospheric Bjerknes feedback. Geophysical Research Letters.", "https://doi.org/10.1029/2024GL109874"),
        ("Kim, Tian & Yu (2025). Western-Pacific warm water and zonal-wind evolution across El Niño events. npj Climate and Atmospheric Science.", "https://doi.org/10.1038/s41612-024-00890-0"),
        ("Jiang et al. (2026). Far-eastern Pacific warming and the atmospheric response during the 2023/24 El Niño. Geophysical Research Letters.", "https://doi.org/10.1029/2026GL123227"),
        ("Fan et al. (2026). Multi-year La Niña influence on different El Niño development pathways. The Innovation Geoscience.", "https://doi.org/10.59717/j.xinn-geo.2026.100224"),
        ("Fang & Mu (2018). Both thermocline state and western-Pacific wind anomalies are necessary for ENSO predictability. Scientific Reports.", "https://www.nature.com/articles/s41598-018-28964-z"),
        ("Lin et al. (2025). North Pacific Meridional Mode conditions and state-dependent westerly wind bursts. npj Climate and Atmospheric Science.", "https://doi.org/10.1038/s41612-025-01013-z"),
        ("Zuo et al. (2019). The ORAS5 ocean synthesis system. Ocean Science, 15, 779–808.", "https://doi.org/10.5194/os-15-779-2019"),
        ("Carton et al. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983.", "https://doi.org/10.1175/JCLI-D-18-0149.1"),
        ("Kirtman et al. (2014). The North American Multimodel Ensemble: Phase-1 seasonal-to-interannual prediction. Bulletin of the American Meteorological Society.", "https://doi.org/10.1175/BAMS-D-12-00050.1"),
        ("Chen et al. (2015). Strong influence of westerly wind bursts on El Niño diversity. Nature Geoscience.", "https://doi.org/10.1038/ngeo2399"),
    ]
    for i, (citation, url) in enumerate(refs, 1):
        add_reference(doc, i, citation, url)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
