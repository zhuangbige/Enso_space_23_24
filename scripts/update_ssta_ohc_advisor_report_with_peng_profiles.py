#!/usr/bin/env python3
"""Update the user's edited SSTA--OHC report with Peng-style SODA3 profiles."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports"
    / "ENSO发展期_SSTA-OHC不同步_阶段性汇报_修订版_20260815.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports"
    / "ENSO发展期_SSTA-OHC不同步_导师汇报_剖面更新版_20260816.docx"
)
FIGURE_DIR = ROOT / "outputs/research_2026/soda3_peng2025_profile_comparison"
FIGURES = {
    "three_stage": FIGURE_DIR
    / "fig_soda3_peng2025_jfm_aproct_ndj_2023_vs_strong_event_composite.png",
    "four_season": FIGURE_DIR
    / "fig_soda3_peng2025_jfm_amj_jas_ond_2023_vs_1982_1997_2015_profiles.png",
    "monthly": FIGURE_DIR
    / "fig_soda3_peng2025_jan_jun_2023_monthly_profile_evolution.png",
}
CJK_FONT = "PingFang SC"


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_text(paragraph: Paragraph, text: str) -> Paragraph:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)
    return paragraph


def paragraph_by_text(doc: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph matching {text!r}, found {len(matches)}")
    return matches[0]


def paragraph_by_prefix(doc: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def paragraph_by_contains(doc: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if text in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph containing {text!r}, found {len(matches)}")
    return matches[0]


def clone_after(anchor: Paragraph, prototype: Paragraph, text: str = "") -> Paragraph:
    element = deepcopy(prototype._p)
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    replace_text(paragraph, text)
    return paragraph


def set_figure(
    paragraph: Paragraph,
    path: Path,
    *,
    width_inches: float,
    alt_text: str,
) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    doc_properties = paragraph._p.xpath(".//*[local-name()='docPr']")
    if doc_properties:
        doc_properties[0].set("descr", alt_text)


def add_figure_after(
    anchor: Paragraph,
    image_prototype: Paragraph,
    caption_prototype: Paragraph,
    path: Path,
    caption: str,
    *,
    width_inches: float,
    alt_text: str,
) -> tuple[Paragraph, Paragraph]:
    image_paragraph = clone_after(anchor, image_prototype)
    set_figure(
        image_paragraph,
        path,
        width_inches=width_inches,
        alt_text=alt_text,
    )
    caption_paragraph = clone_after(image_paragraph, caption_prototype, caption)
    return image_paragraph, caption_paragraph


def empty_headers(doc: Document) -> None:
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            clear_paragraph(paragraph)


def set_document_font(doc: Document) -> None:
    """Use a locally available Office-safe CJK font throughout the report."""
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = CJK_FONT
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            from docx.oxml import OxmlElement

            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{slot}"), CJK_FONT)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = CJK_FONT
            r_pr = run._element.get_or_add_rPr()
            r_fonts = r_pr.find(qn("w:rFonts"))
            if r_fonts is None:
                from docx.oxml import OxmlElement

                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{slot}"), CJK_FONT)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for path in FIGURES.values():
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(SOURCE)
    # Several empty paragraphs exist; use the current Figure 4 image paragraph
    # as the visual prototype instead of relying on the first empty paragraph.
    old_profile_caption = paragraph_by_prefix(doc, "图4 ")
    old_profile_image = Paragraph(old_profile_caption._p.getprevious(), old_profile_caption._parent)
    image_prototype = old_profile_image
    caption_prototype = old_profile_caption
    body_prototype = paragraph_by_prefix(doc, "深度—经度空间显示")
    heading2_prototype = paragraph_by_text(doc, "剖面解释")

    heading = paragraph_by_text(doc, "连续垂向结构：不同步来自哪里")
    replace_text(heading, "SODA3剖面复现：2023的异常纵向结构")
    # The edited source has a standalone page-break paragraph immediately
    # before this heading.  Moving the break onto the heading avoids a blank
    # page when the preceding section already fills its last page.
    previous_break = Paragraph(heading._p.getprevious(), heading._parent)
    if previous_break.text:
        raise ValueError("Expected an empty page-break paragraph before the profile section")
    previous_break._p.getparent().remove(previous_break._p)
    heading.paragraph_format.page_break_before = True
    introduction = paragraph_by_prefix(doc, "深度—经度空间显示")
    replace_text(
        introduction,
        "为检验Peng等（2025）提出的西太平洋次表层暖储备及其东传过程，采用SODA3.15.2重建赤道太平洋温度异常剖面。剖面按2°S—2°N余弦加权，异常相对于1982—2022年逐月气候态计算且不去趋势；黑线和灰色虚线分别表示实际与气候态20°C等温线。",
    )
    set_figure(
        old_profile_image,
        FIGURES["three_stage"],
        width_inches=6.40,
        alt_text="SODA3复现Peng二〇二五年三阶段赤道太平洋温度异常剖面",
    )
    replace_text(
        old_profile_caption,
        "图4  SODA3对Peng等（2025）三阶段剖面诊断的复现。左列为2023/24，中央为1982/83、1997/98和2015/16强事件合成，右列为二者之差；时段依次为JFM、4—10月和NDJ。填色为位温异常，虚线框标示FEP（90°W—80°W）。",
    )

    subheading = paragraph_by_text(doc, "剖面解释")
    replace_text(subheading, "三阶段演变")
    subheading.paragraph_format.page_break_before = True
    jfm = paragraph_by_prefix(doc, "1997。")
    replace_text(
        jfm,
        "JFM。2023年西太平洋120°E—180°、100—220 m暖异常显著强于强事件合成，区域平均差约+2.61°C；与此同时，160°W—90°W、50—180 m比合成低约1.08°C。初始热异常因此主要储存在西部次表层，而非已经建立为全盆地一致增暖。",
    )
    aproct = paragraph_by_prefix(doc, "2015。")
    replace_text(
        aproct,
        "4—10月。暖异常向东、向上扩展并到达东太平洋，但中东太平洋50—180 m仍比强事件合成低约1.44°C；同期FEP上层0—100 m略偏暖。该垂向配置表明，2023年的东部响应更偏浅层，而温跃层附近的整层放大弱于典型强事件。",
    )
    ndj = paragraph_by_prefix(doc, "2023。年初暖量")
    replace_text(
        ndj,
        "NDJ。2023/24东太平洋表层保持显著暖异常，但160°W—90°W、50—180 m仍比强事件合成低约1.84°C，FEP上层0—100 m低约3.02°C。成熟期海表增暖并未对应同等强度的东部次表层热异常和温跃层加深。",
    )
    profile_synthesis = clone_after(
        ndj,
        body_prototype,
        "综合来看，2023区别于典型强事件的关键不是缺少暖水，而是暖异常的空间配置不同：西部次表层储热异常强，东部表层响应较早，东部温跃层附近的增暖却持续偏弱。由此，后续机制检验应分别针对西部100—220 m热储备和中东部50—180 m热量收支展开。",
    )

    seasonal_heading = clone_after(profile_synthesis, heading2_prototype, "四季连续演变")
    seasonal_heading.paragraph_format.page_break_before = True
    _, seasonal_caption = add_figure_after(
        seasonal_heading,
        image_prototype,
        caption_prototype,
        FIGURES["four_season"],
        "图5  2023与强事件合成的JFM、AMJ、JAS和OND连续剖面对比。定义同图4。差值列显示，2023的正异常长期偏留在西部次表层，而中东太平洋温跃层附近持续弱于强事件合成。",
        width_inches=6.35,
        alt_text="二〇二三年与强事件合成的四季SODA3赤道太平洋温度异常剖面",
    )
    seasonal_analysis = clone_after(
        seasonal_caption,
        body_prototype,
        "从JFM到OND，2023的西部次表层正差逐季减弱，中东太平洋负差则延续到成熟期。该结果支持“西部初始暖储备异常强、随后发生东传和上翻”的剖面事实，同时说明东部次表层放大效率明显低于1982、1997和2015的平均状态。仅凭温度剖面尚不能区分水平平流、垂向交换和表面热通量的相对贡献。",
    )

    monthly_heading = clone_after(seasonal_analysis, heading2_prototype, "早期逐月垂向补偿")
    monthly_heading.paragraph_format.page_break_before = True
    monthly_intro = clone_after(
        monthly_heading,
        body_prototype,
        "1—6月逐月剖面用于分辨季节平均中被平滑的早期转换。对照组仍为1982、1997和2015发展年的逐月合成。",
    )
    _, monthly_caption = add_figure_after(
        monthly_intro,
        image_prototype,
        caption_prototype,
        FIGURES["monthly"],
        "图6  2023年1—6月SODA3赤道太平洋逐月温度异常剖面、强事件合成及二者之差。定义同图4。",
        width_inches=6.40,
        alt_text="二〇二三年一月至六月SODA3赤道太平洋逐月温度异常剖面对比",
    )
    clone_after(
        monthly_caption,
        body_prototype,
        "在与剖面一致的2°S—2°N、90°W—80°W口径下，2023年2月约5 m温度异常已转为+0.155°C，而0—300 m OHC仍为−0.249 GJ m⁻²；3月二者共同转正。剖面显示，2月的正异常主要局限于近表层，其下仍存在冷异常，因此早期不同步应解释为浅暖层与冷次表层的垂向补偿，而不是整个东太平洋水柱同步转暖。",
    )

    replace_text(
        paragraph_by_contains(doc, "13次暖事件1—6月FEP"),
        "图7  13次暖事件1—6月FEP OISST SSTA、SODA3 0—300 m OHC及首次转正月差。黑框表示SSTA>0且OHC<0；紫色为2023。",
    )
    replace_text(
        paragraph_by_contains(doc, "五次FEP异号事件的1—6月相空间轨迹"),
        "图8  五次FEP异号事件的1—6月相空间轨迹。横轴为0—300 m OHC异常，纵轴为OISST SSTA；数字1—6表示月份。",
    )
    replace_text(
        paragraph_by_contains(doc, "共同正异常持续性及6月联合暖度"),
        "图9  共同正异常持续性及6月联合暖度与成熟期DJF Niño-3.4异常的关系。置换检验10万次；结果为探索性且未作多重比较校正。",
    )

    replace_text(
        paragraph_by_prefix(doc, "当前证据支持把2023描述为"),
        "当前证据表明，2023的发展并非单一的“表层先暖—整层追赶”。盆地尺度上，JFM存在异常强的西太平洋次表层暖储备，但中东太平洋温跃层仍弱于典型强事件；随后暖异常向东、向上扩展，远东表层较早转暖，而东部次表层放大到成熟期仍明显偏弱。1997体现整层预置后表层响应，2015体现次表层支撑下的表面回冷，2023则体现强西部储热、浅层东部响应与偏弱东部次表层放大的组合。",
    )
    replace_text(
        paragraph_by_prefix(doc, "主样本为13次发展年暖事件"),
        "事件普查主样本为13次发展年暖事件：1982、1986、1991、1994、1997、2002、2004、2006、2009、2015、2018、2019、2023；普查主区域FEP=5°S—5°N、90°W—80°W。Peng型剖面对比单独采用2°S—2°N、120°E—80°W，并以1982、1997和2015构成强事件合成。1987/88作为连续事件敏感性样本，不进入主统计。",
    )
    replace_text(
        paragraph_by_prefix(doc, "海洋："),
        "海洋：SODA3.15.2月平均位温。OHC按ρ₀Cp对0—300 m分层积分；Peng型剖面显示位温异常而非OHC，采用2°S—2°N平均、1982—2022年月气候态且不去趋势，并叠加实际与气候态20°C等温线。",
    )
    replace_text(
        paragraph_by_prefix(doc, "13事件统计属于小样本探索"),
        "13事件统计属于小样本探索，相关检验未作多重比较校正；1982与2023的异号OHC都接近零线。三阶段和四季剖面能够约束暖异常的位置、深度与演变，但不能单独确定风应力、水平平流、垂向交换或表面热通量的因果贡献，机制归因仍需热收支和独立资料闭合。",
    )

    empty_headers(doc)
    set_document_font(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
