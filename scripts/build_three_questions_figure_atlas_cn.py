#!/usr/bin/env python3
"""Build a Chinese DOCX atlas for the curated warm-heat-fate figures."""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_warm_heat_fate_complete_report_cn as style  # noqa: E402


PACKAGE = ROOT / "outputs/research_2026/three_questions_figure_package_20260805"
OUTPUT = PACKAGE / "三个核心问题_图件总览_20260805.docx"

MAIN_FIGURES = [
    (
        "main01",
        "问题一｜暖量是否以及如何到达东太平洋",
        "图1  1997/98、2015/16和2023/24赤道太平洋0–300 m OHC经度—时间演变。",
        "三次事件的暖量都向东延伸，但1997的盆地尺度库存和东传最强；2023的东部表面暖异常没有同等强度的深层支撑。",
        6.15,
    ),
    (
        "main02",
        "问题一—二｜CT/CP两类事件的到达与三箱暖量命运",
        "图2  CT与CP事件跨140°W的东向Q20、东部储热建立及成熟—衰减期三箱暖量收支。",
        "CP也能把暖量送入东太平洋；类型差异主要体现在到达强度、持续性和衰减期分配，而不是简单的到达/未到达。",
        6.15,
    ),
    (
        "main03",
        "问题二｜暖量到达后在输出、散热和储存间如何分配",
        "图3  1997、2015、2023的东部暖量输入、次年5°N输出与海表热损失。",
        "1997为强排放、2015为部分排放、2023出口近闭合，说明暖量是否到达并不足以决定事件命运。",
        6.15,
    ),
    (
        "main04",
        "问题二—三｜5°N暖量出口的Ekman—地转竞争",
        "图4  三事件次年1–5月跨5°N总Q20的Ekman、地转和残差分解。",
        "2015和2023的地转阻力接近，而风驱Ekman明显不同；两事件净出口差异主要来自大气能否提供足够强的风驱输送。",
        6.15,
    ),
    (
        "main05",
        "问题三｜FEP先暖为何不代表相同海洋状态",
        "图5  1997、2015、2023的FEP指标、0–300 m OHC和D20演变。",
        "FEP先暖只是表面起点；盆地尺度次表层暖量库存及温跃层向东调整决定表面异常能否获得持续海洋支撑。",
        6.0,
    ),
    (
        "main06",
        "问题三｜FEP之后的大气闭合为何不同",
        "图6  FEP、OHC、D20、u850和OLR/对流响应的三事件过程比较。",
        "1997和2015较早形成深对流—西风响应；2023到5月才出现较弱接通。决定性分叉位于FEP表面增暖之后、深对流东移之前。",
        6.15,
    ),
    (
        "main07",
        "问题三｜1997与2023海气耦合机制综合对照",
        "图7  1997与2023从OHC/D20背景、SST和风场到对流、气压梯度、风应力及暖量输出的全过程对照。",
        "只有盆地尺度次表层支撑和大气对流接通同时成立，东部表面增暖才会转化为足够强的Ekman输出并打开5°N暖量出口。",
        4.8,
    ),
    (
        "main08",
        "问题三｜13事件统计与风应力替换的定量归因",
        "图8  13个暖事件的FEP—OHC—u850—5°N Ekman联系，以及1997、2015、2023局地风应力替换试验。",
        "2023给定FEP暖度后的中太平洋OHC缺口为13事件中最大；OHC—西风与西风—Ekman两段联系稳健。固定2023海洋状态、换入2015风应力后，重构总出口由+0.10增至+0.88 ZJ。",
        6.15,
    ),
]

SUPPLEMENT_ROWS = [
    ["S1", "东部箱发展—衰减", "CT/CP预算反转", "支撑图2"],
    ["S2", "完整预算剩余项", "混合、夹卷及跨等温面等候选过程", "回应预算闭合问题"],
    ["S3", "CT/CP 5°N流速分解", "复合层面的地转—Ekman竞争", "支撑图4"],
    ["S4", "CT/CP逐月耦合时序", "SST、降水、风应力与出口先后次序", "支撑问题三"],
    ["S5", "CT/CP海气过程序列", "降水—风应力—暖量出口链", "支撑S4"],
    ["S6", "ERA5全风应力旋度", "独立大气资料验证", "稳健性"],
    ["S7", "风应力敏感性", "固定温度条件下改变风应力", "不作严格因果表述"],
    ["S8", "混合层预算", "日界线输送与混合层热收支", "剩余项解释"],
    ["S9", "西部暖量箱预算", "160°E–180°暖量释放与脉冲依赖", "上游过程"],
]


style.FIGURES.update(
    {
        "main01": PACKAGE / "01_main_figures/Fig01_eastward_OHC_arrival_hovmoller.png",
        "main02": PACKAGE / "01_main_figures/Fig02_CT_CP_three_box_arrival_and_fate.png",
        "main03": PACKAGE / "01_main_figures/Fig03_event_arrival_export_surface_loss.png",
        "main04": PACKAGE / "01_main_figures/Fig04_event_5N_dynamical_decomposition.png",
        "main05": PACKAGE / "01_main_figures/Fig05_FEP_subsurface_preconditioning.png",
        "main06": PACKAGE / "01_main_figures/Fig06_FEP_atmospheric_closure.png",
        "main07": PACKAGE / "01_main_figures/Fig07_1997_2023_coupled_mechanism_synthesis.png",
        "main08": PACKAGE / "01_main_figures/Fig08_third_question_coupled_attribution.png",
        "main_sheet": PACKAGE / "01_main_figures/_main_contact_sheet.png",
        "supp_sheet": PACKAGE / "02_supplementary_figures/_supplement_contact_sheet.png",
        "explore_sheet": PACKAGE / "03_exploratory_not_for_manuscript/_exploratory_contact_sheet.png",
    }
)


def set_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(86)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(16)
    style.set_run_font(
        kicker.add_run("论文图件包  |  主文、补充与探索图分级"),
        font=style.HEAD_FONT,
        size=11,
        bold=True,
        color=style.GOLD,
    )
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    style.set_run_font(
        title.add_run("“暖量去哪了”三个核心问题图件总览"),
        font=style.HEAD_FONT,
        size=27,
        bold=True,
        color=style.NAVY,
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_after = Pt(68)
    style.set_run_font(
        subtitle.add_run("8张主文图 · 9张补充图 · 9张探索性图"),
        font=style.HEAD_FONT,
        size=14,
        bold=True,
        color=style.DARK_BLUE,
    )
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Pt(0)
    style.set_run_font(date.add_run("2026年8月5日"), size=10.5, color=style.MUTED)
    doc.add_page_break()


def build_atlas() -> Path:
    doc = Document()
    style.configure_document(doc)
    doc.core_properties.title = "暖量去哪了：三个核心问题图件总览"
    doc.core_properties.subject = "ENSO暖量到达、去向和海气耦合命运的论文图件包"
    doc.core_properties.keywords = "ENSO; SODA3; OHC; H20; Q20; Ekman; figures"
    doc.core_properties.comments = "Curated from verified project figures on 2026-08-05."
    set_cover(doc)

    style.add_section_heading(doc, "1 图序总览与取舍原则")
    style.add_para(
        doc,
        "本图件包从当前研究目录约200张PNG中筛出与‘暖量到达—暖量分配—海气耦合命运’直接相关的证据。主文只保留能够独立承担论点的图；补充材料用于稳健性、预算细节和过程解释；阈值、预测技巧和缩减力平衡等仍在发展的诊断单独归入探索区。",
        indent=False,
    )
    style.add_table(
        doc,
        "表1  三个核心问题与主文图对应关系",
        ["科学问题", "主图", "核心证据", "论文作用"],
        [
            ["1 暖量是否到达", "图1—2", "OHC经度—时间演变；140°W Q20与三箱H20", "回答到达和路径"],
            ["2 到达后去了哪里", "图3—4", "5°N输出、海表热损失；Ekman—地转分解", "回答分配和动力出口"],
            ["3 为什么命运不同", "图5—8", "FEP、OHC/D20、u850/OLR、事件统计和风应力替换", "形成并定量检验海气耦合闭环"],
        ],
        [1900, 1300, 3630, 2530],
        font_size=8.4,
    )

    for key, heading, caption, conclusion, width in MAIN_FIGURES:
        style.add_section_heading(doc, heading, page_break=True)
        style.add_callout(
            doc,
            "一句话结论",
            conclusion,
            fill=style.LIGHT_BLUE,
        )
        style.add_figure(doc, key, caption, width=width)

    style.add_section_heading(doc, "2 补充材料图", page_break=True)
    style.add_para(
        doc,
        "补充图不重复主文结论，而是承担预算细节、类型普适性、资料独立性和敏感性说明。它们已经统一编号为Fig. S1—S9。",
        indent=False,
    )
    style.add_figure(
        doc,
        "supp_sheet",
        "补充材料图总览。原分辨率文件均位于02_supplementary_figures目录。",
        width=6.15,
    )
    style.add_table(
        doc,
        "表2  补充图的功能与建议位置",
        ["编号", "主题", "承担的证据", "建议"],
        SUPPLEMENT_ROWS,
        [850, 2150, 4100, 2260],
        font_size=8.0,
    )

    style.add_section_heading(doc, "3 探索性图：保留，但暂不进入当前论文", page_break=True)
    style.add_callout(
        doc,
        "取舍原则",
        "补偿阈值、缩减力平衡、高频风门、暖库释放和OHC几何等图并非无效，而是目前仍依赖阈值、事件内相关或尚未完成的因果验证。把它们从主文移出，可以避免论文从三个核心问题分叉成多个尚未闭合的新故事。",
        fill="FFF4E5",
        color=style.GOLD,
    )
    style.add_figure(
        doc,
        "explore_sheet",
        "探索性图总览：作为后续试验储备，不在当前论文中承担主结论。",
        width=6.15,
    )
    style.add_para(
        doc,
        "所有原图均保留在原始分析目录；本图件包只做复制、统一编号和论文角色分级。CSV清单figure_manifest.csv记录了统一文件名、对应问题、论文角色、核心用途和原始路径。",
        bold=True,
        indent=False,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_atlas())
