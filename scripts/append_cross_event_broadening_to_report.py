#!/usr/bin/env python3
"""Build the ENSO report with native Word equations and no black-square marks."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_20260818.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "数据公式与跨事件展宽_定义严谨修订版_20260818.docx"
)
FIGURE = (
    ROOT
    / "outputs/research_2026/enso_3d_warm_heat_spreading_decoupling/"
    "enso_cross_event_broadening_jja_coupling_en.png"
)

CHINESE_FONT = "Noto Sans CJK SC"
TEXT_COLOR = "111111"
MUTED_COLOR = "4F4F4F"
CALLOUT_FILL = "F1F4F6"
PAGINATION_TAGS = (
    "keepNext",
    "keepLines",
    "pageBreakBefore",
    "widowControl",
    "suppressLineNumbers",
)


def set_font(run, *, size: float, bold: bool = False, color: str = TEXT_COLOR) -> None:
    run.font.name = CHINESE_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), CHINESE_FONT)
    color_element = rpr.find(qn("w:color"))
    if color_element is None:
        color_element = OxmlElement("w:color")
        rpr.append(color_element)
    color_element.set(qn("w:val"), color)


def replace_normal_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = CHINESE_FONT
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), CHINESE_FONT)


def strip_pagination_properties(paragraph) -> None:
    """Remove Word pagination flags shown as black squares with Show/Hide enabled."""
    ppr = paragraph._p.pPr
    if ppr is not None:
        for tag in PAGINATION_TAGS:
            element = ppr.find(qn(f"w:{tag}"))
            if element is not None:
                ppr.remove(element)
        num_pr = ppr.find(qn("w:numPr"))
        if num_pr is not None:
            ppr.remove(num_pr)
    for element in paragraph._p.xpath('.//*[local-name()="lastRenderedPageBreak"]'):
        element.getparent().remove(element)


def strip_document_pagination_properties(document: Document) -> None:
    containers = [document.paragraphs]
    for section in document.sections:
        containers.extend([section.header.paragraphs, section.footer.paragraphs])
    for paragraphs in containers:
        for paragraph in paragraphs:
            strip_pagination_properties(paragraph)


def add_manual_page_break_at_start(paragraph) -> None:
    """Insert a real page break without paragraph-level pagination properties."""
    run = OxmlElement("w:r")
    break_element = OxmlElement("w:br")
    break_element.set(qn("w:type"), "page")
    run.append(break_element)
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, run)


def add_title(document: Document, text: str, *, new_page: bool = False) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.0
    if new_page:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    set_font(paragraph.add_run(text), size=21, bold=True)
    strip_pagination_properties(paragraph)
    return paragraph


def add_subheading(document: Document, text: str) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    set_font(paragraph.add_run(text), size=13, bold=True)
    strip_pagination_properties(paragraph)
    return paragraph


def add_body(
    document: Document,
    text: str,
    *,
    after: float = 5.0,
    size: float = 10.5,
) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    set_font(paragraph.add_run(text), size=size)
    strip_pagination_properties(paragraph)
    return paragraph


def add_labeled_body(
    document: Document,
    label: str,
    text: str,
    *,
    after: float = 4.0,
    size: float = 10.2,
) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    set_font(paragraph.add_run(label), size=size, bold=True)
    set_font(paragraph.add_run(text), size=size)
    strip_pagination_properties(paragraph)
    return paragraph


def add_caption(document: Document, text: str) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.0
    set_font(paragraph.add_run(text), size=8, color=MUTED_COLOR)
    strip_pagination_properties(paragraph)
    return paragraph


def add_callout(document: Document, text: str) -> object:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.15
    ppr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), CALLOUT_FILL)
    ppr.append(shading)
    set_font(paragraph.add_run(text), size=11.5, bold=True, color="24313A")
    strip_pagination_properties(paragraph)
    return paragraph


def set_picture_alt_text(paragraph, title: str, description: str) -> None:
    drawing = paragraph._p.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("title", title)
        drawing[0].set("descr", description)


# ------------------------ Native Word equation helpers ---------------------


def math_run(text: str, *, plain: bool = False) -> OxmlElement:
    run = OxmlElement("m:r")
    if plain:
        run_properties = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        style.set(qn("m:val"), "p")
        run_properties.append(style)
        run.append(run_properties)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def _math_parts(value) -> list[OxmlElement]:
    if isinstance(value, str):
        return [math_run(value)]
    if hasattr(value, "tag"):
        return [value]
    parts = []
    for item in value:
        parts.extend(_math_parts(item))
    return parts


def math_sub(base, subscript) -> OxmlElement:
    element = OxmlElement("m:sSub")
    base_element = OxmlElement("m:e")
    for child in _math_parts(base):
        base_element.append(child)
    sub_element = OxmlElement("m:sub")
    for child in _math_parts(subscript):
        sub_element.append(child)
    element.append(base_element)
    element.append(sub_element)
    return element


def math_fraction(numerator, denominator) -> OxmlElement:
    element = OxmlElement("m:f")
    numerator_element = OxmlElement("m:num")
    for child in _math_parts(numerator):
        numerator_element.append(child)
    denominator_element = OxmlElement("m:den")
    for child in _math_parts(denominator):
        denominator_element.append(child)
    element.append(numerator_element)
    element.append(denominator_element)
    return element


def math_accent(base, character: str = "¯") -> OxmlElement:
    """Return an OMML accent so the overbar remains editable in Word."""
    element = OxmlElement("m:acc")
    properties = OxmlElement("m:accPr")
    accent_character = OxmlElement("m:chr")
    accent_character.set(qn("m:val"), character)
    properties.append(accent_character)
    base_element = OxmlElement("m:e")
    for child in _math_parts(base):
        base_element.append(child)
    element.append(properties)
    element.append(base_element)
    return element


def add_native_equation(document: Document, parts: Iterable, *, after: float = 4.0) -> object:
    """Append an editable OMML equation, not an image or plain-text formula."""
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(after)
    equation = OxmlElement("m:oMath")
    for part in parts:
        if isinstance(part, str):
            equation.append(math_run(part))
        else:
            equation.append(part)
    paragraph._p.append(equation)
    strip_pagination_properties(paragraph)
    return paragraph


def sub(base: str, index: str) -> OxmlElement:
    return math_sub(base, index)


def three_band_mean(index: str | None = None) -> OxmlElement:
    subscript = "3band" if index is None else f"3band,{index}"
    return math_sub(math_accent("H"), subscript)


def build_front_matter(document: Document, original_first_paragraph) -> None:
    """Create two method pages and move them before the existing report."""
    body = document.element.body
    start_index = len(body) - 1

    add_title(document, "本报告使用的数据与基本公式")
    add_body(
        document,
        "本报告比较1997、2015和2023年1—6月海表温度、上层海洋热储备及风场的演变，"
        "并将三次个例中看到的经向—垂向结构扩展到13次ENSO暖事件。所有异常均按同历月计算，"
        "以避免把季节循环误认为事件信号。",
        after=5,
        size=10.3,
    )
    add_subheading(document, "数据资料")
    add_labeled_body(
        document,
        "海表温度与SST指数。",
        "使用NOAA OISST月平均海温制作SSTA图，并计算Niño3、Niño3.4和Niño4指数。"
        "OISST异常采用1982—2024年同历月线性去趋势基准。",
    )
    add_labeled_body(
        document,
        "海洋三维温度与OHC。",
        "使用SODA3.15.2月平均温度，保留原生上300 m深度层。温度、热含量密度和0—300 m OHC"
        "均采用1980—2024年同历月线性去趋势基准。",
    )
    add_labeled_body(
        document,
        "风场与大气耦合量。",
        "空间图叠加ERA5 850 hPa风异常。跨事件耦合响应使用GPCP降水、ERA5海平面气压和ERA5纬向风应力；"
        "历史预期关系以1982—2014年为训练期，并按同历月逐年留一拟合。",
    )
    add_labeled_body(
        document,
        "事件与时间窗。",
        "个例为1997、2015和2023年。跨事件检验包含1982、1986、1991、1994、1997、2002、2004、"
        "2006、2009、2015、2018、2019和2023年；海洋原因窗为JFM，大气结果窗为JJA。",
        after=5,
    )
    add_subheading(document, "同历月线性去趋势异常")
    add_native_equation(
        document,
        [
            sub("X′", "y,m"),
            "=",
            sub("X", "y,m"),
            "−",
            sub("⟨X⟩", "m"),
            "−",
            sub("b", "m"),
            "(y−ȳ)",
        ],
        after=3,
    )
    add_body(
        document,
        "其中y和m分别表示年份和月份，⟨X⟩ₘ为同历月均值，bₘ为同历月线性趋势。"
        "地图中的SSTA、OHC和风异常以及剖面中的温度异常均使用这一形式。",
        after=5,
        size=9.8,
    )
    add_subheading(document, "热含量密度与单位面积OHC")
    add_native_equation(
        document,
        [
            "q′=",
            sub("ρ", "0"),
            sub("c", "p"),
            "T′",
            "    ;    ",
            sub("H′", "B,L,t"),
            "=",
            sub("ρ", "0"),
            sub("c", "p"),
            math_sub(
                ["⟨", math_sub("∫", "L"), sub("T′", "B"), "(λ,z,t)dz", "⟩"],
                "λ",
            ),
        ],
        after=3,
    )
    add_body(
        document,
        "q′为热含量密度异常，单位MJ m⁻³；H′为纬带B、深度层L内的单位面积OHC异常，单位GJ m⁻²。"
        "计算采用ρ₀=1025 kg m⁻³、cₚ=3990 J kg⁻¹ K⁻¹；纬度平均使用余弦纬度权重。",
        after=0,
        size=9.8,
    )

    add_title(document, "跨事件展宽指标与JJA耦合残差", new_page=True)
    add_subheading(document, "区域定义")
    add_body(
        document,
        "地图范围为30°S—30°N、120°E—80°W；FEP定义为5°S—5°N、90°W—80°W。"
        "经度—深度剖面为5°S—5°N余弦加权平均，纬度—深度剖面对固定120°E—80°W完整经度范围平均。"
        "跨事件计算则跟随每次事件的赤道正OHC暖核质心经度λc，取λc±10°窗口，并划分赤道、南侧和北侧三纬带："
        "5°S—5°N、15°S—5°S和5°N—15°N。",
        after=5,
        size=10.0,
    )
    add_subheading(document, "离赤道暖量与三纬带面积加权平均OHC")
    add_native_equation(
        document,
        [sub("H", "off"), "=", math_fraction([sub("H", "S"), "+", sub("H", "N")], "2")],
        after=3,
    )
    add_native_equation(
        document,
        [
            three_band_mean(),
            "=",
            math_fraction(
                [sub("w", "E"), sub("H", "E"), "+", sub("w", "S"), sub("H", "S"), "+", sub("w", "N"), sub("H", "N")],
                [sub("w", "E"), "+", sub("w", "S"), "+", sub("w", "N")],
            ),
            "    ,    ",
            sub("w", "B"),
            "=sin",
            sub("φ", "2,B"),
            "−sin",
            sub("φ", "1,B"),
        ],
        after=3,
    )
    add_native_equation(
        document,
        [
            sub("H", "off,i"),
            "=α+β",
            three_band_mean("i"),
            "+",
            sub("X", "i"),
        ],
        after=4,
    )
    add_body(
        document,
        "横轴使用残差Xᵢ。Xᵢ>0表示在三纬带面积加权平均OHC相同的条件下，相对更多的0—300 m暖量"
        "配置在赤道两侧；Xᵢ<0表示暖量相对集中于赤道带。该指标描述赤道—离赤道之间的相对分配，"
        "不能把离赤道增加与赤道减少分别识别出来。",
        after=5,
        size=9.8,
    )
    add_subheading(document, "JJA型态调整后综合耦合残差")
    add_body(
        document,
        "三个大气分量分别为：5°S—5°N、180°E—250°E的GPCP降水；130°E—160°E减220°E—250°E的"
        "ERA5海平面气压驱动；5°S—5°N、160°E—220°E的ERA5纬向风应力。每个同历月的大气预期为",
        after=3,
        size=9.8,
    )
    add_native_equation(
        document,
        [
            "A=",
            sub("γ", "0"),
            "+",
            sub("γ", "1"),
            "(year−1998)+",
            sub("γ", "2"),
            "Niño3.4+",
            sub("γ", "3"),
            "(Niño3−Niño4)+ε",
        ],
        after=3,
    )
    add_native_equation(
        document,
        [
            sub("R", "A"),
            "=",
            math_fraction("A−Â", sub("s", "ε")),
            "    ,    ",
            sub("C", "JJA,i"),
            "=",
            math_sub(
                [
                    "⟨",
                    math_fraction(
                        [sub("R", "P"), "+", sub("R", "SLP"), "+", sub("R", "τ")],
                        "3",
                    ),
                    "⟩",
                ],
                "JJA",
            ),
        ],
        after=3,
    )
    add_native_equation(
        document,
        [
            sub("C", "JJA,i"),
            "=",
            sub("δ", "0"),
            "+",
            sub("δ", "1"),
            three_band_mean("i"),
            "+",
            sub("Y", "i"),
            "    ,    ",
            sub("r", "partial"),
            "=corr(",
            sub("X", "i"),
            ",",
            sub("Y", "i"),
            ")",
        ],
        after=4,
    )
    add_body(
        document,
        "Â与sε均来自1982—2014年同历月逐年留一训练，因此Rₐ是交叉验证标准化残差。"
        "Yᵢ为进一步控制JFM三纬带面积加权平均OHC后的JJA耦合残差；因此最终相关比较的是相同三纬带平均暖含量条件下，"
        "早期经向暖量配置与随后JJA耦合响应强弱之间的关系。",
        after=0,
        size=9.8,
    )

    # The existing first heading starts on a new page through a manual break,
    # which does not create the black-square pagination marker in Word.
    add_manual_page_break_at_start(original_first_paragraph)

    front_elements = list(body)[start_index:-1]
    for element in front_elements:
        body.remove(element)
    for position, element in enumerate(front_elements):
        body.insert(position, element)


def append_cross_event_pages(document: Document) -> None:
    add_title(document, "从个例到跨事件：暖量展宽是否影响发展期耦合", new_page=True)
    add_body(
        document,
        "上述剖面对2023给出的核心线索，并不只是“暖量偏南”，而是早期暖异常同时分布在赤道带及其两侧，"
        "并更多保存在次表层，使海洋暖量呈现更宽、更深的空间配置。为检验这一结构是否具有跨事件意义，"
        "下面不直接把南侧暖异常解释为经向输送，而是提出一个更保守、可量化的问题：在三纬带面积加权平均OHC相近时，"
        "配置到赤道两侧的暖量越多，随后海洋储热向赤道表层—对流—风反馈的转化是否越弱？",
        after=4,
    )
    add_body(
        document,
        "对13次ENSO暖事件，采用随赤道正OHC暖核质心经度移动的±10°窗口，计算JFM三个纬带的0—300 m OHC："
        "5°S—5°N、5°—15°S和5°—15°N。离赤道暖量定义为南、北两侧OHC的平均，并剔除三纬带面积加权平均OHC的线性影响。"
        "随后将该指标与JJA型态调整后综合海气耦合残差比较；后者控制Niño3.4强度、Niño3−Niño4空间型态和长期线性变化。",
        after=4,
    )

    figure_paragraph = document.add_paragraph(style="Normal")
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.space_before = Pt(2)
    figure_paragraph.paragraph_format.space_after = Pt(0)
    figure_paragraph.add_run().add_picture(str(FIGURE), width=Inches(6.25))
    set_picture_alt_text(
        figure_paragraph,
        "Early meridional OHC spreading and JJA coupling",
        "Scatterplot for 13 ENSO warm events showing a negative relationship between JFM off-equatorial OHC residual and JJA pattern-adjusted coupling residual.",
    )
    strip_pagination_properties(figure_paragraph)
    add_caption(
        document,
        "图8  早期经向暖量展宽与随后JJA海气耦合响应的跨事件关系。横轴正值表示在相同三纬带面积加权平均OHC下，"
        "相对更多的0—300 m暖量配置在5°—15°S和5°—15°N；纵轴负值表示降水、海平面气压和风应力响应弱于相同SST强度及EP—CP型态下的历史预期。",
    )
    add_body(
        document,
        "在13次暖事件中，JFM离赤道暖量越多，随后JJA型态调整后的海气耦合响应越弱（偏相关r=−0.81，"
        "置换检验p=0.001，FDR q=0.003）。逐事件剔除后，相关仍保持在−0.84至−0.75之间，说明该关系并非由"
        "1997、2015或2023中的某一个事件单独决定。结果表明，除三纬带平均暖含量之外，暖量是否集中在赤道附近"
        "与随后海气耦合响应的相对强弱有关，并与耦合转化效率降低的解释一致。",
        after=0,
    )

    add_title(
        document,
        "2023的定位：发展期暖量更宽、更深，赤道表层可利用比例偏低",
        new_page=True,
    )
    add_subheading(document, "三点结论")
    add_labeled_body(
        document,
        "结论一。",
        "2023年并不缺少海洋暖量，但发展早期暖量主要位于西部、次表层及赤道南侧，与赤道东太平洋表层增暖未能及时重合，"
        "表现为发展期海气耦合建立较慢，而非成熟期持续脱耦。",
        after=6,
    )
    add_labeled_body(
        document,
        "结论二。",
        "与1997和2015年相比，2023年暖异常分布更宽、更深。暖量向赤道两侧和次表层展开后，"
        "赤道表层可直接参与对流—风反馈的暖量比例可能降低。",
        after=6,
    )
    add_labeled_body(
        document,
        "结论三。",
        "13次暖事件中，在控制三纬带面积加权平均OHC后，JFM离赤道暖量越多，随后JJA海气耦合响应越弱"
        "（r=−0.81，p=0.001）。这一结果支持暖量空间集中度影响发展期耦合响应的解释，但尚不能证明经向扩散是直接原因。",
        after=0,
    )


def audit_document(document: Document) -> None:
    prohibited = ("•", "●", "▪", "◦", "‣")
    equation_count = len(document.element.xpath(".//m:oMath"))
    if equation_count < 8:
        raise ValueError(f"Expected at least 8 native Word equations, found {equation_count}")
    for index, paragraph in enumerate(document.paragraphs):
        if any(marker in paragraph.text for marker in prohibited):
            raise ValueError(f"Prohibited bullet marker in paragraph {index}")
        ppr = paragraph._p.pPr
        if ppr is None:
            continue
        if ppr.find(qn("w:numPr")) is not None:
            raise ValueError(f"Unexpected list numbering in paragraph {index}")
        for tag in PAGINATION_TAGS:
            if ppr.find(qn(f"w:{tag}")) is not None:
                raise ValueError(f"Pagination marker property {tag} in paragraph {index}")


def main() -> None:
    for path in (SOURCE, FIGURE):
        if not path.exists():
            raise FileNotFoundError(path)

    document = Document(SOURCE)
    original_first_paragraph = document.paragraphs[0]
    replace_normal_font(document)
    append_cross_event_pages(document)
    build_front_matter(document, original_first_paragraph)
    strip_document_pagination_properties(document)
    audit_document(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
    print(f"paragraphs: {len(document.paragraphs)}")
    print(f"inline shapes: {len(document.inline_shapes)}")
    print(f"native equations: {len(document.element.xpath('.//m:oMath'))}")


if __name__ == "__main__":
    main()
