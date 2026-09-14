#!/usr/bin/env python3
"""Build the Chinese DOCX methods report for the 2026 recharge analysis."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs/research_2026/enso_recharge_event_matching_2026"
OUTPUT = OUT_DIR / "2026_ENSO_recharge_vertical_profile_flow_methods_report_CN.docx"

FIGURES = {
    "group_profiles": ROOT
    / "outputs/research_2026/enso_vertical_flow_official_roni_2026/official_roni_vertical_profiles_2026.png",
    "matching": OUT_DIR / "2026_recharge_event_matching_matrix.png",
    "analog_profiles": OUT_DIR / "2026_closest_analogue_vertical_profiles.png",
    "flow_storage": OUT_DIR / "2026_recharge_analogue_flow_and_storage.png",
}

MATH_XSL = Path("/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl")
MATH_NS = "http://www.w3.org/1998/Math/MathML"

FONT_LATIN = "Times New Roman"
FONT_CJK = "Arial Unicode MS"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = "F2F2F2"


def set_run_font(
    run,
    *,
    size: float = 11,
    italic: bool = False,
    color: RGBColor = BLACK,
    bold: bool = False,
) -> None:
    contains_cjk = any("\u3400" <= character <= "\u9fff" for character in run.text)
    primary_font = FONT_CJK if contains_cjk else FONT_LATIN
    run.font.name = primary_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), primary_font)
    rfonts.set(qn("w:hAnsi"), primary_font)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rfonts.set(qn("w:cs"), primary_font)
    if contains_cjk:
        rfonts.set(qn("w:hint"), "eastAsia")


def set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = 6,
    line: float = 1.3,
    keep_next: bool = False,
    keep_together: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_together


def add_text(
    doc: Document,
    text: str,
    *,
    size: float = 11,
    before: float = 0,
    after: float = 6,
    line: float = 1.3,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    italic: bool = False,
    color: RGBColor = BLACK,
    keep_next: bool = False,
) -> object:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    set_paragraph_spacing(
        paragraph,
        before=before,
        after=after,
        line=line,
        keep_next=keep_next,
    )
    run = paragraph.add_run(text)
    set_run_font(run, size=size, italic=italic, color=color)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1) -> object:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tokens = {
        1: (16, 16, 8),
        2: (13.5, 12, 6),
        3: (12, 9, 4),
    }
    size, before, after = tokens[level]
    set_paragraph_spacing(
        paragraph,
        before=before,
        after=after,
        line=1.15,
        keep_next=True,
    )
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=False)
    return paragraph


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=9, color=GRAY)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3

    for level, size in ((1, 16), (2, 13.5), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("2026年赤道太平洋充电过程：方法与逐事件比较")
    set_run_font(r, size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_title_page(doc: Document) -> None:
    add_text(doc, "海洋与气候诊断报告", size=11, after=28, color=GRAY)
    add_text(
        doc,
        "2026年赤道太平洋充电过程",
        size=22,
        after=5,
        line=1.1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_text(
        doc,
        "Vertical profile、暖水储量与断面输送的逐事件比较",
        size=15,
        after=24,
        line=1.15,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_text(
        doc,
        "本报告系统说明分析所包含的理论模块、权威方法来源、针对2026问题所作的适配，以及由这些模块共同支持的阶段性结论。",
        size=11.5,
        after=54,
        line=1.4,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_text(doc, "资料更新至：2026年6月", size=10.5, after=3, color=GRAY)
    add_text(doc, "报告生成日期：2026年8月", size=10.5, after=0, color=GRAY)
    doc.add_page_break()


def mathml(math_body: str) -> str:
    return f'<math xmlns="{MATH_NS}" display="block">{math_body}</math>'


def add_equation(doc: Document, body: str, *, after: float = 7) -> object:
    if not MATH_XSL.exists():
        raise FileNotFoundError(MATH_XSL)
    transform = etree.XSLT(etree.parse(str(MATH_XSL)))
    root = etree.fromstring(mathml(body).encode("utf-8"))
    omml = transform(root).getroot()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=2, after=after, line=1.1, keep_together=True)
    paragraph._p.append(omml)
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 110) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must sum to 9360, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 9.5,
    first_col_left: bool = True,
) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_properties.append(header_flag)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.15)
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=False)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            p = cells[index].paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if first_col_left and index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            set_paragraph_spacing(p, after=0, line=1.15)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=False)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_table_caption(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, before=4, after=4, line=1.15, keep_next=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5)
    return paragraph


def set_image_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description[:80])


def add_figure(
    doc: Document,
    path: Path,
    caption: str,
    *,
    width: float = 6.55,
    page_break_before: bool = False,
) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if page_break_before:
        doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=4, after=3, line=1.0, keep_next=True)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    set_image_alt_text(shape, caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(cap, before=0, after=8, line=1.15, keep_together=True)
    r = cap.add_run(caption)
    set_run_font(r, size=9.2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    rpr.append(size)
    new_run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_reference(doc: Document, text: str, url: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.3)
    fmt.first_line_indent = Inches(-0.3)
    fmt.space_after = Pt(5)
    fmt.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=9.8)
    if url:
        paragraph.add_run(" ")
        add_hyperlink(paragraph, url, url)


def build_report() -> Document:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "摘要", 1)
    add_text(
        doc,
        "本分析的目标不是仅凭温度剖面推断海流，而是先利用经度—深度温度异常识别2026年暖核的位置、厚度与东传—上倾结构，再由20°C等温面计算暖水储量，并加入纬向流速计算真实断面输送。历史事件依据NOAA CPC Relative Oceanic Niño Index（RONI）划分为中等、强和极强事件。历史场使用SODA3.15.2，2026实时场使用GODAS。最终通过vertical profile、三分区暖水体积V20和三个断面的暖水体积输送M20进行逐事件匹配。",
    )
    add_text(
        doc,
        "结果表明，2026年JFM仍属于混合型早期充电状态：垂向结构最接近2009，暖水储量配置更接近2023/1982；到AMJ，1997在垂向型态、暖水储量演变和断面输送三方面同时成为最稳定类比。该结论说明2026发展路径在春季转入类似1997的快速东向再分配，但不表示其最终成熟强度必然达到1997水平。",
    )

    add_heading(doc, "1. 科学问题与理论框架", 1)
    add_text(
        doc,
        "Jin（1997）的ENSO充电振子理论将赤道太平洋平均温跃层深度或暖含量视为ENSO的海洋记忆。暖事件发生前，赤道暖水库充电；发展和成熟过程中，风应力驱动的海洋调整改变温跃层倾斜，并使暖水向东或向赤道外重新分配；随后赤道暖水库放电。Li et al.（2022）进一步强调，不同经度的温跃层异常对充电与放电阶段具有不同作用，因此仅分析盆地总暖量会损失关键的空间信息。",
    )
    add_text(
        doc,
        "基于这一理论，本报告同时保留三个层次的信息：暖异常的经度—深度型态、暖水体积在西部—中部—东部之间的配置，以及暖水穿过固定经度断面的真实输送。三者分别回答“暖水在哪里”“区域储量如何变化”和“暖水以多大流量向东移动”。",
    )

    add_heading(doc, "2. 分析模块及方法来源", 1)
    add_table_caption(doc, "Table 1. 本次分析所包含的九个模块及其方法来源。")
    module_rows = [
        ("1 充电—放电理论", "赤道暖含量、温跃层调整与ENSO相位转换", "Jin（1997）；Li et al.（2022）", "直接采用理论框架"),
        ("2 事件强度分类", "中等、强、极强暖事件", "NOAA CPC RONI", "采用官方阈值"),
        ("3 资料与异常", "SODA3历史场、GODAS 2026实时场", "Carton et al.（2018）；NOAA GODAS", "同历月异常与去趋势"),
        ("4 Vertical profile", "5°S—5°N经度—深度温度异常", "温跃层诊断；Taylor（2001）", "逐事件型态相关"),
        ("5 D20与V20", "20°C等温面和暖水体积", "Meinen and McPhaden（2000）", "增加三个经度分区"),
        ("6 储量变化率", "dV20/dt与区域暖水收支", "Meinen and McPhaden（2001）", "月际有限差分"),
        ("7 M20输送", "20°C以上水体的真实断面流量", "Meinen and McPhaden（2001）", "增加三个内部断面"),
        ("8 Q20输送", "相对20°C的暖水热输送", "标准海洋平流热通量", "仅作敏感性检验"),
        ("9 事件匹配", "相关、标准化距离与等权排序", "Taylor型态统计；本研究构建", "不是既有ENSO指数"),
    ]
    add_table(
        doc,
        ["模块", "核心内容", "权威来源", "本次处理"],
        module_rows,
        [1800, 3000, 2200, 2360],
        font_size=8.8,
    )

    add_heading(doc, "3. 事件分类与资料体系", 1)
    add_heading(doc, "3.1 NOAA CPC RONI强度分类", 2)
    add_text(
        doc,
        "RONI以ERSST.v5 Niño3.4三个月滑动平均为基础，减去20°S—20°N热带平均SST异常，并调整方差使其与原Niño3.4指数尺度可比。NOAA CPC采用0.5°C递增阈值划分强度：1.0—1.5°C为中等，1.5—2.0°C为强，达到或超过2.0°C为极强。该分类提供事件强度标签，但不参与2026相似度计算。",
    )
    add_table_caption(doc, "Table 2. 本次逐事件比较采用的官方RONI暖事件分组。")
    add_table(
        doc,
        ["类别", "峰值RONI范围", "事件年份"],
        [
            ("中等", "1.0—1.5°C", "1994、2002、2023"),
            ("强", "1.5—2.0°C", "1986、2009"),
            ("极强", "≥2.0°C", "1982、1991、1997、2015"),
        ],
        [1700, 2500, 5160],
        font_size=9.5,
    )
    add_text(
        doc,
        "2026在分析中作为目标发展年，不预先赋予中等、强或极强标签。这样可以避免先验强度分类决定海洋过程结论。",
    )

    add_heading(doc, "3.2 SODA3历史场与GODAS 2026实时场", 2)
    add_text(
        doc,
        "历史事件使用SODA3.15.2温度和纬向流速。SODA3基于MOM5海洋模式，并利用温盐剖面、卫星SST等资料进行顺序资料同化（Carton et al., 2018）。历史异常使用1982—2024同历月线性去趋势，即每个网格上的1月只与其他年份1月比较。",
    )
    add_equation(
        doc,
        "<mrow><msubsup><mi>T</mi><mrow><mi>y</mi><mo>,</mo><mi>m</mi></mrow><mo>′</mo></msubsup><mo>=</mo><msub><mi>T</mi><mrow><mi>y</mi><mo>,</mo><mi>m</mi></mrow></msub><mo>−</mo><mrow><mo>(</mo><msub><mi>a</mi><mi>m</mi></msub><mo>+</mo><msub><mi>b</mi><mi>m</mi></msub><mrow><mo>(</mo><mi>y</mi><mo>−</mo><mover><mi>y</mi><mo>¯</mo></mover><mo>)</mo></mrow><mo>)</mo></mrow></mrow>",
    )
    add_text(
        doc,
        "式中，T为月平均温度，y为年份，m为同历月，aₘ和bₘ分别为1982—2024同历月回归截距和趋势，ȳ为基准期平均年份。",
        size=10.2,
    )
    add_text(
        doc,
        "2026使用NOAA/NCEP GODAS potential temperature和zonal current，异常相对于GODAS 1991—2020同历月气候态。由于SODA与GODAS属于不同资料系统，后续流量和储量比较增加跨产品校准；vertical profile则主要依靠型态相关，而不是直接比较绝对幅度。",
    )

    add_heading(doc, "4. Vertical profile模块", 1)
    add_heading(doc, "4.1 纬度平均与垂向结构", 2)
    add_text(
        doc,
        "对5°S—5°N、120°E—80°W、0—300 m的温度异常进行余弦纬度加权平均，保留月份、经度和深度三个维度。该处理能够识别暖核位于西部还是东部、暖异常处于表层还是次表层、暖层是否向东上倾，以及东太平洋冷异常是否消失。",
    )
    add_equation(
        doc,
        "<mrow><msup><mi>T</mi><mo>′</mo></msup><mrow><mo>(</mo><mi>λ</mi><mo>,</mo><mi>z</mi><mo>,</mo><mi>m</mi><mo>)</mo></mrow><mo>=</mo><mfrac><mrow><msubsup><mi>∫</mi><mrow><mn>5</mn><mo>°</mo><mi>S</mi></mrow><mrow><mn>5</mn><mo>°</mo><mi>N</mi></mrow></msubsup><msup><mi>T</mi><mo>′</mo></msup><mrow><mo>(</mo><mi>φ</mi><mo>,</mo><mi>λ</mi><mo>,</mo><mi>z</mi><mo>,</mo><mi>m</mi><mo>)</mo></mrow><mi>cos</mi><mi>φ</mi><mi>d</mi><mi>φ</mi></mrow><mrow><msubsup><mi>∫</mi><mrow><mn>5</mn><mo>°</mo><mi>S</mi></mrow><mrow><mn>5</mn><mo>°</mo><mi>N</mi></mrow></msubsup><mi>cos</mi><mi>φ</mi><mi>d</mi><mi>φ</mi></mrow></mfrac></mrow>",
    )
    add_text(
        doc,
        "其中，φ为纬度，λ为经度，z为深度，m为月份。余弦权重用于修正经纬网格在不同纬度上的面积差异。",
        size=10.2,
    )

    add_heading(doc, "4.2 型态统计", 2)
    add_text(
        doc,
        "2026与每个历史事件的完整month—longitude—depth场计算Pearson型态相关，同时报告RMSE和RMS幅度比。Taylor（2001）指出，相关、均方根差和方差幅度能够从互补角度描述复杂气候场的型态一致性。本分析主排序使用相关系数，RMSE和幅度比用于解释，从而降低SODA—GODAS产品幅度差异对剖面排序的影响。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>r</mi><mrow><mi>T</mi><mo>,</mo><mi>i</mi></mrow></msub><mo>=</mo><mi>corr</mi><mrow><mo>[</mo><msubsup><mi>T</mi><mn>2026</mn><mo>′</mo></msubsup><mrow><mo>(</mo><mi>λ</mi><mo>,</mo><mi>z</mi><mo>,</mo><mi>m</mi><mo>)</mo></mrow><mo>,</mo><msubsup><mi>T</mi><mi>i</mi><mo>′</mo></msubsup><mrow><mo>(</mo><mi>λ</mi><mo>,</mo><mi>z</mi><mo>,</mo><mi>m</mi><mo>)</mo></mrow><mo>]</mo></mrow></mrow>",
    )

    add_figure(
        doc,
        FIGURES["group_profiles"],
        "Fig. 1. 2026年JFM和AMJ经度—深度温度异常与官方RONI中等、强和极强事件合成的比较。历史场为SODA3 1982—2024同历月去趋势异常，2026为GODAS相对于1991—2020同历月气候态的异常。三条虚线表示180°、160°W和140°W断面。分组比较显示，2026完整1—6月型态与极强组合成相关最高，但分组平均不能确定具体历史类比，因此后续进一步开展逐事件比较。",
        width=6.7,
        page_break_before=False,
    )

    add_heading(doc, "5. D20、暖水体积与储量变化率", 1)
    add_heading(doc, "5.1 20°C等温面深度", 2)
    add_text(
        doc,
        "Meinen and McPhaden（2000）将5°S—5°N、120°E—80°W中温度不低于20°C的水体定义为赤道太平洋暖水体积。20°C等温面通常位于主温跃层中部，因此D20可以作为温跃层深度代理。在相邻温度层之间采用线性插值。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>D</mi><mn>20</mn></msub><mo>=</mo><msub><mi>z</mi><mn>1</mn></msub><mo>+</mo><mfrac><mrow><mn>20</mn><mo>−</mo><msub><mi>T</mi><mn>1</mn></msub></mrow><mrow><msub><mi>T</mi><mn>2</mn></msub><mo>−</mo><msub><mi>T</mi><mn>1</mn></msub></mrow></mfrac><mrow><mo>(</mo><msub><mi>z</mi><mn>2</mn></msub><mo>−</mo><msub><mi>z</mi><mn>1</mn></msub><mo>)</mo></mrow></mrow>",
    )
    add_text(
        doc,
        "式中，z₁和z₂为包围20°C等温面的相邻深度，T₁和T₂为对应温度。",
        size=10.2,
    )

    add_heading(doc, "5.2 暖水体积与三个经度分区", 2)
    add_equation(
        doc,
        "<mrow><msub><mi>V</mi><mn>20</mn></msub><mo>=</mo><msub><mi>∬</mi><mi>A</mi></msub><msub><mi>D</mi><mn>20</mn></msub><mrow><mo>(</mo><mi>λ</mi><mo>,</mo><mi>φ</mi><mo>)</mo></mrow><mi>d</mi><mi>A</mi></mrow>",
    )
    add_text(
        doc,
        "球面面积权重按R²[sin(φ₂)−sin(φ₁)]Δλ计算。为了区分暖水停留在西部、进入中部或到达东部，将盆地划为120°E—180°、180°—140°W和140°W—80°W三个箱体。该具体分区是针对2026暖水东传建立的诊断；其物理思路受Li et al.（2022）关于西部、中部和东部温跃层作用差异的启发。",
    )

    add_heading(doc, "5.3 储量变化率及其物理边界", 2)
    add_equation(
        doc,
        "<mrow><mfrac><mrow><mi>d</mi><msub><mi>V</mi><mn>20</mn></msub></mrow><mrow><mi>d</mi><mi>t</mi></mrow></mfrac><mo>≈</mo><mfrac><mrow><msub><mi>V</mi><mn>20</mn></msub><mrow><mo>(</mo><msub><mi>t</mi><mn>2</mn></msub><mo>)</mo></mrow><mo>−</mo><msub><mi>V</mi><mn>20</mn></msub><mrow><mo>(</mo><msub><mi>t</mi><mn>1</mn></msub><mo>)</mo></mrow></mrow><mrow><msub><mi>t</mi><mn>2</mn></msub><mo>−</mo><msub><mi>t</mi><mn>1</mn></msub></mrow></mfrac></mrow>",
    )
    add_text(
        doc,
        "dV20/dt可换算为Sv，但它表示区域内部储量增长或减少，不是某一边界上的真实流量。完整暖水体积收支还包含东、西、南、北边界的水平输送，以及跨20°C等温面的混合、上升流和水团转化（Meinen and McPhaden, 2001）。",
    )

    add_heading(doc, "6. 暖水体积输送与热输送", 1)
    add_heading(doc, "6.1 真实断面暖水体积输送M20", 2)
    add_text(
        doc,
        "Vertical profile能够确定断面上哪些网格属于20°C暖水层，但不能单独给出真实海流。必须加入纬向流速u，才能计算穿过固定经度断面的暖水体积输送。该定义是Meinen and McPhaden（2001）暖水体积收支方法的直接延伸。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>M</mi><mn>20</mn></msub><mo>=</mo><mi>∬</mi><mi>u</mi><mi>I</mi><mrow><mo>(</mo><mi>T</mi><mo>≥</mo><mn>20</mn><mo>°</mo><mi>C</mi><mo>)</mo></mrow><mi>d</mi><mi>z</mi><mi>d</mi><mi>y</mi></mrow>",
    )
    add_text(
        doc,
        "I为示性函数；温度不低于20°C时I=1，否则I=0。积分范围为5°S—5°N、0—300 m，固定断面为180°、160°W和140°W。正值表示向东，单位为Sv。三个内部断面分别判断暖水是否越过日界线、进入中太平洋并最终到达东部储量箱。",
    )

    add_heading(doc, "6.2 相对20°C的暖水热输送Q20", 2)
    add_text(
        doc,
        "标准海洋平流热输送为ρcp∫u(T−Tr)dA（Zheng and Giese, 2009）。本分析取参考温度Tr=20°C，并只计算20°C以上的超额热量。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>Q</mi><mn>20</mn></msub><mo>=</mo><msub><mi>ρ</mi><mn>0</mn></msub><msub><mi>c</mi><mi>p</mi></msub><mi>∬</mi><mi>u</mi><mi>max</mi><mrow><mo>(</mo><mi>T</mi><mo>−</mo><mn>20</mn><mo>°</mo><mi>C</mi><mo>,</mo><mn>0</mn><mo>)</mo></mrow><mi>d</mi><mi>z</mi><mi>d</mi><mi>y</mi></mrow>",
    )
    add_text(
        doc,
        "计算采用ρ₀=1025 kg m⁻³、cp=3990 J kg⁻¹ K⁻¹，结果单位为PW。Q20不是全水深绝对海洋热输送，而是20°C以上暖水的参考温度热输送。由于它同时包含M20和温度幅度信息，本分析仅将其作为敏感性检验，不进入主排序。",
    )

    add_heading(doc, "7. 跨产品校准与验证", 1)
    add_text(
        doc,
        "历史事件来自SODA，2026来自GODAS。为减少产品系统性幅度差异，利用1997、2002、2015和2023四个重叠事件建立SODA对GODAS的线性映射。断面输送验证包含4个事件×6个月×3个断面，共72个样本。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>X</mi><mi>SODA</mi></msub><mo>=</mo><mi>α</mi><mo>+</mo><mi>β</mi><msub><mi>X</mi><mi>GODAS</mi></msub></mrow>",
    )
    add_text(
        doc,
        "验证中，M20跨产品相关为0.781，Q20为0.797；盆地V20异常相关为0.944。由此判断，暖水储量的跨产品比较最稳健，流量的符号、时间演变和相对经度结构较可靠，但流量绝对幅度仍具有较大产品不确定性。Vertical profile未进行逐格点幅度回归，因此使用型态相关作为主判据。",
    )

    add_heading(doc, "8. 逐事件匹配模块", 1)
    add_heading(doc, "8.1 三个等权分量", 2)
    add_text(
        doc,
        "主匹配由三个等权分量构成。第一分量是month—longitude—depth温度异常型态相关；第二分量是西部、中部和东部V20逐月演变；第三分量是180°、160°W和140°W断面M20逐月演变。V20和M20同时评价型态相关和幅度标准化距离。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>d</mi><mi>i</mi></msub><mo>=</mo><msqrt><mfrac><mn>1</mn><mi>K</mi></mfrac><munderover><mi>∑</mi><mrow><mi>k</mi><mo>=</mo><mn>1</mn></mrow><mi>K</mi></munderover><msup><mrow><mo>[</mo><mfrac><mrow><msub><mi>X</mi><mrow><mn>2026</mn><mo>,</mo><mi>k</mi></mrow></msub><mo>−</mo><msub><mi>X</mi><mrow><mi>i</mi><mo>,</mo><mi>k</mi></mrow></msub></mrow><msub><mi>s</mi><mi>k</mi></msub></mfrac><mo>]</mo></mrow><mn>2</mn></msup></msqrt></mrow>",
    )
    add_text(
        doc,
        "式中，i为历史事件，k为月份—区域或月份—断面组合，sₖ为九个历史事件在该特征上的标准差。d越小表示幅度越接近；相关系数越高表示逐月、逐区域演变越一致。",
    )

    add_heading(doc, "8.2 综合排序的解释", 2)
    add_text(
        doc,
        "对于V20和M20，先平均相关排名与标准化距离排名，再与vertical-profile排名等权平均。最终图中的combined similarity只是把九个候选事件的平均名次线性转换到0—1，不能解释成概率或物理相似百分比。该综合排序是本研究为2026跨产品逐事件比较建立的透明诊断，不是Li et al.（2022）或其他文献中的现成ENSO指数。",
    )
    add_equation(
        doc,
        "<mrow><msub><mi>S</mi><mi>i</mi></msub><mo>=</mo><mn>1</mn><mo>−</mo><mfrac><mrow><msub><mover><mi>R</mi><mo>¯</mo></mover><mi>i</mi></msub><mo>−</mo><mn>1</mn></mrow><mrow><mi>N</mi><mo>−</mo><mn>1</mn></mrow></mfrac></mrow>",
    )

    add_figure(
        doc,
        FIGURES["matching"],
        "Fig. 2. 2026与九个历史暖事件的阶段性匹配矩阵。Vertical-profile列给出温度异常型态相关；V20和M20列同时给出型态相关r与标准化距离d；颜色表示每个阶段内部的相对排名，combined similarity不是概率。JFM没有唯一类比，而AMJ的1997在三个主模块中同时排名第一。",
        width=6.65,
        page_break_before=False,
    )

    add_heading(doc, "9. 结果：2026充电过程与哪些事件相似", 1)
    add_heading(doc, "9.1 JFM：混合型早期充电", 2)
    add_text(
        doc,
        "JFM的不同量指向不同事件。经度—深度温度异常最接近2009，型态相关为0.735；三分区V20配置最接近2023，型态相关为0.816、标准化距离为0.640；1982在剖面、储量和流量三个方面表现最均衡。综合排名前三为2023、1982和2009。",
    )
    add_text(
        doc,
        "因此，JFM不适合写成“2026早期就是某一个经典事件”。更准确的表述是：2026具有类似2009的西—中太平洋深层暖水结构，同时其暖水储量配置更接近2023/1982，早期仍处于多种结构特征并存的混合型充电阶段。",
    )

    add_heading(doc, "9.2 AMJ：转入1997型东向再分配", 2)
    add_text(
        doc,
        "到AMJ，1997成为最稳定类比：vertical profile相关为0.824，V20演变相关为0.953，M20演变相关为0.904，Q20敏感性相关为0.786。2026与1997共同表现为暖异常向东、向上发展，西部和中部储量减少，东部储量增加，同时断面暖水输送由西向东增强。",
    )

    add_figure(
        doc,
        FIGURES["analog_profiles"],
        "Fig. 3. 2026阶段性垂向结构类比。上排比较2026 JFM与JFM垂向型态最接近的2009；下排比较2026 AMJ与综合过程最接近的1997；右列为2026减去对应历史事件。三条虚线为180°、160°W和140°W。JFM相似性主要来自西—中太平洋次表层暖水和东部冷异常结构；AMJ相似性来自暖层快速向东上倾并到达东太平洋。",
        width=6.65,
        page_break_before=False,
    )

    add_heading(doc, "9.3 断面输送和储量变化的量化", 2)
    add_table_caption(doc, "Table 3. AMJ三个断面的平均M20异常（Sv）。")
    add_table(
        doc,
        ["事件", "180°", "160°W", "140°W"],
        [
            ("2026", "+10.6", "+17.6", "+25.2"),
            ("1997", "+8.4", "+25.2", "+36.1"),
            ("2023", "+10.0", "+14.6", "+20.3"),
        ],
        [1800, 2520, 2520, 2520],
        font_size=9.5,
    )
    add_text(
        doc,
        "2026和1997都表现为越向东输送越强，但2026中、东太平洋的输送幅度低于1997，更接近“1997路径的较弱版本”。2023也具有向东增强的梯度，但整体幅度较弱。",
    )
    add_table_caption(doc, "Table 4. 4—6月三分区V20异常变化（10¹³ m³）。")
    add_table(
        doc,
        ["事件", "西部", "中部", "东部"],
        [
            ("2026", "−5.9", "−2.9", "+8.4"),
            ("1997", "−5.2", "−8.2", "+6.3"),
            ("2023", "−5.1", "−2.2", "+2.2"),
        ],
        [1800, 2520, 2520, 2520],
        font_size=9.5,
    )
    add_text(
        doc,
        "2026西部暖水消耗接近1997和2023，中部消耗弱于1997，但东部暖水增长显著大于2023。这说明2026在AMJ发生了强烈的东向再分配，而不仅是整个赤道盆地同步充电。",
    )

    add_figure(
        doc,
        FIGURES["flow_storage"],
        "Fig. 4. 2026与主要历史类比事件的暖水体积输送和分区储量演变。上排为180°、160°W和140°W断面M20异常；下排为西部、中部和东部V20异常。2026黑线为根据重叠事件回归换算后的SODA等效幅度。2026在4月前具有很强的向东输送，随后流量减弱；与此同时西部和中部储量下降、东部储量继续增加，表明发展过程从盆地充电转入快速东向再分配。",
        width=6.65,
        page_break_before=False,
    )

    add_heading(doc, "10. 方法适用边界", 1)
    add_heading(doc, "10.1 Vertical profile不能单独给出真实流量", 2)
    add_text(
        doc,
        "温度剖面只能确定20°C暖水层的位置和厚度，并由此计算D20、V20和dV20/dt。真实断面输送必须加入流速u。当前M20代表5°S—5°N、0—300 m内部经度断面的净东向暖水输送，不能替代包含南、北、东、西边界和跨等温面水团转化的完整暖水收支。",
    )
    add_heading(doc, "10.2 跨产品幅度仍有不确定性", 2)
    add_text(
        doc,
        "SODA与GODAS的V20一致性较高，但断面流量标准差不同。回归校准降低了平均幅度偏差，却不能完全消除流速结构、垂向网格、资料同化和月平均uT协方差的差异。因此，事件排序应优先解释为型态、方向和阶段演变的一致性。",
    )
    add_heading(doc, "10.3 相似路径不等于相同最终强度", 2)
    add_text(
        doc,
        "当前资料只覆盖2026年1—6月。AMJ与1997相似表示发展期海洋暖水东向再分配路径相似，并不保证2026成熟期SST、海气耦合或遥相关达到1997水平。尤其是2026在4月达到断面流量峰值后已经减弱，后续强度仍取决于风应力、对流反馈和跨赤道外暖水交换。",
    )

    add_heading(doc, "11. 结论", 1)
    add_text(
        doc,
        "第一，整套分析以Jin充电振子、Meinen and McPhaden暖水体积与输送定义，以及Li et al.温跃层空间非对称思想为物理基础；RONI分组来自NOAA CPC，资料来自SODA3和GODAS。",
    )
    add_text(
        doc,
        "第二，JFM的2026属于混合型早期充电状态：垂向结构偏2009，储量配置偏2023/1982，尚不存在唯一历史类比。",
    )
    add_text(
        doc,
        "第三，到AMJ，1997在vertical profile、V20和M20三方面同时成为最接近事件。2026表现为西部和中部暖水减少、东部增加，并伴随由西向东增强的暖水输送，因此其发展过程更接近1997式快速东向再分配，但输送幅度整体弱于1997。",
    )
    add_text(
        doc,
        "第四，跨资料回归和综合排名是本研究为2026问题新建的诊断，而不是任何单篇文章的现成方法。其作用是透明整合不同物理量，不应解释为概率或最终ENSO强度预测。",
    )

    doc.add_page_break()
    add_heading(doc, "参考文献", 1)
    add_reference(
        doc,
        "Behringer, D. W., and Y. Xue, 2004: Evaluation of the Global Ocean Data Assimilation System at NCEP: The Pacific Ocean. Eighth Symposium on Integrated Observing and Assimilation Systems for Atmosphere, Oceans, and Land Surface, AMS.",
        "https://www.cpc.ncep.noaa.gov/products/people/yxue/pub/13.pdf",
    )
    add_reference(
        doc,
        "Carton, J. A., G. A. Chepurin, and L. Chen, 2018: SODA3: A New Ocean Climate Reanalysis. Journal of Climate, 31, 6967–6983.",
        "https://doi.org/10.1175/JCLI-D-18-0149.1",
    )
    add_reference(
        doc,
        "Jin, F.-F., 1997: An Equatorial Ocean Recharge Paradigm for ENSO. Part I: Conceptual Model. Journal of the Atmospheric Sciences, 54, 811–829.",
        "https://doi.org/10.1175/1520-0469(1997)054%3C0811:AEORPF%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "Li, X., Z.-Z. Hu, S. Zhao, R. Ding, and B. Zhang, 2022: On the Asymmetry of the Tropical Pacific Thermocline Fluctuation Associated With ENSO Recharge and Discharge. Geophysical Research Letters, 49, e2022GL099242.",
        "https://doi.org/10.1029/2022GL099242",
    )
    add_reference(
        doc,
        "Meinen, C. S., and M. J. McPhaden, 2000: Observations of Warm Water Volume Changes in the Equatorial Pacific and Their Relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559.",
        "https://doi.org/10.1175/1520-0442(2000)013%3C3551:OOWWVC%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "Meinen, C. S., and M. J. McPhaden, 2001: Interannual Variability in Warm Water Volume Transports in the Equatorial Pacific during 1993–99. Journal of Physical Oceanography, 31, 1324–1345.",
        "https://doi.org/10.1175/1520-0485(2001)031%3C1324:IVIWWV%3E2.0.CO;2",
    )
    add_reference(
        doc,
        "Taylor, K. E., 2001: Summarizing Multiple Aspects of Model Performance in a Single Diagram. Journal of Geophysical Research: Atmospheres, 106, 7183–7192.",
        "https://doi.org/10.1029/2000JD900719",
    )
    add_reference(
        doc,
        "Zheng, Y., and B. S. Giese, 2009: Ocean Heat Transport in Simple Ocean Data Assimilation: Structure and Mechanisms. Journal of Geophysical Research: Oceans, 114, C11009.",
        "https://doi.org/10.1029/2008JC005190",
    )
    add_reference(
        doc,
        "NOAA Climate Prediction Center, 2026: Relative Oceanic Niño Index (RONI): Historical El Niño/La Niña Episodes and Official Strength Thresholds. Accessed August 2026.",
        "https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/enso/roni/",
    )
    add_reference(
        doc,
        "NOAA Physical Sciences Laboratory, 2026: NCEP Global Ocean Data Assimilation System (GODAS). Accessed August 2026.",
        "https://psl.noaa.gov/data/gridded/data.godas.html",
    )
    add_reference(
        doc,
        "University of Maryland SODA Project, 2026: SODA3.15.2 Reanalysis Documentation. Accessed August 2026.",
        "https://www.soda.umd.edu/soda3/",
    )

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, path in FIGURES.items():
        if not path.exists():
            raise FileNotFoundError(f"Missing figure {name}: {path}")
    document = build_report()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
