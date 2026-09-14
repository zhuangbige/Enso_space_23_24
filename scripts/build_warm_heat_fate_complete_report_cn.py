#!/usr/bin/env python3
"""Build the complete Chinese DOCX report for the ENSO warm-heat-fate study."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "outputs/research_2026/reports"
OUTPUT = REPORT_DIR / "厄尔尼诺暖量去向与海气耦合闭环_完整中文报告_20260804.docx"

FIGURES = {
    "three_box": ROOT
    / "outputs/research_2026/soda3_three_box_warm_heat_fate/soda3_three_box_warm_heat_fate.png",
    "coupling": ROOT
    / "outputs/research_2026/soda3_air_sea_coupled_fate/soda3_air_sea_coupled_warm_heat_mechanism.png",
    "curl": ROOT
    / "outputs/research_2026/era5_full_stress_curl_warm_heat/era5_full_stress_curl_warm_heat.png",
    "decomposition": ROOT
    / "outputs/research_2026/soda3_5n_current_decomposition/soda3_5n_current_decomposition.png",
    "timing": ROOT
    / "outputs/research_2026/soda3_5n_coupled_timing/soda3_5n_coupled_timing.png",
}

THREE_BOX_SUMMARY = ROOT / "data/derived/research_2026/soda3_three_box_warm_heat_fate_cumulative_summary.csv"
COUPLING_SUMMARY = ROOT / "data/derived/research_2026/soda3_air_sea_coupled_warm_heat_summary.csv"
CURL_SUMMARY = ROOT / "data/derived/research_2026/era5_full_stress_curl_warm_heat_summary.csv"
DECOMP_SUMMARY = ROOT / "data/derived/research_2026/soda3_5n_current_decomposition_summary.csv"
TIMING_SUMMARY = ROOT / "data/derived/research_2026/soda3_5n_coupled_timing_summary.csv"

# Design preset: narrative_proposal, with editorial_cover first-page pattern.
# Named override CJK_SAFE_FONT: use the explicit macOS Simplified-Chinese
# family in every script slot so both Word and the headless QA renderer resolve
# the same glyphs (font aliases such as STSong are not stable in LibreOffice).
BODY_FONT = "PingFang SC"
HEAD_FONT = "PingFang SC"
NAVY = "203748"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "111111"
MUTED = "5B6573"
LIGHT_FILL = "F4F6F9"
LIGHT_BLUE = "E8EEF5"
GRID = "B8C2CC"
WHITE = "FFFFFF"
GOLD = "9A6A00"
RED = "9B1C1C"
GREEN = "1F6B4F"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    font: str = BODY_FONT,
    size: float = 11,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str = INK,
) -> None:
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:shd"))
    if old is not None:
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_margins(cell, *, top=80, bottom=80, start=120, end=120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = tblpr.find(qn(tag))
        if existing is not None:
            tblpr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tblpr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    tblpr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color: str = GRID, size: int = 5) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def mark_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def prevent_row_split(row) -> None:
    """Keep a table row intact when Word/LibreOffice paginates it."""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, font=BODY_FONT, size=8.5, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_numbering(doc: Document, *, numbered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if numbered else "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if numbered else "•")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.extend([start, fmt, text, suff])
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "279")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def configure_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    # Named override CHINESE_PROSE_INDENT: two-character first-line indent.
    normal.paragraph_format.first_line_indent = Pt(22)

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), HEAD_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), HEAD_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)

    caption = doc.styles["Caption"]
    caption.font.name = BODY_FONT
    caption._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.first_line_indent = Pt(0)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("ENSO暖量去向研究报告"), size=8.5, bold=True, color=MUTED)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(footer_table, [7200, 2160], indent_dxa=0)
    left, right = footer_table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    lp = left.paragraphs[0]
    lp.paragraph_format.first_line_indent = Pt(0)
    lp.paragraph_format.space_after = Pt(0)
    set_run_font(lp.add_run("厄尔尼诺暖量去向与海气耦合闭环"), size=8.5, color=MUTED)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.first_line_indent = Pt(0)
    rp.paragraph_format.space_after = Pt(0)
    set_run_font(rp.add_run("第 "), size=8.5, color=MUTED)
    add_page_field(rp)
    set_run_font(rp.add_run(" 页"), size=8.5, color=MUTED)

    props = doc.core_properties
    props.title = "厄尔尼诺暖量去了哪里？——CT与CP事件的海气耦合闭环"
    props.subject = "基于SODA3、ERA5、OISST与GPCP的暖量收支与动力分解"
    props.author = "Daiyue"
    props.keywords = "ENSO; CT El Niño; CP El Niño; SODA3; warm heat; Ekman; geostrophic; air-sea coupling"
    props.comments = "Generated from verified project analyses on 2026-08-04."
    return create_numbering(doc, numbered=False), create_numbering(doc, numbered=True)


def add_para(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    indent: bool = True,
    color: str = INK,
    after: float = 8,
    keep: bool = False,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22 if indent else 0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    p.paragraph_format.keep_with_next = keep
    set_run_font(p.add_run(text), size=11, bold=bold, italic=italic, color=color)


def add_mixed_para(doc: Document, parts: list[tuple[str, bool]], *, indent=True, after=8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22 if indent else 0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    for text, bold in parts:
        set_run_font(p.add_run(text), size=11, bold=bold)


def add_list_item(doc: Document, text: str, num_id: int, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, num])
    ppr.append(numpr)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=11, bold=True)
        set_run_font(p.add_run(text[len(bold_prefix) :]), size=11)
    else:
        set_run_font(p.add_run(text), size=11)


def add_callout(doc: Document, label: str, text: str, *, fill=LIGHT_FILL, color=DARK_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Match the table indent to the callout cell's start margin so that the
    # visible content edge remains aligned with the document text grid.
    set_table_geometry(table, [TABLE_WIDTH_DXA], indent_dxa=180)
    trpr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trpr.append(cant_split)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, val, size in (("left", "single", 20), ("top", "nil", 0), ("bottom", "nil", 0), ("right", "nil", 0)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)
        borders.append(node)
    tcpr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.20
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(label + "  "), font=HEAD_FONT, size=11, bold=True, color=color)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 9,
) -> None:
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Pt(0)
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.keep_with_next = True
    set_run_font(cp.add_run(caption), size=9.5, bold=True, color=INK)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    mark_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        set_run_font(p.add_run(header), font=HEAD_FONT, size=font_size, bold=True, color=DARK_BLUE)
    for row in rows:
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_run_font(p.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, key: str, caption: str, *, width=6.25, page_break=False) -> None:
    path = FIGURES[key]
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph(style="Caption")
    set_run_font(cp.add_run(caption), size=9, color=MUTED)


def add_equation(doc: Document, equation: str, explanation: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(equation), size=11, italic=True, color=NAVY)
    add_para(doc, explanation, indent=False, color=MUTED, after=7)


def add_reference(doc: Document, number: int, citation: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(20)
    p.paragraph_format.first_line_indent = Pt(-20)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(f"[{number}] {citation} "), size=9.5)
    add_hyperlink(p, "DOI/链接", url)


def add_section_heading(doc: Document, text: str, *, page_break=False) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = page_break
    set_run_font(p.add_run(text), font=HEAD_FONT, size=16, bold=True, color=BLUE)


def summary_value(frame: pd.DataFrame, **filters) -> pd.Series:
    selected = frame
    for key, value in filters.items():
        selected = selected.loc[selected[key].eq(value)]
    if len(selected) != 1:
        raise ValueError(f"Expected one row for {filters}, found {len(selected)}")
    return selected.iloc[0]


def build_report() -> Path:
    three = pd.read_csv(THREE_BOX_SUMMARY)
    coupling = pd.read_csv(COUPLING_SUMMARY)
    curl = pd.read_csv(CURL_SUMMARY)
    decomp = pd.read_csv(DECOMP_SUMMARY)
    timing = pd.read_csv(TIMING_SUMMARY)

    doc = Document()
    bullet_id, number_id = configure_document(doc)
    next_step_number_id = create_numbering(doc, numbered=True)

    # Cover: editorial_cover header pattern.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(92)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run("研究报告  |  2026年8月"), font=HEAD_FONT, size=11, bold=True, color=GOLD)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("厄尔尼诺暖量去了哪里？"), font=HEAD_FONT, size=30, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_after = Pt(6)
    set_run_font(subtitle.add_run("CT与CP事件的暖量收支、动力分解与海气耦合闭环"), font=HEAD_FONT, size=15, bold=True, color=DARK_BLUE)
    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.first_line_indent = Pt(0)
    desc.paragraph_format.space_after = Pt(66)
    set_run_font(
        desc.add_run("基于 SODA3.15.2、ERA5、OISST 与 GPCP 的多资料诊断"),
        size=11,
        italic=True,
        color=MUTED,
    )
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.first_line_indent = Pt(0)
    status.paragraph_format.space_after = Pt(5)
    set_run_font(status.add_run("研究状态：观测—再分析诊断闭环已完成"), font=HEAD_FONT, size=11, bold=True, color=GREEN)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Pt(0)
    set_run_font(date.add_run("2026年8月4日"), size=10.5, color=MUTED)
    doc.add_page_break()

    add_section_heading(doc, "摘要")
    add_para(
        doc,
        "本报告围绕一个具体问题展开：传统冷舌型（CT）厄尔尼诺发展期输送到东太平洋的异常暖量，在成熟—衰减期究竟去了哪里，为什么中央太平洋型（CP）没有表现出相同的暖量排放路径？研究首先复现Chakravorty等（2021）的CT/CP充放电框架，随后利用SODA3建立西—中—东三箱H20暖量收支，并联合OISST、GPCP和ERA5检验海气响应，最终在5°N、140°W—80°W断面上将总Q20分解为地转、Ekman和残差项。",
    )
    add_para(
        doc,
        "结果表明，Jun(0)—Nov(0)期间CT和CP均把暖量向东输送，但CT跨140°W进入东部的暖量和东部储存增长更强。Dec(0)—May(+1)期间，CT东部暖量显著下降；暖量一部分跨5°N向北输出，另一部分通过海表热通量散失。CT的原生SODA3北向Q20为+2.77 ZJ，CP为−1.00 ZJ，CT−CP差异为+3.77 ZJ。动力分解显示，该差异由地转+2.10 ZJ、Ekman+1.77 ZJ和不显著残差−0.10 ZJ组成。对于CT本身，显著的北向直接项是Ekman（+2.18 ZJ，约占总出口79%）；CT/CP的地转差异主要来自CP的显著南向地转异常。SODA3自身风应力与独立ERA5估算的CT−CP Ekman差异分别为+1.77和+1.67 ZJ，量级高度一致。",
    )
    add_para(
        doc,
        "逐月时序进一步显示，Dec(0)先出现东部H20和降水差异，Jan(+1)东部SST及总Q20差异变得可辨识，Feb西向风应力、Ekman和地转输送均显著为正，Mar风驱出口仍持续。因此，当前已经形成观测—再分析意义上的海气耦合诊断闭环，但尚不能替代风应力扰动或耦合模式敏感性试验意义上的严格因果闭环。",
    )
    add_callout(
        doc,
        "核心判断",
        "CT东太平洋暖量和对流增强后，局地西向风应力驱动北向Ekman出口；海洋SSH—密度压力场调整同时参与CT/CP差异。文章应写成“风驱Ekman与压力场调整共同控制暖量北出口”，而不是“纯Ekman北逃”。",
        fill=LIGHT_BLUE,
    )
    add_para(
        doc,
        "关键词：厄尔尼诺多样性；冷舌型；中央太平洋型；暖水体积；H20；Q20；Ekman输送；地转输送；海气耦合",
        indent=False,
        bold=True,
        color=DARK_BLUE,
    )

    doc.add_heading("内容结构", level=2)
    for item in (
        "研究背景与科学问题",
        "已有研究与本研究定位",
        "资料、事件和方法",
        "三箱暖量收支：暖量到达东部以后发生了什么",
        "海气耦合证据：对流、风应力与表面散热",
        "5°N动力分解：地转、Ekman与残差",
        "逐月时序与闭环判断",
        "创新性、局限与论文发展方案",
    ):
        add_list_item(doc, item, number_id)

    add_section_heading(doc, "1 研究背景与科学问题", page_break=True)
    doc.add_heading("1.1 从“暖水体积”到“暖量命运”", level=2)
    add_para(
        doc,
        "经典recharge–discharge oscillator认为，厄尔尼诺发生前赤道太平洋暖水体积（WWV）逐步充电，成熟期之后通过经向质量和热量输送放电，从而为相位反转创造条件。WWV通常定义为5°S—5°N、120°E—80°W范围内20°C等温面以上的水体体积。已有理论和观测已证明WWV与ENSO演变密切相关，但“WWV减少”本身没有回答暖水具体从哪里离开、以何种动力形式离开，以及海表热通量是否同时参与。",
    )
    add_para(
        doc,
        "本研究把问题从一个盆地平均指数推进为区域能量问题：发展期向东输送的20°C以上异常暖量，到达140°W以东后，是跨经向边界离开、通过海表散热消失、向深层转移，还是滞留在东部？进一步地，CT和CP事件在这一分配上是否存在系统差异？",
    )
    doc.add_heading("1.2 CT与CP的定义", level=2)
    add_para(
        doc,
        "CT（Cold Tongue El Niño）指主要增暖中心进入东太平洋冷舌区的传统厄尔尼诺；CP（Central Pacific El Niño）指增暖中心主要停留在日界线附近中央太平洋的事件。本文沿用Chakravorty等（2021）的分类逻辑，并在SODA3、OISST、GPCP和ERA5共同覆盖期内选择4个CT和4个CP事件。CT事件为1982、1986、1997和2006；CP事件为1991、1994、2002和2004。",
    )
    add_table(
        doc,
        "表1 事件、阶段与分析区域",
        ["项目", "定义", "用途"],
        [
            ["CT事件", "1982、1986、1997、2006", "传统冷舌型合成"],
            ["CP事件", "1991、1994、2002、2004", "中央太平洋型合成"],
            ["发展期", "Jun(0)—Nov(0)", "诊断暖量向东输送与东部储存"],
            ["成熟—衰减期", "Dec(0)—May(+1)", "诊断暖量出口、散热与动力分解"],
            ["三箱", "10°S—5°N，0—300 m；160°E—180°、180°—140°W、140°W—80°W", "追踪西—中—东暖量迁移"],
            ["主断面", "5°N，140°W—80°W，0—300 m", "分解北向Q20"],
        ],
        [1550, 4310, 3500],
        font_size=8.8,
    )
    doc.add_heading("1.3 核心科学问题", level=2)
    for text in (
        "发展期CT与CP分别把多少暖量送入东太平洋？",
        "成熟—衰减期东部暖量通过北界、南界和海表热通量分别损失多少？",
        "为什么CT跨5°N北向输出显著，而CP没有？",
        "北向总Q20中有多少可以归因于Ekman输送，有多少对应SSH—密度压力场控制的地转调整？",
        "东部暖量、SST/对流、风应力和海洋出口是否具有物理一致的先后顺序？",
    ):
        add_list_item(doc, text, bullet_id)

    add_section_heading(doc, "2 已有研究与本研究定位", page_break=True)
    doc.add_heading("2.1 已有工作已经建立的事实", level=2)
    add_para(
        doc,
        "与本研究最接近的工作可分为四类。第一类是CT/CP充放电差异：Chakravorty等（2021）利用SODA证明CT成熟期存在显著极向放电，而CP的赤道充放电循环较弱。第二类是地转—Ekman分解：Bosc和Delcroix（2008）用高度计与现场观测证明，5°N—5°S的WWV变化是经向Ekman与地转暖水输送相互抵消后的残差；Li等（2022）进一步用GODAS讨论Sverdrup、Ekman和地转输送对WWV和ENSO的作用。第三类是单事件闭合预算：Verdy等（2023）用TPOSE发现2015/16厄尔尼诺恢复期5°N的年龄转/近Ekman输送主导暖表层水北向出口。第四类是热源与热汇：Sun（2000）和Huguenin等（2020）分别从个例与模式角度强调极向热输送和非绝热过程共同终止厄尔尼诺。",
    )
    add_table(
        doc,
        "表2 最接近的已有研究与本研究的区别",
        ["研究", "已完成的内容", "本研究进一步推进"],
        [
            ["Chakravorty et al. (2021)", "SODA中的CT/CP充放电与经向暖水输送", "东部H20能量去向、表面散热、5°N动力分解和大气闭环"],
            ["Bosc & Delcroix (2008)", "观测WWV及地转—Ekman暖水输送", "CT/CP多事件和东部箱体ZJ级能量收支"],
            ["Li et al. (2022)", "GODAS全盆地经向输送及Sverdrup分解", "事件类型差异、北界限定与海气响应时序"],
            ["Verdy et al. (2023)", "2015/16单事件的闭合温盐体积预算", "4+4事件、SODA3、Q20及CT/CP组间机制"],
            ["Sun (2000); Huguenin et al. (2020)", "极向热输送、海表热通量和非绝热WWV过程", "动力出口与热力散失在同一事件框架中闭合"],
        ],
        [2200, 3480, 3680],
        font_size=8.3,
    )
    doc.add_heading("2.2 可以主张与不能主张的创新", level=2)
    add_callout(
        doc,
        "不能主张",
        "不能把“厄尔尼诺暖水跨5°N极向排出”或“地转与Ekman共同控制WWV”写成首次发现；这些机制已有明确文献基础。",
        fill="FFF4E5",
        color=GOLD,
    )
    add_callout(
        doc,
        "可主张的新增贡献",
        "在统一SODA3暖量框架中比较CT与CP到达东太平洋后的能量命运，并同时闭合三箱H20、海表散热、跨5°N总Q20、地转—Ekman分量以及OISST/GPCP/ERA5大气响应。该贡献属于已有理论上的新综合与新定量证据。",
        fill=LIGHT_BLUE,
    )

    add_section_heading(doc, "3 资料、事件和方法")
    doc.add_heading("3.1 资料", level=2)
    for text in (
        "SODA3.15.2：月平均温度、盐度、海表高度、三维流速、纬向风应力和海表热通量；主分析期1980—2022年。",
        "OISST：东太平洋海表温度异常，用于连接海洋暖量与表面增暖。",
        "GPCP：月降水，用于刻画对流中心是否进入东太平洋。",
        "ERA5：纬向/经向风应力及短波、长波、潜热、感热通量，用于独立检验动力和热力反馈。",
    ):
        add_list_item(doc, text, bullet_id)
    doc.add_heading("3.2 暖量与输送定义", level=2)
    add_equation(
        doc,
        "H20 = ρ₀ cₚ ∭ max(T − 20°C, 0) dV",
        "H20以ZJ表示20°C以上相对于20°C参考温度的暖量，而非仅以D20或WWV表示水体几何体积。",
    )
    add_equation(
        doc,
        "Q20 = ρ₀ cₚ ∬ v · max(T − 20°C, 0) dz dx",
        "Q20以PW表示穿过边界的暖量输送；5°N断面规定向北为正。",
    )
    add_equation(
        doc,
        "dH20/dt = C_zonal + C_meridional + Q_surface + R",
        "三箱预算由纬向辐合、经向辐合、海表暖量通量及未解析残差组成。海表向下为正，报告中的“热损失”统一取负号。",
    )
    doc.add_heading("3.3 地转—Ekman分解", level=2)
    add_para(
        doc,
        "在5°N、140°W—80°W断面，原生SODA3 v给出总Q20。地转流由海表高度的表面地转速度和温盐密度梯度的热成风积分重建；密度主结果使用UNESCO EOS-80，线性状态方程作为敏感性。Ekman项由SODA3自身纬向风应力和0—50 m暖层平均温度折算。残差严格定义为总量减去地转和Ekman，不预先命名为某一种洋流。",
    )
    add_equation(
        doc,
        "v_g(0) = (g/f) ∂SSH/∂x",
        "在北半球5°N，北向为正；表面地转速度由沿断面的SSH梯度确定。",
    )
    add_equation(
        doc,
        "v_g(d) = v_g(0) + [g/(ρ₀f)] ∫₀ᵈ ∂ρ/∂x dd",
        "深度d向下为正；第二项表示由温盐密度场给出的热成风剪切。",
    )
    add_equation(
        doc,
        "M_y = −τ_x/(ρ₀f)",
        "异常西向风应力（τx<0）在5°N对应北向Ekman输送。",
    )
    doc.add_heading("3.4 异常、平滑与显著性", level=2)
    add_para(
        doc,
        "所有时间序列先按日历月在1982—2022年线性去趋势。三箱与海气耦合主诊断使用中心3个月平均；仅有Dec—May断面的动力分解使用未平滑月异常，以避免把May与下一年Dec错误连接。统计推断不把单个月份当作独立样本，而以整事件为单位进行20,000次bootstrap。主要结果报告CT、CP和CT−CP的均值及95%区间。",
    )
    doc.add_heading("3.5 质量控制", level=2)
    for text in (
        "1980—2022年Dec—May共258个月，断面维度为20层×120经度格点，无缺月和缺测。",
        "重建总Q20与既有三箱5°N边界逐月结果的最大误差为2.75×10⁻⁹ PW，相关系数为1.000。",
        "地转+Ekman与8个事件总Q20的相关为0.995；逐一删除任一事件后仍为0.991—0.997。",
        "全部月—深度—经度格点的总流速异常与重建地转流异常相关为0.76。",
        "90°W、85°W和80°W三个东界以及EOS-80/线性状态方程给出一致结论。",
    ):
        add_list_item(doc, text, bullet_id)

    add_section_heading(doc, "4 三箱暖量收支：暖量到达东部以后发生了什么", page_break=True)
    doc.add_heading("4.1 发展期：暖量由西向东转移", level=2)
    add_para(
        doc,
        "Jun(0)—Nov(0)期间，CT跨140°W向东输送+6.25 ZJ，CP为+2.97 ZJ；两者均显著。东部箱体储存增长在CT和CP中分别为+7.81和+3.71 ZJ，CT−CP为+4.10 ZJ，95%区间[+1.23,+7.00]。这说明CT与CP都能把暖量送到东部，但CT到达的暖量更强、更深且更靠东。",
    )
    doc.add_heading("4.2 成熟—衰减期：CT东部暖量快速释放", level=2)
    add_para(
        doc,
        "Dec(0)—May(+1)期间，东部H20在CT和CP中分别减少−7.86和−3.03 ZJ。CT跨5°N北向输出+2.74 ZJ，而CP为−0.96 ZJ，即CP合成表现为向东部箱体的净输入。CT−CP北向输出差异为+3.70 ZJ，区间[+1.65,+5.65]。南边界没有发现同样稳健的异常极向输出，因此关键通道位于北侧5°N。",
    )
    add_para(
        doc,
        "与此同时，SODA3东部海表暖量损失在CT和CP中分别为3.69和1.76 ZJ。东部暖量衰减不是单一出口造成，而是北向动力输出和海气热交换同时发生。",
    )
    add_figure(
        doc,
        "three_box",
        "图1  SODA3三箱H20暖量收支。发展期CT把更多暖量跨140°W送入东部；成熟—衰减期CT的北向输出和东部储存下降明显强于CP。",
        width=6.15,
    )

    add_section_heading(doc, "5 海气耦合证据：对流、风应力与表面散热", page_break=True)
    doc.add_heading("5.1 东太平洋对流响应", level=2)
    precip_ct = summary_value(coupling, contrast="CT", phase="mature_decay_Dec0_to_May1", metric="east_precip_mean_mmday")
    precip_cp = summary_value(coupling, contrast="CP", phase="mature_decay_Dec0_to_May1", metric="east_precip_mean_mmday")
    add_para(
        doc,
        f"成熟—衰减期东部GPCP降水异常在CT为{precip_ct['mean']:+.2f} mm day⁻¹，95%区间[{precip_ct['ci_low']:+.2f},{precip_ct['ci_high']:+.2f}]；CP为{precip_cp['mean']:+.2f} mm day⁻¹且区间跨零。CT−CP整个阶段的降水差异下界略低于零，因此应作为机制线索而非最强统计结论；但空间分布和逐月结果显示，CT把对流中心带入东部，而CP主要停留在中太平洋。",
    )
    doc.add_heading("5.2 5°N风应力与完整旋度", level=2)
    tau_diff = summary_value(coupling, contrast="CT_minus_CP", phase="mature_decay_Dec0_to_May1", metric="tau5n_westward_mean_nm2")
    curl_diff = summary_value(curl, contrast="CT_minus_CP", metric="curl_full_mean_1e8_nm3")
    add_para(
        doc,
        f"以向西风应力为正，CT−CP在5°N、140°W—80°W的异常为{tau_diff['mean']:+.4f} N m⁻²，区间[{tau_diff['ci_low']:+.4f},{tau_diff['ci_high']:+.4f}]。完整ERA5风应力旋度在0—10°N、140°W—80°W为{curl_diff['mean']:+.3f}×10⁻⁸ N m⁻³，区间[{curl_diff['ci_low']:+.3f},{curl_diff['ci_high']:+.3f}]；其中−∂τx/∂y是主贡献，经向应力梯度提供较小但同号的补充。",
    )
    doc.add_heading("5.3 海表热力反馈", level=2)
    qnet_ct = summary_value(coupling, contrast="CT", phase="mature_decay_Dec0_to_May1", metric="era5_qnet_loss_zj")
    qnet_cp = summary_value(coupling, contrast="CP", phase="mature_decay_Dec0_to_May1", metric="era5_qnet_loss_zj")
    sw_diff = summary_value(coupling, contrast="CT_minus_CP", phase="mature_decay_Dec0_to_May1", metric="era5_shortwave_loss_zj")
    latent_diff = summary_value(coupling, contrast="CT_minus_CP", phase="mature_decay_Dec0_to_May1", metric="era5_latent_loss_zj")
    add_para(
        doc,
        f"ERA5东部净向下海表热通量异常换算成海洋热损失后，CT和CP分别为{qnet_ct['mean']:+.2f}和{qnet_cp['mean']:+.2f} ZJ。CT中短波减少和潜热散失均重要：CT−CP短波损失差异为{sw_diff['mean']:+.2f} ZJ且显著，符合云和深对流增强；潜热差异为{latent_diff['mean']:+.2f} ZJ但区间跨零，不能单独强调为类型差异。SODA3和ERA5的事件热损失相关为0.94，说明热力结论不依赖单一产品。",
    )
    add_figure(
        doc,
        "coupling",
        "图2  东太平洋暖量去向的海气耦合证据。图中连接东部暖量、SST/降水、5°N风应力、北向输出和海表散热。",
        width=6.25,
    )
    add_figure(
        doc,
        "curl",
        "图3  ERA5完整风应力旋度与一阶Ekman暖量估算。独立ERA5应力支持CT东部大气反馈打开5°N暖量出口。",
        width=6.25,
    )

    add_section_heading(doc, "6 5°N动力分解：地转、Ekman与残差", page_break=True)
    doc.add_heading("6.1 总量与分量", level=2)
    result_rows = []
    for group, label in (("CT", "CT"), ("CP", "CP"), ("CT_minus_CP", "CT−CP")):
        values = []
        for metric in (
            "total_220_280e_zj",
            "geo_eos80_220_280e_zj",
            "ekman_220_280e_zj",
            "residual_eos80_220_280e_zj",
        ):
            row = summary_value(decomp, contrast=group, metric=metric)
            values.append(f"{row['mean']:+.2f} [{row['ci_low']:+.2f},{row['ci_high']:+.2f}]")
        result_rows.append([label, *values])
    add_table(
        doc,
        "表3 Dec(0)—May(+1)跨5°N暖量输送分解（ZJ；括号为95%区间）",
        ["对比", "总Q20", "地转", "Ekman", "残差"],
        result_rows,
        [1050, 2140, 2140, 2140, 1890],
        font_size=8.3,
    )
    add_para(
        doc,
        "CT的总北向输出为+2.77 ZJ，其中Ekman为+2.18 ZJ且显著，事件均值约占总量79%；CT地转项为+0.87 ZJ，但区间跨零。CP总量为−1.00 ZJ，其中显著成分是−1.23 ZJ的南向地转异常，Ekman项不显著。因而，如果回答“CT暖水如何向北离开”，第一主角是风驱Ekman；如果回答“为什么CT与CP不同”，地转和Ekman的组间差异都重要。",
    )
    add_para(
        doc,
        "CT−CP总差异为+3.77 ZJ，地转和Ekman均显著，按均值分别约占56%和47%，残差约−3%。两者百分比略超过100%是因为残差为小幅负值。这里的“贡献”是运动学分解，不表示SSH—密度压力场和风应力是彼此独立的因果强迫；压力场本身也会响应风和热力过程。",
    )
    doc.add_heading("6.2 独立资料和敏感性", level=2)
    add_para(
        doc,
        "SODA3自身风应力得到的CT−CP Ekman暖量差异为+1.77 ZJ，独立ERA5一阶估算为+1.67 ZJ，二者95%区间几乎重合。将断面东端分别设为90°W、85°W和80°W时，总差异为+3.51、+3.55和+3.77 ZJ；地转项为+1.95、+1.98和+2.10 ZJ；Ekman项为+1.60、+1.68和+1.77 ZJ。线性状态方程给出的80°W地转差异为+2.19 ZJ，与EOS-80的+2.10 ZJ一致。",
    )
    add_figure(
        doc,
        "decomposition",
        "图4  SODA3 5°N总流速、重建地转流速及Q20分解。右下展示90°W、85°W和80°W三种东界敏感性。",
        width=6.25,
    )

    add_section_heading(doc, "7 逐月时序与闭环判断", page_break=True)
    doc.add_heading("7.1 月尺度先后顺序", level=2)
    sequence_metrics = [
        (11, "east_h20_zj_anom", "Dec(0)东部H20", "ZJ"),
        (11, "gpcp_east_precip_anom", "Dec(0)东部降水", "mm day⁻¹"),
        (12, "oisst_east_sst_anom", "Jan(+1)东部SST", "°C"),
        (12, "total_220_280e_anom_pw", "Jan(+1)总Q20", "PW"),
        (13, "tau_west_anom_nm2", "Feb西向风应力", "N m⁻²"),
        (13, "ekman_220_280e_anom_pw", "Feb Ekman Q20", "PW"),
        (14, "ekman_220_280e_anom_pw", "Mar Ekman Q20", "PW"),
    ]
    timing_rows = []
    for month, metric, label, unit in sequence_metrics:
        row = summary_value(timing, relative_month=month, contrast="CT_minus_CP", metric=metric)
        timing_rows.append(
            [label, f"{row['mean']:+.3f}", f"[{row['ci_low']:+.3f},{row['ci_high']:+.3f}]", unit]
        )
    add_table(
        doc,
        "表4 CT−CP逐月关键差异",
        ["时间与变量", "均值", "95%区间", "单位"],
        timing_rows,
        [3300, 1500, 2800, 1760],
        font_size=8.7,
    )
    add_para(
        doc,
        "统计可辨识的顺序为：Dec(0)先出现东部暖量和降水差异；Jan(+1)东部SST和总Q20显著；Feb西向风应力以及Ekman、地转Q20均显著为正；Mar Ekman和总出口仍持续。该顺序与“暖量/对流—风应力—动力出口”的物理链一致，但月平均资料无法区分同月内谁先谁后。",
    )
    add_figure(
        doc,
        "timing",
        "图5  CT−CP从东部暖量、SST和降水到5°N风应力、Ekman/地转及总Q20的逐月演变。阴影为整事件bootstrap 95%区间。",
        width=6.25,
    )
    doc.add_heading("7.2 现在是否已经闭环", level=2)
    closure_rows = [
        ["暖量到达东部", "140°W东向Q20与东部H20增长", "已闭合"],
        ["东部海气响应", "OISST增暖、GPCP对流位置、ERA5应力", "机制支持；阶段平均降水差异偏临界"],
        ["动力出口", "SODA3总Q20、地转—Ekman分解、ERA5独立应力", "已闭合"],
        ["热力去向", "SODA3/ERA5海表热损失与通量分量", "已闭合"],
        ["时间顺序", "Dec—Mar逐月事件bootstrap", "已闭合至月尺度"],
        ["严格因果", "尚无风应力关闭或耦合模式扰动试验", "未闭合"],
    ]
    add_table(
        doc,
        "表5 海气耦合闭环的证据状态",
        ["环节", "证据", "判断"],
        closure_rows,
        [1950, 5000, 2410],
        font_size=8.7,
    )
    add_callout(
        doc,
        "闭环结论",
        "观测—再分析诊断闭环已完成；尚缺风应力关闭或耦合模式扰动试验，因此不能宣称严格因果闭环。",
        fill=LIGHT_BLUE,
    )

    add_section_heading(doc, "8 科学解释与现实意义")
    doc.add_heading("8.1 为什么CT与CP的暖量命运不同", level=2)
    add_para(
        doc,
        "CT发展期形成更深、更靠东的暖量储存，使东太平洋SST和对流异常足以改变局地风应力。恢复中的东风/西向应力在5°N驱动北向Ekman输送，完整正风应力旋度则对应浅层经向翻转和压力场调整。CP虽然也能在发展期把部分暖量送入东部，但对流和风应力中心更偏中太平洋，东部5°N没有形成同强度的直接动力出口；其合成反而表现为南向地转异常。因此，CT/CP差异不是“有没有暖量”这么简单，而是暖量到达位置、对流阈值和海洋压力场共同决定的后续分配。",
    )
    doc.add_heading("8.2 现实意义", level=2)
    for text in (
        "解释ENSO相位转换：北向暖量输出和海表散热帮助东太平洋热含量回落，为CT之后的冷却或La Niña转变创造条件。",
        "解释事件持续时间：如果风驱北向出口较弱或压力场补偿不同，东部暖量可能保留更久，从而影响厄尔尼诺衰减速度。",
        "改善模式诊断：模式即使能模拟Niño3.4 SST，也可能错误表示5°N经向暖量出口、应力旋度或表面散热；这些是更严格的过程检验量。",
        "连接区域影响：排出的暖量进入热带北太平洋上层环流，可能影响海洋热浪、热带气旋潜在强度和后续季节的背景海温，但这些远端影响不是本文已经验证的结论。",
    ):
        add_list_item(doc, text, bullet_id)

    add_section_heading(doc, "9 论文定位、核心主张与图件方案", page_break=True)
    doc.add_heading("9.1 建议的论文定位", level=2)
    add_para(
        doc,
        "最合适的定位是“观测—再分析资料驱动的ENSO暖量命运与海气耦合机制诊断”，而不是预报技巧比较论文。文章主问题应保持简洁：为什么CT厄尔尼诺到达东太平洋的暖量在成熟后显著跨5°N向北离开，而CP没有？SODA3的独特优势在于可以同时获得温盐、SSH、三维流速、风应力和热通量，从而把暖量收支与海洋动力结构放在统一框架中。",
    )
    doc.add_heading("9.2 可直接使用的核心主张", level=2)
    for text in (
        "CT和CP发展期都向东输送暖量，但CT在东太平洋形成更强的暖量储存。",
        "CT成熟—衰减期东部暖量通过5°N北向出口和海表散热共同释放；CP北界表现为净输入。",
        "CT本身的直接北向动力出口主要由风驱Ekman解释，而CT/CP差异同时包含显著地转调整。",
        "SODA3和ERA5对CT−CP Ekman差异给出近乎相同的能量量级。",
        "月尺度时序支持东部暖量/对流先于风应力和动力出口变得可辨识。",
    ):
        add_list_item(doc, text, bullet_id)
    doc.add_heading("9.3 建议标题", level=2)
    for text in (
        "厄尔尼诺暖量去了哪里？CT与CP事件东太平洋暖量命运的差异",
        "Wind-driven Ekman export and pressure-field adjustment govern the contrasting warm-heat fate of cold-tongue and central-Pacific El Niño",
        "Contrasting eastern-Pacific warm-heat discharge pathways between cold-tongue and central-Pacific El Niño",
    ):
        add_list_item(doc, text, bullet_id)
    doc.add_heading("9.4 主文图件建议", level=2)
    add_table(
        doc,
        "表6 建议的论文主图结构",
        ["图", "内容", "回答的问题"],
        [
            ["Fig. 1", "CT/CP事件与三箱、边界、阶段定义", "研究对象和分析框架是什么？"],
            ["Fig. 2", "发展期西—中—东H20及Q20", "暖量怎样到达东部？"],
            ["Fig. 3", "成熟—衰减期东部收支", "暖量通过哪些路径离开？"],
            ["Fig. 4", "SST/降水、应力和完整旋度", "大气如何响应东部增暖？"],
            ["Fig. 5", "5°N总流—地转流断面及Q20分解", "北向出口由什么动力组成？"],
            ["Fig. 6", "Dec—May逐月耦合时序", "证据链是否具有正确先后顺序？"],
        ],
        [1000, 4220, 4140],
        font_size=8.8,
    )

    add_section_heading(doc, "10 局限与下一步", page_break=True)
    doc.add_heading("10.1 当前局限", level=2)
    for text in (
        "样本只有4个CT和4个CP。1986 CT是北向总Q20为负的例外，因此文章应写成合成倾向，而不是所有CT事件的确定规律。",
        "CT−CP成熟—衰减期整体降水差异接近但略未达到95%门槛；空间位置和逐月证据支持机制，但不宜把降水差异写成最强结论。",
        "5°N接近赤道，科氏参数较小，SSH梯度和热成风重建容易放大噪声。当前已用多东界、两种状态方程和流速相关验证，但仍应保留方法限定。",
        "地转压力场会受到风应力和热力过程影响，因此地转与Ekman分量是运动学分解，不是彼此独立的因果试验。",
        "SODA3属于同化产品；虽然总Q20由原生v得到，分量和表面通量仍应尽可能用独立资料或模式进行复核。",
    ):
        add_list_item(doc, text, bullet_id)
    doc.add_heading("10.2 最有价值的下一步", level=2)
    for text in (
        "事件异质性：解释1986 CT为何没有表现出群体平均的北向出口，并检验事件强度、暖量到达经度和对流阈值的作用。",
        "严格因果试验：在海洋模式或耦合模式中关闭/替换5°N附近异常风应力，检验北向Q20是否相应减弱。",
        "独立海洋产品：用ORAS5、GODAS或TPOSE复核5°N总量和地转—年龄转分量；特别关注2015/16单事件与多事件合成的一致性。",
        "模式过程评估：把5°N北向Q20、完整应力旋度和东部表面热损失作为CMIP/季节预报模式的过程指标，而非立即转向预报排名。",
    ):
        add_list_item(doc, text, next_step_number_id)
    add_callout(
        doc,
        "建议",
        "当前已可进入SODA3诊断论文写作；若要把支持的机制提升为严格因果机制，优先补风应力敏感性试验。",
        fill="FFF4E5",
        color=GOLD,
    )

    add_section_heading(doc, "参考文献")
    references = [
        (
            "Jin, F.-F. (1997). An equatorial ocean recharge paradigm for ENSO. Part I: Conceptual model. Journal of the Atmospheric Sciences, 54, 811–829.",
            "https://doi.org/10.1175/1520-0469(1997)054%3C0811:AEORPF%3E2.0.CO;2",
        ),
        (
            "Meinen, C. S., & McPhaden, M. J. (2000). Observations of warm water volume changes in the equatorial Pacific and their relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559.",
            "https://doi.org/10.1175/1520-0442(2000)013%3C3551:OOWWVC%3E2.0.CO;2",
        ),
        (
            "Sun, D.-Z. (2000). The heat sources and sinks of the 1986–87 El Niño. Journal of Climate, 13, 3533–3550.",
            "https://doi.org/10.1175/1520-0442(2000)013%3C3533:THSASO%3E2.0.CO;2",
        ),
        (
            "Bosc, C., & Delcroix, T. (2008). Observed equatorial Rossby waves and ENSO-related warm water volume changes in the equatorial Pacific Ocean. Journal of Geophysical Research: Oceans, 113, C06003.",
            "https://doi.org/10.1029/2007JC004613",
        ),
        (
            "Huguenin, M. F., Holmes, R. M., & England, M. H. (2020). Key role of diabatic processes in regulating warm water volume variability over ENSO events. Journal of Climate, 33, 9783–9801.",
            "https://doi.org/10.1175/JCLI-D-20-0198.1",
        ),
        (
            "Chakravorty, S., Perez, R. C., Gnanaseelan, C., & Anderson, B. T. (2021). Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 126, e2020JC017075.",
            "https://doi.org/10.1029/2020JC017075",
        ),
        (
            "Li, X., Hu, Z.-Z., Huang, B., & Jin, F.-F. (2022). Oceanic meridional transports and their roles in warm water volume variability and ENSO in the tropical Pacific. Climate Dynamics, 59, 245–261.",
            "https://doi.org/10.1007/s00382-021-06124-w",
        ),
        (
            "Verdy, A., Mazloff, M. R., Cornuelle, B. D., & Kim, S. B. (2023). Balancing volume, temperature, and salinity budgets during 2014–2018 in the Tropical Pacific Ocean State Estimate. Journal of Geophysical Research: Oceans, 128, e2022JC019576.",
            "https://doi.org/10.1029/2022JC019576",
        ),
        (
            "Carton, J. A., Chepurin, G. A., & Chen, L. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983.",
            "https://doi.org/10.1175/JCLI-D-18-0149.1",
        ),
        (
            "Hersbach, H., et al. (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146, 1999–2049.",
            "https://doi.org/10.1002/qj.3803",
        ),
        (
            "Huang, B., et al. (2021). Improvements of the Daily Optimum Interpolation Sea Surface Temperature (DOISST) version 2.1. Journal of Climate, 34, 2923–2939.",
            "https://doi.org/10.1175/JCLI-D-20-0166.1",
        ),
        (
            "Adler, R. F., et al. (2018). The Global Precipitation Climatology Project monthly analysis: An update. Atmosphere, 9, 138.",
            "https://doi.org/10.3390/atmos9040138",
        ),
    ]
    for number, (citation, url) in enumerate(references, start=1):
        add_reference(doc, number, citation, url)

    add_section_heading(doc, "附录A 结果与可复现产物", page_break=True)
    add_para(
        doc,
        "以下文件构成当前报告的直接计算依据。报告中的数字来自相应CSV汇总表，图件来自同一分析脚本输出。",
    )
    artifact_rows = [
        ["三箱暖量去向", "outputs/research_2026/soda3_three_box_warm_heat_fate/", "scripts/analyze_soda3_three_box_warm_heat_fate.py"],
        ["海气耦合闭环", "outputs/research_2026/soda3_air_sea_coupled_fate/", "scripts/analyze_soda3_air_sea_coupled_warm_heat_fate.py"],
        ["完整风应力旋度", "outputs/research_2026/era5_full_stress_curl_warm_heat/", "scripts/analyze_era5_full_stress_curl_warm_heat_export.py"],
        ["5°N动力分解", "outputs/research_2026/soda3_5n_current_decomposition/", "scripts/analyze_soda3_5n_current_decomposition.py"],
        ["逐月耦合时序", "outputs/research_2026/soda3_5n_coupled_timing/", "scripts/analyze_soda3_5n_coupled_timing.py"],
        ["SODA3断面提取", "data/interim/research_2026/soda3_5n_geostrophic_sections_1980_2022.nc", "scripts/fetch_soda3_5n_geostrophic_sections.py"],
    ]
    add_table(
        doc,
        "表A1 主要产物与脚本",
        ["模块", "报告/数据目录", "主脚本"],
        artifact_rows,
        [1900, 4050, 3410],
        font_size=7.8,
    )
    add_para(
        doc,
        "最终质量控制包括：5个新增动力分解单元测试、13个既有暖量/耦合测试，共18项全部通过；DOCX生成后另行执行逐页渲染和视觉检查。",
        bold=True,
        indent=False,
    )

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
