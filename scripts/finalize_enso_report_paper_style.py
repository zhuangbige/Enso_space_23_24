#!/usr/bin/env python3
"""Create the paper-style final ENSO report.

The final pass removes figure-level suptitles and source-note strips from the
embedded raster figures, transfers those details into external Fig. captions,
and removes bold formatting from all editable Word text. Heading hierarchy is
retained through font size alone.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "数据公式与跨事件展宽_定义严谨修订版_20260818.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "论文图注无粗体终版_20260818.docx"
)

FONT = "PingFang SC"
CAPTION_COLOR = "4F4F4F"

# Pixel crop boxes, ordered by the 11 inline figures. Only the global figure
# title and the small source/method strip are removed; axes, panel labels,
# colorbars, and their units remain inside the scientific graphic.
CROP_BOXES = {
    1: (0, 140, 4376, 1975),
    2: (0, 140, 4138, 1875),
    3: (0, 140, 4376, 1975),
    4: (0, 140, 4138, 1875),
    5: (0, 140, 4376, 1975),
    6: (0, 140, 4138, 1875),
    7: (0, 115, 6147, 2620),
    8: (0, 150, 4156, 2168),
    9: (0, 150, 4156, 2168),
    10: (0, 150, 4156, 2168),
    11: (0, 180, 3360, 1840),
}

CAPTIONS = {
    30: (
        "Fig. 1(a): 1997年1—6月OISST海表温度异常（填色）与ERA5 850 hPa风异常（箭头）。"
        "范围为30°S—30°N、120°E—80°W；OISST采用1982—2024年、ERA5风采用1980—2024年同历月线性去趋势异常。"
        "各面板风矢量标尺为4 m s⁻¹；虚线框标示Niño3、Niño3.4、EP和FEP诊断区。"
    ),
    32: (
        "Fig. 1(b): 1997年1—6月SODA3 0—300 m单位面积OHC异常（填色）与ERA5 850 hPa风异常（箭头）。"
        "范围同Fig. 1(a)；OHC和风均采用1980—2024年同历月线性去趋势异常。三次事件统一使用±3 GJ m⁻²色标，"
        "各面板风矢量标尺为4 m s⁻¹。"
    ),
    38: (
        "Fig. 2(a): 2015年1—6月OISST海表温度异常与ERA5 850 hPa风异常。填色、箭头、诊断框、空间范围、"
        "距平基准及风矢量标尺同Fig. 1(a)。"
    ),
    40: (
        "Fig. 2(b): 2015年1—6月SODA3 0—300 m单位面积OHC异常与ERA5 850 hPa风异常。"
        "空间范围、距平基准、风矢量标尺和统一色标同Fig. 1(b)。"
    ),
    46: (
        "Fig. 3(a): 2023年1—6月OISST海表温度异常与ERA5 850 hPa风异常。2月远东FEP已出现局地暖异常；"
        "填色、箭头、诊断框、空间范围、距平基准及风矢量标尺同Fig. 1(a)。"
    ),
    48: (
        "Fig. 3(b): 2023年1—6月SODA3 0—300 m单位面积OHC异常与ERA5 850 hPa风异常。"
        "空间范围、距平基准、风矢量标尺和统一色标同Fig. 1(b)。"
    ),
    54: (
        "Fig. 4: 1997、2015和2023年1—6月赤道太平洋热含量密度异常q′=ρ₀cₚT′的经度—深度剖面。"
        "各行依次为1997、2015和2023年，各列为1—6月；采用5°S—5°N余弦纬度加权平均、SODA3原生0—300 m深度层，"
        "经度范围为120°E—80°W。异常为1980—2024年同历月线性去趋势结果，共用色标，单位MJ m⁻³。"
    ),
    63: (
        "Fig. 5: 1997年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。范围为30°S—30°N、0—300 m；"
        "采用SODA3温度和1980—2024年同历月线性去趋势异常，六个月共用±6 MJ m⁻³色标。"
    ),
    68: (
        "Fig. 6: 2015年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。空间范围、异常定义及统一色标同Fig. 5。"
    ),
    73: (
        "Fig. 7: 2023年1—6月固定120°E—80°W平均的纬度—深度热含量密度异常。空间范围、异常定义及统一色标同Fig. 5。"
    ),
    82: (
        "Fig. 8: 13次ENSO暖事件中，JFM暖核附近的经向OHC配置与随后JJA发展期海气响应的关系。"
        "横轴为离赤道0—300 m OHC在控制三纬带面积加权平均OHC后的残差；纵轴为控制相同海洋量后的JJA型态调整耦合残差。"
        "海洋资料为SODA3.15.2，窗口跟随暖核质心经度±10°，纬度范围15°S—15°N，异常基准为1980—2024年同历月线性去趋势。"
        "耦合残差控制SST幅度和EP—CP型态；灰色阴影为事件重采样回归线的95%区间。"
    ),
}

PAGINATION_TAGS = (
    "keepNext",
    "keepLines",
    "pageBreakBefore",
    "widowControl",
    "suppressLineNumbers",
)


def set_run_font(run, size: float, color: str = CAPTION_COLOR) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = None
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT)
    color_element = rpr.find(qn("w:color"))
    if color_element is None:
        color_element = OxmlElement("w:color")
        rpr.append(color_element)
    color_element.set(qn("w:val"), color)


def strip_pagination(paragraph) -> None:
    ppr = paragraph._p.pPr
    if ppr is None:
        return
    for tag in PAGINATION_TAGS:
        element = ppr.find(qn(f"w:{tag}"))
        if element is not None:
            ppr.remove(element)
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def replace_captions(document: Document) -> None:
    for paragraph_index, caption in CAPTIONS.items():
        paragraph = document.paragraphs[paragraph_index]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.12
        run = paragraph.add_run(caption)
        set_run_font(run, 8.8)
        strip_pagination(paragraph)


def remove_all_bold(document: Document) -> None:
    """Remove bold from body text, captions, styles, headers, and footers."""
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            element = getattr(part, "_element", None)
        if element is None or not hasattr(element, "xpath"):
            continue
        for bold_element in list(element.xpath(".//w:b | .//w:bCs")):
            parent = bold_element.getparent()
            if parent is not None:
                parent.remove(bold_element)


def use_installed_regular_font(document: Document) -> None:
    """Replace the unavailable CJK family while retaining every existing size."""
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            element = getattr(part, "_element", None)
        if element is None or not hasattr(element, "xpath"):
            continue
        for fonts in element.xpath(".//w:rFonts"):
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                key = qn(f"w:{attribute}")
                if fonts.get(key) == "Noto Sans CJK SC":
                    fonts.set(key, FONT)


def crop_embedded_images(source_docx: Path) -> tuple[dict[str, bytes], dict[int, tuple[int, int]]]:
    replacements: dict[str, bytes] = {}
    dimensions: dict[int, tuple[int, int]] = {}
    with ZipFile(source_docx) as archive:
        for image_index, box in CROP_BOXES.items():
            member = f"word/media/image{image_index}.png"
            with Image.open(BytesIO(archive.read(member))) as image:
                if image.size != (box[2], box[3] if box[1] == 0 else image.size[1]):
                    expected_width = box[2]
                    if image.width != expected_width:
                        raise ValueError(
                            f"Unexpected width for image{image_index}: {image.width}, expected {expected_width}"
                        )
                cropped = image.crop(box)
                output = BytesIO()
                cropped.save(output, format="PNG", optimize=True)
                replacements[member] = output.getvalue()
                dimensions[image_index] = cropped.size
    return replacements, dimensions


def resize_inline_shapes(document: Document, dimensions: dict[int, tuple[int, int]]) -> None:
    if len(document.inline_shapes) != len(CROP_BOXES):
        raise ValueError(
            f"Expected {len(CROP_BOXES)} inline figures, found {len(document.inline_shapes)}"
        )
    for image_index, shape in enumerate(document.inline_shapes, start=1):
        width_px, height_px = dimensions[image_index]
        shape.height = int(shape.width * height_px / width_px)


def write_docx_with_replaced_media(
    intermediate_docx: Path, output_docx: Path, replacements: dict[str, bytes]
) -> None:
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        dir=output_docx.parent, suffix=".docx", delete=False
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)
    try:
        with ZipFile(intermediate_docx) as source, ZipFile(
            temporary_path, "w", compression=ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                data = replacements.get(info.filename, source.read(info.filename))
                copied_info = ZipInfo(info.filename, date_time=info.date_time)
                copied_info.compress_type = ZIP_DEFLATED
                copied_info.external_attr = info.external_attr
                copied_info.internal_attr = info.internal_attr
                copied_info.flag_bits = info.flag_bits
                copied_info.create_system = info.create_system
                target.writestr(copied_info, data)
        temporary_path.replace(output_docx)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def audit(document: Document) -> None:
    if len(document.inline_shapes) != 11:
        raise ValueError("The final report must contain 11 inline figures")
    if len(document.element.xpath(".//m:oMath")) != 8:
        raise ValueError("The final report must retain 8 native Word equations")
    if len(document.element.xpath(".//w:b | .//w:bCs")):
        raise ValueError("Bold formatting remains in the document body")
    captions = [p.text for p in document.paragraphs if p.text.startswith("Fig. ")]
    if len(captions) != 11:
        raise ValueError(f"Expected 11 Fig. captions, found {len(captions)}")
    for paragraph in document.paragraphs:
        strip_pagination(paragraph)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    replacements, cropped_dimensions = crop_embedded_images(SOURCE)
    document = Document(SOURCE)
    replace_captions(document)
    resize_inline_shapes(document, cropped_dimensions)
    use_installed_regular_font(document)
    remove_all_bold(document)
    for paragraph in document.paragraphs:
        strip_pagination(paragraph)
    audit(document)

    with NamedTemporaryFile(
        dir=OUTPUT.parent, suffix=".docx", delete=False
    ) as intermediate_file:
        intermediate_path = Path(intermediate_file.name)
    try:
        document.save(intermediate_path)
        write_docx_with_replaced_media(intermediate_path, OUTPUT, replacements)
    finally:
        if intermediate_path.exists():
            intermediate_path.unlink()

    final_document = Document(OUTPUT)
    audit(final_document)
    print(OUTPUT)
    print(f"paragraphs: {len(final_document.paragraphs)}")
    print(f"inline figures: {len(final_document.inline_shapes)}")
    print(f"native equations: {len(final_document.element.xpath('.//m:oMath'))}")
    print(f"Fig. captions: {sum(p.text.startswith('Fig. ') for p in final_document.paragraphs)}")


if __name__ == "__main__":
    main()
