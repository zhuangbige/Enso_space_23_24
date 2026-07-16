#!/usr/bin/env python3
"""Build the revised ENSO comparison DOCX with strip-format figures."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/research_2026/reports/enso_event_comparison_unified_1980_2014_en_final_event_groups.docx"
FIG_DIR = ROOT / "outputs/research_2026/revised_event_figures"
INPUT_DIR = ROOT / "data/interim/research_2026/monthly_panel_inputs"
PMM_CSV = ROOT / "data/derived/research_2026/pmm_like_index_oisst_era5_1982_2026.csv"
PMM_FIG = ROOT / "outputs/research_2026/pmm_index/fig_pmm_like_eof_patterns_and_index.png"
PMM_FOCUS = ROOT / "outputs/research_2026/pmm_index/fig_pmm_like_index_2026_jan_may.png"
GROUP_ASSIGNMENTS_CSV = FIG_DIR / "enso_event_group_assignments.csv"

EVENTS = {
    "2015/16": {
        "npz": INPUT_DIR / "monthly_2015_detrended_fields_1980_2014.npz",
        "year": 2015,
        "label": "2015/16 development year",
        "ssta": FIG_DIR / "fig_2015_revised_ssta_u850_detrended_1980_2014.png",
        "precip": FIG_DIR / "fig_2015_revised_era5_precip_u850_detrended_1980_2014.png",
        "d20": FIG_DIR / "fig_2015_revised_d20_u850_detrended_1980_2014.png",
    },
    "2023/24": {
        "npz": INPUT_DIR / "monthly_2023_detrended_fields_1980_2014.npz",
        "year": 2023,
        "label": "2023/24 development year",
        "ssta": FIG_DIR / "fig_2023_revised_ssta_u850_detrended_1980_2014.png",
        "precip": FIG_DIR / "fig_2023_revised_era5_precip_u850_detrended_1980_2014.png",
        "d20": FIG_DIR / "fig_2023_revised_d20_u850_detrended_1980_2014.png",
    },
    "2026": {
        "npz": INPUT_DIR / "monthly_2026_detrended_fields_1980_2014.npz",
        "year": 2026,
        "label": "2026 January-May",
        "ssta": FIG_DIR / "fig_2026_revised_ssta_u850_detrended_1980_2014.png",
        "precip": FIG_DIR / "fig_2026_revised_era5_precip_u850_detrended_1980_2014.png",
        "d20": FIG_DIR / "fig_2026_revised_d20_u850_detrended_1980_2014.png",
    },
}

CLASSIC_EVENT_FIGS = {
    "ssta": FIG_DIR / "fig_classic_strong_events_revised_ssta_u850_detrended_1980_2014.png",
    "precip": FIG_DIR / "fig_classic_strong_events_revised_era5_precip_u850_detrended_1980_2014.png",
    "d20": FIG_DIR / "fig_classic_strong_events_revised_d20_u850_detrended_1980_2014.png",
}

SENSITIVITY_EVENT_FIGS = {
    "ssta": FIG_DIR / "fig_extended_sensitivity_events_revised_ssta_u850_detrended_1980_2014.png",
    "precip": FIG_DIR / "fig_extended_sensitivity_events_revised_era5_precip_u850_detrended_1980_2014.png",
    "d20": FIG_DIR / "fig_extended_sensitivity_events_revised_d20_u850_detrended_1980_2014.png",
}

REGIONS = {
    "Nino3.4": (-5, 5, 190, 240),
    "CWP": (-5, 5, 160, 220),
    "CP": (-5, 5, 180, 220),
    "EP": (-5, 5, 220, 260),
    "FEP": (-5, 5, 270, 280),
}
D20_REGIONS = {"West": (130, 180), "Central": (180, 220), "East": (220, 275)}


def set_run_font(run, size: float | None = None, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str = "FFFFFF") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def area_mean(field: np.ndarray, lat: np.ndarray, lon: np.ndarray, box: tuple[float, float, float, float]) -> np.ndarray:
    lat0, lat1, lon0, lon1 = box
    lm = (lat >= lat0) & (lat <= lat1)
    om = (lon >= lon0) & (lon <= lon1)
    sub = field[:, lm, :][:, :, om]
    weights = np.cos(np.deg2rad(lat[lm]))
    return np.nansum(sub * weights[None, :, None], axis=(1, 2)) / (np.sum(weights) * om.sum())


def d20_mean(field: np.ndarray, lon: np.ndarray, bounds: tuple[float, float]) -> np.ndarray:
    mask = (lon >= bounds[0]) & (lon <= bounds[1])
    return np.nanmean(field[:, mask], axis=1)


def event_group_labels(group: str) -> list[str]:
    if not GROUP_ASSIGNMENTS_CSV.exists():
        return []
    df = pd.read_csv(GROUP_ASSIGNMENTS_CSV)
    subset = df[df["group"] == group]
    return subset["event_label"].astype(str).tolist()


def load_event_metrics() -> dict[str, dict[str, float]]:
    metrics: dict[str, dict[str, float]] = {}
    for event, info in EVENTS.items():
        z = np.load(info["npz"])
        lat = z["lat"].astype(float)
        lon = z["lon"].astype(float)
        slat = z["sst_lat"].astype(float)
        slon = z["sst_lon"].astype(float)
        dlon = z["d20_lon"].astype(float)
        out: dict[str, float] = {}
        for region, box in REGIONS.items():
            out[f"sst_{region}"] = float(np.nanmean(area_mean(z["sst"], slat, slon, box)))
            out[f"u_{region}"] = float(np.nanmean(area_mean(z["u850"], lat, lon, box)))
            out[f"pr_{region}"] = float(np.nanmean(area_mean(z["precip"], lat, lon, box)))
        for region, bounds in D20_REGIONS.items():
            out[f"d20_{region}"] = float(np.nanmean(d20_mean(z["d20"], dlon, bounds)))
        metrics[event] = out

    pmm = pd.read_csv(PMM_CSV)
    for event, info in EVENTS.items():
        subset = pmm[(pmm["year"] == info["year"]) & (pmm["month"].between(1, 5))]
        metrics[event]["pmm_joint"] = float(subset["pmm_joint_sst_wind"].mean())
        metrics[event]["pmm_sst"] = float(subset["pmm_sst_only"].mean())
    return metrics


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)

    for style_name, size, bold, before, after in [
        ("Normal", 11.2, False, 0, 5),
        ("Heading 1", 16.0, True, 13, 6),
        ("Heading 2", 13.6, True, 9, 5),
        ("Heading 3", 12.2, True, 7, 4),
        ("Caption", 9.2, False, 1, 6),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.08
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), "Times New Roman")


def add_paragraph(doc: Document, text: str = "", style: str | None = None, size: float = 11.2) -> None:
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(5)
    para.paragraph_format.line_spacing = 1.08
    run = para.add_run(text)
    set_run_font(run, size=size)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    para = doc.add_paragraph(style=f"Heading {level}")
    para.paragraph_format.keep_with_next = True
    run = para.add_run(text)
    set_run_font(run, size={1: 16.0, 2: 13.6, 3: 12.2}.get(level, 12.2), bold=True)


def add_caption(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="Caption")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(1)
    para.paragraph_format.space_after = Pt(7)
    run = para.add_run(text)
    set_run_font(run, size=9.2, italic=True)


def add_figure(
    doc: Document,
    image: Path,
    caption: str,
    width: float = 7.30,
    *,
    keep_with_caption: bool = False,
) -> None:
    if not image.exists():
        raise FileNotFoundError(image)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.keep_with_next = keep_with_caption
    para.add_run().add_picture(str(image), width=Inches(width))
    add_caption(doc, caption)


def add_metric_table(doc: Document, metrics: dict[str, dict[str, float]]) -> None:
    add_heading(doc, "Regional Mean Diagnostics under the Unified Climatology", level=2)
    add_paragraph(
        doc,
        "The table reports January-May means. SST is in deg C, u850 is in m/s, precipitation is in mm/day, and D20 is in m. "
        "Positive u850 denotes westerly anomalies. The main classic reference and extended sensitivity events are shown in "
        "separate figures rather than merged into a single table column.",
    )
    rows = [
        ("Niño-3.4 SST", "sst_Nino3.4", "{:.2f}"),
        ("CP SST (180E-140W)", "sst_CP", "{:.2f}"),
        ("EP SST (140W-100W)", "sst_EP", "{:.2f}"),
        ("FEP SST (90W-80W)", "sst_FEP", "{:.2f}"),
        ("CWP u850", "u_CWP", "{:.2f}"),
        ("CP precipitation", "pr_CP", "{:.2f}"),
        ("D20 west / central / east", None, "{:.1f} / {:.1f} / {:.1f}"),
        ("PMM joint index", "pmm_joint", "{:.2f}"),
        ("PMM SST-only index", "pmm_sst", "{:.2f}"),
    ]
    headers = ["Metric", "2015/16", "2023/24", "2026"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for label, key, fmt in rows:
        cells = table.add_row().cells
        cells[0].text = label
        for i, event in enumerate(headers[1:], start=1):
            if key is None:
                vals = (metrics[event]["d20_West"], metrics[event]["d20_Central"], metrics[event]["d20_East"])
                cells[i].text = fmt.format(*vals)
            else:
                cells[i].text = fmt.format(metrics[event][key])
    set_table_width(table, [2.35, 1.55, 1.55, 1.55])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "FFFFFF")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    set_run_font(run, size=8.5, bold=(row is table.rows[0]))


def add_event_group_table(
    doc: Document,
    classic_text: str,
    sensitivity_text: str,
    target_text: str,
) -> None:
    add_heading(doc, "Event Groups and Development-Phase Alignment", level=2)
    add_paragraph(
        doc,
        "Every comparison is aligned to January-May of the developing calendar year, representing the early formation stage "
        "of each El Niño event. The following-year January-May period and the source paper's June-to-following-January event "
        "mean are not used in these development-phase maps.",
    )
    headers = ["Group", "Events", "Aligned window", "Role in this report"]
    rows = [
        (
            "Main classic strong reference",
            classic_text,
            "Jan-May 1982 and Jan-May 1997",
            "Primary historical reference for all three target events; 2015/16 is excluded from the reference.",
        ),
        (
            "Extended sensitivity",
            sensitivity_text,
            "Jan-May 2002 and Jan-May 2009",
            "Robustness check only; not mixed into the main classic reference.",
        ),
        (
            "Target events",
            target_text,
            "Jan-May 2015, 2023, and 2026",
            "Events being compared; their ENSO status is not re-audited here.",
        ),
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    set_table_width(table, [1.55, 1.30, 1.75, 2.40])
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=100, start=110, bottom=100, end=110)
            set_cell_shading(cell, "F2F2F2" if row_idx == 0 else "FFFFFF")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in {1, 2} else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_run_font(run, size=8.1, bold=(row_idx == 0))


def build_doc() -> None:
    metrics = load_event_metrics()
    classic_text = ", ".join(event_group_labels("main_classic_strong_reference"))
    sensitivity_text = ", ".join(event_group_labels("extended_sensitivity"))
    target_text = ", ".join(event_group_labels("target_events"))
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Comparison of the 2015/16, 2023/24, and 2026 ENSO Development Phases")
    set_run_font(run, size=20.5, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(
        "Early-formation January-May SST-wind, precipitation-wind, D20-wind, PMM, and group-separated historical diagnostics"
    )
    set_run_font(run, size=11.2)

    add_heading(doc, "Data and Methods", level=1)
    add_paragraph(
        doc,
        "The analysis compares January-May diagnostics for the early formation stage of three ENSO target events and two "
        "separately defined historical groups. "
        "SST is from NOAA OISST, 850 hPa winds and precipitation are from ERA5, and D20 is from ORAS5. "
        "For SST, winds, precipitation, and D20, anomalies are calculated relative to the same-calendar-month climatology "
        "and then linearly detrended. The common baseline is 1980-2014 for ERA5/ORAS5 and 1982-2014 for OISST, "
        "reflecting OISST data availability.",
    )
    add_paragraph(
        doc,
        f"The main classic strong reference group is limited to {classic_text}. The 2015/16 event is retained only as a target "
        "event and is excluded from the reference, preventing the same event from appearing on both sides of the comparison. "
        f"The extended sensitivity group ({sensitivity_text}) is plotted separately and is not mixed into the primary reference.",
    )
    add_paragraph(
        doc,
        "Event labels follow the developing year: 1982/83 maps use January-May 1982, 1997/98 maps use January-May 1997, "
        "2002/03 maps use January-May 2002, 2009/10 maps use January-May 2009, and the target maps use January-May 2015, "
        "2023, and 2026. This calendar alignment isolates early formation rather than the following mature or decay phase.",
    )
    add_paragraph(
        doc,
        "The PMM-like diagnostic is constructed independently using an EOF trained over 1982-2025; January-May 2026 is "
        "projected onto this fixed EOF mode and is not used in the EOF training. The revised map figures use a single-row "
        "strip layout for the three target events and separate multi-row layouts for the classic and sensitivity groups. Precipitation-wind "
        "panels use bright magenta wind vectors with a thicker white underlay to improve visibility over the precipitation shading.",
    )

    add_event_group_table(doc, classic_text, sensitivity_text, target_text)
    doc.add_page_break()
    add_metric_table(doc, metrics)

    add_heading(doc, "Results Overview", level=1)
    add_paragraph(
        doc,
        "Under the unified climatology, the 2026 case shows more than a weak eastern-Pacific surface warming signal. "
        "Its Niño-3.4, CP, and EP SST anomalies are stronger than in the earlier non-unified comparison, and its CWP/CP "
        "lower-tropospheric westerlies and eastern Pacific D20 are also enhanced. In dynamical terms, 2026 is closer to "
        "2015 than to 2023, although the convective response remains weaker than the 2015-like ocean and wind signals.",
    )
    add_paragraph(
        doc,
        "The main classic reference isolates the two strongest historical eastern-Pacific development cases, 1982/83 and "
        "1997/98. The 2002/03 and 2009/10 cases remain useful for testing whether a conclusion depends on choosing only very "
        "strong events, but they are treated as an extended sensitivity group. The 2015/16 event is discussed only alongside "
        "2023/24 and 2026 as a target event.",
    )

    add_heading(doc, "SST Anomalies and 850 hPa Winds", level=1)
    add_paragraph(
        doc,
        "The revised SST-wind panels emphasize the longitudinal evolution of the warm center during the early formation stage. "
        "Compared with the main classic historical cases, 2015 develops a strong central to central-eastern Pacific warming pattern with persistent westerly "
        "anomalies. The 2023 event remains more displaced toward the far eastern Pacific. By April-May 2026, central Pacific "
        "SST and westerly anomalies begin to connect with the eastern Pacific warming signal.",
    )
    add_figure(doc, EVENTS["2015/16"]["ssta"], "Figure 1. January-May 2015 early formation: OISST SSTA with ERA5 850 hPa wind anomalies.")
    add_figure(doc, EVENTS["2023/24"]["ssta"], "Figure 2. January-May 2023 early formation: OISST SSTA with ERA5 850 hPa wind anomalies.")
    add_figure(doc, EVENTS["2026"]["ssta"], "Figure 3. January-May 2026 early formation: OISST SSTA with ERA5 850 hPa wind anomalies.")
    add_figure(
        doc,
        CLASSIC_EVENT_FIGS["ssta"],
        f"Figure 4. Main classic strong reference ({classic_text}), development-year January-May, plotted individually: OISST SSTA with ERA5 850 hPa wind anomalies.",
    )
    add_figure(
        doc,
        SENSITIVITY_EVENT_FIGS["ssta"],
        f"Figure 5. Extended sensitivity group ({sensitivity_text}), development-year January-May, plotted separately: OISST SSTA with ERA5 850 hPa wind anomalies.",
    )

    add_heading(doc, "Precipitation and 850 hPa Winds", level=1)
    add_paragraph(
        doc,
        "The precipitation-wind panels show the maturity of atmospheric coupling. In 2015, the central Pacific convective "
        "response is early and coherent. In 2023, central Pacific precipitation remains weaker during the far-eastern Pacific "
        "warming phase. In 2026, the westerly anomalies already resemble 2015, but central Pacific precipitation remains "
        "suppressed from January to April and only recovers locally in May. The classic and sensitivity groups are shown "
        "separately so that this lag is not obscured by averaging events with different strengths.",
    )
    add_figure(doc, EVENTS["2015/16"]["precip"], "Figure 6. January-May 2015 early formation: ERA5 precipitation with ERA5 850 hPa wind anomalies.")
    add_figure(doc, EVENTS["2023/24"]["precip"], "Figure 7. January-May 2023 early formation: ERA5 precipitation with ERA5 850 hPa wind anomalies.")
    add_figure(doc, EVENTS["2026"]["precip"], "Figure 8. January-May 2026 early formation: ERA5 precipitation with ERA5 850 hPa wind anomalies.")
    add_figure(
        doc,
        CLASSIC_EVENT_FIGS["precip"],
        f"Figure 9. Main classic strong reference ({classic_text}), development-year January-May, plotted individually: ERA5 precipitation with ERA5 850 hPa wind anomalies.",
        width=6.70,
    )
    add_figure(
        doc,
        SENSITIVITY_EVENT_FIGS["precip"],
        f"Figure 10. Extended sensitivity group ({sensitivity_text}), development-year January-May, plotted separately: ERA5 precipitation with ERA5 850 hPa wind anomalies.",
    )

    add_heading(doc, "D20 and Equatorial u850", level=1)
    add_paragraph(
        doc,
        "D20 remains the clearest evidence for the similarity between 2015 and 2026. Both events show relative shoaling in "
        "the western Pacific and deepening in the central-eastern Pacific, consistent with a tilt or transport of warm water "
        "toward the central-eastern Pacific. The separately plotted classic and sensitivity groups show that this west-shallow/east-deep structure is "
        "not uniform across early El Niño development cases. The 2023 event has strong eastern Pacific D20 deepening but lacks the "
        "same clear western Pacific shoaling, making its subsurface structure less similar to 2015/2026.",
    )
    add_figure(doc, EVENTS["2015/16"]["d20"], "Figure 11. January-May 2015 early formation: ORAS5 equatorial D20 anomalies with ERA5 equatorial u850.")
    add_figure(doc, EVENTS["2023/24"]["d20"], "Figure 12. January-May 2023 early formation: ORAS5 equatorial D20 anomalies with ERA5 equatorial u850.")
    add_figure(doc, EVENTS["2026"]["d20"], "Figure 13. January-May 2026 early formation: ORAS5 equatorial D20 anomalies with ERA5 equatorial u850.")
    add_figure(
        doc,
        CLASSIC_EVENT_FIGS["d20"],
        f"Figure 14. Main classic strong reference ({classic_text}), development-year January-May, plotted individually: ORAS5 equatorial D20 anomalies with ERA5 equatorial u850.",
    )
    add_figure(
        doc,
        SENSITIVITY_EVENT_FIGS["d20"],
        f"Figure 15. Extended sensitivity group ({sensitivity_text}), development-year January-May, plotted separately: ORAS5 equatorial D20 anomalies with ERA5 equatorial u850.",
    )

    add_heading(doc, "PMM: North Pacific Subtropical Background", level=1)
    add_paragraph(
        doc,
        "The PMM-like EOF represents the leading coupled mode of North Pacific subtropical SST and near-surface winds after "
        "the linear influence of Niño-3.4 is removed. Its positive phase corresponds to warm northeastern Pacific SST, weakened "
        "trade winds, and reduced evaporative cooling. The joint SST+wind PMM index is positive during January-May in both "
        "2015 and 2026, whereas it is weak or negative in 2023. This indicates that the subtropical coupled state in 2026 is "
        "closer to 2015 than to 2023.",
    )
    add_figure(doc, PMM_FIG, "Figure 16. PMM-like EOF diagnostics from OISST SST and ERA5 10 m wind. EOF training period: 1982-2025.")
    add_figure(doc, PMM_FOCUS, "Figure 17. 2026 January-May PMM-like joint SST+wind and SST-only indices.")

    add_heading(doc, "Summary", level=1)
    add_paragraph(
        doc,
        "The three individual events differ mainly in the maturity and spatial organization of air-sea coupling. The 2015/16 "
        "event is the most strongly coupled case, with central-eastern Pacific SST warming, strong westerlies, positive central "
        "Pacific precipitation, and a clear west-shallow/east-deep D20 structure. The 2023/24 development phase is more strongly "
        "led by far-eastern Pacific warming and has weaker central Pacific and PMM support.",
    )
    add_paragraph(
        doc,
        "The 2026 event retains some 2023-like east-Pacific-led surface warming characteristics, but its D20, u850, and PMM "
        "signals are more similar to 2015. Relative to the two-event main classic reference, 2026 already shows a strong "
        "dynamical preconditioning signal, while the precipitation and deep-convective response still lags. Its subsequent "
        "evolution will depend on whether central Pacific precipitation and the Walker circulation continue to respond after summer.",
    )
    add_paragraph(
        doc,
        "The extended 2002/03 and 2009/10 sensitivity cases should be used only to test robustness across weaker or differently "
        "organized early development pathways. They are not members of the main classic strong reference and do not change the "
        "target-only status of 2015/16.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
