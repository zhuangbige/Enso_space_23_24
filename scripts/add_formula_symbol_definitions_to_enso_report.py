#!/usr/bin/env python3
"""Add complete symbol definitions below every methods formula group."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "统一1982_2024基准_暖核公式补充版_20260818.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "统一1982_2024基准_暖核公式与完整符号说明版_20260818.docx"
)

FONT = "PingFang SC"


DETREND_DEFINITION = (
    "式中：X(y,m)为变量X在年份y、月份m的原始月值；X′(y,m)为去除同历月平均和"
    "线性趋势后的异常；y为年份，m为历月；⟨X⟩m为1982—2024年所有m月样本的平均值；"
    "bm为仅用m月样本拟合的线性趋势斜率，单位为X的单位每年；ȳ为基准期平均年份，"
    "1982—2024年时为2003；bm(y−ȳ)为相对平均年份的趋势项；尖括号⟨·⟩表示平均。"
)

OHC_DEFINITION = (
    "式中：q′为体积热含量密度异常，在SI单位下为J m⁻³，图中换算为MJ m⁻³；"
    "ρ₀=1025 kg m⁻³为参考海水密度；cₚ=3990 J kg⁻¹ K⁻¹为海水定压比热；"
    "T′为同历月线性去趋势温度异常，用K或°C表示时数值相同；H′(B,L,t)为时刻t、"
    "纬带B、深度层L内的单位水平面积OHC异常，SI单位为J m⁻²，本报告用GJ m⁻²；"
    "λ为经度，z为向下为正的深度；T′B(λ,z,t)为纬带B内余弦纬度加权后的温度异常；"
    "∫L·dz表示在深度层L内的垂向积分；⟨·⟩λ表示在所选经度窗内的平均。"
)

WARM_CORE_DEFINITION = (
    "式中：q′j,k为5°S—5°N余弦平均后、深度层j和经度格点k处的热含量密度异常；"
    "q′⁺j,k=max(q′j,k,0)为其正值部分，冷异常的权重因此为0；j为深度层索引，k为经度格点索引；"
    "Δzj为第j层的厚度，Δxk为第k个经度格点对应的东西向宽度；aj,k=ΔzjΔxk为经度—深度"
    "剖面网格面积；Wj,k=q′⁺j,kaj,k为质心计算权重，若q′用MJ m⁻³、a用m²，则W的单位为MJ m⁻¹，"
    "但质心中只使用其相对比例；λk和zj分别为经度格点中心和深度层中心；λc和zc分别为暖核的加权"
    "质心经度和深度；∑j,k表示对120°E—80°W、0—300 m域内所有j、k网格求和；"
    "λ∈[λc−10°,λc+10°]为后续三纬带OHC计算的跟随暖核移动窗。"
)

THREE_BAND_DEFINITION = (
    "式中：HE、HS和HN分别为事件暖核λc±10°窗内、5°S—5°N赤道带、15°S—5°S南侧带和"
    "5°N—15°N北侧带的JFM平均、0—300 m单位面积OHC异常，单位为GJ m⁻²；"
    "Hoff=(HS+HN)/2为南、北两侧带OHC的算术平均；H̄3band为三纬带的面积加权平均OHC，"
    "不是整个海域的体积积分总热量；B∈{E,S,N}为纬带标记；wB=sinφ2,B−sinφ1,B为纬带B的相对球面面积权重，"
    "φ1,B和φ2,B分别为该纬带的南、北边界纬度；i为13次ENSO暖事件的事件索引；α为截距，"
    "β为H̄3band对Hoff的线性回归系数；Xi为第i次事件在扣除三纬带平均OHC线性影响后的离赤道OHC残差，"
    "单位仍为GJ m⁻²。"
)

ATMOSPHERE_MODEL_DEFINITION = (
    "式中：A为某一同历月的单个大气响应量，即GPCP赤道太平洋降水、ERA5东西海平面气压驱动"
    "或ERA5中西太平洋纬向风应力；γ0为回归截距；γ1为长期线性变化系数；"
    "year为样本年份，year−1998为以1982—2014年训练期中心年份1998为零点的时间项；"
    "Niño3.4为ENSO海温强度指标；Niño3−Niño4为EP—CP空间型态指标；两者均来自去趋势OISST异常，"
    "单位为°C；γ2和γ3分别为A对海温强度和EP—CP型态的回归系数；ε为模型未解释残差。"
)

COUPLING_DEFINITION = (
    "式中：Â为从1982—2014年同历月逐年留一训练得到的A预测值；sε为对应训练样本回归残差ε的"
    "标准差；RA=(A−Â)/sε为大气量A的交叉验证标准化残差；RP、RSLP和Rτ分别为降水、"
    "海平面气压驱动和纬向风应力的标准化残差；⟨·⟩JJA表示对6—8月取平均；"
    "CJJA,i为第i次事件的JJA三分量等权综合耦合残差，单位为z；δ0为截距，δ1为H̄3band对CJJA的回归系数；"
    "Yi为第i次事件进一步扣除JFM三纬带平均OHC线性影响后的JJA耦合残差；"
    "corr(Xi,Yi)为13次事件上Xi与Yi的Pearson相关，rpartial为该偏相关系数。CJJA或Yi>0表示大气响应"
    "强于同等SST强度和型态下的历史预期，<0表示弱于历史预期。"
)


def set_run_font(run, size: float = 9.0) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = None
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT)


def replace_paragraph_text(paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run, 9.0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.1


def make_definition(document: Document, text: str):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text), 8.8)
    return paragraph


def find_one(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def insert_before(anchor, paragraph) -> None:
    anchor._p.addprevious(paragraph._p)


def audit(source: Document, output: Document) -> None:
    if len(output.inline_shapes) != len(source.inline_shapes):
        raise ValueError("The edit changed the number of figures")
    source_equations = len(source.element.xpath(".//m:oMath"))
    output_equations = len(output.element.xpath(".//m:oMath"))
    if output_equations != source_equations:
        raise ValueError("The edit changed the native equations")
    if len(output.element.xpath(".//w:b | .//w:bCs")):
        raise ValueError("Bold formatting was introduced")
    if len(output.element.xpath(".//w:numPr")):
        raise ValueError("List numbering was introduced")
    definition_count = sum(p.text.startswith("式中：") for p in output.paragraphs)
    if definition_count != 6:
        raise ValueError(f"Expected six complete symbol-definition paragraphs, found {definition_count}")


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_document = Document(SOURCE)
    document = Document(SOURCE)

    replace_paragraph_text(
        find_one(document, "其中y和m分别表示年份和月份"),
        DETREND_DEFINITION,
    )
    replace_paragraph_text(
        find_one(document, "q′为热含量密度异常"),
        OHC_DEFINITION,
    )

    warm_definition = make_definition(document, WARM_CORE_DEFINITION)
    insert_before(
        find_one(document, "离赤道暖量与三纬带面积加权平均OHC"),
        warm_definition,
    )

    three_band_definition = make_definition(document, THREE_BAND_DEFINITION)
    insert_before(find_one(document, "横轴使用残差X"), three_band_definition)

    atmosphere_definition = make_definition(document, ATMOSPHERE_MODEL_DEFINITION)
    coupling_definition = make_definition(document, COUPLING_DEFINITION)
    final_note = find_one(document, "Â与sε均来自1982—2014年")
    insert_before(final_note, atmosphere_definition)
    insert_before(final_note, coupling_definition)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    final_document = Document(OUTPUT)
    audit(source_document, final_document)
    print(OUTPUT)
    print(f"figures: {len(final_document.inline_shapes)}")
    print(f"native equations: {len(final_document.element.xpath('.//m:oMath'))}")
    print(f"definition paragraphs: {sum(p.text.startswith('式中：') for p in final_document.paragraphs)}")
    print(f"bold tags: {len(final_document.element.xpath('.//w:b | .//w:bCs'))}")
    print(f"numbering tags: {len(final_document.element.xpath('.//w:numPr'))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
