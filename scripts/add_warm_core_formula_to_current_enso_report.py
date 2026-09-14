#!/usr/bin/env python3
"""Add only the warm-core centroid formulas to the user's current ENSO report."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "统一1982_2024基准_论文图注无粗体版_20260818.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO发展期_SSTA-OHC不同步_阶段性汇报_赤道带与纬向剖面修正版_"
    "统一1982_2024基准_暖核公式补充版_20260818.docx"
)

FONT = "PingFang SC"
PAGINATION_TAGS = (
    "keepNext",
    "keepLines",
    "pageBreakBefore",
    "widowControl",
    "suppressLineNumbers",
)


def set_font(run, size: float) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = None
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT)


def strip_pagination(paragraph) -> None:
    ppr = paragraph._p.pPr
    if ppr is None:
        return
    for tag in PAGINATION_TAGS:
        element = ppr.find(qn(f"w:{tag}"))
        if element is not None:
            ppr.remove(element)
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def math_run(text: str) -> OxmlElement:
    run = OxmlElement("m:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "18")
    properties.append(size)
    properties.append(size_cs)
    run.append(properties)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_parts(value) -> list[OxmlElement]:
    if isinstance(value, str):
        return [math_run(value)]
    if hasattr(value, "tag"):
        return [value]
    output: list[OxmlElement] = []
    for item in value:
        output.extend(math_parts(item))
    return output


def math_sub(base, subscript) -> OxmlElement:
    element = OxmlElement("m:sSub")
    base_element = OxmlElement("m:e")
    for child in math_parts(base):
        base_element.append(child)
    sub_element = OxmlElement("m:sub")
    for child in math_parts(subscript):
        sub_element.append(child)
    element.append(base_element)
    element.append(sub_element)
    return element


def math_fraction(numerator, denominator) -> OxmlElement:
    element = OxmlElement("m:f")
    numerator_element = OxmlElement("m:num")
    for child in math_parts(numerator):
        numerator_element.append(child)
    denominator_element = OxmlElement("m:den")
    for child in math_parts(denominator):
        denominator_element.append(child)
    element.append(numerator_element)
    element.append(denominator_element)
    return element


def make_lead_body(document: Document, label: str, text: str):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 0.98
    set_font(paragraph.add_run(label), 10.5)
    set_font(paragraph.add_run(text), 8.5)
    strip_pagination(paragraph)
    return paragraph


def make_body(document: Document, text: str):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.12
    set_font(paragraph.add_run(text), 9.5)
    strip_pagination(paragraph)
    return paragraph


def make_equation(document: Document, parts: Iterable):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    equation = OxmlElement("m:oMath")
    for part in parts:
        for child in math_parts(part):
            equation.append(child)
    paragraph._p.append(equation)
    strip_pagination(paragraph)
    return paragraph


def insert_after(anchor, paragraphs: list) -> None:
    current = anchor._p
    for paragraph in paragraphs:
        current.addnext(paragraph._p)
        current = paragraph._p


def audit(source: Document, output: Document) -> None:
    if len(output.inline_shapes) != len(source.inline_shapes):
        raise ValueError("The edit changed the number of figures")
    source_equations = len(source.element.xpath(".//m:oMath"))
    output_equations = len(output.element.xpath(".//m:oMath"))
    if output_equations != source_equations + 2:
        raise ValueError(
            f"Expected two added native equations: {source_equations} -> {output_equations}"
        )
    if len(output.element.xpath(".//w:b | .//w:bCs")):
        raise ValueError("Bold formatting was introduced")
    if len(output.element.xpath(".//w:numPr")):
        raise ValueError("List numbering was introduced")
    if sum(p.text.startswith("暖核质心与移动窗口。") for p in output.paragraphs) != 1:
        raise ValueError("Warm-core formula heading is missing or duplicated")


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_document = Document(SOURCE)
    document = Document(SOURCE)
    targets = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("地图范围为30°S—30°N")
    ]
    if len(targets) != 1:
        raise ValueError(f"Expected one region-definition paragraph, found {len(targets)}")

    introduction = make_lead_body(
        document,
        "暖核质心与移动窗口。",
        "基于5°S—5°N余弦平均剖面，以正热含量密度异常与网格面积加权；"
        "j、k分别表示深度与经度，无额外峰值阈值。",
    )
    weight_equation = make_equation(
        document,
        [
            math_sub("q′⁺", "j,k"),
            " = max(",
            math_sub("q′", "j,k"),
            ", 0),    ",
            math_sub("a", "j,k"),
            " = Δ",
            math_sub("z", "j"),
            " Δ",
            math_sub("x", "k"),
            ",    ",
            math_sub("W", "j,k"),
            " = ",
            math_sub("q′⁺", "j,k"),
            math_sub("a", "j,k"),
            ".",
        ],
    )
    centroid_equation = make_equation(
        document,
        [
            math_sub("λ", "c"),
            " = (",
            math_sub("∑", "j,k"),
            " ",
            math_sub("λ", "k"),
            " ",
            math_sub("W", "j,k"),
            ") / (",
            math_sub("∑", "j,k"),
            " ",
            math_sub("W", "j,k"),
            "),        ",
            math_sub("z", "c"),
            " = (",
            math_sub("∑", "j,k"),
            " ",
            math_sub("z", "j"),
            " ",
            math_sub("W", "j,k"),
            ") / (",
            math_sub("∑", "j,k"),
            " ",
            math_sub("W", "j,k"),
            "),        ",
            "λ ∈ [",
            math_sub("λ", "c"),
            " − 10°, ",
            math_sub("λ", "c"),
            " + 10°].",
        ],
    )
    insert_after(
        targets[0],
        [
            introduction,
            weight_equation,
            centroid_equation,
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    final_document = Document(OUTPUT)
    audit(source_document, final_document)
    print(OUTPUT)
    print(f"figures: {len(final_document.inline_shapes)}")
    print(f"native equations: {len(final_document.element.xpath('.//m:oMath'))}")
    print(f"bold tags: {len(final_document.element.xpath('.//w:b | .//w:bCs'))}")
    print(f"numbering tags: {len(final_document.element.xpath('.//w:numPr'))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
