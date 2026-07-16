from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/Users/daiyue/Documents/Maryland")
OUT_DIR = ROOT / "outputs/research_2026/reports"

FIG = {
    "fep": ROOT / "outputs/research_2026/figures/fig1_2026_sst_fep_residual.png",
    "sst_precip": ROOT / "outputs/research_2026/spatial_meteorology/fig4_2026_sst_precip_spatial.png",
    "olr": ROOT / "outputs/research_2026/spatial_meteorology/fig6_2026_olr_convection_spatial.png",
    "era5_surface": ROOT / "outputs/research_2026/era5_spatial/fig8_2026_era5_surface_wind_pressure.png",
    "era5_precip": ROOT / "outputs/research_2026/era5_spatial/fig9_2026_era5_precip_u850_omega.png",
    "era5_hov": ROOT / "outputs/research_2026/era5_spatial/fig10_2026_era5_equatorial_hovmoller.png",
}

COLORS = {
    "ink": "1C1B18",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "teal": "2C7773",
    "rust": "C85B45",
    "light_blue": "E8EEF5",
    "light_teal": "E7F1EF",
    "light_rust": "F5E4DE",
    "table_header": "F2F4F7",
    "line": "D8D1C2",
    "muted": "5B5A55",
}


def ensure_inputs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for path in FIG.values():
        if not path.exists():
            raise FileNotFoundError(path)


def set_run_font(run, lang: str, size: float | None = None, bold: bool | None = None, color: str | None = None):
    font = "PingFang SC" if lang == "cn" else "Calibri"
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, lang: str, size: float, color: str = COLORS["ink"], bold: bool = False):
    font = "PingFang SC" if lang == "cn" else "Calibri"
    style.font.name = font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def style_document(doc: Document, lang: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], lang, 11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    set_style_font(styles["Heading 1"], lang, 16, COLORS["blue"], True)
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    set_style_font(styles["Heading 2"], lang, 13, COLORS["blue"], True)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)
    set_style_font(styles["Heading 3"], lang, 12, COLORS["dark_blue"], True)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ENSO/FEP 2026 diagnostic | Yue Dai | 2026-06-22")
    set_run_font(run, lang, 8.5, color=COLORS["muted"])


def add_title(doc: Document, title: str, subtitle: str, lang: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_run_font(run, lang, 22, True, COLORS["ink"])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(subtitle)
    set_run_font(run, lang, 12, False, COLORS["muted"])
    add_horizontal_rule(doc, COLORS["line"])


def add_horizontal_rule(doc: Document, color: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(table, top=80, bottom=80, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for key, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]], widths: list[int], lang: str) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_fixed(table, widths)
    set_cell_margins(table)

    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, lang, 9.2 if len(text) > 90 else 10, bold=(r == 0), color=COLORS["ink"])
            if r == 0:
                shade_cell(cell, COLORS["table_header"])
            elif r % 2 == 1:
                shade_cell(cell, "F8F6F0")
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str, lang: str, fill: str = COLORS["light_teal"]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_fixed(table, [9360])
    set_cell_margins(table, top=120, bottom=120, start=180, end=180)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, lang, 11.5, True, COLORS["teal"])
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, lang, 10.5, False, COLORS["ink"])
    doc.add_paragraph()


def add_bullet(doc: Document, text: str, lang: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, lang, 10.8)


def add_caption(doc: Document, text: str, lang: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, lang, 9.5, False, COLORS["muted"])


def add_figure(doc: Document, image_path: Path, caption: str, note: str, lang: str, width: float = 6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption, lang)
    add_callout(doc, "图件解读" if lang == "cn" else "Figure reading", note, lang, COLORS["light_blue"])


def add_source_text(doc: Document, text: str, lang: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, lang, 9.5, False, COLORS["muted"])


def build_cn() -> Path:
    doc = Document()
    style_document(doc, "cn")
    add_title(
        doc,
        "2026 年当前 ENSO/FEP 事件的实时气象诊断小报告",
        "基于 ERA5、OISST、GPCP 和 CPC OLR 的阶段性结果 | 2026-06-22",
        "cn",
    )

    doc.add_heading("一、当前核心判断", level=1)
    add_callout(
        doc,
        "一句话结论",
        "2026 当前事件不能简单写成“海温强、大气完全不响应”。更准确的说法是：海温/FEP 先行增暖，1-2 月风压和降水明显不同步，3-5 月风压分支开始滞后恢复，而降水/对流仍最滞后、最不稳定。",
        "cn",
    )
    for item in [
        "海温端：OISST 显示 2026 年 2 月以后 FEP 和东太平洋暖异常快速增强。",
        "风压端：ERA5 显示 1-2 月中太平洋仍偏东风、SOI-like 压差偏 La Niña-like；但 MAM 和 Apr-May 风场与海平面气压开始转向 El Niño-like。",
        "降水/对流端：GPCP 与 CPC OLR 均显示响应中心偏西、偏北或空间破碎，赤道中东太平洋尚未形成稳定连续的强对流/降水响应。",
        "omega850 解释：omega850 是 850 hPa 垂直压力速度，负值表示偏上升，正值表示偏下沉；本报告中正的 omega850 异常意味着相对气候态更偏下沉，不利于深对流增强。",
    ]:
        add_bullet(doc, item, "cn")

    doc.add_heading("二、研究背景与问题来源", level=1)
    p = doc.add_paragraph()
    r = p.add_run(
        "前期工作从 Peng et al. (2025) 关于 2023/24 强厄尔尼诺的研究出发：该事件海洋动力增暖很强，但同期风场、气压、降水和遥相关响应偏弱。当前 2026 事件提供了一个实时对照案例：如果东太平洋/FEP 海温再次率先增暖，大气分支是否同步跟上？"
    )
    set_run_font(r, "cn", 11)
    p = doc.add_paragraph()
    r = p.add_run(
        "因此，本阶段不是判断 2026 是否已经成为强厄尔尼诺，而是检查“海温、风压、降水/对流是否同步”的早期信号。"
    )
    set_run_font(r, "cn", 11)

    doc.add_heading("三、数据链与异常定义", level=1)
    rows = [
        ["变量", "数据集", "当前用途", "可用时段"],
        ["SST 海表温度", "NOAA OISST v2.1", "判断 FEP、Niño 区和热带海盆暖异常", "到 2026-05"],
        ["降水", "GPCP monthly precipitation", "检验降水是否与 SST 同步增强", "到 2026-05"],
        ["对流", "CPC Blended OLR", "独立检验深对流位置；OLR 负值代表对流增强", "到 2026-06-15"],
        ["风压", "ERA5 single levels", "10 m 风、平均海平面气压、总降水", "1982-2022 与 2026-01 至 2026-05"],
        ["低层环流", "ERA5 pressure levels", "850 hPa 风和 omega850", "1982-2022 与 2026-01 至 2026-05"],
    ]
    add_table(doc, rows, [1600, 2500, 3500, 1760], "cn")
    p = doc.add_paragraph()
    r = p.add_run(
        "异常定义：2026 anomaly = 2026 某月或某时段平均值 − 1982-2022 同月气候态。MAM 表示 3-5 月，Apr-May 表示 4-5 月；完整 AMJ 仍需等待 2026 年 6 月月平均数据。"
    )
    set_run_font(r, "cn", 10.8, False, COLORS["muted"])

    doc.add_heading("四、图件证据链", level=1)
    add_figure(
        doc,
        FIG["fep"],
        "图 1. OISST 指数与 FEP residual。数据：NOAA OISST v2.1 月平均海表温度异常。",
        "该图确认海温端已经先行变化：2026 年 2 月后 FEP 和 Niño1+2 附近快速转暖，FEP residual 为正，表示沿岸远东太平洋暖异常相对热带太平洋整体更突出。",
        "cn",
    )
    add_figure(
        doc,
        FIG["sst_precip"],
        "图 2. OISST 海温异常与 GPCP 降水异常空间分布。",
        "MAM 和 Apr-May 的东太平洋/FEP 偏暖明显，但 GPCP 降水增强中心偏向西太平洋和赤道以北；赤道中东太平洋没有形成连续强正降水带。这说明 SST 增暖与降水响应之间仍存在空间错位。",
        "cn",
    )
    add_figure(
        doc,
        FIG["olr"],
        "图 3. CPC Blended OLR 对流异常。",
        "OLR 负异常表示深对流增强，正异常表示对流偏弱。2026 春季对流增强更偏向西太平洋/中太平洋北侧，而不是沿赤道中东太平洋大范围增强。这与 GPCP 的“降水响应错位”判断一致。",
        "cn",
    )
    add_figure(
        doc,
        FIG["era5_surface"],
        "图 4. ERA5 10 m 风异常与平均海平面气压异常。",
        "ERA5 补齐了前一版 NCEP 只到 2 月的限制。MAM 与 Apr-May 显示赤道中西太平洋逐渐出现偏西风异常，中东太平洋气压偏低，说明风压分支在春季开始恢复到更 El Niño-like 的方向。",
        "cn",
    )
    add_figure(
        doc,
        FIG["era5_precip"],
        "图 5. ERA5 总降水、u850 和 omega850 异常。",
        "u850 在 MAM/Apr-May 已经转为偏西风，但降水增强中心仍偏西/偏北，中太平洋 omega850 没有表现为稳定强上升响应。这里 omega850 的负值才代表上升，正值代表下沉。",
        "cn",
    )
    add_figure(
        doc,
        FIG["era5_hov"],
        "图 6. ERA5 赤道经度-时间 Hovmoller 图。",
        "1-2 月赤道中太平洋仍为偏东风异常；3-5 月中西太平洋风场逐渐转偏西。降水在 1-4 月仍有大范围偏弱，5 月才出现更多局地恢复。这是“阶段性耦合恢复”的直接证据。",
        "cn",
    )

    doc.add_page_break()
    doc.add_heading("五、关键数值", level=1)
    rows = [
        ["时段", "指标", "异常值", "含义"],
        ["Jan", "中太平洋 u10", "-1.15 m/s", "年初仍是偏东风，不是典型 El Niño 偏西风响应"],
        ["Feb", "中太平洋 u850", "-2.95 m/s", "850 hPa 偏东风很强，低层大气尚未跟上海温"],
        ["Feb", "SOI-like MSLP", "+2.63 hPa", "气压梯度偏 La Niña-like"],
        ["MAM", "CWP u850", "+0.98 m/s", "春季中西太平洋低层风开始偏西"],
        ["MAM", "中太平洋降水", "-0.87 mm/day", "降水仍弱于气候态"],
        ["Apr-May", "CWP u850", "+1.47 m/s", "风场恢复更明显"],
        ["Apr-May", "SOI-like MSLP", "-1.16 hPa", "气压分支也转向 El Niño-like"],
        ["Apr-May", "中太平洋降水", "-0.26 mm/day", "降水接近恢复，但仍不是强正异常"],
    ]
    add_table(doc, rows, [1200, 2300, 1600, 4260], "cn")

    doc.add_heading("六、阶段性解释与边界", level=1)
    for item in [
        "可以说：2026 年初存在明显海气不同步，尤其表现为海温转暖时风场和降水仍偏 La Niña-like 或弱响应。",
        "也可以说：ERA5 1-5 月显示春季风压分支开始恢复，因此不能再用“持续完全脱耦”概括 2026 当前事件。",
        "当前最值得继续追踪的是降水、OLR 与垂直运动，因为它们相对于风压更滞后，也更容易受到背景态和季节内扰动影响。",
        "当前还不能说：2026 已经形成强 El Niño，或已经完整复现 2023/24 的海气脱耦。完整 AMJ 和 JJA 结论需要后续月平均数据。",
    ]:
        add_bullet(doc, item, "cn")

    doc.add_heading("七、下一步工作", level=1)
    for item in [
        "等待 2026 年 6 月 ERA5 月平均，正式构造 AMJ 2026，而不是用 MAM 或 Apr-May 近似。",
        "把 ERA5 风压结果接入之前 2023/24 的 response-ratio / observed-vs-expected 框架，形成 2023 与 2026 的同口径对比。",
        "用同一套 ERA5 变量回算 2023/24 AMJ/JJA，避免不同资料集造成的解释差异。",
        "下载或计算 200 hPa divergence / velocity potential，补 Walker circulation 上层分支。",
        "继续使用 GPCP 与 OLR 双约束，检查降水/对流中心是否从偏北/偏西转向赤道中东太平洋。",
    ]:
        add_bullet(doc, item, "cn")

    doc.add_heading("八、数据来源与本地文件", level=1)
    for text in [
        "ERA5 monthly single levels: https://cds.climate.copernicus.eu/datasets/reanalysis-era5-single-levels-monthly-means",
        "ERA5 monthly pressure levels: https://cds.climate.copernicus.eu/datasets/reanalysis-era5-pressure-levels-monthly-means",
        "NOAA OISST v2.1: https://www.ncei.noaa.gov/products/optimum-interpolation-sst",
        "GPCP monthly precipitation: https://psl.noaa.gov/data/gridded/data.gpcp.html",
        "CPC Blended OLR: https://psl.noaa.gov/data/gridded/data.cpc_blended_olr-1deg.html",
        "本地报告：reports/2026_ERA5再分析气象图说明.md",
        "本地表格：data/derived/research_2026/current_2026_era5_region_summary.csv",
    ]:
        add_source_text(doc, text, "cn")

    out = OUT_DIR / "2026_ENSO_FEP_ERA5_Current_Diagnosis_Report_CN.docx"
    doc.save(out)
    return out


def build_en() -> Path:
    doc = Document()
    style_document(doc, "en")
    add_title(
        doc,
        "Real-Time Meteorological Diagnosis of the 2026 ENSO/FEP Event",
        "A short report based on ERA5, OISST, GPCP, and CPC OLR | 22 June 2026",
        "en",
    )

    doc.add_heading("1. Current Takeaway", level=1)
    add_callout(
        doc,
        "One-sentence conclusion",
        "The current 2026 case should not be described as “strong SST with no atmospheric response.” A more accurate reading is that SST/FEP warming leads first, wind-pressure and rainfall are clearly asynchronous in January-February, wind-pressure recovers with lag in March-May, and rainfall/convection remains the slowest and least stable branch.",
        "en",
    )
    for item in [
        "Ocean branch: OISST shows rapid FEP and eastern-Pacific warming after February 2026.",
        "Wind-pressure branch: ERA5 shows central-Pacific easterlies and a La Nina-like SOI-like pressure contrast in January-February, but MAM and Apr-May shift toward a more El Nino-like wind-pressure response.",
        "Rainfall/convection branch: GPCP and CPC OLR both show a displaced, fragmented response, with enhanced rainfall/convection located more westward or north of the equator rather than along the equatorial central-eastern Pacific.",
        "omega850 note: omega850 is vertical pressure velocity at 850 hPa. Negative values indicate anomalous upward motion, while positive values indicate anomalous downward motion.",
    ]:
        add_bullet(doc, item, "en")

    doc.add_heading("2. Research Context", level=1)
    p = doc.add_paragraph()
    r = p.add_run(
        "This work follows the 2023/24 El Nino question raised by Peng et al. (2025): the event was ocean-dynamically strong, but the atmospheric response in wind, pressure, rainfall, and teleconnections was weaker than expected. The ongoing 2026 case provides a real-time comparison: if eastern-Pacific/FEP SST warms first again, does the atmosphere respond synchronously?"
    )
    set_run_font(r, "en", 11)
    p = doc.add_paragraph()
    r = p.add_run(
        "At this stage, the goal is not to classify 2026 as a strong El Nino. The goal is to diagnose whether SST, wind-pressure, rainfall, and convection evolve at the same time."
    )
    set_run_font(r, "en", 11)

    doc.add_heading("3. Data Chain and Anomaly Definition", level=1)
    rows = [
        ["Variable", "Dataset", "Use in this report", "Availability"],
        ["SST", "NOAA OISST v2.1", "Diagnose FEP, Nino-region, and tropical-basin warming", "Through May 2026"],
        ["Rainfall", "GPCP monthly precipitation", "Test whether rainfall strengthens with SST", "Through May 2026"],
        ["Convection", "CPC Blended OLR", "Independent deep-convection constraint; negative OLR means stronger convection", "Through 15 June 2026"],
        ["Wind-pressure", "ERA5 single levels", "10 m wind, mean sea-level pressure, total precipitation", "1982-2022 and Jan-May 2026"],
        ["Low-level circulation", "ERA5 pressure levels", "850 hPa wind and omega850", "1982-2022 and Jan-May 2026"],
    ]
    add_table(doc, rows, [1500, 2500, 3600, 1760], "en")
    p = doc.add_paragraph()
    r = p.add_run(
        "Anomaly definition: 2026 anomaly = 2026 monthly or period mean minus the 1982-2022 same-month climatology. MAM means March-May; Apr-May means April-May. Full AMJ still requires the June 2026 monthly mean."
    )
    set_run_font(r, "en", 10.8, False, COLORS["muted"])

    doc.add_heading("4. Figure-Based Evidence Chain", level=1)
    add_figure(
        doc,
        FIG["fep"],
        "Figure 1. OISST indices and FEP residual. Dataset: NOAA OISST v2.1 monthly SST anomaly.",
        "The ocean branch moves first: after February 2026, FEP and Nino1+2 warming intensifies rapidly. The positive FEP residual indicates that coastal far-eastern Pacific warming is stronger than the broader tropical-Pacific background.",
        "en",
    )
    add_figure(
        doc,
        FIG["sst_precip"],
        "Figure 2. OISST SST anomaly and GPCP precipitation anomaly.",
        "MAM and Apr-May show clear FEP/eastern-Pacific warming. However, GPCP rainfall enhancement is displaced toward the western Pacific and north of the equator, and a continuous positive equatorial central-eastern Pacific rainfall band is absent.",
        "en",
    )
    add_figure(
        doc,
        FIG["olr"],
        "Figure 3. CPC Blended OLR convection anomaly.",
        "Negative OLR indicates stronger deep convection, while positive OLR indicates suppressed convection. Spring 2026 convection is enhanced more over the western/northern central Pacific than over the equatorial central-eastern Pacific, supporting the GPCP-based displaced-rainfall interpretation.",
        "en",
    )
    add_figure(
        doc,
        FIG["era5_surface"],
        "Figure 4. ERA5 10 m wind anomaly and mean sea-level pressure anomaly.",
        "ERA5 extends the wind-pressure diagnosis beyond the earlier NCEP-only January-February view. In MAM and Apr-May, central-western equatorial Pacific winds turn westerly, and pressure lowers over the central-eastern tropical Pacific, indicating a lagged recovery toward an El Nino-like state.",
        "en",
    )
    add_figure(
        doc,
        FIG["era5_precip"],
        "Figure 5. ERA5 total precipitation, u850, and omega850 anomalies.",
        "u850 becomes westerly in MAM/Apr-May, but rainfall remains displaced westward/northward, and central-Pacific omega850 does not show a stable strong upward-motion response. For omega850, negative values indicate upward motion and positive values indicate downward motion.",
        "en",
    )
    add_figure(
        doc,
        FIG["era5_hov"],
        "Figure 6. ERA5 equatorial longitude-time Hovmoller diagnostics.",
        "January-February shows clear central-Pacific easterly anomalies. From March to May, central-western Pacific winds gradually turn westerly. Rainfall remains broadly weak from January to April and becomes more locally positive only in May. This is direct evidence for staged coupling recovery.",
        "en",
    )

    doc.add_page_break()
    doc.add_heading("5. Key Numbers", level=1)
    rows = [
        ["Period", "Metric", "Anomaly", "Interpretation"],
        ["Jan", "Central-Pacific u10", "-1.15 m/s", "Early easterlies; not the canonical El Nino westerly response"],
        ["Feb", "Central-Pacific u850", "-2.95 m/s", "Strong lower-tropospheric easterly anomaly"],
        ["Feb", "SOI-like MSLP", "+2.63 hPa", "Pressure gradient remains La Nina-like"],
        ["MAM", "CWP u850", "+0.98 m/s", "Spring low-level wind begins to recover westerly"],
        ["MAM", "Central-Pacific rainfall", "-0.87 mm/day", "Rainfall remains below climatology"],
        ["Apr-May", "CWP u850", "+1.47 m/s", "Wind recovery becomes clearer"],
        ["Apr-May", "SOI-like MSLP", "-1.16 hPa", "Pressure branch shifts toward El Nino-like"],
        ["Apr-May", "Central-Pacific rainfall", "-0.26 mm/day", "Rainfall approaches recovery but is not strongly positive"],
    ]
    add_table(doc, rows, [1200, 2300, 1600, 4260], "en")

    doc.add_heading("6. Interpretation and Boundaries", level=1)
    for item in [
        "It is reasonable to state that early 2026 shows clear air-sea asynchrony, especially when SST warms while wind and rainfall remain La Nina-like or weak.",
        "ERA5 through May shows spring recovery of the wind-pressure branch, so the current case should not be summarized as persistent full decoupling.",
        "Rainfall, OLR, and vertical motion remain the most important branches to track because they lag behind wind-pressure recovery and are more sensitive to background state and intraseasonal variability.",
        "It is not yet justified to claim that 2026 has become a strong El Nino or that it fully reproduces the 2023/24 decoupling. Full AMJ and JJA diagnostics require later monthly means.",
    ]:
        add_bullet(doc, item, "en")

    doc.add_heading("7. Next Steps", level=1)
    for item in [
        "Wait for the June 2026 ERA5 monthly mean and construct formal AMJ 2026 diagnostics.",
        "Insert ERA5 wind-pressure diagnostics into the previous 2023/24 response-ratio / observed-versus-expected framework for same-metric comparison.",
        "Recalculate 2023/24 AMJ/JJA using the same ERA5 variables to avoid dataset-dependent interpretation.",
        "Add 200 hPa divergence or velocity potential to diagnose the upper branch of the Walker circulation.",
        "Keep both GPCP and OLR constraints to test whether rainfall/convection shifts from west/north displacement toward the equatorial central-eastern Pacific.",
    ]:
        add_bullet(doc, item, "en")

    doc.add_heading("8. Data Sources and Local Files", level=1)
    for text in [
        "ERA5 monthly single levels: https://cds.climate.copernicus.eu/datasets/reanalysis-era5-single-levels-monthly-means",
        "ERA5 monthly pressure levels: https://cds.climate.copernicus.eu/datasets/reanalysis-era5-pressure-levels-monthly-means",
        "NOAA OISST v2.1: https://www.ncei.noaa.gov/products/optimum-interpolation-sst",
        "GPCP monthly precipitation: https://psl.noaa.gov/data/gridded/data.gpcp.html",
        "CPC Blended OLR: https://psl.noaa.gov/data/gridded/data.cpc_blended_olr-1deg.html",
        "Local report: reports/2026_ERA5再分析气象图说明.md",
        "Local table: data/derived/research_2026/current_2026_era5_region_summary.csv",
    ]:
        add_source_text(doc, text, "en")

    out = OUT_DIR / "2026_ENSO_FEP_ERA5_Current_Diagnosis_Report_EN.docx"
    doc.save(out)
    return out


def main() -> None:
    ensure_inputs()
    cn = build_cn()
    en = build_en()
    print(cn)
    print(en)


if __name__ == "__main__":
    main()
