#!/usr/bin/env python3
"""Remove the cross-event correlation analysis from the English ENSO report."""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization_No_Correlation.docx"
)


REPLACEMENTS = {
    1: (
        "This report compares the January-June evolution of sea surface "
        "temperature, upper-ocean heat storage, and wind fields in 1997, "
        "2015, and 2023."
    ),
    5: (
        "Wind fields. Spatial maps are overlaid with ERA5 850-hPa wind "
        "anomalies, defined using same-calendar-month linear detrending over "
        "1982-2024."
    ),
    6: (
        "Events and time windows. The case studies are 1997, 2015, and 2023, "
        "with January-June used to compare their development-phase evolution."
    ),
    15: (
        "Maps cover 30°S-30°N, 120°E-80°W; FEP is defined as 5°S-5°N, "
        "90°W-80°W. Longitude-depth sections are cosine-weighted averages over "
        "5°S-5°N, whereas latitude-depth sections are averaged over the fixed "
        "full-longitude range 120°E-80°W."
    ),
    73: (
        "Thus, the distinctive feature of 2023 was not merely a southward "
        "displacement of heat. Rather, heat was not concentrated near the "
        "equatorial warm core but spread to both sides of the equator, "
        "especially southward, producing a broader and deeper heat "
        "distribution. This configuration may reduce the fraction of heat in "
        "the equatorial surface layer that can directly participate in the "
        "SST-convection-wind feedback."
    ),
    81: "The analysis supports two conclusions:",
}


# Original paragraph indices. These include formula/image-only paragraphs whose
# visible text is empty in python-docx.
DELETE_INDICES = (
    {13}
    | set(range(16, 34))
    | set(range(74, 81))
    | {84}
)


def replace_text_preserving_first_run(paragraph, text: str) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    original_paragraphs = list(document.paragraphs)
    if len(original_paragraphs) != 85:
        raise ValueError(
            f"Expected 85 source paragraphs, found {len(original_paragraphs)}"
        )

    for index, replacement in REPLACEMENTS.items():
        replace_text_preserving_first_run(original_paragraphs[index], replacement)
    for index in sorted(DELETE_INDICES, reverse=True):
        delete_paragraph(original_paragraphs[index])

    document.save(OUTPUT)

    check = Document(OUTPUT)
    visible_text = "\n".join(paragraph.text for paragraph in check.paragraphs)
    prohibited = (
        "Cross-Event",
        "cross-event",
        "partial r",
        "permutation p",
        "FDR q",
        "Fig. 8",
        "13 ENSO warm events",
        "JJA Pattern-Adjusted Composite Coupling Residual",
        "Off-Equatorial Heat Content and the Three-Band",
    )
    remaining = [term for term in prohibited if term in visible_text]
    if remaining:
        raise ValueError(f"Correlation-analysis text remains: {remaining}")
    if len(check.inline_shapes) != 10:
        raise ValueError(
            f"Expected 10 retained figures after removing Fig. 8, found "
            f"{len(check.inline_shapes)}"
        )
    if "The analysis supports two conclusions:" not in visible_text:
        raise ValueError("Two-conclusion transition was not retained")
    print(OUTPUT)
    print(f"paragraphs={len(check.paragraphs)} inline_shapes={len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
