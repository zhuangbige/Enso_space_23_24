#!/usr/bin/env python3
"""Add concise strength-class and pattern-similarity material to the ENSO report."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization_"
    "Strength_Profile_2026_Revision.docx"
)
OUTPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO_Development_Phase_SSTA_OHC_Desynchronization_"
    "Strength_Profile_2026_Pattern_Similarity_Revision.docx"
)
PATTERN_FIGURE = (
    ROOT
    / "outputs/research_2026/enso_strength_vertical_profiles_2026/"
    "fig_2026_vertical_pattern_similarity_monthly_only.png"
)


def first_text_run(paragraph: Paragraph):
    for run in paragraph.runs:
        if run.text:
            return run
    return paragraph.runs[0] if paragraph.runs else None


def replace_text(paragraph: Paragraph, text: str, run_template=None) -> Paragraph:
    if run_template is None:
        run_template = first_text_run(paragraph)
    rpr = (
        deepcopy(run_template._r.rPr)
        if run_template is not None and run_template._r.rPr is not None
        else None
    )
    paragraph.clear()
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return paragraph


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix!r}")


def paragraph_index(paragraphs: list[Paragraph], target: Paragraph) -> int:
    for index, paragraph in enumerate(paragraphs):
        if paragraph._p is target._p:
            return index
    raise ValueError("Paragraph is not in the document")


def clone_after(
    reference: Paragraph, template: Paragraph, text: str, run_template=None
) -> Paragraph:
    element = deepcopy(template._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    return replace_text(paragraph, text, run_template or first_text_run(template))


def empty_clone_after(reference: Paragraph, template: Paragraph) -> Paragraph:
    element = deepcopy(template._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    reference._p.addnext(element)
    return Paragraph(element, reference._parent)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def strip_pagination_properties(paragraph: Paragraph) -> None:
    properties = paragraph._p.pPr
    if properties is None:
        return
    for tag in ("keepNext", "keepLines", "pageBreakBefore", "widowControl"):
        for element in list(properties.findall(qn(f"w:{tag}"))):
            properties.remove(element)


def strip_manual_breaks(paragraph: Paragraph) -> None:
    for break_element in list(paragraph._p.xpath(".//w:br")):
        parent = break_element.getparent()
        if parent is not None:
            parent.remove(break_element)


def math_run(text: str) -> OxmlElement:
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_parts(value) -> list[OxmlElement]:
    if isinstance(value, str):
        return [math_run(value)]
    if hasattr(value, "tag"):
        return [value]
    output: list[OxmlElement] = []
    for item in value:
        output.extend(math_parts(item))
    return output


def math_sub(base, subscript) -> OxmlElement:
    element = OxmlElement("m:sSub")
    base_element = OxmlElement("m:e")
    for child in math_parts(base):
        base_element.append(child)
    sub_element = OxmlElement("m:sub")
    for child in math_parts(subscript):
        sub_element.append(child)
    element.append(base_element)
    element.append(sub_element)
    return element


def math_sup(base, superscript) -> OxmlElement:
    element = OxmlElement("m:sSup")
    base_element = OxmlElement("m:e")
    for child in math_parts(base):
        base_element.append(child)
    superscript_element = OxmlElement("m:sup")
    for child in math_parts(superscript):
        superscript_element.append(child)
    element.append(base_element)
    element.append(superscript_element)
    return element


def math_fraction(numerator, denominator) -> OxmlElement:
    element = OxmlElement("m:f")
    numerator_element = OxmlElement("m:num")
    for child in math_parts(numerator):
        numerator_element.append(child)
    denominator_element = OxmlElement("m:den")
    for child in math_parts(denominator):
        denominator_element.append(child)
    element.append(numerator_element)
    element.append(denominator_element)
    return element


def math_radical(base) -> OxmlElement:
    element = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    hide_degree = OxmlElement("m:degHide")
    hide_degree.set(qn("m:val"), "1")
    properties.append(hide_degree)
    degree = OxmlElement("m:deg")
    base_element = OxmlElement("m:e")
    for child in math_parts(base):
        base_element.append(child)
    element.append(properties)
    element.append(degree)
    element.append(base_element)
    return element


def equation_after(reference: Paragraph, template: Paragraph, parts) -> Paragraph:
    paragraph = empty_clone_after(reference, template)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    strip_pagination_properties(paragraph)
    equation_paragraph = OxmlElement("m:oMathPara")
    equation = OxmlElement("m:oMath")
    for part in math_parts(parts):
        equation.append(part)
    equation_paragraph.append(equation)
    paragraph._p.append(equation_paragraph)
    return paragraph


def insert_methods(document: Document) -> None:
    replace_text(
        document.paragraphs[1],
        "This report compares January-June SSTA, upper-ocean heat storage, and winds "
        "in 1997, 2015, and 2023, then uses strength-conditioned vertical profiles "
        "to assess whether 2026 follows a strong or moderate El Niño pathway.",
    )
    events = find_paragraph(document, "Events and time windows.")
    heading_template = find_paragraph(
        document, "Same-Calendar-Month Linearly Detrended Anomalies"
    )
    body_template = find_paragraph(document, "Here, X(y,m)")
    equation_template = document.paragraphs[
        paragraph_index(document.paragraphs, heading_template) + 1
    ]

    replace_text(
        events,
        "Events and time windows. The case studies are 1997, 2015, and 2023. "
        "The strength composites use developing-year JFM and AMJ, and the 2026 "
        "pattern comparison is evaluated monthly from January through June.",
    )

    reference = clone_after(
        events,
        heading_template,
        "Strength Classes and Spatial Pattern Correlation",
    )
    strip_pagination_properties(reference)
    reference = clone_after(
        reference,
        body_template,
        "Class intensity uses the peak same-calendar-month detrended OISST Niño3.4 "
        "anomaly during the first development winter:",
    )
    reference = equation_after(
        reference,
        equation_template,
        [
            math_sub("A", "y"),
            "=",
            math_sub("max", "m∈DJFᵧ"),
            math_sub("Niño3.4′", "y,m"),
        ],
    )
    reference = equation_after(
        reference,
        equation_template,
        [
            math_sub("C", "y"),
            ":  strong if ",
            math_sub("A", "y"),
            "≥1.5°C;   moderate if 0.5°C≤",
            math_sub("A", "y"),
            "<1.5°C.",
        ],
    )
    reference = clone_after(
        reference,
        body_template,
        "Retained strong years are 1982, 1991, 1997, 2009, 2015, and 2023; moderate "
        "years are 1986, 1994, and 2002, with 1986 classified by its first winter. "
        "For period p, xᵢ is the 2026 GODAS field and yᵢ the SODA3 class composite "
        "on Ω (5°S-5°N, 121°E-81°W, 0-300 m):",
    )

    reference = equation_after(
        reference,
        equation_template,
        [
            math_sub("r", "c,p"),
            "=",
            math_fraction(
                [
                    math_sub("Σ", "i∈Ω"),
                    "(",
                    math_sub("x", "i"),
                    "−x̄)(",
                    math_sub("y", "i"),
                    "−ȳ)",
                ],
                math_radical(
                    [
                        math_sub("Σ", "i∈Ω"),
                        math_sup(["(", math_sub("x", "i"), "−x̄)"], "2"),
                        math_sub("Σ", "i∈Ω"),
                        math_sup(["(", math_sub("y", "i"), "−ȳ)"], "2"),
                    ]
                ),
            ),
        ],
    )
    reference = clone_after(
        reference,
        body_template,
        "The valid-cell correlation measures spatial pattern rather than amplitude.",
    )
    strip_pagination_properties(reference)

    # Keep the method page compact enough that the existing event-section page break
    # does not strand the final OHC-definition lines on an otherwise empty page.
    ohc_heading = find_paragraph(document, "Heat-Content Density and OHC per Unit Area")
    ohc_definition = document.paragraphs[
        paragraph_index(document.paragraphs, ohc_heading) + 2
    ]
    replace_text(
        ohc_definition,
        "Here, q′ is volumetric heat-content density (J m⁻³; shown in MJ m⁻³); "
        "ρ₀=1025 kg m⁻³ and cₚ=3990 J kg⁻¹ K⁻¹; T′ is the detrended "
        "temperature anomaly; and H′B,L,t is OHC per unit area (J m⁻²; shown in GJ "
        "m⁻²). λ is longitude, z is depth positive downward, T′B is the cosine-latitude-"
        "weighted band mean, and ⟨·⟩λ is the longitude-window mean.",
        first_text_run(body_template),
    )
    strip_pagination_properties(ohc_definition)


def add_pattern_figure_after(reference: Paragraph, image_template: Paragraph) -> Paragraph:
    if not PATTERN_FIGURE.exists():
        raise FileNotFoundError(PATTERN_FIGURE)
    paragraph = empty_clone_after(reference, image_template)
    strip_pagination_properties(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(PATTERN_FIGURE), width=Inches(6.15))
    shape._inline.docPr.set("name", "Figure 9 monthly vertical pattern similarity")
    shape._inline.docPr.set(
        "descr",
        "January-June spatial pattern correlations between the 2026 vertical "
        "temperature anomaly and historical strong- and moderate-event composites",
    )
    return paragraph


def revise_strength_results(document: Document) -> None:
    heading = find_paragraph(
        document, "Strength-Conditioned Vertical Structure and the 2026 Development Path"
    )
    strip_pagination_properties(heading)
    strip_manual_breaks(heading)

    intro = find_paragraph(document, "The case-study sections identify")
    setup = find_paragraph(document, "The comparison is organized as strong composite")
    caption = find_paragraph(document, "Fig. 8:")
    depth_result = find_paragraph(document, "Initial reservoir depth is not")
    transfer_result = find_paragraph(document, "The subsequent eastward transfer")
    decisive_result = find_paragraph(document, "The decisive separation appears")

    replace_text(
        intro,
        "Figure 8 compares the JFM and AMJ equatorial vertical-temperature structures "
        "of the strong and moderate composites with 2026.",
    )
    replace_text(
        setup,
        "The diagnosis uses initial reservoir depth, JFM-to-AMJ centroid displacement, "
        "and warming of the eastern-Pacific 50-100 m layer.",
    )
    replace_text(
        caption,
        "Fig. 8: JFM and AMJ equatorial-Pacific (5°S-5°N) potential-temperature "
        "anomalies for the strong composite, moderate composite, their difference, "
        "and 2026. Historical SODA3 anomalies use a 1980-2024 same-month detrended "
        "baseline; 2026 GODAS anomalies use its 1991-2020 climatology. The dashed "
        "box marks the eastern-Pacific 50-100 m layer.",
    )
    replace_text(
        depth_result,
        "The JFM west-central reservoir depth is moderate-like in 2026 (137.2 m; "
        "moderate 137.4 m; strong 155.6 m). Its subsequent transfer is strong-like: "
        "the positive-anomaly centroid shifts 56.8° eastward from JFM to AMJ, compared "
        "with 42.4° for the strong composite and 7.2° for the moderate composite.",
    )
    replace_text(
        transfer_result,
        "The AMJ eastern-Pacific 50-100 m anomaly reaches +2.78°C in 2026, compared "
        "with +1.58°C in the strong composite and approximately 0°C in the moderate "
        "composite. Thus, 2026 combines a moderate-like initial depth with unusually "
        "rapid eastward transfer and shoaling into the key eastern layer.",
    )
    replace_text(
        decisive_result,
        "Whole-field similarity was then evaluated by correlating the 2026 "
        "longitude-depth anomaly with each historical class composite month by month "
        "(Fig. 9).",
    )

    image_template = document.paragraphs[
        paragraph_index(document.paragraphs, caption) - 1
    ]
    new_image = add_pattern_figure_after(decisive_result, image_template)
    new_caption = clone_after(
        new_image,
        caption,
        "Fig. 9: January-June spatial pattern correlations between the 2026 GODAS "
        "longitude-depth temperature anomaly and the historical strong (red) and "
        "moderate (blue) SODA3 composites over the common 5°S-5°N, 121°E-81°W, "
        "0-300 m grid.",
    )
    new_result = clone_after(
        new_caption,
        depth_result,
        "The two patterns are least distinguishable in March (r=0.41 for strong and "
        "0.37 for moderate), but separate sharply thereafter. Strong-versus-moderate "
        "correlations are 0.82 versus 0.26 in April, 0.82 versus -0.08 in May, and "
        "0.88 versus 0.12 in June. The AMJ correlations are 0.88 and 0.06, confirming "
        "that 2026 enters the strong-event vertical pathway after March.",
    )
    for paragraph in (
        intro,
        setup,
        caption,
        depth_result,
        transfer_result,
        decisive_result,
        new_image,
        new_caption,
        new_result,
    ):
        strip_pagination_properties(paragraph)
        strip_manual_breaks(paragraph)


def revise_synthesis(document: Document) -> None:
    conclusion = find_paragraph(document, "Conclusion 3.")
    replace_text(
        conclusion,
        "Conclusion 3. As of June 2026, the initial reservoir depth is moderate-like, "
        "whereas eastward transfer, eastern 50-100 m warming, and April-June whole-field "
        "pattern correlations are strong-like. This supports entry into a strong-event "
        "development pathway, while the cross-product comparison and incomplete 2026/27 "
        "life cycle preclude a final intensity forecast.",
    )


def clean_added_section_markers(document: Document) -> None:
    start = paragraph_index(
        document.paragraphs,
        find_paragraph(
            document,
            "Strength-Conditioned Vertical Structure and the 2026 Development Path",
        ),
    )
    for paragraph in document.paragraphs[start:]:
        strip_manual_breaks(paragraph)
        strip_pagination_properties(paragraph)
        if paragraph.text.startswith(("•", "·", "▪", "◦")):
            replace_text(paragraph, paragraph.text.lstrip("•·▪◦ "))


def audit(document: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = (
        "Strength Classes and Spatial Pattern Correlation",
        "Fig. 9:",
        "0.82 versus -0.08 in May",
        "enters the strong-event vertical pathway after March",
    )
    for phrase in required:
        if phrase not in text:
            raise ValueError(f"Required content missing: {phrase}")
    if len(document.inline_shapes) != 12:
        raise ValueError(f"Expected 12 inline figures, found {len(document.inline_shapes)}")
    if len(document.element.xpath(".//m:oMathPara")) != 5:
        raise ValueError("Expected the original two plus three added display equations")
    target = find_paragraph(
        document, "Strength-Conditioned Vertical Structure and the 2026 Development Path"
    )
    if target._p.xpath(".//w:br"):
        raise ValueError("Manual line break remains in the revised section heading")
    properties = target._p.pPr
    if properties is not None and properties.find(qn("w:pageBreakBefore")) is not None:
        raise ValueError("Page-break-before marker remains on the revised section heading")
    for paragraph in document.paragraphs[paragraph_index(document.paragraphs, target) :]:
        if paragraph.text.startswith(("•", "·", "▪", "◦")):
            raise ValueError(f"Stray bullet-like marker remains: {paragraph.text[:30]!r}")


def strip_embedded_fonts(path: Path) -> None:
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(path) as source, ZipFile(
            temporary, "w", compression=ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                if info.filename.startswith("word/fonts/"):
                    continue
                if info.filename == "word/_rels/fontTable.xml.rels":
                    continue
                data = source.read(info.filename)
                if info.filename == "word/fontTable.xml":
                    root = etree.fromstring(data)
                    for element in root.xpath(
                        ".//*[starts-with(local-name(), 'embed')]",
                        namespaces=namespaces,
                    ):
                        parent = element.getparent()
                        if parent is not None:
                            parent.remove(element)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone="yes",
                    )
                target.writestr(info, data)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    insert_methods(document)
    revise_strength_results(document)
    revise_synthesis(document)
    clean_added_section_markers(document)
    document.core_properties.subject = (
        "Strength-conditioned vertical profiles and 2026 pattern similarity"
    )
    audit(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    strip_embedded_fonts(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
