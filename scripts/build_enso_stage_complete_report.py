#!/usr/bin/env python3
"""Build the complete Chinese ENSO stage report as a polished DOCX."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data/derived/research_2026"
FIG_BASE = ROOT / "outputs/research_2026"
OUTPUT = (
    FIG_BASE
    / "reports/ENSO阶段完整报告_海气接通与类型转换_20260731.docx"
)

FIGURES = {
    "definition": (
        FIG_BASE
        / "complete_enso_type_research/"
        "fig1_type_definition_and_transition_checkpoint.png"
    ),
    "nmme": (
        FIG_BASE
        / "complete_enso_type_research/"
        "fig2_nmme_barrier_and_coupling_update.png"
    ),
    "pathways": (
        FIG_BASE
        / "complete_enso_type_research/"
        "fig3_event_pathway_comparison.png"
    ),
    "event_matrix": (
        FIG_BASE
        / "perturbation_attribution/"
        "fig1_event_factor_and_response_matrix.png"
    ),
    "single_effects": (
        FIG_BASE
        / "perturbation_attribution/"
        "fig2_single_factor_perturbation_effects.png"
    ),
    "replacements": (
        FIG_BASE
        / "perturbation_attribution/"
        "fig3_2023_counterfactual_replacements.png"
    ),
    "ne_bridge": (
        FIG_BASE
        / "perturbation_attribution/"
        "fig4_fep_ne_wind_bridge.png"
    ),
    "factorial": (
        FIG_BASE
        / "perturbation_attribution/"
        "fig5_next_factorial_experiment_matrix.png"
    ),
}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203748"
MUTED = "5F6B76"
GOLD = "A87919"
LIGHT_FILL = "F4F6F9"
BLUE_FILL = "E8EEF5"
WHITE = "FFFFFF"
GRID = "C7D1DC"
GREEN_FILL = "E7F2EC"
AMBER_FILL = "FFF4D6"
RED_FILL = "FBE9E7"


def set_run_font(
    run,
    *,
    latin: str = "Calibri",
    east_asia: str = "Songti SC",
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    *,
    top: int = 80,
    bottom: int = 80,
    start: int = 120,
    end: int = 120,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", start),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRID, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_in: list[float]) -> None:
    total_dxa = int(round(sum(widths_in) * 1440))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width_dxa = int(round(widths_in[index] * 1440))
            cell.width = Inches(widths_in[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def mark_table_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def paragraph_border_bottom(paragraph, color: str = BLUE, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Songti SC")
    r_pr.extend((fonts, color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((r_pr, text_node))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.line_spacing = 1.0

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def set_running_header_footer(section, *, first_page: bool = False) -> None:
    section.different_first_page_header_footer = first_page
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [3.8, 2.7])
    mark_table_header_row(table.rows[0])
    set_table_borders(table, color=WHITE, size=0)
    left = table.cell(0, 0).paragraphs[0]
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(
        left.add_run("ENSO TYPE-TRANSITION STUDY"),
        size=8.5,
        bold=True,
        color=MUTED,
    )
    set_run_font(
        right.add_run("阶段报告 · 2026年7月31日"),
        size=8.5,
        color=MUTED,
    )
    rule = header.add_paragraph()
    rule.paragraph_format.space_after = Pt(0)
    paragraph_border_bottom(rule, color=GRID, size=6)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run("研究阶段报告  |  "), size=9, color=MUTED)
    add_page_field(p)


def add_body(
    document: Document,
    text: str,
    *,
    bold_prefix: str | None = None,
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text.split("：", 1)
        set_run_font(
            paragraph.add_run(first + "："),
            size=11,
            bold=True,
            color=INK,
        )
        set_run_font(paragraph.add_run(rest), size=11)
    else:
        set_run_font(paragraph.add_run(text), size=11)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        set_run_font(paragraph.add_run(item), size=11)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        set_run_font(paragraph.add_run(item), size=11)


def add_callout(
    document: Document,
    title: str,
    body: str,
    *,
    fill: str = LIGHT_FILL,
    title_color: str = DARK_BLUE,
) -> None:
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    mark_table_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.paragraphs[0].paragraph_format.space_after = Pt(4)
    set_run_font(
        cell.paragraphs[0].add_run(title),
        size=11,
        bold=True,
        color=title_color,
    )
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.20
    set_run_font(paragraph.add_run(body), size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    source: str | None = None,
    font_size: float = 8.6,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    mark_table_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, LIGHT_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(
            paragraph.add_run(header),
            size=font_size,
            bold=True,
            color=INK,
        )
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            set_run_font(
                paragraph.add_run(str(value)),
                size=font_size,
                color="222222",
            )
    set_table_geometry(table, widths)
    if source:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        set_run_font(
            paragraph.add_run(source),
            size=8.2,
            italic=True,
            color=MUTED,
        )


def add_figure(
    document: Document,
    path: Path,
    caption: str,
    *,
    width: float = 6.35,
) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(
        str(path),
        width=Inches(width),
    )
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split("。", 1)[0])
    caption_paragraph = document.add_paragraph(style="Caption")
    set_run_font(caption_paragraph.add_run(caption), size=9, color=MUTED)


def add_section_break(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True


def add_reference(
    document: Document,
    number: int,
    citation: str,
    url: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(5)
    set_run_font(paragraph.add_run(f"[{number}] {citation} "), size=9.5)
    add_hyperlink(paragraph, url, url)


def build_report() -> None:
    for figure in FIGURES.values():
        if not figure.exists():
            raise FileNotFoundError(figure)

    partial = pd.read_csv(
        DERIVED / "enso_perturbation_partial_correlations.csv"
    )
    interactions = pd.read_csv(
        DERIVED / "enso_next_factorial_interactions.csv"
    )
    factorial_skill = pd.read_csv(
        DERIVED / "enso_next_factorial_loeo_skill.csv"
    )
    events = pd.read_csv(
        DERIVED / "enso_ne_wind_bridge_event_indices.csv"
    ).set_index("year")

    document = Document()
    configure_document(document)
    set_running_header_footer(document.sections[0], first_page=True)

    # Editorial cover pattern.
    for _ in range(5):
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_run_font(
        kicker.add_run("RESEARCH STAGE REPORT"),
        size=10.5,
        bold=True,
        color=GOLD,
    )
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_run_font(
        title.add_run("ENSO海气接通与EP/CP类型转换"),
        size=28,
        bold=True,
        color=INK,
        east_asia="Heiti SC",
    )
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(
        subtitle.add_run("阶段完整报告与下一步实验矩阵"),
        size=17,
        bold=True,
        color=DARK_BLUE,
        east_asia="Heiti SC",
    )
    detail = document.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_after = Pt(38)
    set_run_font(
        detail.add_run(
            "SODA3 / ORAS5 / ERA5 / ERSSTv5 / GPCP / NOAA NMME"
        ),
        size=10.5,
        color=GOLD,
    )
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    set_run_font(
        date.add_run("2026年7月31日"),
        size=12,
        bold=True,
        color=INK,
    )
    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        audience.add_run(
            "研究阶段版本：观测诊断与统计扰动已完成；动力因果试验待运行"
        ),
        size=9.5,
        italic=True,
        color=MUTED,
    )

    add_section_break(document)

    document.add_heading("摘要", level=1)
    add_callout(
        document,
        "本阶段最重要的结论",
        (
            "FEP异常偏高不是东太平洋型厄尔尼诺形成的充分条件。"
            "更完整的分岔由两道连续关口控制：第一，春季暖水是否越过"
            "140°W并在中东太平洋形成持续D20/SSH和上层储热；第二，"
            "降水—气压—赤道西风应力的Bjerknes反馈是否随后接上。"
            "2023表现为强FEP表面路线、冷NPMM、少WWB、偏强东北风入侵"
            "和弱赤道纬向风反馈，因此不能视为1997的弱化复制。"
        ),
        fill=GREEN_FILL,
    )
    add_body(
        document,
        (
            "本报告整合13个暖事件（1982—2023）的海洋热含量、D20、"
            "暖水输送、SST型态、低层风、风应力、降水、气压、OLR和"
            "NMME回报结果。所有事件统计统一采用1982—2024气候态或在"
            "验证折内重估的异常和趋势；类型同时使用Niño3−Niño4与正交"
            "E/C坐标，避免把强度误当成暖中心位置。"
        ),
    )
    add_body(
        document,
        (
            "观测归因中，只有FEP经向降水偶极在控制4月海洋门控、早期"
            "型态和FEP幅度后，与MJJ综合接通显著相关（partial r=+0.72，"
            "95% CI +0.47至+0.91，FDR q=0.024）。但该信号主要体现在"
            "降水和SLP支路，纬向风应力支路区间仍跨零，所以它更像"
            "“对流—气压已经响应”的中介状态，而不是足以推动暖水东传"
            "的独立触发器。"
        ),
    )
    add_body(
        document,
        (
            "新完成的27格观测约束因子矩阵显示，NPMM×WWB以及"
            "“海洋向东门控×有利风场”对接通的交互点估计为正，但"
            "bootstrap区间都略跨零；没有任何格点使E/C与Niño3−Niño4"
            "同时、稳健地转成EP。因而本阶段已经闭合的是证据链和可证伪"
            "实验设计，而不是最终因果定律。下一步应优先运行四格AGCM"
            "的Indo-Atlantic×residual-FEP矩阵，以及九格耦合模式的"
            "ocean-gate×wind-regime矩阵。"
        ),
    )

    document.add_heading("决策摘要", level=2)
    add_table(
        document,
        ["问题", "当前答案", "证据等级", "下一步判据"],
        [
            [
                "FEP高是否必然形成EP？",
                "否；必要性也非普适。",
                "强",
                "FEP on/off的AGCM成对强迫。",
            ],
            [
                "目前最稳定的海洋量？",
                "3—4月140°W Q20+D20海洋门控。",
                "中—强",
                "独立再分析及耦合初值扰动。",
            ],
            [
                "大气何时算“接上”？",
                "降水、SLP与赤道纬向应力需同时响应。",
                "中",
                "集合均值三支路同向且区间不跨零。",
            ],
            [
                "找到普适类型转折点了吗？",
                "没有；4月是候选检查点，不是硬阈值。",
                "中",
                "两种分型指标同时过零。",
            ],
            [
                "当前可否作结？",
                "可作阶段结论；不可宣布单因子因果。",
                "明确边界",
                "完成P0动力矩阵后再升级结论。",
            ],
        ],
        [1.45, 2.05, 0.85, 2.15],
        source="证据等级依据：跨资料一致性、留一事件验证、事件bootstrap及定义敏感性。",
    )

    document.add_heading("报告结构", level=2)
    add_table(
        document,
        ["部分", "内容"],
        [
            ["1—3", "研究问题、数据、统一计算框架"],
            ["4—6", "海洋门控、大气接通、2023事件归因"],
            ["7—8", "因子扰动、最新27格实验结果"],
            ["9", "下一步AGCM与耦合模式实验矩阵"],
            ["10—12", "预报意义、文献定位、结论边界与产物"],
        ],
        [1.1, 5.4],
    )

    add_section_break(document)
    document.add_heading("1 研究问题与逻辑演化", level=1)
    add_body(
        document,
        (
            "研究起点是2023/24事件中的“东部海温很强、但风和SOI响应"
            "偏弱”。随着SODA3热含量和输送、ORAS5 D20、ERA5风场、"
            "GPCP降水以及NMME回报被纳入，问题从“为什么FEP会先暖”"
            "推进为一个更严格的条件问题：在海洋已经具备向东发展的"
            "条件后，什么因素决定大气反馈能否接上，并最终决定EP或CP"
            "空间型态？"
        ),
    )
    document.add_heading("1.1 三个不能混为一谈的目标", level=2)
    add_numbered(
        document,
        [
            "发生：当年是否会发展为El Niño；",
            "分型：在暖事件已发生的条件下，成熟暖中心偏东还是偏中；",
            "放大：已有海洋/表面异常能否通过降水—气压—风应力形成正反馈。",
        ],
    )
    add_body(
        document,
        (
            "本报告的主要统计目标是第二和第三个问题。TSI、FEP残差、"
            "NPMM、WWB和接通残差都不应被宣传成独立判断El Niño是否"
            "发生的通用指数。全年度走步检验已表明，经典海洋门控量在"
            "非暖年并不能稳定超过4月SST基准。"
        ),
    )
    document.add_heading("1.2 预先提出的竞争机制", level=2)
    add_table(
        document,
        ["机制", "物理角色", "可检验预期"],
        [
            [
                "多年La Niña遗产",
                "西部充热，同时可留下西暖—中冷背景。",
                "既可能增强海洋准备度，也可能削弱典型西风响应。",
            ],
            [
                "Indo-Atlantic加热",
                "通过Walker环流远程影响太平洋对流。",
                "去掉外海盆暖异常后，太平洋风压响应增强。",
            ],
            [
                "NPMM",
                "改变热带—副热带SST梯度、MJO与WWB概率。",
                "冷NPMM时WWB减少；暖NPMM放大WWB效应。",
            ],
            [
                "FEP残余与经向降水偶极",
                "组织赤道增强、北侧抑制的对流结构。",
                "先增强降水/SLP，再可能通过东北风抑制赤道西风。",
            ],
            [
                "随机/状态依赖WWB",
                "激发Kelvin型调整和暖量东传。",
                "风领先140°W输送，且效应依赖海洋与NPMM背景。",
            ],
        ],
        [1.25, 2.4, 2.85],
    )

    document.add_heading("2 数据、样本与统一处理", level=1)
    add_table(
        document,
        ["资料", "主要变量", "本研究用途", "时间与分辨率"],
        [
            [
                "SODA3.15.2",
                "OHC、Q20、D20、SSH、热收支、风应力",
                "海洋门控与5日传播",
                "1980—2024；月/5日",
            ],
            [
                "ORAS5",
                "D20、风应力",
                "海洋储存走廊独立检查",
                "1980—2026年6月；月",
            ],
            [
                "ERA5",
                "u850/v850、SLP、风应力、OLR",
                "大气响应与东北风桥梁",
                "1980—2026；月",
            ],
            [
                "ERSSTv5 / OISST",
                "SST、Niño指数、空间型态",
                "类型定义与FEP残差",
                "1980—2024/2026；月",
            ],
            [
                "GPCP",
                "降水",
                "经向降水偶极与对流支路",
                "1982—2024；月",
            ],
            [
                "NOAA NMME",
                "季节SST回报",
                "固定SON目标的春季障碍与更新",
                "1991—2020；月",
            ],
            [
                "TAO 0°, 140°W",
                "流速、温度剖面、D20",
                "SODA3 Q20/D20现场观测核验",
                "共同可用年份；日/月",
            ],
        ],
        [1.15, 1.65, 2.25, 1.45],
        source=(
            "异常主气候态为1982—2024；SODA3过程图使用1980—2024同月异常。"
            "所有预测验证在训练折内重估趋势、均值和尺度。"
        ),
        font_size=8.2,
    )
    document.add_heading("2.1 事件样本与类型定义", level=2)
    add_body(
        document,
        (
            "暖事件为1982、1986、1991、1994、1997、2002、2004、"
            "2006、2009、2015、2018、2019和2023，共13个。成熟期"
            "定义为发展年12月至次年2月。Niño3−Niño4保持与传统箱区"
            "标签一致；正交E/C坐标则把东部和中部型从幅度中尽量分离。"
            "只有两种指标同号时才给出稳定的EP/CP描述；符号相反时记为"
            "混合或定义敏感事件。"
        ),
    )
    add_figure(
        document,
        FIGURES["definition"],
        (
            "图1  类型定义、4月候选检查点与稳健性。Niño3−Niño4与强度"
            "相关较高；正交E/C暴露2015和2023的混合型态。"
        ),
    )

    document.add_heading("3 计算框架与公式", level=1)
    document.add_heading("3.1 海洋门控", level=2)
    add_callout(
        document,
        "输送—储存指数（TSI）",
        (
            "TSI = 0.5 × [ z(Q20 at 140°W, Mar–Apr) + "
            "z(D20 at 140°W, Mar–Apr) ]。Q20回答暖量是否越过140°W，"
            "D20回答输送是否转化为当地温跃层加深和储热。该量是条件式"
            "春季分型更新量，不是普适El Niño发生指数。"
        ),
    )
    document.add_heading("3.2 期望响应与接通残差", level=2)
    add_callout(
        document,
        "逐月留一事件期望关系",
        (
            "对每个月m和响应支路k，用除目标事件外的年份拟合："
            "Y(k) = a(m) + b(m)·trend + c(m)·Niño3.4 + "
            "d(m)·SST-type + ε。响应残差R为观测减期望后再除以训练"
            "残差标准差。MJJ接通分数是降水、东西向SLP驱动和中西太平洋"
            "纬向风应力三个标准化残差的平均。报告同时保留三支路，"
            "防止一个综合量掩盖“对流响应但西风没有接上”。"
        ),
    )
    document.add_heading("3.3 FEP、降水偶极、东北风与WWB", level=2)
    add_table(
        document,
        ["量", "定义", "解释"],
        [
            [
                "FEP residual",
                "FEP SST减去同期趋势、Niño3.4和SST型态可解释部分",
                "东部表面路线的额外暖异常",
            ],
            [
                "Rain dipole",
                "P(5°S—5°N) − P(6°—10°N)，再移除SST期望",
                "赤道增强、北侧抑制的对流结构",
            ],
            [
                "NE850",
                "−(u850+v850)/√2，3°—7°N、180°—130°W",
                "正值为向赤道/向西的东北风入侵",
            ],
            [
                "WWB impulse",
                "Σ max[τx月内扰动 − 1σ, 0]，135°E—180°",
                "五日高频正风应力脉冲代理",
            ],
        ],
        [1.15, 3.25, 2.1],
    )
    document.add_heading("3.4 统计识别与防过拟合", level=2)
    add_bullets(
        document,
        [
            "偏相关同时控制4月海洋门控、4月E/C与同期FEP残差，并用30,000次置换和10,000次事件bootstrap。",
            "多重比较使用Benjamini–Hochberg FDR；预测使用留一事件或年代顺序扩展窗口。",
            "所有2023反事实曲面排除2023训练；岭回归惩罚在0.25—5之间做敏感性。",
            "删除1982、1997和2015三个极端事件，检查结论是否只由强EP杠杆维持。",
            "统计扰动仅表示观测约束敏感性，不等价于AGCM或耦合模式因果干预。",
        ],
    )

    add_section_break(document)
    document.add_heading("4 海洋端：暖水多不等于一定东部型", level=1)
    add_body(
        document,
        (
            "西太平洋总暖水体积说明盆地被“充电”，但不能说明暖量最终"
            "停在中太平洋还是继续进入东部。13事件中，越过日界线的"
            "Q20与成熟型态关系较弱；到160°W增强；到140°W最强。"
            "因此日界线更像进入中太平洋的入口，140°W则是目前样本中"
            "最接近EP/CP分叉的过程检查带，而不是一条不变的物理墙。"
        ),
    )
    add_table(
        document,
        ["海洋量", "与成熟型态r", "FDR q", "解释"],
        [
            ["西部WWV", "+0.389", "0.219", "背景充热，不决定位置"],
            ["Q20跨日界线", "+0.169", "0.582", "只说明进入中太平洋"],
            ["Q20跨160°W", "+0.514", "0.079", "过渡带"],
            ["Q20跨140°W", "+0.765", "0.020", "最强的东部闸门信号"],
            ["4月Q20跨140°W", "+0.822", "0.009", "候选动力检查点"],
        ],
        [1.55, 1.1, 0.8, 3.05],
        source="成熟型态为连续Niño3−Niño4；13个暖事件，事件层FDR。",
    )
    add_body(
        document,
        (
            "5日资料给出了动力顺序：西太平洋正风应力扰动与随后140°W"
            "暖水输送的最大正相关出现在风领先约20天（r=+0.172，"
            "事件块bootstrap CI +0.117至+0.259）。日界线D20到140°W"
            "通常领先15—20天，表观速度与赤道Kelvin型调整相容。"
            "但CP事件也能检测到传播，所以“有Kelvin波”本身仍不是"
            "分型判据；到达后的幅度、持续时间和储存才是第二个海洋条件。"
        ),
    )
    add_figure(
        document,
        FIGURES["pathways"],
        (
            "图2  1997、2015和2023的三条发展路径。三者暖水库接近，"
            "但向东重分配、FEP表面路线和大气兑现明显不同。"
        ),
    )
    document.add_heading("4.1 海洋门控的预测含义与边界", level=2)
    add_body(
        document,
        (
            "在暖事件条件样本中，4月加入140°W Q20+D20后，留一事件"
            "分型技巧明显提高，且ORAS5 D20与TAO 140°W流速/D20提供"
            "独立支持。然而，删除三个极端EP后技巧明显下降，2023又是"
            "TSI的主要失效案例。因此海洋门控适合回答“经典深层路线"
            "是否已实现”，不能单独覆盖FEP表面先行或混合型事件。"
        ),
    )

    document.add_heading("5 大气端：何时才算反馈接上", level=1)
    add_body(
        document,
        (
            "经典Bjerknes反馈必须区分至少三条支路：SST引起对流/降水"
            "加热异常，进而改变海平面气压梯度，最后形成能推动上层海洋"
            "的赤道纬向风应力。2023的关键异常是前两条支路可部分响应，"
            "但第三条偏弱；把它们简单平均会掩盖海气脱耦的具体位置。"
        ),
    )
    coupling_rows = []
    selected = partial.loc[
        partial["response"].eq("MJJ atmosphere coupling residual")
    ]
    for _, row in selected.iterrows():
        coupling_rows.append(
            [
                row["factor"],
                f"{row['partial_r']:+.2f}",
                (
                    f"[{row['bootstrap_ci_low']:+.2f}, "
                    f"{row['bootstrap_ci_high']:+.2f}]"
                ),
                f"{row['fdr_q']:.3f}",
            ]
        )
    add_table(
        document,
        ["候选因素", "partial r", "95% CI", "FDR q"],
        coupling_rows,
        [2.45, 0.9, 1.75, 1.0],
        source="控制：4月海洋门控、4月E/C、MJJ residual FEP；n=13。",
    )
    add_body(
        document,
        (
            "FEP经向降水偶极是唯一通过FDR的量，但其与降水响应"
            "partial r=+0.79、与SLP梯度+0.71，而与纬向风应力仅+0.46"
            "且区间跨零。它可识别局地对流结构，却不能证明赤道西风已经"
            "形成。五个候选因素同时进入模型时，留一RMSE skill仅+0.165；"
            "只加降水偶极为+0.359，显示小样本全因子模型存在过拟合。"
        ),
    )
    add_figure(
        document,
        FIGURES["event_matrix"],
        (
            "图3  13个暖事件的背景因子、海洋门控、接通和成熟型态"
            "标准化矩阵。2023的背景组合与1997/2015明显不同。"
        ),
    )

    document.add_heading("6 2023为何出现强FEP但弱西风", level=1)
    add_table(
        document,
        ["状态量", "2023标准差", "物理含义"],
        [
            ["多年La Niña记忆", "+1.62", "强西部暖水遗产"],
            ["Indo-Atlantic总加热", "+2.27", "强远程背景强迫候选"],
            ["NPMM", "−1.07", "不利MJO/WWB背景"],
            ["FEP residual", "+1.93", "表面东部路线很强"],
            ["经向降水偶极", "+1.22", "对流空间结构已出现"],
            ["高频WWB", "−0.66", "正风应力脉冲偏少"],
            ["东北风入侵", "+1.39", "抑制中太平洋赤道西风"],
            ["MJJ接通", "−0.24", "综合大气兑现偏弱"],
        ],
        [2.2, 1.2, 3.1],
        source="标准化基于13个暖事件；正负号按有利于所列物理量的方向解释。",
    )
    add_body(
        document,
        (
            "2023同时具备强海洋预条件和强FEP残余，但冷NPMM、少WWB、"
            "强东北风入侵与外海盆增暖构成了不利于赤道西风接通的组合。"
            "因此“暖水多→低压→风吹向暖水”只能描述局地压力响应，"
            "不能替代整个Walker环流、纬向SST梯度和对流位置的共同约束。"
        ),
    )
    add_figure(
        document,
        FIGURES["ne_bridge"],
        (
            "图4  FEP—降水偶极—东北风—纬向风应力桥梁。2023个例"
            "方向与Jiang机制一致，但13事件bootstrap区间均跨零。"
        ),
    )
    add_body(
        document,
        (
            "跨事件检验中，FEP residual→东北风为r=+0.05，降水偶极"
            "→东北风为+0.51，东北风→纬向风应力为−0.28，东北风→综合"
            "接通为−0.06。方向上支持“北侧降水抑制/东北风入侵削弱"
            "西风”，但没有一条路径在13事件中达到稳健显著。最合理的"
            "表述是：Jiang机制对2023有解释力，但还需要特定背景态配合，"
            "不是所有暖事件的普适主导因子。"
        ),
    )

    add_section_break(document)
    document.add_heading("7 单因子扰动与2023条件替换", level=1)
    add_body(
        document,
        (
            "第一轮观测扰动使用岭回归结构模拟器，沿“背景→梯度/降水"
            "偶极/WWB→接通→成熟型态”传播。每次增加一个训练样本标准差"
            "后，所有因素对接通和成熟型态的95%区间都跨零。降水偶极对"
            "接通的点估计最大（+0.49 SD），但不能单独宣布因果。"
        ),
    )
    add_figure(
        document,
        FIGURES["single_effects"],
        (
            "图5  +1 SD单因子观测扰动。误差线为10,000次事件bootstrap"
            "区间；没有一个因素的接通效应区间完全排除零。"
        ),
    )
    document.add_heading("7.1 2023替换是否会翻型", level=2)
    add_body(
        document,
        (
            "模拟器排除2023训练，依次把多年La Niña、Indo-Atlantic、"
            "NPMM、降水偶极和WWB替换为中性值或1997/2015平均。七组"
            "单因子/组合替换在E/C和Niño3−Niño4上均未出现过零，"
            "翻型记录为0。这不是“因素不重要”，而是说明在当前观测"
            "样本和线性结构下，成熟型态受海洋位置和多因子共同约束，"
            "不能靠单一背景量完成反事实翻转。"
        ),
    )
    add_figure(
        document,
        FIGURES["replacements"],
        (
            "图6  2023条件替换。所有成熟型态变化均未在两种指标上"
            "形成一致过零。"
        ),
    )

    document.add_heading("8 最新完成：27格因子交互实验", level=1)
    add_body(
        document,
        (
            "第二轮把问题从单因子推进到预先指定的交互矩阵。三个3×3"
            "曲面均以2023为反事实参考，但只用其余12个暖事件训练："
            "A为NPMM×WWB，B为Indo-Atlantic×residual-FEP，C为"
            "4月海洋门控×风场状态。风场状态的正端同时表示WWB增强和"
            "3°—7°N东北风减弱。"
        ),
    )
    interaction_rows = []
    labels = {
        "A_NPMM_x_WWB": "NPMM × WWB",
        "B_IA_x_FEP": "Indo-Atlantic × residual FEP",
        "C_OCEAN_x_WIND": "Ocean gate × favorable wind",
    }
    selected_interactions = interactions.loc[
        interactions["target"].eq("coupling_mjj")
    ]
    for _, row in selected_interactions.iterrows():
        skill_row = factorial_skill.loc[
            factorial_skill["matrix_id"].eq(row["matrix_id"])
            & factorial_skill["target"].eq("coupling_mjj")
        ].iloc[0]
        interaction_rows.append(
            [
                labels[row["matrix_id"]],
                f"{row['interaction_coefficient_response_sd']:+.2f}",
                (
                    f"[{row['bootstrap_ci_low']:+.2f}, "
                    f"{row['bootstrap_ci_high']:+.2f}]"
                ),
                f"{skill_row['incremental_rmse_skill']:+.3f}",
                "否",
            ]
        )
    add_table(
        document,
        ["交互矩阵", "接通交互/SD", "95% CI", "LOEO增量skill", "稳健翻型"],
        interaction_rows,
        [2.05, 1.05, 1.35, 1.15, 0.9],
        source=(
            "LOEO增量skill相对于同一矩阵的控制变量模型；"
            "稳健翻型要求E/C和Niño3−Niño4同时过零且区间不跨零。"
        ),
        font_size=8.2,
    )
    add_figure(
        document,
        FIGURES["factorial"],
        (
            "图7  围绕2023的下一阶段因子交互曲面。上排为接通变化；"
            "下排为成熟E/C与Niño3−Niño4。×表示相应95%区间包含零。"
        ),
    )
    document.add_heading("8.1 如何解读新结果", level=2)
    add_numbered(
        document,
        [
            "NPMM×WWB对接通的交互为+0.42 SD，95% CI [−0.04,+0.71]；点估计符合“暖NPMM使WWB更有效”，但还未排除零。",
            "海洋门控×有利风场对接通的交互为+0.74 SD，95% CI [−0.07,+0.89]；在不同岭参数下保持正号，删除极端事件后仍为+0.14至+0.20，但事件bootstrap仍不够稳健。",
            "Indo-Atlantic×FEP在留一事件中对接通的增量skill为−0.684，说明它可作为动力试验的因果问题，却不是可靠的观测经验预测器。",
            "成熟型态曲面的留一增量skill均为负；九格ocean×wind矩阵中，6格两种指标都偏CP，3格两指标符号不同，没有稳健EP格点。",
            "因此观测矩阵定位了最值得干预的交互，却没有找到可直接业务化的类型阈值。"
        ],
    )
    add_callout(
        document,
        "新的阶段判断",
        (
            "最接近“转折点”的不是某一个FEP、NPMM或WWB数值，而是"
            "海洋暖量位置与风场状态的联合门槛；但当前13事件不能估计"
            "稳定非线性边界。动力模式必须用固定初值和成对强迫检验"
            "“改变一个条件是否真的让类型翻转”。"
        ),
        fill=AMBER_FILL,
        title_color=GOLD,
    )

    add_section_break(document)
    document.add_heading("9 下一步动力实验矩阵", level=1)
    add_body(
        document,
        (
            "下一步不是继续无限增加相关量，而是把最小因果对照做完整。"
            "矩阵分为AGCM诊断层和耦合模式类型层：前者回答给定SST时"
            "为什么不出现西风，后者回答改变海洋门控或风场后成熟型态"
            "是否真的转换。"
        ),
    )
    add_table(
        document,
        ["优先级", "矩阵", "规模", "核心对照", "判定"],
        [
            [
                "P0",
                "AGCM: IA on/off × FEP on/off",
                "4格×30成员",
                "外海盆与局地FEP的主效应/交互",
                "降水、SLP、赤道τx三支路",
            ],
            [
                "P0",
                "Coupled: ocean gate × wind regime",
                "9格×15成员",
                "低/中/高海洋门控 × 不利/中性/有利风",
                "DJF两种类型指标同时过零",
            ],
            [
                "P1",
                "Coupled: NPMM × WWB",
                "4格×20成员",
                "冷/中性NPMM × 无/加WWB",
                "三个WWB起始日同号",
            ],
            [
                "P1",
                "AGCM机制加试",
                "2格×30成员",
                "NPMM-neutral；取消北侧降水亏损",
                "拆分Lin与Jiang路径",
            ],
        ],
        [0.65, 1.95, 1.05, 1.8, 1.05],
        source=(
            "完整19单元矩阵见 experiments/enso_type_transition_matrix/"
            "experiment_matrix.csv 与 experiment_matrix.yaml。"
        ),
        font_size=8.1,
    )
    document.add_heading("9.1 P0-A：四格AGCM矩阵", level=2)
    add_body(
        document,
        (
            "用2023全球SST作为控制，分别去掉热带印度洋—大西洋异常和"
            "residual-FEP分量。预先声明交互对比为："
            "[IA on,FEP on − IA on,FEP off] − "
            "[IA off,FEP on − IA off,FEP off]。若外海盆确实抑制FEP"
            "引起的赤道西风，该对比应在纬向风应力上为负，并同时看到"
            "降水/SLP路径变化。只在降水显著而τx不显著时，仍不能称为"
            "Bjerknes反馈接上。"
        ),
    )
    document.add_heading("9.2 P0-B：九格耦合矩阵", level=2)
    add_body(
        document,
        (
            "以2023年3月1日海气状态为基准，将3—4月140°W Q20/D20"
            "门控设为−1、0、+1个事件标准差，并与不利/中性/有利风场"
            "交叉。有利风场为+1 SD WWB和−1 SD东北风；不利风场相反。"
            "海洋扰动应以盆地内热量重分配实现，不能凭空增加总热含量。"
            "核心检验是“高海洋+有利风”是否比两个单因素高值都更偏EP。"
        ),
    )
    document.add_heading("9.3 预注册判据", level=2)
    add_table(
        document,
        ["假设", "支持条件", "反证条件"],
        [
            [
                "海洋和风是联合门槛",
                "双高格显著超过两个单高格，且两种类型指标同向。",
                "只有海洋主效应或只有风主效应；交互不稳定。",
            ],
            [
                "外海盆抑制FEP风响应",
                "去IA后τx/SLP响应增强，IA×FEP交互符合预设方向。",
                "只改变降水而τx无响应，或不同成员符号随机。",
            ],
            [
                "NPMM调制WWB效率",
                "同一WWB在neutral NPMM下产生更强Kelvin波/Q20与EP型。",
                "三种WWB起始日结果不一致或NPMM主效应为零。",
            ],
            [
                "Jiang经向偶极是抑制桥梁",
                "取消北侧降水亏损后NE减弱、赤道西风增强。",
                "NE与τx不随加热修正改变。",
            ],
        ],
        [1.35, 2.6, 2.55],
        font_size=8.2,
    )
    add_callout(
        document,
        "停止规则",
        (
            "先完成P0 ocean-gate×wind-regime矩阵。若没有任何预先"
            "声明的组合使两种类型指标同向翻转，就不继续扩大到更多"
            "背景因子，而应回到初值平衡、扰动幅度和模式系统偏差检查。"
        ),
        fill=RED_FILL,
        title_color="9B1C1C",
    )

    document.add_heading("10 对春季预报障碍与业务预测的意义", level=1)
    add_body(
        document,
        (
            "NMME固定SON目标检验显示，暖中心位置技巧落后于Niño3.4"
            "幅度技巧，但EP还是CP更受春季障碍影响会随类型定义改变。"
            "按Niño3−Niño4，EP春季惩罚点估计约+0.30°C、CP约+0.01°C；"
            "按正交E/C，同一事件分组会得到不同排序。因此稳妥结论是"
            "“春季存在类型定位困难”，而不是所有EP必然更难预测。"
        ),
    )
    add_figure(
        document,
        FIGURES["nmme"],
        (
            "图8  NMME春季障碍与观测耦合更新。5—6月接通残差是当前"
            "最有希望的春末更新量，但年代顺序置信区间仍跨零。"
        ),
    )
    add_body(
        document,
        (
            "在暖事件样本中，早春SON类型绝对误差与5—6月接通残差的"
            "偏相关为r=0.79。将5—6月观测接通加入3—4月NMME预报，"
            "留一事件RMSE约改善27%；年代顺序扩展窗口的点估计同向，"
            "但95%区间跨零。业务上最合理的流程是：先由动力模式判断"
            "暖事件概率，再由4月海洋门控更新经典EP路线概率，最后在"
            "5—6月用真实接通状态校正成熟暖中心位置。"
        ),
    )

    add_section_break(document)
    document.add_heading("11 与权威研究的关系与本研究增量", level=1)
    add_table(
        document,
        ["研究", "已有结论", "本项目的检验与增量"],
        [
            [
                "Peng et al. (2025)",
                "2023/24强增暖主要由海洋动力生成，风/SOI偏弱。",
                "把期望响应拆到季节、降水、SLP、τx及类型预测更新。",
            ],
            [
                "Fan et al. (2026)",
                "多年La Niña使增暖西伸并削弱SO响应。",
                "事件方向支持，但MLN独立partial r不显著。",
            ],
            [
                "Lin et al. (2025)",
                "冷NPMM抑制MJO/WWB链。",
                "2023冷NPMM、少WWB；NPMM×WWB接通交互为正但区间跨零。",
            ],
            [
                "Jiang et al. (2026)",
                "residual FEP产生经向降水偶极与东北风。",
                "2023个例链完整；13事件桥梁方向一致但不显著。",
            ],
            [
                "Li et al. (2024)",
                "SST→降水加热→SLP→风应力的过程分解。",
                "构造三支路接通残差，识别“对流接上而西风未接上”。",
            ],
            [
                "Kim et al. (2025)",
                "西部暖水多但持续东风可抑制发展。",
                "将暖水库、向东输送和大气兑现分成连续关口。",
            ],
            [
                "Chen et al. (2015)",
                "WWB效应依赖初始海洋态。",
                "SODA3核验风领先140°W输送约20天，并设计联合门控试验。",
            ],
            [
                "Zuo et al. (2019)",
                "ORAS5提供同化海洋状态重建。",
                "以ORAS5 D20独立支持SODA3的东移储存走廊。",
            ],
        ],
        [1.35, 2.25, 2.9],
        font_size=8.0,
    )
    document.add_heading("11.1 本研究真正新增的内容", level=2)
    add_bullets(
        document,
        [
            "把FEP表面路线与Q20/D20经典深层路线放进同一套13事件坐标，而不是只比较1997、2015和2023三个个例。",
            "用两种成熟型态定义明确暴露2015和2023的分类敏感性，避免用强度代替位置。",
            "把大气反馈拆为降水、SLP和纬向风应力支路，并引入FEP降水偶极与3°—7°N东北风桥梁。",
            "以留一事件、年代顺序、FDR、事件bootstrap、删除极端事件和独立海洋资料构成可重复的反证框架。",
            "从单因子替换推进到27格交互曲面，并生成19单元动力模式预注册矩阵。",
        ],
    )

    document.add_heading("12 可作结、不可作结与论文主线", level=1)
    add_table(
        document,
        ["可作结", "仍不可作结"],
        [
            [
                "FEP偏高不是EP形成的充分条件。",
                "尚无适用于所有暖事件的单一数值阈值。",
            ],
            [
                "类型形成至少包含充热、向东重分配、储存和大气兑现。",
                "13事件不能稳定估计高维非线性边界。",
            ],
            [
                "4月140°W Q20/D20是经典EP路线的候选检查点。",
                "140°W不是固定不变的动力分界线。",
            ],
            [
                "2023是强FEP表面路线+弱纬向风反馈+混合成熟型。",
                "观测回归不能证明MLN、IA、NPMM或WWB谁是唯一主因。",
            ],
            [
                "海洋门控×风场状态是下一步最高优先级交互。",
                "未运行耦合矩阵前不能声称人为改变条件会导致翻型。",
            ],
        ],
        [3.25, 3.25],
        font_size=8.6,
    )
    add_callout(
        document,
        "建议的论文核心句",
        (
            "FEP异常增暖可以启动或标记东部表面路线，但EP型的形成"
            "取决于春季上层暖量是否向东输送并储存在中东太平洋，以及"
            "随后降水—气压—赤道西风应力反馈能否共同接上。2023说明"
            "这两道关口可以不同步；因而ENSO类型转换更接近条件链和"
            "状态依赖交互，而不是单一阈值。"
        ),
        fill=GREEN_FILL,
    )
    document.add_heading("12.1 建议的论文结构", level=2)
    add_numbered(
        document,
        [
            "引言：FEP不充分、2023弱耦合与EP/CP类型预测难题；",
            "资料与方法：双类型定义、SODA3/ORAS5海洋门控、三支路接通残差；",
            "结果一：140°W输送—储存走廊与4月候选检查点；",
            "结果二：2023的FEP—降水偶极—东北风与弱纬向应力；",
            "结果三：单因子和交互扰动均未给出稳健翻型；",
            "讨论：与Peng、Fan、Lin、Jiang、Li、Kim的对应及因果矩阵；",
            "结论：双关口框架、预测更新意义和明确的不确定性边界。",
        ],
    )

    document.add_heading("13 可复现产物", level=1)
    add_table(
        document,
        ["类别", "文件"],
        [
            [
                "阶段统计报告",
                "reports/ENSO大气接通与类型转换扰动实验.md",
            ],
            [
                "事件因子表",
                "data/derived/research_2026/enso_ne_wind_bridge_event_indices.csv",
            ],
            [
                "偏相关与FDR",
                "data/derived/research_2026/enso_perturbation_partial_correlations.csv",
            ],
            [
                "27格矩阵预测",
                "data/derived/research_2026/enso_next_factorial_matrix_predictions.csv",
            ],
            [
                "交互与LOEO",
                "enso_next_factorial_interactions.csv / enso_next_factorial_loeo_skill.csv",
            ],
            [
                "动力矩阵",
                "experiments/enso_type_transition_matrix/experiment_matrix.csv",
            ],
            [
                "动力预注册",
                "experiments/enso_type_transition_matrix/experiment_matrix.yaml",
            ],
            [
                "主分析代码",
                "scripts/run_enso_perturbation_attribution.py",
            ],
            [
                "交互矩阵代码",
                "scripts/run_enso_next_factorial_matrix.py",
            ],
        ],
        [1.35, 5.15],
        font_size=8.0,
    )

    add_section_break(document)
    document.add_heading("参考文献与资料", level=1)
    references = [
        (
            "Peng et al. (2025), Strong 2023–2024 El Niño generated by ocean dynamics, Nature Geoscience.",
            "https://www.nature.com/articles/s41561-025-01700-9",
        ),
        (
            "Fan et al. (2026), Multi-year La Niña dampens the Southern Oscillation during subsequent El Niño, The Innovation Geoscience.",
            "https://doi.org/10.59717/j.xinn-geo.2026.100224",
        ),
        (
            "Lin et al. (2025), Diverse NPMM conditions deviate the 2023/24 El Niño from the 1997/98 and 2015/16 events.",
            "https://doi.org/10.1038/s41612-025-01013-z",
        ),
        (
            "Jiang et al. (2026), Geophysical Research Letters, doi:10.1029/2026GL123227.",
            "https://doi.org/10.1029/2026GL123227",
        ),
        (
            "Li et al. (2024), Process-separated atmospheric Bjerknes feedback, Geophysical Research Letters.",
            "https://doi.org/10.1029/2024GL109874",
        ),
        (
            "Kim, Tian & Yu (2025), npj Climate and Atmospheric Science.",
            "https://doi.org/10.1038/s41612-024-00890-0",
        ),
        (
            "Chen et al. (2015), Strong influence of westerly wind bursts on El Niño diversity, Nature Geoscience.",
            "https://doi.org/10.1038/ngeo2399",
        ),
        (
            "Takahashi et al. (2011), Introducing the eastern and central Pacific ENSO indices, Geophysical Research Letters.",
            "https://doi.org/10.1029/2011GL047364",
        ),
        (
            "Zuo et al. (2019), The ECMWF ocean reanalysis system 5 (ORAS5), Ocean Science.",
            "https://doi.org/10.5194/os-15-779-2019",
        ),
        (
            "Fang & Mu (2018), A perfect-model perspective on ENSO predictability.",
            "https://www.nature.com/articles/s41598-018-28964-z",
        ),
        (
            "SODA: Simple Ocean Data Assimilation, University of Maryland.",
            "https://www.soda.umd.edu/",
        ),
        (
            "NOAA Climate Prediction Center, North American Multi-Model Ensemble.",
            "https://www.cpc.ncep.noaa.gov/products/NMME/data.html",
        ),
    ]
    for index, (citation, url) in enumerate(references, start=1):
        add_reference(document, index, citation, url)

    document.add_heading("附录A 关键数值核对", level=1)
    event_rows = []
    for year in (1997, 2015, 2023):
        row = events.loc[year]
        event_rows.append(
            [
                str(year),
                f"{row['mature_ec_flavor']:+.2f}",
                f"{row['mature_box_type']:+.2f}",
                f"{row['east_ocean_state_z_z']:+.2f}",
                f"{row['fep_residual_mjj_z']:+.2f}",
                f"{row['coupling_mjj_z']:+.2f}",
            ]
        )
    add_table(
        document,
        ["事件", "E/C", "N3−N4", "海洋门控z", "FEP residual z", "接通z"],
        event_rows,
        [0.75, 0.8, 0.9, 1.25, 1.45, 1.0],
        source="1997、2015和2023的统一事件表；N3−N4单位为°C。",
    )
    add_body(
        document,
        (
            "1997在两种类型定义上均为强EP，海洋门控与接通都强；"
            "2015的Niño3−Niño4偏EP而E/C接近中性，提示混合型；"
            "2023两种指标符号相反，海洋门控仅略正而FEP residual极强，"
            "接通却偏弱。三者共同证明“暖事件强度”“暖中心位置”和"
            "“大气反馈效率”必须分开描述。"
        ),
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
