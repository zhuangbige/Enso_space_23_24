#!/usr/bin/env python3
"""Create a Chinese-safe PDF from the complete ENSO stage DOCX.

LibreOffice on the current macOS runtime drops CJK glyphs during headless
DOCX-to-PDF conversion.  This script reads the already-built DOCX in document
order and lays it out with ReportLab using embedded Chinese fonts.
"""

from __future__ import annotations

import html
import io
import re
from pathlib import Path

from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
INPUT = (
    ROOT
    / "outputs/research_2026/reports/"
    "ENSO阶段完整报告_海气接通与类型转换_20260731.docx"
)
OUTPUT = (
    ROOT
    / "output/pdf/"
    "ENSO阶段完整报告_海气接通与类型转换_20260731.pdf"
)

SONGTI = "/System/Library/Fonts/Supplemental/Songti.ttc"
HEITI = "/System/Library/Fonts/STHeiti Medium.ttc"

BLUE = colors.HexColor("#2E74B5")
DARK_BLUE = colors.HexColor("#1F4D78")
INK = colors.HexColor("#203748")
MUTED = colors.HexColor("#5F6B76")
GOLD = colors.HexColor("#A87919")
LIGHT_FILL = colors.HexColor("#F4F6F9")
GREEN_FILL = colors.HexColor("#E7F2EC")
GRID = colors.HexColor("#C7D1DC")


def register_fonts() -> None:
    pdfmetrics.registerFont(
        TTFont("SongtiSC", SONGTI, subfontIndex=6)
    )
    pdfmetrics.registerFont(
        TTFont("SongtiSC-Bold", SONGTI, subfontIndex=1)
    )
    pdfmetrics.registerFont(
        TTFont("HeitiSC", HEITI, subfontIndex=1)
    )
    pdfmetrics.registerFontFamily(
        "SongtiSC",
        normal="SongtiSC",
        bold="SongtiSC-Bold",
        italic="SongtiSC",
        boldItalic="SongtiSC-Bold",
    )


def paragraph_text(paragraph: DocxParagraph) -> str:
    return "".join(
        node.text or "" for node in paragraph._p.xpath(".//w:t")
    )


def safe_markup(text: str) -> str:
    url_pattern = re.compile(r"https?://[^\s]+")
    output: list[str] = []
    position = 0
    for match in url_pattern.finditer(text):
        output.append(html.escape(text[position : match.start()]))
        url = match.group(0)
        trailing = ""
        while url and url[-1] in ".,;)]":
            trailing = url[-1] + trailing
            url = url[:-1]
        output.append(
            f'<link href="{html.escape(url)}" '
            f'color="#2E74B5">{html.escape(url)}</link>'
        )
        output.append(html.escape(trailing))
        position = match.end()
    output.append(html.escape(text[position:]))
    return "".join(output).replace("\n", "<br/>")


def styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle(
            "BodyCN",
            parent=base["BodyText"],
            fontName="SongtiSC",
            fontSize=10.2,
            leading=15.0,
            textColor=colors.HexColor("#222222"),
            alignment=TA_JUSTIFY,
            wordWrap="CJK",
            spaceAfter=7,
        ),
        "h1": ParagraphStyle(
            "H1CN",
            parent=base["Heading1"],
            fontName="HeitiSC",
            fontSize=16,
            leading=21,
            textColor=BLUE,
            spaceBefore=12,
            spaceAfter=8,
            keepWithNext=1,
            wordWrap="CJK",
        ),
        "h2": ParagraphStyle(
            "H2CN",
            parent=base["Heading2"],
            fontName="HeitiSC",
            fontSize=13,
            leading=18,
            textColor=BLUE,
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=1,
            wordWrap="CJK",
        ),
        "h3": ParagraphStyle(
            "H3CN",
            parent=base["Heading3"],
            fontName="HeitiSC",
            fontSize=11.5,
            leading=16,
            textColor=DARK_BLUE,
            spaceBefore=7,
            spaceAfter=4,
            keepWithNext=1,
            wordWrap="CJK",
        ),
        "caption": ParagraphStyle(
            "CaptionCN",
            parent=base["BodyText"],
            fontName="SongtiSC",
            fontSize=8.3,
            leading=11.2,
            textColor=MUTED,
            alignment=TA_CENTER,
            wordWrap="CJK",
            spaceBefore=3,
            spaceAfter=8,
        ),
        "list": ParagraphStyle(
            "ListCN",
            parent=base["BodyText"],
            fontName="SongtiSC",
            fontSize=10.1,
            leading=14.2,
            textColor=colors.HexColor("#222222"),
            leftIndent=18,
            firstLineIndent=-12,
            spaceAfter=4,
            wordWrap="CJK",
        ),
        "table_source": ParagraphStyle(
            "TableSourceCN",
            parent=base["BodyText"],
            fontName="SongtiSC",
            fontSize=7.6,
            leading=10,
            textColor=MUTED,
            spaceBefore=3,
            spaceAfter=5,
            wordWrap="CJK",
        ),
        "cover_kicker": ParagraphStyle(
            "CoverKickerCN",
            parent=base["BodyText"],
            fontName="HeitiSC",
            fontSize=10.5,
            leading=14,
            textColor=GOLD,
            alignment=TA_CENTER,
            spaceAfter=16,
        ),
        "cover_title": ParagraphStyle(
            "CoverTitleCN",
            parent=base["Title"],
            fontName="HeitiSC",
            fontSize=28,
            leading=36,
            textColor=INK,
            alignment=TA_CENTER,
            spaceAfter=8,
            wordWrap="CJK",
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitleCN",
            parent=base["Title"],
            fontName="HeitiSC",
            fontSize=17,
            leading=23,
            textColor=DARK_BLUE,
            alignment=TA_CENTER,
            spaceAfter=10,
            wordWrap="CJK",
        ),
        "cover_meta": ParagraphStyle(
            "CoverMetaCN",
            parent=base["BodyText"],
            fontName="SongtiSC",
            fontSize=10,
            leading=14,
            textColor=MUTED,
            alignment=TA_CENTER,
            spaceAfter=5,
            wordWrap="CJK",
        ),
    }


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield DocxParagraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield DocxTable(child, document)


def has_section_break(paragraph: DocxParagraph) -> bool:
    return bool(paragraph._p.xpath("./w:pPr/w:sectPr"))


def image_blob(paragraph: DocxParagraph) -> bytes | None:
    blips = paragraph._p.xpath(".//a:blip")
    if not blips:
        return None
    relation_id = blips[0].get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    if not relation_id:
        return None
    return paragraph.part.related_parts[relation_id].blob


def reportlab_image(blob: bytes) -> Image:
    with PILImage.open(io.BytesIO(blob)) as image:
        width_px, height_px = image.size
    max_width = 6.35 * inch
    max_height = 7.15 * inch
    scale = min(max_width / width_px, max_height / height_px)
    width = width_px * scale
    height = height_px * scale
    return Image(io.BytesIO(blob), width=width, height=height)


def table_column_widths(table: DocxTable) -> list[float]:
    widths = []
    for cell in table.rows[0].cells:
        tc_w = cell._tc.tcPr.tcW
        try:
            widths.append(float(tc_w.w))
        except (TypeError, ValueError, AttributeError):
            widths.append(1.0)
    total = sum(widths)
    if total <= 0:
        widths = [1.0] * len(widths)
        total = float(len(widths))
    return [6.5 * inch * width / total for width in widths]


def convert_table(
    table: DocxTable,
    style_map: dict[str, ParagraphStyle],
) -> Table | KeepTogether:
    column_count = len(table.columns)
    if column_count == 1:
        texts = [
            paragraph_text(paragraph).strip()
            for paragraph in table.cell(0, 0).paragraphs
            if paragraph_text(paragraph).strip()
        ]
        title = safe_markup(texts[0]) if texts else ""
        body = safe_markup(" ".join(texts[1:])) if len(texts) > 1 else ""
        callout_style = ParagraphStyle(
            "CalloutCN",
            parent=style_map["body"],
            fontSize=9.5,
            leading=14,
            textColor=INK,
            spaceAfter=0,
        )
        content = Paragraph(
            f'<font name="SongtiSC-Bold" color="#1F4D78">'
            f"{title}</font><br/>{body}",
            callout_style,
        )
        output = Table(
            [[content]],
            colWidths=[6.5 * inch],
            hAlign="LEFT",
        )
        output.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), GREEN_FILL),
                    ("BOX", (0, 0), (-1, -1), 0.6, GRID),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        return KeepTogether([output, Spacer(1, 5)])

    font_size = 8.3
    if column_count == 4:
        font_size = 7.7
    elif column_count >= 5:
        font_size = 7.1
    cell_style = ParagraphStyle(
        f"TableCell{column_count}",
        parent=style_map["body"],
        fontSize=font_size,
        leading=font_size + 2.2,
        alignment=TA_LEFT,
        spaceAfter=0,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        f"TableHeader{column_count}",
        parent=cell_style,
        fontName="SongtiSC-Bold",
        textColor=INK,
    )
    rows = []
    for row_index, row in enumerate(table.rows):
        current = []
        for cell in row.cells:
            text = "<br/>".join(
                safe_markup(paragraph_text(paragraph).strip())
                for paragraph in cell.paragraphs
                if paragraph_text(paragraph).strip()
            )
            current.append(
                Paragraph(
                    text or " ",
                    header_style if row_index == 0 else cell_style,
                )
            )
        rows.append(current)
    output = Table(
        rows,
        colWidths=table_column_widths(table),
        repeatRows=1,
        hAlign="LEFT",
    )
    output.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), LIGHT_FILL),
                ("GRID", (0, 0), (-1, -1), 0.45, GRID),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return output


def draw_cover_page(canvas, document) -> None:
    canvas.saveState()
    canvas.setFillColor(MUTED)
    canvas.setFont("SongtiSC", 8)
    canvas.drawCentredString(
        letter[0] / 2,
        0.52 * inch,
        "观测诊断与统计扰动完成 · 动力因果试验待运行",
    )
    canvas.restoreState()


def draw_content_page(canvas, document) -> None:
    canvas.saveState()
    width, height = letter
    canvas.setStrokeColor(GRID)
    canvas.setLineWidth(0.5)
    canvas.line(0.78 * inch, height - 0.55 * inch, width - 0.78 * inch, height - 0.55 * inch)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.drawString(0.78 * inch, height - 0.42 * inch, "ENSO TYPE-TRANSITION STUDY")
    canvas.setFont("SongtiSC", 7.5)
    canvas.drawRightString(
        width - 0.78 * inch,
        height - 0.42 * inch,
        "阶段报告 · 2026年7月31日",
    )
    canvas.setFont("SongtiSC", 8)
    canvas.drawRightString(
        width - 0.78 * inch,
        0.48 * inch,
        f"研究阶段报告  |  {canvas.getPageNumber()}",
    )
    canvas.restoreState()


def build_pdf() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)
    register_fonts()
    style_map = styles()
    document = Document(INPUT)
    story = [
        Spacer(1, 1.45 * inch),
        Paragraph("RESEARCH STAGE REPORT", style_map["cover_kicker"]),
        Paragraph(
            "ENSO海气接通与EP/CP类型转换",
            style_map["cover_title"],
        ),
        Paragraph(
            "阶段完整报告与下一步实验矩阵",
            style_map["cover_subtitle"],
        ),
        Spacer(1, 0.12 * inch),
        Paragraph(
            "SODA3 / ORAS5 / ERA5 / ERSSTv5 / GPCP / NOAA NMME",
            style_map["cover_meta"],
        ),
        Spacer(1, 0.65 * inch),
        Paragraph(
            '<font name="SongtiSC-Bold" size="12" color="#203748">'
            "2026年7月31日</font>",
            style_map["cover_meta"],
        ),
        Paragraph(
            "研究阶段版本：观测诊断与统计扰动已完成；动力因果试验待运行",
            style_map["cover_meta"],
        ),
        PageBreak(),
    ]

    after_cover = False
    list_number = 0
    pending_image: Image | None = None
    for block in iter_blocks(document):
        if isinstance(block, DocxParagraph):
            if not after_cover:
                if has_section_break(block):
                    after_cover = True
                continue
            if has_section_break(block):
                if pending_image is not None:
                    story.append(pending_image)
                    pending_image = None
                story.append(PageBreak())
                list_number = 0
                continue
            blob = image_blob(block)
            if blob is not None:
                if pending_image is not None:
                    story.append(pending_image)
                pending_image = reportlab_image(blob)
                continue
            text = paragraph_text(block).strip()
            if not text:
                continue
            style_name = block.style.name if block.style else "Normal"
            if style_name == "Caption":
                caption = Paragraph(
                    safe_markup(text),
                    style_map["caption"],
                )
                if pending_image is not None:
                    story.append(
                        KeepTogether([pending_image, caption])
                    )
                    pending_image = None
                else:
                    story.append(caption)
                continue
            if pending_image is not None:
                story.append(pending_image)
                pending_image = None
            if style_name == "Heading 1":
                story.append(Paragraph(safe_markup(text), style_map["h1"]))
                list_number = 0
            elif style_name == "Heading 2":
                story.append(Paragraph(safe_markup(text), style_map["h2"]))
                list_number = 0
            elif style_name == "Heading 3":
                story.append(Paragraph(safe_markup(text), style_map["h3"]))
                list_number = 0
            elif style_name.startswith("List Bullet"):
                story.append(
                    Paragraph(
                        "• " + safe_markup(text),
                        style_map["list"],
                    )
                )
                list_number = 0
            elif style_name.startswith("List Number"):
                list_number += 1
                story.append(
                    Paragraph(
                        f"{list_number}. " + safe_markup(text),
                        style_map["list"],
                    )
                )
            else:
                list_number = 0
                source_like = (
                    text.startswith("证据等级依据")
                    or text.startswith("异常主气候态")
                    or text.startswith("控制：")
                    or text.startswith("成熟型态为")
                    or text.startswith("标准化基于")
                    or text.startswith("LOEO增量")
                    or text.startswith("完整19单元")
                    or text.startswith("1997、2015和2023")
                )
                story.append(
                    Paragraph(
                        safe_markup(text),
                        (
                            style_map["table_source"]
                            if source_like
                            else style_map["body"]
                        ),
                    )
                )
        else:
            if not after_cover:
                continue
            if pending_image is not None:
                story.append(pending_image)
                pending_image = None
            story.append(convert_table(block, style_map))
            story.append(Spacer(1, 4))
            list_number = 0
    if pending_image is not None:
        story.append(pending_image)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(OUTPUT),
        pagesize=letter,
        rightMargin=0.78 * inch,
        leftMargin=0.78 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.70 * inch,
        title="ENSO海气接通与EP/CP类型转换：阶段完整报告与下一步实验矩阵",
        author="Yue Dai",
        subject="ENSO type transition, air-sea coupling, and perturbation matrix",
    )
    pdf.build(
        story,
        onFirstPage=draw_cover_page,
        onLaterPages=draw_content_page,
    )
    print(OUTPUT)


if __name__ == "__main__":
    build_pdf()
