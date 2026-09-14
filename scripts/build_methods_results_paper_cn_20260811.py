#!/usr/bin/env python3
"""Build the Chinese methods-and-results manuscript from the current ENSO deck.

Named visual override: ``journal_manuscript_cn``.  The document uses a plain
black-and-white journal layout on A4 paper, with no cover page, bullets,
decorative callouts, muted text, or slide-style takeaway boxes.
"""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs/research_2026/papers"
OUTPUT = OUT_DIR / "厄尔尼诺异常暖量到达分配与事件命运_中文小论文_方法结果版_20260811.docx"
FIG_DIR = ROOT / "outputs/research_2026/three_paper_english_figures_20260811"

BODY_FONT = "Arial Unicode MS"
HEAD_FONT = "Arial Unicode MS"
LATIN_FONT = "Arial Unicode MS"
MATH_FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, *, chinese: str = BODY_FONT, latin: str = LATIN_FONT,
                 size: float = 11.0, bold: bool = False, italic: bool = False) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), chinese)
    rfonts.set(qn("w:cs"), latin)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, width: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), width)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9.0)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.35)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.45)
    section.right_margin = Cm(2.45)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Title", 18, 0, 20),
        ("Heading 1", 14, 16, 8),
        ("Heading 2", 12, 12, 5),
        ("Heading 3", 11, 9, 4),
    ):
        style = styles[name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = 1.15

    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.keep_with_next = True

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    add_field_page_number(footer_para)


def add_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(20)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, chinese=HEAD_FONT, size=18, bold=True)


def add_heading(doc: Document, text: str, level: int, *, first: bool = False) -> None:
    paragraph = doc.add_heading(text, level=level)
    if first and level == 1:
        paragraph.paragraph_format.page_break_before = False
    for run in paragraph.runs:
        set_run_font(run, chinese=HEAD_FONT,
                     size={1: 14, 2: 12, 3: 11}.get(level, 11), bold=True)


def add_body(doc: Document, text: str, *, indent: bool = True) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.first_line_indent = Pt(22) if indent else Pt(0)
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_equation(doc: Document, equation: str, number: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(8.05), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES
    )
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.05), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    run = paragraph.add_run("\t" + equation)
    set_run_font(run, chinese=BODY_FONT, latin=MATH_FONT, size=11.2, italic=True)
    number_run = paragraph.add_run(f"\t（{number}）")
    set_run_font(number_run, size=11.0)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]],
              widths_cm: list[float], font_size: float = 9.5) -> None:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(5)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    run = cap.add_run(caption)
    set_run_font(run, size=10.5)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = 1
    table.autofit = False
    set_table_borders(table)
    table.rows[0].height = Cm(0.85)
    repeat_table_header(table.rows[0])
    keep_table_row(table.rows[0])
    for index, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        r = paragraph.add_run(text)
        set_run_font(r, chinese=HEAD_FONT, size=font_size, bold=True)

    for values in rows:
        row = table.add_row()
        keep_table_row(row)
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            r = paragraph.add_run(str(text))
            set_run_font(r, size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.first_line_indent = Pt(0)


def add_figure(doc: Document, filename: str, caption: str, *, width: float = 6.25) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    set_run_font(r, size=10.0)


def add_reference(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.hanging_indent = Pt(18)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.8)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "厄尔尼诺异常暖量在东太平洋的到达、分配及事件命运差异"
    props.subject = "中文版数据、方法与结果稿"
    props.keywords = "ENSO; SODA3; OHC; heat transport; air-sea coupling"
    props.comments = ""

    add_title(doc, "厄尔尼诺异常暖量在东太平洋的到达、分配及事件命运差异")

    add_heading(doc, "1 数据与方法", 1, first=True)
    add_heading(doc, "1.1 数据资料与研究事件", 2)
    add_body(doc, "海洋资料采用Simple Ocean Data Assimilation version 3.15.2（SODA3.15.2）再分析资料[1]。月平均位温、三维流速和净海表热通量用于计算1980—2024年热带太平洋0—300 m海洋热含量、固定控制体边界焓输送和海表热收支；1997、2015和2023三个事件另使用SODA3五日平均位温及u、v、w流速进行离线反向粒子追踪。1980—2022年的月资料由APDRC镜像读取，2023—2024年使用马里兰大学发布的SODA3年度文件补充。密度和定压比热分别取ρ₀=1025 kg m⁻³和cₚ=3990 J kg⁻¹ K⁻¹。")
    add_body(doc, "大气资料采用ERA5月平均925 hPa和850 hPa风场及表面风应力[2]。MJJ东北风指标沿用L. Jiang等[5]给出的区域，即3°—7°N、180°E—130°W；主分析使用925 hPa东北风，850 hPa指标用于层间一致性和三事件时序对照。东部海表温度取自SODA3上层温度场，并以相同的逐月气候态和趋势方法计算异常。")
    add_body(doc, "事件统计固定采用13个暖事件发展年：1982、1986、1991、1994、1997、2002、2004、2006、2009、2015、2018、2019和2023。该事件集合保证SODA3固定层预算与ERA5状态量在同一时间窗内均有完整资料。1997、2015和2023分别作为强到达、强耦合但到达量较小，以及最东部先暖且后续释放较弱的典型对照。所有事件均以发展年记号表示，例如2023代表2023/24事件。")

    add_table(
        doc,
        "表1  数据资料及其在本文中的用途",
        ["资料", "时段与分辨率", "主要变量", "用途"],
        [
            ["SODA3.15.2月资料", "1980—2024，月平均", "位温、u、v、w、净海表热通量", "固定层OHC、边界输送、三箱预算和东部SST"],
            ["SODA3.15.2五日资料", "1997、2015、2023年1—6月，五日平均", "位温、u、v、w", "正温度异常加权的三维反向轨迹"],
            ["ERA5", "1960—2024资料中的共同事件时段，月平均", "925 hPa和850 hPa风、表面风应力", "MJJ东北风及次年5°N动力状态"],
        ],
        [3.2, 4.3, 4.0, 4.6],
        9.3,
    )

    add_heading(doc, "1.2 空间范围、时间窗口与异常定义", 2)
    add_body(doc, "经度—时间传播诊断覆盖10°S—5°N、130°E—80°W。固定控制体的东箱覆盖10°S—5°N、140°W—80°W和0—300 m；西箱与中箱分别定义为160°E—180°和180°—140°W，纬度与深度范围与东箱一致。发展年1—6月用于衡量跨140°W到达量及东箱储热增量，发展年JFM用于表示早期北向路由，发展年MJJ用于表示随后出现的东北风状态，次年1—5月用于诊断成熟后东箱热量损失和各项去向。")
    add_body(doc, "异常采用逐日历月气候态和线性趋势同时去除。OHC经度—时间图以1980—2014年为基准，以保持三个典型事件的传播比较与现有图件一致；控制体储量、边界输送、SST、风应力和第三问统计统一使用1980—2022年基准。粒子释放温度异常的背景值在1980—2024年同月资料上拟合，并在每次计算中排除相应事件年，以避免事件自身进入背景估计。不同基准仅用于不同诊断模块，所有跨事件统计均在模块内部使用一致口径。")

    add_table(
        doc,
        "表2  主要诊断窗口及其物理含义",
        ["阶段", "时间窗口", "主要量", "用途"],
        [
            ["到达", "发展年1月15日至6月15日", "140°W累计输入、东箱OHC增量", "确定到达规模"],
            ["早期路由", "发展年JFM", "东箱5°N北向输出", "识别到达后的早期分流"],
            ["大气标记", "发展年MJJ", "Jiang区域东北风", "检验海洋先行量与后续大气状态的对应"],
            ["后续命运", "次年1月15日至5月15日", "边界输出、海表散热和残差", "量化成熟后的热量分配"],
        ],
        [3.0, 4.1, 4.3, 4.7],
        9.3,
    )

    add_heading(doc, "1.3 固定层OHC及其经向平均", 2)
    add_body(doc, "固定层海洋热含量先在每个网格点对0—300 m位温进行垂向积分，再按纬度余弦进行面积权重。该定义不设置温度阈值，也不引入暖水体积或等温面掩膜。")
    add_equation(doc, "H₃₀₀(λ,φ,t)=ρ₀cₚ∫₀³⁰⁰T(λ,φ,z,t)dz", 1)
    add_equation(doc, "H′₃₀₀=H₃₀₀−Ĉₘ−bₘ(y−ȳₘ)", 2)
    add_body(doc, "式（2）中的Ĉₘ和bₘ分别为基准期同一日历月的平均值和线性趋势，y为年份。经度—时间图对10°S—5°N的H′₃₀₀按cosφ加权平均，结果以GJ m⁻²表示。东箱储量则对控制体内的温度异常进行体积分：H′ₑ=ρ₀cₚ∭ₑT′dV，并以ZJ表示。")

    add_heading(doc, "1.4 正温度异常加权的反向粒子轨迹", 2)
    add_body(doc, "为检验欧拉OHC东移是否具有可解析的水体路径，本文参考拉格朗日热收支研究中将轨迹与能量收支分开检验的处理方式[4]，在东太平洋释放粒子并向后积分。释放区为5°S—5°N、140°W—100°W和15—160 m，释放时刻为发展年6月末，反向积分至1月初。粒子温度异常及其统计权重定义为")
    add_equation(doc, "T′ₚ=Tₚ−T̂ₘ,−e，    ωₚ=max(T′ₚ,0)cosφₚΔzₚ", 3)
    add_body(doc, "其中T̂ₘ,−e为排除事件年e后得到的同月线性背景，只保留释放时T′ₚ>0.05 °C的粒子。ωₚ用于计算轨迹位移的加权分位数和活动轨迹占比，不被解释为守恒能量。粒子位置按下式演化：")
    add_equation(doc, "dλ/dt=u/(Rcosφ)，    dφ/dt=v/R，    dz/dt=−w", 4)
    add_body(doc, "SODA3五日场之间采用线性时间插值，积分算法为中点二阶Runge—Kutta方法。主试验采用SODA3海洋学约定的向上为正w，因此深度坐标满足dz/dt=−w；同时进行不含w及w反号试验，检验东移方向对垂向速度处理的敏感性。")

    add_heading(doc, "1.5 质量守恒的边界焓输送与东箱OHC预算", 2)
    add_body(doc, "开放控制体的单边热输送会受到净体积输送和参考温度的影响。本文先对每个箱体四个侧边界施加同一法向速度校正，使月平均净体积输送严格为零，再计算异常焓输送。设第i个边界的原始体积输送和湿面积分别为Vᵢ和Aᵢ，则")
    add_equation(doc, "δu=(ΣᵢVᵢ)/(ΣᵢAᵢ)，    Vᵢ*=Vᵢ−δuAᵢ", 5)
    add_body(doc, "校正后的边界焓输送采用当月控制体四边界平均温度的同月线性期望值T̄ᵦ作为动态参考：")
    add_equation(doc, "QΓ=ρ₀cₚ∬Γvₙ*(T−T̄ᵦ)dA", 6)
    add_body(doc, "对东箱西边界140°W，正值表示由西向东进入东箱。发展年1—6月累计到达量为Aᴶᵃⁿ–ᴶᵘⁿ=∫Q′₁₄₀Wdt；东箱储热增量为ΔH′ₑ=H′ₑ(Jun)−H′ₑ(Jan)。月样本采用梯形积分，1 PW·s=10⁻⁶ ZJ。质量校正后东箱四边界净体积误差小于4×10⁻¹⁴ Sv，总边界收敛及预算残差不依赖常数参考温标；各单边分量仍依赖所采用的参考定义，因此本文把单边结果用于当前口径下的分配诊断，而把总收敛和残差用于闭合判断。")
    add_body(doc, "次年1—5月东箱固定层能量预算写为")
    add_equation(doc, "ΔH′ₑ=C′ₑ+Q′ₛ+R′，    C′ₑ=Q′W+Q′E+Q′N+Q′S", 7)
    add_body(doc, "Q′W、Q′E、Q′N和Q′S均以进入东箱为正，Q′ₛ为向海洋输入的净海表热通量，R′为端点储量变化减去侧边界总收敛和海表通量后的剩余。结果中统一将−ΔH′ₑ、−Q′W、−Q′N、−Q′S和−Q′ₛ分别表示为东箱损失、向西返回、向北输出、向南输出和海表散热。R′不被命名为具体物理通道；其可能来源包括300 m底边界垂向平流、垂向混合与扩散、SODA同化或分析增量，以及月平均积分和端点离散误差。")

    add_heading(doc, "1.6 第三问的条件相关与样本外检验", 2)
    add_body(doc, "第三问不预设完整的中间机制，而是在统一控制到达量后分别检验三组关系：发展年JFM北向输出与随后MJJ东北风，次年1—5月5°N西向应力与北向输出，以及同期东箱SST与海表散热。对任意变量X和Y及控制量A，先分别进行普通最小二乘残差化，再计算偏相关：")
    add_equation(doc, "Xᵣ=X−(aₓ+bₓA)，    Yᵣ=Y−(aᵧ+bᵧA)，    rₓᵧ·A=corr(Xᵣ,Yᵣ)", 8)
    add_body(doc, "JFM北向输出与MJJ东北风的主检验控制同一JFM的140°W输入；晚期动力和热力关系控制发展年1—6月140°W累计输入。显著性采用20,000次事件置换的双侧检验。早期路由主检验采用10,000次事件bootstrap计算95%置信区间，统一稳健性矩阵采用20,000次bootstrap。每组关系同时报告Pearson偏相关、Spearman秩相关、删除1997、2015和2023后的相关系数以及逐一留一事件后的符号范围。同一检验族内的多重比较使用Benjamini—Hochberg方法控制假发现率。")
    add_body(doc, "样本外检验采用逐一留一事件法。对第i个事件，仅使用其余12个事件拟合普通最小二乘模型，再预测被留出的事件：")
    add_equation(doc, "ŷ₋ᵢ=β̂₀,−i+Σⱼβ̂ⱼ,−ixᵢⱼ，    RMSEz={n⁻¹Σᵢ(ŷ₋ᵢ−yᵢ)²}¹ᐟ²/sᵧ", 9)
    add_body(doc, "北向输出比较仅含到达量的基线模型与加入次年5°N应力的状态模型；海表散热比较仅含到达量的基线模型与加入东箱SST的状态模型。标准化RMSE越小表示未见事件的预测误差越小。所有统计结果均解释为事件间的时间顺序、条件关系或样本外信息增量，不作为耦合模式意义上的单向因果证明。")

    add_heading(doc, "2 结果", 1)
    add_heading(doc, "2.1 0—300 m OHC异常的东传及其路径", 2)
    add_body(doc, "三个典型事件均在发展期形成由中西太平洋向140°W以东扩展的正H′₃₀₀信号，但强度和连贯性不同。1997的正异常最早建立且幅度最大，春季以后在140°W以东形成连续而强的热储存。2015的东传结构较弱，但中东太平洋正异常仍清晰可辨。2023在发展初期表现为西太平洋和最东部正异常并存，中太平洋连接较迟；到发展年春末和初夏，140°W以东仍形成持续的正OHC异常。因此，2023的主要特征并不是东部缺少固定层储热，而是东传结构较1997弱且建立方式不同。")
    add_figure(doc, "q1_ohc300_hovmoller.png", "图1  1997/98、2015/16和2023/24事件0—300 m OHC异常的经度—时间演变。异常以1980—2014年同月气候态和趋势为基准，虚线表示140°W，三幅图使用同一色标。", width=6.25)
    add_body(doc, "反向轨迹给出了与欧拉OHC东移一致的路径证据。正温度异常加权的东向位移中位数在1997、2015和2023中分别为30.0°、21.0°和17.7°；积分至1月初时仍保持可解析状态的释放权重占比分别为62%、64%和75%。不含垂向速度及垂向速度反号试验均保持三次事件的东向位移，说明路径方向不依赖单一的w符号设定，但位移幅度仍对三维积分细节敏感。轨迹诊断由此确认暖异常具有从西侧来源区向东太平洋移动的可解析路径，但不承担热量守恒闭合。")
    add_figure(doc, "q1_lagrangian_summary.png", "图2  三个典型事件正温度异常加权轨迹的东向位移中位数及活动轨迹权重占比。权重只用于轨迹统计，不表示守恒能量。", width=6.1)

    add_heading(doc, "2.2 跨140°W输入与东箱储热共同确认暖量到达", 2)
    add_body(doc, "质量守恒边界输送和独立储量变化给出了相互一致的到达量估计。发展年1—6月，1997跨140°W累计进入东箱的异常焓为16.96 ZJ，东箱OHC增加17.32 ZJ；2015相应为8.49和8.60 ZJ；2023为9.90和11.12 ZJ。1997的到达量约为后两次事件的两倍，而2023的边界输入和箱内储量增量均不低于2015。输入量与储量增量并不要求完全相等，因为同一时段仍包含东、南、北边界交换、海表热通量和预算剩余，但两项独立量在三次事件中保持一致的正号和相近数量级。")
    add_figure(doc, "q1_arrival_storage.png", "图3  1997、2015和2023发展年1—6月跨140°W累计输入与东箱0—300 m OHC增量。东箱范围为10°S—5°N、140°W—80°W。", width=6.0)
    add_body(doc, "图1—3分别从空间传播、三维路径和控制体收支三个角度回答了第一个问题。三次事件的异常暖量均到达东太平洋；2015与2023的差异不能归结为2023的东部输入不足。")

    add_heading(doc, "2.3 到达后的东箱热量分配", 2)
    add_body(doc, "次年1—5月，1997、2015和2023东箱OHC分别减少21.76、19.89和15.95 ZJ。按当前动态边界温度参考计算，向西返回在三次事件中均为最大的已识别水平分量，分别为9.52、8.65和9.45 ZJ，事件间差异很小。北向输出则由1997的5.70 ZJ降至2015的2.99 ZJ和2023的2.17 ZJ；海表散热由4.91 ZJ依次降至1.51和0.47 ZJ。南向分量较小，且2023为−0.72 ZJ，表示当前符号下南边界对东箱存在净输入。")
    add_figure(doc, "q2_fate_partition.png", "图4  1997、2015和2023次年1—5月东箱OHC损失及其分配。正值表示离开东箱；单边分量采用逐月动态边界温度参考，预算残差按原符号单列。", width=6.3)

    add_table(
        doc,
        "表3  三个典型事件次年1—5月东箱热量收支（ZJ）",
        ["事件", "东箱损失", "向西返回", "向北输出", "向南输出", "海表散热", "未闭合汇"],
        [
            ["1997", "21.76", "9.52", "5.70", "0.98", "4.91", "0.52（2.4%）"],
            ["2015", "19.89", "8.65", "2.99", "0.05", "1.51", "6.39（32.1%）"],
            ["2023", "15.95", "9.45", "2.17", "−0.72", "0.47", "4.43（27.7%）"],
        ],
        [1.6, 2.2, 2.2, 2.2, 2.2, 2.2, 2.6],
        9.1,
    )

    add_heading(doc, "2.4 预算残差及可解释范围", 2)
    add_body(doc, "1997的未闭合汇仅为0.52 ZJ，占东箱损失的2.4%，其预算可视为近似闭合。2015和2023的未闭合汇分别为6.39和4.43 ZJ，占东箱损失的32.1%和27.7%。在13个事件中，|R′|/|ΔH′ₑ|的中位数约为26%，|R′|的中位数为2.84 ZJ，11个事件表现为未闭合汇。净海表热通量已作为显式项进入式（7），因此残差不能再归因于遗漏Qnet；现阶段更需要检验300 m底通量、垂向混合和扩散、SODA分析增量以及月平均离散误差。")
    add_body(doc, "较大的残差不改变三个事件均发生明显东箱热量损失，也不改变向西返回为当前参考下最大已识别水平项的结果；但它限制了对2015和2023全部损失的精确归因。因而第二个问题可以回答为：三次事件都存在相近的向西返回，稳定的事件差异主要体现在北向输出和海表散热；尚不能把每一焦耳东箱损失都归入具体物理通道。")

    add_heading(doc, "2.5 早期北向路由与随后东北风", 2)
    add_body(doc, "在固定发展年JFM跨140°W到达量后，JFM北向输出与随后MJJ Jiang区域925 hPa东北风之间的偏相关为r=0.68，事件置换p=0.007，bootstrap 95%置信区间为0.23—0.95；Spearman秩相关为0.56。删除1997、2015或2023后，偏相关分别为0.74、0.60和0.57，逐一留一事件的相关系数范围为0.57—0.87，说明该关系并非由单个典型事件独立维持。")
    add_figure(doc, "q3_early_north_to_ne.png", "图5  发展年JFM北向输出与随后MJJ东北风的条件关系。（a）固定JFM跨140°W到达量后的残差关系；（b）1997、2015和2023的JFM北向输出、JFM NE850和MJJ NE850标准化对照。", width=6.3)
    add_body(doc, "三个典型事件的时序对照进一步显示，1997和2015的JFM北向输出分别为−0.71和−1.08 ZJ，而2023为+0.95 ZJ。2023的JFM NE850仍为−0.74 m s⁻¹，到MJJ转为+1.39个标准差；2015的MJJ NE850则为−2.21个标准差。该结果提供了“早期海洋路由在前、东北风异常在后”的跨事件时序约束，但未检验两者之间的具体中间过程，因此本文不将其写成已经闭合的因果链。")

    add_heading(doc, "2.6 晚期动力与热力状态对事件命运的约束", 2)
    add_body(doc, "固定发展年1—6月到达量后，次年1—5月5°N西向应力与同期北向输出的偏相关为r=0.74，置换p=0.005，95%置信区间为0.29—0.93，Spearman相关为0.68。分别删除1997、2015和2023后，偏相关为0.73、0.74和0.74，表明局地动力状态包含到达量之外的稳定信息。该关系只描述同一后续时段内动力状态与北向输出的对应，不用于补写从MJJ东北风到次年应力的中间机制。")
    add_body(doc, "热力分支同样表现出稳定关系。固定到达量后，东箱SST与海表散热的偏相关为r=0.65，经同一检验族假发现率校正后的q=0.016，95%置信区间为0.23—0.91，Spearman相关为0.64。删除三个典型事件后的相关系数分别为0.66、0.67和0.63。相比之下，5°N应力与向西返回的偏相关仅为r=−0.31，p=0.306，置信区间跨零，因此当前资料不能把相近的向西返回直接归因于局地5°N应力；这一分量仍需结合SSH、密度结构和盆地尺度调整进一步诊断。")
    add_figure(doc, "q3_thermo_return.png", "图6  固定发展年1—6月跨140°W到达量后的后续状态关系。（a）东箱SST与海表散热；（b）5°N应力与向西返回。仅（a）通过当前显著性与稳健性检验。", width=6.15)

    add_table(
        doc,
        "表4  三组保留关系的事件稳健性",
        ["关系", "Pearson偏相关", "Spearman", "95%置信区间", "删除1997/2015/2023"],
        [
            ["JFM北向输出—MJJ NE925", "0.68", "0.56", "0.23—0.95", "0.74 / 0.60 / 0.57"],
            ["5°N应力—晚期北向输出", "0.74", "0.68", "0.29—0.93", "0.73 / 0.74 / 0.74"],
            ["东箱SST—海表散热", "0.65", "0.64", "0.23—0.91", "0.66 / 0.67 / 0.63"],
        ],
        [5.2, 3.0, 2.6, 3.0, 3.6],
        9.2,
    )

    add_heading(doc, "2.7 逐一留一事件的样本外检验", 2)
    add_body(doc, "逐一留一事件检验表明，状态变量的统计信息能够传递到未参与拟合的事件。对北向输出，仅使用到达量时的标准化RMSE为0.75，预测相关为0.63；加入次年5°N应力后，RMSE降至0.55，预测相关升至0.82。对海表散热，仅使用到达量时的RMSE为0.93，预测相关为0.35；加入东箱SST后，RMSE降至0.77，预测相关升至0.62。样本量仍只有13个事件，因此这些改进应理解为到达量之外的样本外信息增量，而不是稳定的业务预报模型。")
    add_figure(doc, "q3_loeo_rmse.png", "图7  北向输出和海表散热模型的逐一留一事件标准化RMSE。每次预测均由其余12个事件拟合，较小的RMSE表示更低的样本外误差。", width=5.8)
    add_body(doc, "综合三个问题的结果，1997、2015和2023均有明确的暖量到达证据，且2023的发展年边界输入和东箱储热不低于2015。三次事件的向西返回相近，事件差异主要表现为北向输出和海表散热的递减。控制到达量后，JFM北向路由与MJJ东北风、次年5°N应力与北向输出以及东箱SST与海表散热分别构成三组独立统计约束。2023由此表现为较强的早期北向分流，随后出现东北风增强，并在次年具有较弱的北向输出和海表散热；这一表述保留了时间顺序和统计证据，同时不越过现阶段尚未检验的中间机制。")

    add_heading(doc, "参考文献", 1)
    add_reference(doc, "[1] Carton J A, Chepurin G A, Chen L. SODA3: A new ocean climate reanalysis. Journal of Climate, 2018, 31: 6967—6983.")
    add_reference(doc, "[2] Hersbach H, Bell B, Berrisford P, et al. The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 2020, 146: 1999—2049.")
    add_reference(doc, "[3] Chakravorty S, Perez R C, Gnanaseelan C, Anderson B T. Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 2021, 126: e2020JC017075.")
    add_reference(doc, "[4] Jiang W, Forget G, Song Y, Haine T W N. Coupled air-sea interaction drove and sustained the 2013—2016 North Pacific marine heatwave. Research Square, 2025. https://doi.org/10.21203/rs.3.rs-7762242/v1.")
    add_reference(doc, "[5] Jiang L, Chen H-C, Zhang W, Tseng Y-H. Early establishment of far eastern Pacific warming shapes distinctive air-sea coupling in the 2023—2024 El Niño. Geophysical Research Letters, 2026, 53: e2026GL123227.")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
