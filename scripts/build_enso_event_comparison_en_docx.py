#!/usr/bin/env python3
"""Build the English January-June ENSO event-comparison report."""

from __future__ import annotations

import numpy as np

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_enso_event_comparison_docx import (
    EVENTS,
    ORDINARY_FIGURES,
    ROOT,
    add_figure,
    add_heading,
    add_paragraph,
    configure_styles,
    format_compact_table,
    load_event_metrics,
    load_ordinary_analysis,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
    set_table_width,
)


OUT = ROOT / "outputs/research_2026/reports/enso_event_comparison_unified_1980_2014_jan_jun_2026_en.docx"


def format_table(table) -> None:
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    set_table_width(table, [2.2, 1.43, 1.43, 1.43])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, size=8.6, bold=(row is table.rows[0]))


def add_metric_table(doc: Document, metrics: dict[str, dict[str, float]]) -> None:
    add_heading(doc, "Table 1. Key Regional Metrics", level=2)
    add_paragraph(
        doc,
        "Each cell gives the January-June mean / June value. Positive Peng CWP u10 denotes a near-surface westerly anomaly.",
        size=9.2,
    )
    rows = [
        ("Niño3.4 SST (°C)", "sst_Nino3.4", "{:+.2f}"),
        ("Peng CWP u10 (m s⁻¹)", "u_CWP", "{:+.2f}"),
        ("CP precipitation (mm day⁻¹)", "pr_CP", "{:+.2f}"),
        ("Western D20 (m)", "d20_West", "{:+.1f}"),
        ("Eastern D20 (m)", "d20_East", "{:+.1f}"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Metric", "2015/16", "2023/24", "2026"]):
        cell.text = label
    for label, key, fmt in rows:
        cells = table.add_row().cells
        cells[0].text = label
        for index, event in enumerate(["2015/16", "2023/24", "2026"], start=1):
            cells[index].text = f"{fmt.format(metrics[event][key])} / {fmt.format(metrics[event][key + '_june'])}"
    format_table(table)


def add_expected_response_table(doc: Document, ordinary) -> None:
    expected = ordinary["expected"]
    means = (
        expected.groupby(["target_label", "response"], as_index=False)[["observed", "expected", "residual"]]
        .mean(numeric_only=True)
        .set_index(["target_label", "response"])
    )
    add_heading(doc, "Table 2. Observed versus SST-Conditioned Response", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Response (Jan-Jun mean)", "2015/16", "2023/24", "2026 state"]):
        cell.text = label
    labels = {2015: "2015/16", 2023: "2023/24", 2026: "2026 developing state"}
    for metric_label, response in (("Peng CWP u10 (m/s)", "cwp_u10"), ("CP precipitation (mm/day)", "cp_precip")):
        cells = table.add_row().cells
        cells[0].text = metric_label
        for index, year in enumerate((2015, 2023, 2026), start=1):
            row = means.loc[(labels[year], response)]
            cells[index].text = f"{row.observed:+.2f} / {row.expected:+.2f} ({row.residual:+.2f})"
    format_compact_table(table, [2.05, 1.48, 1.48, 1.48])
    add_paragraph(
        doc,
        "Each cell gives observed / historical expectation (residual). Calendar-month Niño3 regressions are fitted over "
        "1982-2014 with all three target years excluded.",
        size=9.2,
    )


def add_similarity_table(doc: Document, ordinary) -> None:
    similarity = ordinary["similarity"]
    add_heading(doc, "Table 3. Layer-Resolved Similarity of 2026 to 2015 and 2023", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, label in zip(
        table.rows[0].cells,
        ["Physical layer", "Candidate", "Correlation r", "Centered RMSD", "Std. ratio"],
    ):
        cell.text = label
    layer_labels = {
        "Tropical-Pacific SST": "SST",
        "850 hPa zonal wind": "u850",
        "Precipitation": "Precipitation",
        "Equatorial D20": "D20",
    }
    for layer in ("Tropical-Pacific SST", "850 hPa zonal wind", "Precipitation", "Equatorial D20"):
        subset = similarity.loc[similarity["layer"].eq(layer)].sort_values("candidate_year")
        for row in subset.itertuples(index=False):
            cells = table.add_row().cells
            cells[0].text = layer_labels[layer]
            cells[1].text = str(row.candidate_label)
            cells[2].text = f"{row.pattern_correlation:.3f}"
            cells[3].text = f"{row.centered_rmsd:.2f} {row.unit}"
            cells[4].text = f"{row.standard_deviation_ratio:.3f}"
    format_compact_table(table, [1.25, 1.35, 1.1, 1.55, 1.15])
    add_paragraph(
        doc,
        "Taylor statistics constrain pattern correlation, centered error, and relative amplitude separately; no aggregate "
        "similarity score is constructed.",
        size=9.2,
    )


def build_doc() -> None:
    metrics = load_event_metrics()
    analysis = load_ordinary_analysis()
    expected = analysis["expected"]
    wind_2026 = expected.loc[
        expected["target_year"].eq(2026) & expected["response"].eq("cwp_u10")
    ].sort_values("month")
    precipitation_2026 = expected.loc[
        expected["target_year"].eq(2026) & expected["response"].eq("cp_precip")
    ].sort_values("month")
    if wind_2026["month"].astype(int).tolist() != list(range(1, 7)):
        raise ValueError("Incomplete January-June 2026 Peng CWP u10 response")
    wind_observed = wind_2026["observed"].to_numpy(float)
    wind_expected = wind_2026["expected"].to_numpy(float)
    wind_residual = wind_2026["residual"].to_numpy(float)
    wind_mean_residual = float(np.mean(wind_residual))
    precip_mean_residual = float(precipitation_2026["residual"].mean())
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "Air-Sea Coupling Diagnosis of the 2026 ENSO Event"
    doc.core_properties.subject = "January-June comparison with 2015/16 and 2023/24"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Air-Sea Coupling Diagnosis of the 2026 ENSO Event")
    set_run_font(run, size=20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Comparison with 2015/16 and 2023/24; observations through June 2026")
    set_run_font(run, size=10.5)

    add_heading(doc, "Abstract", level=1)
    add_paragraph(
        doc,
        "Early-2026 SST retained the eastern-leading geometry of 2023, but by June the equatorial thermocline and low-level "
        "zonal wind were closer to the strongly coupled 2015 state. Calendar-month Niño3 regressions show that January-June "
        f"2026 Peng CWP u10 had a mean residual of {wind_mean_residual:+.2f} m s⁻¹, whereas the CP-precipitation residual was "
        f"{precip_mean_residual:+.2f} mm day⁻¹. The observed monthly sequence places thermocline adjustment and near-surface "
        "westerly development before the May-June convective recovery.",
    )

    add_heading(doc, "Study Positioning and Methodological Basis", level=1)
    add_paragraph(
        doc,
        "This report does not introduce a new aggregate coupling index. It applies established process diagnostics within a "
        "common framework for the developing phases of 2015/16, 2023/24, and 2026. Peng et al. (2025) treated Niño3 SST as "
        "the oceanic state, estimated the conditional central-western Pacific zonal-wind response from historical regressions, "
        "and compared the observed 2023/24 response with that expectation. Jiang et al. (2026) jointly examined eastern-Pacific "
        "SST, central-Pacific winds, and precipitation using historical-event comparisons and regression reconstructions. Li et al. "
        "(2024) separated the atmospheric Bjerknes feedback into SST-gradient, precipitation-heating, pressure-gradient, and "
        "zonal-wind-stress links. The observed-versus-expected and layer-resolved diagnostics used here therefore have direct "
        "methodological precedents; the present extension applies them to monthly 2026 observations and compares the result "
        "separately with 2015/16 and 2023/24.",
    )
    add_paragraph(
        doc,
        "This report follows the Peng et al. (2025) CWP definition exactly: area-mean 10 m zonal wind (u10) over 5°S–5°N, "
        "140°E–160°W, used to measure the near-surface atmospheric response to Niño3 SST. CP precipitation is calculated over "
        "5°S–5°N, 180°–140°W, matching the central-Pacific diagnostic box of Sun et al. (2016). CWP u10 and CP precipitation "
        "are coordinate-defined research diagnostics rather than operational monitoring indices such as Niño3 or Niño3.4; the "
        "same boxes, baselines, and anomaly processing are used for all events. The spatial maps retain 850 hPa wind vectors to "
        "show lower-tropospheric circulation structure, but those vectors are not substituted for the Peng CWP u10 metric.",
    )

    add_heading(doc, "1. Data and Methods", level=1)
    add_paragraph(
        doc,
        "SST is from NOAA OISST; 10 m wind, 850 hPa wind, and precipitation are from ERA5; and 20°C isotherm depth (D20) is "
        "from ORAS5. June 2026 u10 is the mean of all 720 hourly fields from the public ARCO-ERA5T mirror; the remaining u10 "
        "values use ERA5 monthly means. ERA5 850 hPa wind, precipitation, and ORAS5 anomalies use 1980-2014, whereas OISST and Peng CWP u10 use "
        "1982-2014. Calendar-month climatologies and linear trends are removed. "
        "All events are compared over January-June of the developing year.",
    )
    add_paragraph(
        doc,
        "Atmospheric conditional responses are obtained from 1982-2014 calendar-month regressions on Niño3 SST, excluding "
        "2015, 2023, and 2026. Historical similarity is evaluated with the area-weighted pattern correlation, centered RMSD, "
        "and standard-deviation ratio of Taylor (2001), separately for SST, u850, precipitation, and D20. The three shaded "
        "month groups in Figure 11 are descriptive summaries of the observed sequence, not formal coupling indices, statistically "
        "identified change points, or event-strength forecasts; they are used to assess consistency with progressive feedback development, not causal closure.",
    )
    add_metric_table(doc, metrics)

    add_heading(doc, "2. Evolution of the Three Events", level=1)
    add_heading(doc, "2.1 SST and 850 hPa Wind", level=2)
    add_paragraph(
        doc,
        "In 2015, central-eastern Pacific warming developed with coherent westerly anomalies. The 2023 warming center was farther "
        "east and the westerly response weaker. In 2026, eastern warming led during January-March, expanded into the central "
        "Pacific after April, and coincided with markedly stronger 850 hPa westerly anomalies in June. These maps describe the "
        "lower-tropospheric circulation; the CWP observed-versus-expected calculation uses 10 m wind.",
    )
    add_figure(doc, EVENTS["2015/16"]["ssta"], "Figure 1. OISST SSTA and ERA5 850 hPa wind anomalies during the 2015/16 development year.")
    add_figure(doc, EVENTS["2023/24"]["ssta"], "Figure 2. OISST SSTA and ERA5 850 hPa wind anomalies during the 2023/24 development year.")
    add_figure(doc, EVENTS["2026"]["ssta"], "Figure 3. OISST SSTA and ERA5 850 hPa wind anomalies, January-June 2026.")

    add_heading(doc, "2.2 Precipitation and 850 hPa Wind", level=2)
    add_paragraph(
        doc,
        "Central-Pacific precipitation and westerlies strengthened together in 2015, whereas convection remained weak in 2023. "
        "In 2026, CP precipitation was suppressed through April, turned positive in May, and strengthened in June, but remained "
        "below the response expected from contemporaneous SST.",
    )
    add_figure(doc, EVENTS["2015/16"]["precip"], "Figure 4. ERA5 precipitation and 850 hPa wind anomalies during the 2015/16 development year.")
    add_figure(doc, EVENTS["2023/24"]["precip"], "Figure 5. ERA5 precipitation and 850 hPa wind anomalies during the 2023/24 development year.")
    add_figure(doc, EVENTS["2026"]["precip"], "Figure 6. ERA5 precipitation and 850 hPa wind anomalies, January-June 2026.")

    add_heading(doc, "2.3 D20 and Equatorial u850", level=2)
    add_paragraph(
        doc,
        "June 2026 D20 reached -12.1 m in the western Pacific and +46.6 m in the east, establishing a clear shallow-west/deep-east "
        "structure. Western D20 in June 2023 was only -2.7 m. The basin-scale thermocline adjustment in 2026 was therefore "
        "substantially closer to 2015.",
    )
    add_figure(doc, EVENTS["2015/16"]["d20"], "Figure 7. ORAS5 equatorial D20 anomalies and ERA5 equatorial u850 during the 2015/16 development year.")
    add_figure(doc, EVENTS["2023/24"]["d20"], "Figure 8. ORAS5 equatorial D20 anomalies and ERA5 equatorial u850 during the 2023/24 development year.")
    add_figure(doc, EVENTS["2026"]["d20"], "Figure 9. ORAS5 equatorial D20 anomalies and ERA5 equatorial u850, January-June 2026.")

    add_heading(doc, "3. SST-Conditioned Atmospheric Response", level=1)
    add_expected_response_table(doc, analysis)
    add_figure(
        doc,
        ORDINARY_FIGURES["expected"],
        "Figure 10. Observed atmospheric responses and Niño3 SST-conditioned historical expectations for the three events; regression period: 1982-2014.",
    )
    add_paragraph(
        doc,
        "Westerlies were substantially stronger than expected in 2015, whereas both wind and precipitation were weaker in 2023. "
        f"The 2026 response is differentiated by process: the January-June mean residual is {wind_mean_residual:+.2f} m s⁻¹ "
        f"for Peng CWP u10 and {precip_mean_residual:+.2f} mm day⁻¹ for CP precipitation. This comparison separates the "
        "near-surface wind response from the deep-convective response; it is not inferred from 850 hPa wind.",
    )

    add_heading(doc, "4. Coupling Sequence in 2026", level=1)
    add_figure(
        doc,
        ORDINARY_FIGURES["staged"],
        "Figure 11. January-June 2026 D20, Niño3 SST, Peng CWP u10, and CP precipitation, including SST-conditioned atmospheric expectations. Shading denotes descriptive month groups, not formal stages or statistical change points.",
    )
    add_heading(doc, "4.1 Thermocline Adjustment (Figure 11a)", level=2)
    add_paragraph(
        doc,
        "Eastern D20 was already +21.0 m in January, while western D20 was only -2.3 m, indicating a pronounced eastern "
        "subsurface warm state before a fully developed basin-scale tilt. Eastern D20 remained between +22.1 and +30.7 m "
        "during February-April as western D20 decreased from -4.6 to -7.9 m. By May, western D20 reached -18.3 m and eastern "
        "D20 +48.5 m, producing an east-west D20 contrast of 66.8 m; the shallow-west/deep-east structure persisted in June "
        "at -12.1 and +46.6 m. This evolution indicates basin-scale redistribution of warm water toward the central-eastern Pacific.",
    )
    add_heading(doc, "4.2 Eastern-Pacific Surface Warming (Figure 11b)", level=2)
    add_paragraph(
        doc,
        "Niño3 SST increased from +0.05°C in January to +0.63°C in February, remained between +0.49 and +0.61°C during "
        "March-April, and then accelerated to +1.25°C in May and +1.66°C in June. The second warming acceleration coincided "
        "with the sharp May increase in thermocline tilt. Eastern thermocline deepening limits the entrainment of cold subsurface "
        "water, while western shoaling is consistent with eastward warm-water redistribution; the joint D20-SST evolution is "
        "therefore consistent with thermocline feedback sustaining eastern-Pacific warming.",
    )
    add_heading(doc, "4.3 Low-Level Zonal-Wind Response (Figure 11c)", level=2)
    add_paragraph(
        doc,
        f"Peng CWP u10 was {wind_observed[0]:+.2f} m s⁻¹ in January, compared with an SST-conditioned expectation of "
        f"{wind_expected[0]:+.2f} m s⁻¹. February observed and expected values were {wind_observed[1]:+.2f} and "
        f"{wind_expected[1]:+.2f} m s⁻¹. Observed March-May values were {wind_observed[2]:+.2f}, {wind_observed[3]:+.2f}, "
        f"and {wind_observed[4]:+.2f} m s⁻¹. June reached {wind_observed[5]:+.2f} m s⁻¹, with a residual of "
        f"{wind_residual[5]:+.2f} m s⁻¹. This series represents the near-surface wind response directly; 850 hPa wind is retained "
        "only as spatial circulation context. The co-evolution of near-surface westerlies, thermocline adjustment, and eastern "
        "warming is consistent with progressive Bjerknes-feedback development, but the historical regression does not establish causal closure.",
    )
    add_heading(doc, "4.4 Deep-Convective Response (Figure 11d)", level=2)
    add_paragraph(
        doc,
        "CP precipitation remained negative from January through April (-1.24, -1.46, -2.14, and -2.04 mm day⁻¹) and below "
        "its SST-conditioned expectation, with a residual of -3.87 mm day⁻¹ in April. Precipitation turned positive in May "
        "(+0.42 mm day⁻¹) but remained 2.22 mm day⁻¹ below expectation. By June it increased to +2.40 mm day⁻¹, close to the "
        "+3.02 mm day⁻¹ expectation, and the residual narrowed to -0.61 mm day⁻¹. Deep convection therefore joined the evolution "
        "after D20, SST, and westerly adjustment. By June, thermocline adjustment and eastern warming co-occurred, warming coincided "
        "with stronger central-western Pacific westerlies, and recovering central-Pacific convection coexisted with the westerly anomaly. "
        "This joint evolution is consistent with progressive closure of the positive air-sea feedback, but monthly observations and "
        "historical regressions do not establish completed causal closure; convection had also not fully reached its historical conditional amplitude.",
    )

    add_heading(doc, "5. Historical Similarity and Evolution", level=1)
    add_similarity_table(doc, analysis)
    add_paragraph(
        doc,
        "The 2026 SST pattern is closer to 2023 (r=0.683; centered RMSD=0.51°C), whereas u850 and D20 are closer to 2015. Their "
        "correlations with 2015 are 0.504 and 0.880, respectively, and the D20 standard-deviation ratio is 0.991. Precipitation "
        "does not map robustly onto either event. Thus, 2026 began with eastern-leading surface warming and subsequently entered "
        "a basin-scale dynamical coupling regime; it is not a replica of either analogue.",
    )
    add_paragraph(
        doc,
        "Observations through June support further development toward a strongly coupled, basin-scale El Niño, but analogue "
        "statistics do not constrain peak amplitude. Independently, the 9 July 2026 NOAA CPC outlook assigned 97% probability of "
        "persistence into early spring 2027 and 81% probability of a very strong event in October-December 2026. The 3 July WMO "
        "update projected rapid strengthening to strong El Niño conditions during July-September.",
    )

    add_heading(doc, "6. Conclusions", level=1)
    add_paragraph(
        doc,
        "By June 2026, D20 adjustment, the Peng CWP near-surface wind response, and the 850 hPa circulation all evolved toward a "
        "coupled state, while SST retained "
        "the eastern-leading geometry of 2023. Deep convection recovered during May-June but remained weak relative to its SST-"
        "conditioned expectation. The summer-autumn evolution should therefore be evaluated through the persistence of central-Pacific "
        "convection and the continued coherence of westerly anomalies with ocean adjustment.",
    )

    add_heading(doc, "Method and Data References", level=1)
    add_paragraph(
        doc,
        "Peng et al. (2025), Nature Geoscience, doi:10.1038/s41561-025-01700-9: CWP (5°S–5°N, 140°E–160°W) u10, Niño3-conditioned regression, and actual-versus-expected response.\n"
        "Jiang et al. (2026), Geophysical Research Letters, doi:10.1029/2026GL123227: eastern-Pacific-leading warming, central-Pacific wind-precipitation response, and cross-event regression comparison.\n"
        "Li et al. (2024), Geophysical Research Letters, doi:10.1029/2024GL109874: process-separated atmospheric Bjerknes feedback from SST through precipitation heating and pressure to wind stress.\n"
        "Sun et al. (2016), Advances in Atmospheric Sciences, doi:10.1007/s00376-015-5018-6: central-Pacific precipitation and wind-stress responses to Niño3 SST and the 5°S–5°N, 180°–140°W diagnostic box.\n"
        "Kim, Tian & Yu (2025), npj Climate and Atmospheric Science, doi:10.1038/s41612-024-00890-0: D20, western-Pacific warm water, and western-central Pacific zonal-wind evolution across events.\n"
        "Taylor (2001), Journal of Geophysical Research, doi:10.1029/2000JD900719: pattern correlation, centered RMSD, and standard-deviation ratio.\n"
        "Zuo et al. (2019), Ocean Science, doi:10.5194/os-15-779-2019: the ORAS5 system, data assimilation, and ocean-state reconstruction.\n"
        "NOAA CPC ENSO Diagnostic Discussion, 9 July 2026; WMO El Niño Update, 3 July 2026.",
        size=9.2,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
