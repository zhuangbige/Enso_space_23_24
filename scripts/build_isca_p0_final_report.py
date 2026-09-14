#!/usr/bin/env python3
"""Build the final Chinese Isca P0 report from audit-complete artifacts.

The builder deliberately refuses incomplete or unverified model output.  It is
not a generic log dump: all reported effects come from the paired MJJ summary,
with OLR transformed into the common physical orientation used by the project.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from scripts.build_enso_stage_complete_report import (
    add_hyperlink,
    add_page_field,
    mark_table_header_row,
    set_cell_margins,
    set_run_font,
    set_table_borders,
    set_table_geometry,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RESULTS = ROOT / "outputs/research_2026/isca_p0_final/results/p0_screening"
DEFAULT_OUTPUT = (
    ROOT
    / "outputs/research_2026/reports"
    / "ENSO大气接通门控_Isca_P0四格试验_最终中文报告.docx"
)
FIG_ROOT = ROOT / "outputs/research_2026"

LATIN_BODY = "Times New Roman"
CJK_BODY = "Songti SC"
LATIN_HEAD = "Arial"
CJK_HEAD = "PingFang SC"
NAVY = "17365D"
BLUE = "1F4E79"
MID_BLUE = "365F91"
MUTED = "666666"
GRID = "AAB8C5"
HEADER_FILL = "EAF0F6"
LIGHT_FILL = "F4F7FA"

METRIC_CN = {
    "precip_tp_mm_day": "热带太平洋降水",
    "precip_cp_mm_day": "中太平洋降水",
    "olr_cp_w_m2": "中太平洋 -OLR",
    "peng_cwp_u10_m_s": "Peng CWP纬向风",
    "tau_cwp_ocean_n_m2": "海洋所受向东风应力",
    "slp_drive_w_minus_e_hpa": "西减东SLP驱动",
}
CONTRAST_CN = {
    "FEP_at_IA0": "FEP | IA0",
    "FEP_at_IA1": "FEP | IA1",
    "IA_at_FEP0": "IA | FEP0",
    "IA_at_FEP1": "IA | FEP1",
    "IA_x_FEP": "IA×FEP",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", type=Path, default=DEFAULT_RESULTS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--billing-cny", type=float)
    parser.add_argument("--billing-hours", type=float)
    return parser.parse_args()


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def parse_key_values(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines()[1:]:
        if "=" in line:
            key, value = line.split("=", 1)
            result[key.strip()] = value.strip()
    return result


def load_complete_results(result_dir: Path) -> dict[str, object]:
    required = (
        "p0_verification.json",
        "march_aligned_pool_validation.json",
        "factorial_sst_validation.json",
        "parallel_benchmark_validation.json",
        "p0_mjj_factorial_summary.csv",
        "p0_mjj_factorial_summary_oriented.csv",
        "p0_scientific_interpretation.json",
        "pipeline_summary.txt",
    )
    missing = [name for name in required if not (result_dir / name).is_file()]
    if missing:
        raise SystemExit(f"Missing final P0 artifacts: {missing}")

    verification = read_json(result_dir / "p0_verification.json")
    if not (
        verification.get("complete") is True
        and verification.get("expected_files") == 240
        and verification.get("valid_files") == 240
        and verification.get("invalid_files") == 0
        and verification.get("errors") == []
    ):
        raise SystemExit("P0 verification does not prove 240/240 completion")
    march = read_json(result_dir / "march_aligned_pool_validation.json")
    if not (
        march.get("complete") is True
        and march.get("valid_state_hashes") == 10
        and march.get("unique_state_hashes") == 10
        and march.get("errors") == []
    ):
        raise SystemExit("March initial-condition pool is not audit complete")

    oriented = pd.read_csv(result_dir / "p0_mjj_factorial_summary_oriented.csv")
    if len(oriented) != 30 or not oriented["n_members"].eq(10).all():
        raise SystemExit("Expected 30 complete oriented factorial rows with n=10")
    if set(oriented["metric"]) != set(METRIC_CN):
        raise SystemExit("Unexpected metric set in oriented summary")
    interpretation = read_json(result_dir / "p0_scientific_interpretation.json")
    pipeline = parse_key_values(result_dir / "pipeline_summary.txt")
    if "p0_wall_seconds" not in pipeline:
        raise SystemExit("Pipeline summary lacks P0 wall time")

    figures = {
        "framework": FIG_ROOT / "paper_cn/fig01_enso_two_gate_framework_cn.png",
        "pathways": (
            FIG_ROOT
            / "complete_enso_type_research/fig3_event_pathway_comparison.png"
        ),
        "forcing": FIG_ROOT / "p0_agcm_prior/fig6_p0_agcm_forcing_components.png",
        "monthly": result_dir / "figures/fig_p0_monthly_paired_response.png",
        "mjj": result_dir / "figures/fig_p0_mjj_factorial_effects.png",
    }
    absent_figures = [str(path) for path in figures.values() if not path.is_file()]
    if absent_figures:
        raise SystemExit(f"Missing report figures: {absent_figures}")
    return {
        "verification": verification,
        "march": march,
        "sst": read_json(result_dir / "factorial_sst_validation.json"),
        "parallel": read_json(result_dir / "parallel_benchmark_validation.json"),
        "oriented": oriented,
        "interpretation": interpretation,
        "pipeline": pipeline,
        "figures": figures,
    }


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_BODY
    fonts = normal._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), LATIN_BODY)
    fonts.set(qn("w:hAnsi"), LATIN_BODY)
    fonts.set(qn("w:eastAsia"), CJK_BODY)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.first_line_indent = Pt(21)

    heading_tokens = (
        ("Heading 1", 15, BLUE, 14, 7),
        ("Heading 2", 12.5, BLUE, 10, 5),
        ("Heading 3", 11, MID_BLUE, 7, 3),
    )
    for name, size, color, before, after in heading_tokens:
        style = doc.styles[name]
        style.font.name = LATIN_HEAD
        fonts = style._element.get_or_add_rPr().rFonts
        fonts.set(qn("w:ascii"), LATIN_HEAD)
        fonts.set(qn("w:hAnsi"), LATIN_HEAD)
        fonts.set(qn("w:eastAsia"), CJK_HEAD)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)

    caption = doc.styles["Caption"]
    caption.font.name = LATIN_BODY
    fonts = caption._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), LATIN_BODY)
    fonts.set(qn("w:hAnsi"), LATIN_BODY)
    fonts.set(qn("w:eastAsia"), CJK_BODY)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string("222222")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.line_spacing = 1.15

    doc.core_properties.title = "Isca P0四格规定SST集合试验最终报告"
    doc.core_properties.subject = "ENSO大气接通门控与IA×residual-FEP交互"
    doc.core_properties.author = "Daiyue"
    doc.core_properties.keywords = "ENSO; Isca; FEP; Indo-Atlantic; prescribed SST"


def set_header_footer(section) -> None:
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(
        p.add_run("ENSO类型形成与海气脱耦"),
        latin=LATIN_HEAD,
        east_asia=CJK_HEAD,
        size=9,
        color=MUTED,
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    add_page_field(footer)


def add_para(
    doc: Document,
    text: str,
    *,
    indent: bool = True,
    bold: bool = False,
    color: str = "111111",
    size: float = 10.5,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21 if indent else 0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    set_run_font(
        p.add_run(text),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=size,
        bold=bold,
        color=color,
    )


def set_cell_text(cell, text: str, *, bold=False, center=False, size=8.2) -> None:
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(
        p.add_run(str(text)),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=size,
        bold=bold,
        color="111111",
    )


def add_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    numeric_columns: set[int] | None = None,
    size: float = 8.2,
) -> None:
    numeric_columns = numeric_columns or set()
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Pt(0)
    cp.paragraph_format.space_before = Pt(5)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.keep_with_next = True
    set_run_font(
        cp.add_run(caption),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=9,
        bold=True,
        color="111111",
    )
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, widths)
    set_table_borders(table, color=GRID, size=5)
    mark_table_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), HEADER_FILL)
        cell._tc.get_or_add_tcPr().append(shading)
        set_cell_margins(cell, top=90, bottom=90, start=120, end=120)
        set_cell_text(cell, header, bold=True, center=True, size=size)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index], top=90, bottom=90, start=120, end=120)
            set_cell_text(
                cells[index],
                value,
                center=index in numeric_columns,
                size=size,
            )
    set_table_geometry(table, widths)
    set_table_borders(table, color=GRID, size=5)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width_cm: float = 15.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Cm(width_cm))
    picture._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph(style="Caption")
    set_run_font(
        cp.add_run(caption),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=9,
        color="222222",
    )


def add_equation(doc: Document, equation: str, explanation: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(
        p.add_run(equation),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=10.5,
        italic=True,
        color="111111",
    )
    add_para(doc, explanation)


def seconds_text(value: str | int | float) -> str:
    seconds = float(value)
    return f"{seconds:.0f} s（{seconds / 3600:.2f} h）"


def effect_rows(oriented: pd.DataFrame) -> list[list[str]]:
    selected = oriented.loc[
        oriented["contrast"].isin(("FEP_at_IA0", "FEP_at_IA1", "IA_x_FEP"))
    ].copy()
    order = {name: index for index, name in enumerate(METRIC_CN)}
    contrast_order = {"FEP_at_IA0": 0, "FEP_at_IA1": 1, "IA_x_FEP": 2}
    selected["metric_order"] = selected["metric"].map(order)
    selected["contrast_order"] = selected["contrast"].map(contrast_order)
    selected = selected.sort_values(["metric_order", "contrast_order"])
    rows = []
    for row in selected.itertuples(index=False):
        rows.append(
            [
                METRIC_CN[row.metric],
                CONTRAST_CN[row.contrast],
                f"{row.oriented_mean:+.3g}",
                f"[{row.oriented_ci95_low:+.3g}, {row.oriented_ci95_high:+.3g}]",
                f"{row.same_sign_fraction:.0%}",
                str(row.unit),
            ]
        )
    return rows


def cautious_result_text(oriented: pd.DataFrame, interpretation: dict) -> str:
    strong = int(interpolation := interpretation["strong_suppressive_count"])
    suggestive = int(interpretation["suggestive_or_strong_suppressive_count"])
    interaction = oriented.loc[oriented["contrast"].eq("IA_x_FEP")]
    strong_metrics = interaction.loc[
        (interaction["oriented_ci95_high"] < 0)
        & (interaction["same_sign_fraction"] >= 0.8),
        "metric",
    ].tolist()
    labels = "、".join(METRIC_CN[name] for name in strong_metrics) or "无"
    if strong >= 3:
        status = "多条大气支路共同支持IA对FEP响应的抑制调制"
    elif strong >= 1 or suggestive >= 3:
        status = "存在部分支路支持，但还不足以宣称整条Bjerknes链一致受抑"
    else:
        status = "10成员P0未给出足以支持稳健IA抑制的因果证据"
    return (
        f"在统一物理方向后，6个指标中有{strong}个同时满足IA×FEP交互"
        f"95%区间低于0且同号成员不低于80%；对应指标为{labels}。"
        f"按预注册门槛，这意味着{status}。该判断不依赖单一均值或"
        "单一图形，而是同时使用成员配对、bootstrap区间和同号成员比例。"
    )


def add_reference(doc: Document, number: int, citation: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    set_run_font(
        p.add_run(f"[{number}] {citation} "),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=9,
        color="222222",
    )
    add_hyperlink(p, "DOI/链接", url)


def build(args: argparse.Namespace) -> None:
    data = load_complete_results(args.results)
    oriented: pd.DataFrame = data["oriented"]
    interpretation: dict = data["interpretation"]
    pipeline: dict[str, str] = data["pipeline"]
    figures: dict[str, Path] = data["figures"]
    args.output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    set_header_footer(doc.sections[0])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(14)
    set_run_font(
        kicker.add_run("大气机制因果筛选报告"),
        latin=LATIN_HEAD,
        east_asia=CJK_HEAD,
        size=10,
        bold=True,
        color=MID_BLUE,
    )
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    set_run_font(
        title.add_run("ENSO大气接通门控\nIsca P0 IA×residual-FEP四格试验"),
        latin=LATIN_HEAD,
        east_asia=CJK_HEAD,
        size=20,
        bold=True,
        color=NAVY,
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_after = Pt(38)
    set_run_font(
        subtitle.add_run("从2023年FEP表层增暖到降水—气压—西风反馈的配对集合检验"),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=12,
        color=MUTED,
    )
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.first_line_indent = Pt(0)
    meta.paragraph_format.space_after = Pt(0)
    set_run_font(
        meta.add_run("2026年8月 | 中文最终版"),
        latin=LATIN_BODY,
        east_asia=CJK_BODY,
        size=10,
        color=MUTED,
    )
    doc.add_page_break()

    doc.add_heading("摘要", level=1)
    add_para(
        doc,
        "本研究检验一个限定明确的大气门控问题：在2023年观测SST背景已给定时，"
        "远东太平洋residual-FEP分量能否独立引发中太平洋降水—气压—西风响应，"
        "以及印度洋—大西洋（IA）海温背景是否调制该响应。试验使用Isca T42L40、"
        "真实大陆与地形、RRTM辐射和Full Betts–Miller对流，对2023年3—8月进行"
        "IA 0/1×FEP 0/1四格规定SST积分。每格10个独立大气初值，同编号成员使用"
        "同一重启，共40条六个月积分链和240个月平均文件。"
    )
    add_para(doc, cautious_result_text(oriented, interpretation))
    add_para(
        doc,
        "P0结果只约束给定SST后的大气响应，不能被解释为海洋暖量已经向东输送，"
        "也不能单独宣称EP/CP成熟型态已发生转换。其价值在于把Jiang的residual-FEP"
        "局地响应与Peng的外海盆抑制假说组织为同一个可证伪因子试验，并为后续"
        "海洋门控×风场耦合矩阵提供筛选依据。"
    )
    add_para(
        doc,
        "关键词：ENSO类型；海气脱耦；residual-FEP；Indo–Atlantic；Isca；"
        "规定SST集合",
        indent=False,
        bold=True,
    )

    doc.add_heading("1 科学问题与逻辑链", level=1)
    add_para(
        doc,
        "前期SODA3、ORAS5和ERA5诊断将ENSO类型形成拆成三个阶段：西太平洋暖量蓄积、"
        "暖量向东重分配—储存，以及大气Bjerknes反馈兑现。FEP表层海温异常只能直接"
        "描述第二阶段在远东表层的一部分表现，它既不等同于盆地暖量输送完成，"
        "也不保证中太平洋深对流和赤道西风已经建立。"
    )
    add_figure(doc, figures["framework"], "图1  ENSO类型形成的三阶段和海洋、大气两道关口。")
    add_para(
        doc,
        "2023年同时出现正residual-FEP、显著外海盆增暖和相对有限的早期赤道西风响应，"
        "因而适合构造反事实问题：如果从同一观测SST场中系统移除或加回IA与FEP分量，"
        "大气对FEP的响应效率是否改变？"
    )
    add_figure(doc, figures["pathways"], "图2  1997、2015与2023年暖量、FEP和大气接通路径的观测对比。")

    doc.add_heading("2 实验设计", level=1)
    doc.add_heading("2.1 SST因子分解", level=2)
    add_para(
        doc,
        "SST源为NOAA ERSSTv5。对1982—2024年按日历月拟合截距、线性趋势、Niño3.4和"
        "E−C型态，构造2023年3—8月的residual-FEP。FEP种子区为5°S—5°N、90°W—80°W；"
        "用排除2023年的历史残差回归得到东太平洋空间足迹。IA主实验分量为2023年热带"
        "印度洋和热带北大西洋相对1982—2024逐月气候态的总异常；边界使用余弦渐变。"
    )
    add_table(
        doc,
        "表1  P0四格SST反事实构造",
        ["格点", "SST构成", "科学含义"],
        [
            ["IA0_FEP0", "2023观测减IA、FEP", "两个候选分量均移除"],
            ["IA0_FEP1", "基础场+FEP", "FEP单独引起的大气响应"],
            ["IA1_FEP0", "基础场+IA", "IA背景的独立响应"],
            ["IA1_FEP1", "完整2023观测SST", "真实联合状态"],
        ],
        [1.25, 2.1, 2.95],
        numeric_columns={0},
    )
    add_para(
        doc,
        "四格SST边界的IA×FEP交互逐格为零。因而若降水、OLR、SLP或风应力出现非零"
        "交互，该信号来自大气动力—物理的非线性，不是预先写入SST的交叉项。"
    )
    add_figure(doc, figures["forcing"], "图3  2023年3—8月IA与residual-FEP规定SST分量及四格构造。")

    doc.add_heading("2.2 Isca配置与集合初值", level=2)
    add_para(
        doc,
        "模式使用T42水平分辨率、40个垂直层、真实大陆与地形、RRTM辐射、Full Betts–Miller"
        "对流和AMIP海冰气候态。每个模式月为30天，每条积分链从2023年3月1日运行至"
        "8月底。输出网格为128×64×40。"
    )
    add_para(
        doc,
        "10个独立大气种子先统一日期至2022年12月1日，再在完全相同的12—2月气候态SST和"
        "海冰下独立积分三个月。最终10个重启的内部日期均为2023-03-01 00:00:00，且"
        "`atmosphere.res.nc`状态SHA-256为10/10唯一。四情景的同编号成员使用同一重启，"
        "从而形成配对反事实设计。"
    )

    doc.add_heading("2.3 诊断量、因子对比与统计", level=2)
    add_para(
        doc,
        "主要响应量为5°S—5°N的热带太平洋降水（180°—250°E）、中太平洋降水和OLR"
        "（180°—220°E）、Peng CWP 10 m纬向风（140°—200°E）、海洋所受向东风应力"
        "（160°—220°E），以及西部130°—160°E减东部220°—250°E的SLP驱动。纬度权重为"
        "cos(lat)。Isca `flux_u`是海面给大气的动量通量，因此海洋所受东向应力定义为"
        "`-flux_u`。"
    )
    add_equation(
        doc,
        "IA×FEP = (Y₁₁ − Y₁₀) − (Y₀₁ − Y₀₀)",
        "其中Yᵢⱼ为IA=i、FEP=j情景的同成员响应。负的物理定向交互表示IA背景减弱FEP促进大气接通的效应。",
    )
    add_para(
        doc,
        "逐月保留3—8月演变，主统计窗口预注册为MJJ。先对每个成员计算MJJ平均对比，"
        "再在10个配对成员上进行20,000次bootstrap，报告均值、标准差、95%区间和同号成员比例。"
    )

    doc.add_heading("3 工程与科学验证", level=1)
    verification_rows = [
        ["Linux环境", "Isca固定提交；隔离Conda/NetCDF/MPI", "通过"],
        ["单月端到端", "SST、重启、辐射、对流与诊断输出", "通过"],
        ["六月稳定性", "6/6月文件与6/6重启；有限变量", "通过"],
        ["SST四格", "仅3—8月有差异；边界交互逐格为0", "通过"],
        ["并发基准", "6个4核作业；正式`slp`存在且有限", "通过"],
        ["共同日期初值", "日期10/10正确；状态10/10唯一", "通过"],
        ["P0主矩阵", "240/240月文件；128×64×40；错误0", "通过"],
    ]
    add_table(
        doc,
        "表2  P0完成门槛与权威证据",
        ["验证项", "实际覆盖范围", "结果"],
        verification_rows,
        [1.55, 3.65, 1.1],
        numeric_columns={2},
    )
    add_para(
        doc,
        "最终审计不以主脚本返回0为充分条件。`p0_verification.json`逐文件检查尺寸、所需变量和"
        "NaN/infinity；只有`expected_files=valid_files=240`、`invalid_files=0`和空错误列表同时成立，"
        "才允许进入统计和作图。"
    )

    doc.add_heading("4 P0结果", level=1)
    doc.add_heading("4.1 3—8月配对响应", level=2)
    add_figure(doc, figures["monthly"], "图4  P0四情景相对IA0_FEP0的3—8月配对响应。线为集合均值，细线为成员。")
    add_para(
        doc,
        "月度曲线用于区分即时响应、季节演变和延迟接通。所有对比都在同成员内完成，"
        "因而不将不同随机天气状态误解为SST强迫效应。正值表示更有利于大气接通；OLR在"
        "作图和定量判读中均使用负号。"
    )

    doc.add_heading("4.2 MJJ因子效应与交互", level=2)
    add_figure(doc, figures["mjj"], "图5  MJJ配对因子效应。点为10成员均值，误差线为成员bootstrap 95%区间。")
    add_table(
        doc,
        "表3  MJJ统一物理方向的主要对比",
        ["响应量", "对比", "均值", "95%区间", "同号", "单位"],
        effect_rows(oriented),
        [1.55, 0.85, 0.85, 1.3, 0.75, 1.0],
        numeric_columns={1, 2, 3, 4, 5},
        size=7.4,
    )
    add_para(doc, cautious_result_text(oriented, interpretation))
    add_para(
        doc,
        "交互项的物理意义必须与FEP在IA0和IA1背景下的条件效应联合判读。交互为负只表示"
        "IA减弱了FEP效应，并不自动表示FEP在IA1下已经完全无效；同样，某一条降水支路"
        "显著也不等于赤道风应力支路已经闭合。"
    )

    doc.add_heading("5 讨论", level=1)
    doc.add_heading("5.1 对“FEP高不是EP充分条件”的贡献", level=2)
    add_para(
        doc,
        "观测证据已表明FEP可以领先盆地Niño3.4演变，而P0提供的是不同类型证据：它在同一"
        "大气初值和同一其他SST背景下，直接改变FEP与IA分量。因此它可以回答“给定"
        "SST后，外海盆是否改变FEP到降水—气压—风应力的大气转换效率”，而不是再做"
        "一次跨年相关。"
    )
    doc.add_heading("5.2 与Peng、Jiang和Kim等研究的关系", level=2)
    add_para(
        doc,
        "Jiang的residual-FEP框架为局地远东增暖及其大气响应提供可分离坐标；Peng的AGCM"
        "试验强调印度洋、大西洋和慢变背景可减弱热带太平洋风响应；Kim等强调"
        "多年La Niña后的西太平洋暖量、D20和纬向风演变。P0的新增量是将前两条大气"
        "假说组合成一个预注册四格，同时把Kim所代表的动力海洋支路明确留给下一阶段。"
    )
    doc.add_heading("5.3 限制与外推边界", level=2)
    add_para(
        doc,
        "Isca是机制层级大气模式；T42分辨率、Full Betts–Miller对流和单模式物理都可影响"
        "定量幅度。10成员是P0初筛而不是最终多模式归因。试验只固定2023年3—8月的一套"
        "空间型，因而不能直接推广到所有EP、CP或混合事件。若P0信号稳定，应将主要"
        "情景扩展到30成员，并在CAM/ECHAM类全面AGCM中复核。"
    )

    doc.add_heading("6 性能、并发与费用", level=1)
    p0_wall = float(pipeline["p0_wall_seconds"])
    alignment_wall = float(pipeline.get("march_alignment_wall_seconds", 0))
    price = float(pipeline.get("price_per_hour_cny", 1.77))
    performance_rows = [
        ["六月单链基准", "5149 s（1.43 h）", "4 MPI", "2.53元（与其他作业重叠）"],
        ["6作业单月基准", "907 s（0.25 h）", "6×4 MPI", "0.45元"],
        ["10初值共同日期对齐", seconds_text(alignment_wall), "6后4", f"{alignment_wall / 3600 * price:.2f}元"],
        ["P0四格主矩阵", seconds_text(p0_wall), pipeline.get("selected_parallel", "6"), f"{p0_wall / 3600 * price:.2f}元"],
    ]
    if args.billing_cny is not None:
        performance_rows.append(
            [
                "AutoDL实际账单",
                f"{args.billing_hours:.2f} h" if args.billing_hours is not None else "平台计时",
                "25 vCPU / 90 GB",
                f"{args.billing_cny:.2f}元",
            ]
        )
    add_table(
        doc,
        "表4  云端运行性能与费用",
        ["阶段", "墙钟", "并发", "按1.77元/小时估算"],
        performance_rows,
        [2.0, 1.45, 1.1, 1.75],
        numeric_columns={1, 2, 3},
        size=7.8,
    )
    add_para(
        doc,
        "六月单链基准与初值生成在部分时段并行，因此不能将各阶段独立估算费用简单相加当作"
        "平台总账单。最终若报告AutoDL实际费用，以平台计费页为准；表中阶段估算用于再现"
        "计算成本和评估扩展至30成员的资源需求。"
    )

    doc.add_heading("7 结论与下一阶段", level=1)
    add_para(doc, cautious_result_text(oriented, interpretation))
    add_para(
        doc,
        "不论P0是支持、部分支持还是不支持IA抑制，其结论都只针对大气接通支路。"
        "下一阶段需要在耦合模式中使用守恒盆地总上层热量的东西重分配，构造海洋门控"
        "低/中/高×风场不利/中性/有利的3×3矩阵。只有DJF E/C和Niño3−Niño4在集合中同时、"
        "稳健过零，才可宣称事件类型发生因果转换。NPMM×WWB和WWB起始日期敏感性应在主"
        "矩阵出现稳定翻型后再加入，避免在小集合中一次估计不可识别的高阶交互。"
    )

    doc.add_heading("可复现性声明", level=1)
    add_para(
        doc,
        "最终结果包包含240文件验收JSON、月度指标、配对月效应、MJJ因子汇总、符号统一表、"
        "两张PNG/PDF主图、性能日志和运行脚本。归档内`MANIFEST.sha256`对每个载荷计算哈希，"
        "下载到本地后再由独立验收脚本逐项重算，并重新读取240/240和10/10共同初值门槛。"
    )

    doc.add_heading("参考文献", level=1)
    references = [
        (
            "Peng et al. (2025). Strong 2023–2024 El Niño generated by ocean dynamics. Nature Geoscience.",
            "https://www.nature.com/articles/s41561-025-01700-9",
        ),
        (
            "Jiang et al. (2026). Far-eastern Pacific warming and the atmospheric response during the 2023/24 El Niño. Geophysical Research Letters.",
            "https://doi.org/10.1029/2026GL123227",
        ),
        (
            "Kim, Tian & Yu (2025). Western-Pacific warm water and zonal-wind evolution across El Niño events. npj Climate and Atmospheric Science.",
            "https://doi.org/10.1038/s41612-024-00890-0",
        ),
        (
            "Li et al. (2024). Process-separated atmospheric Bjerknes feedback. Geophysical Research Letters.",
            "https://doi.org/10.1029/2024GL109874",
        ),
        (
            "Zuo et al. (2019). The ORAS5 ocean synthesis system. Ocean Science, 15, 779–808.",
            "https://doi.org/10.5194/os-15-779-2019",
        ),
        (
            "Carton et al. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983.",
            "https://doi.org/10.1175/JCLI-D-18-0149.1",
        ),
    ]
    for number, (citation, url) in enumerate(references, 1):
        add_reference(doc, number, citation, url)

    doc.save(args.output)
    print(args.output)


def main() -> None:
    build(parse_args())


if __name__ == "__main__":
    main()
