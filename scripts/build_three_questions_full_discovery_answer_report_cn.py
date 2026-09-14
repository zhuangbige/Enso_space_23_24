#!/usr/bin/env python3
"""Build a complete Chinese report from literature discovery to three answers."""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import build_three_core_questions_methods_results_figures_cn as base  # noqa: E402
import build_warm_heat_fate_complete_report_cn as style  # noqa: E402


OUT_DIR = ROOT / "outputs/research_2026/reports"
OUTPUT = OUT_DIR / "三个核心问题_发现提出回答_完整研究报告_公式严谨性修订版_20260807.docx"
ASSET_DIR = ROOT / "outputs/research_2026/three_questions_full_report_assets_20260805"
FIG_PACKAGE = ROOT / "outputs/research_2026/three_questions_figure_package_20260805"


style.FIGURES.update(
    {
        "literature_chakravorty": ASSET_DIR / "chakravorty2021_figure8_excerpt.png",
        "literature_fan": ROOT
        / "outputs/manual-20260523-el-nino/article_figures/fan2026_the_innovation/figure_1.jpg",
        "literature_zhang": ROOT
        / "outputs/research_2026/article_figures/precip_wind_reference_figures/zhang_2025_fig1_full.png",
        "literature_peng_replot": ROOT
        / "outputs/research_2026/figures_detrended/fig_peng2025_expected_response_lines.png",
        "question_preview": FIG_PACKAGE
        / "01_main_figures/Fig03_event_arrival_export_surface_loss.png",
        "data_stack": FIG_PACKAGE
        / "02_supplementary_figures/FigS04_CT_CP_coupled_timing.png",
        "method_decomp": FIG_PACKAGE
        / "02_supplementary_figures/FigS03_CT_CP_5N_current_decomposition.png",
        "limitations": FIG_PACKAGE
        / "02_supplementary_figures/FigS02_full_budget_residual_decomposition.png",
    }
)


def find_paragraph(doc: Document, exact: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact}")


def replace_paragraph_text(doc: Document, old: str, new: str) -> None:
    paragraph = find_paragraph(doc, old)
    paragraph.text = new
    if paragraph.style.name == "Heading 1":
        style.set_run_font(
            paragraph.runs[0],
            font=style.HEAD_FONT,
            size=16,
            bold=True,
            color=style.BLUE,
        )
    elif paragraph.style.name == "Heading 2":
        style.set_run_font(
            paragraph.runs[0],
            font=style.HEAD_FONT,
            size=13,
            bold=True,
            color=style.BLUE,
        )


def move_new_body_elements_before(doc: Document, marker, anchor) -> None:
    """Move elements appended after marker to immediately before anchor."""
    body = doc._element.body
    children = list(body)
    marker_index = children.index(marker)
    appended = []
    for element in children[marker_index + 1 :]:
        if element.tag.endswith("}sectPr"):
            continue
        appended.append(element)
    body.remove(marker)
    for element in appended:
        body.remove(element)
        anchor.addprevious(element)


def iter_table_paragraphs(table):
    """Yield paragraphs in a table, including any nested tables."""
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_all_paragraphs(doc: Document):
    """Yield body, table, header and footer paragraphs for text normalization."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from iter_table_paragraphs(table)
    seen_parts = set()
    for section in doc.sections:
        for part in (section.header, section.footer):
            part_id = id(part._element)
            if part_id in seen_parts:
                continue
            seen_parts.add(part_id)
            yield from part.paragraphs
            for table in part.tables:
                yield from iter_table_paragraphs(table)


def normalize_warm_transport_notation(doc: Document) -> None:
    """Reserve Qwarm,20 for threshold-weighted warm-water energy transport."""
    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            if "Q20" in run.text:
                run.text = run.text.replace("Q20", "Qwarm,20")


def add_front_matter(doc: Document, bullet_id: int, number_id: int) -> None:
    style.add_section_heading(doc, "1 报告摘要：从充放电框架到暖量命运闭环")
    style.add_para(
        doc,
        "本报告整理的是一条连续演化的研究线：从Chakravorty等（2021）关于冷舌型（CT）和中央太平洋型（CP）厄尔尼诺充放电差异的复合研究出发，先复现其170°E、5°N和10°S暖水体积输送逻辑，再追问暖量是否真正到达东太平洋、到达后如何分配，以及为什么1997、2015和2023在相似FEP先暖条件下形成不同的海气耦合和暖量出口。",
    )
    style.add_para(
        doc,
        "文献已经证明赤道太平洋暖水体积是ENSO记忆的重要载体，也证明CT事件通常具有更清晰的充放电循环；但“归一化的暖水体积输送”并不能回答东部箱体获得了多少暖热量、多少从南北边界离开、多少由海表散失、多少仍留在海洋中，更不能解释相似表面增暖为什么会产生不同出口。因此，本研究把问题从WWV相位关系推进到H20/Qwarm,20阈值暖量收支、5°N动力分解和海气耦合过程归因。",
    )
    style.add_callout(
        doc,
        "整篇报告的一句话结论",
        "暖量到达东太平洋只是必要条件，不是命运的充分条件；最终分配由盆地尺度次表层支撑、大气对流—风场接通以及5°N处Ekman与地转输送的竞争共同决定。",
        fill=style.LIGHT_BLUE,
    )
    style.add_figure(
        doc,
        "literature_chakravorty",
        "文献图L1  Chakravorty等（2021）Figure 8：CT与CP事件在170°E、5°N、10°S的归一化暖水输送及WWV演变。原图来源：Chakravorty et al., Journal of Geophysical Research: Oceans, 126, e2020JC017075, https://doi.org/10.1029/2020JC017075。此处按学术研究用途引用，版权归原作者/期刊。",
        width=6.15,
    )

    style.add_section_heading(doc, "2 文献发现：别人已经回答了什么", page_break=True)
    doc.add_heading("2.1 充放电、WWV与ENSO记忆", level=2)
    style.add_para(
        doc,
        "Jin（1997）的recharge oscillator把纬向平均温跃层深度与风应力驱动的经向暖水输送失衡视为ENSO慢记忆：发展前赤道带被充水，暖事件成熟后暖水向极放出。Meinen和McPhaden（2000）的观测进一步表明，赤道太平洋WWV异常通常领先Niño-3.4 SST数月，支持WWV作为海洋预条件而非同期表面温度的替代量。",
    )
    style.add_para(
        doc,
        "后续研究把这个框架拆得更细。Bosc和Delcroix（2008）指出5°N—5°S之间的WWV变化是Ekman与地转经向输送相互补偿后的残差；Sun（2000）及Huguenin等（2020）强调暖事件终止同时涉及极向热输送、海表热通量和非绝热过程；Verdy等（2023）则用TPOSE闭合2014—2018体积、温盐收支，说明单个事件的北界输出能够被定量追踪。它们共同说明“暖水向极排放”是真过程，但也说明只画WWV或单一断面还不足以回答每个事件的暖量去向。",
    )

    doc.add_heading("2.2 Chakravorty等（2021）的方法、结论与留下的缺口", level=2)
    style.add_para(
        doc,
        "Chakravorty等主要使用SODA2.1.6（1958—2008），并以SODA2.2.4作长时段稳健性检验。所有月序列去季节气候态和线性趋势后施加3个月滑动平均；事件识别不是只用固定Niño-3/Niño-4阈值，而是结合surface center-of-heat index（CHI）区分CT和CP。作者比较SST、风应力/旋度、次表层热含量、WWV—Niño-3.4相图、surface/subsurface CHI、D20 EOF，以及从海表积分到逐月20°C等温面深度的170°E纬向、5°N和10°S经向暖水体积输送。",
    )
    style.add_para(
        doc,
        "其核心结论是：CT事件在前期出现西太平洋暖水积累、沿赤道向东输送和成熟后的向极放电，WWV对后续SST具有较清楚的先导；CP事件的次表层预兆和充放电循环较弱，因此不能用同样的赤道温跃层前兆解释。Figure 8尤其显示CT在成熟—衰减期5°N北侧向极排放增强，而10°S仍可维持一段时间的向赤道输送。",
    )
    style.add_para(
        doc,
        "它留下了四个直接缺口：第一，输送被标准化，缺少ZJ级能量量纲；第二，没有设置西—中—东固定控制体，无法区分“经过断面”和“进入东部后储存”；第三，没有把南北输出、海表散热和剩余项放在同一预算；第四，没有把5°N总输送拆成地转、Ekman和残差，也没有解释2015/2023这类相似输入却不同出口的海气耦合分叉。",
    )
    style.add_table(
        doc,
        "表L1  从经典文献到本研究的递进关系",
        ["文献", "已经回答", "使用的关键证据", "仍未回答"],
        [
            ["Jin (1997); Meinen & McPhaden (2000)", "WWV构成ENSO慢记忆并领先SST", "理论模型；观测WWV—SST相位关系", "到达东部后的能量归宿"],
            ["Bosc & Delcroix (2008); Verdy et al. (2023)", "经向输送存在地转—Ekman补偿；个例预算可闭合", "高度计/现场观测；TPOSE", "CT/CP群体差异和三强事件对照"],
            ["Chakravorty et al. (2021)", "CT有清晰充放电，CP较弱", "SODA复合、CHI、D20 EOF、归一化输送", "ZJ级H20/Q20去向与动力分解"],
            ["本研究", "到达—分配—耦合命运的一条证据链", "SODA3三箱预算、独立大气资料、13事件与换风试验", "上游外强迫的严格因果份额"],
        ],
        [1900, 2450, 2460, 2550],
        font_size=8.0,
    )

    doc.add_heading("2.3 2023/24把问题从‘暖不暖’推向‘海气是否接通’", level=2)
    style.add_para(
        doc,
        "Peng等（2025）指出，2023/24具有强海洋增暖，却只有偏弱的Southern Oscillation和热带太平洋风异常；其大气模式试验显示，2023年印度洋和大西洋增暖以及慢变SST背景趋势削弱了热带太平洋风响应，而海洋后报试验在4月后关闭风应力反馈仍可重现6—12月增暖的87%。这意味着2023不是“暖量不足”，而是海洋记忆很强、经典Bjerknes正反馈却偏弱。",
    )
    style.add_figure(
        doc,
        "literature_peng_replot",
        "文献图L2  依据Peng等（2025）Figure 1及其公开source archive重绘的历史期望线：2023/24风和气压响应只有历史Niño-3回归期望的约27%和31%。数据来源：Peng et al., Nature Geoscience 18, 471–478, https://doi.org/10.1038/s41561-025-01700-9；公开代码与数据：https://doi.org/10.5281/zenodo.15074285。本图为本研究重绘。",
        width=6.15,
    )
    style.add_para(
        doc,
        "Zhang等（2025）进一步表明，2023/24虽然是卫星时代第四强的El Niño之一，但热带太平洋降水异常约为0.75 mm d⁻¹，仅为1982/83、1997/98和2015/16三次超级事件的约三分之一；AGCM试验把弱降水和弱遥相关与印度洋、大西洋异常增暖联系起来。该结果把“风弱”延伸为“对流加热源也弱”，说明海气接通必须同时检查SST、降水/OLR、SLP和低层风。",
    )
    style.add_figure(
        doc,
        "literature_zhang",
        "文献图L3  Zhang等（2025）Figure 1：2023/24的SST很强，但热带太平洋降水、u850及PNA响应明显低于历史强事件关系。原图来源：Zhang et al., Communications Earth & Environment, 6, 595, https://doi.org/10.1038/s43247-025-02584-8（开放获取，CC BY 4.0）。",
        width=5.9,
    )
    style.add_para(
        doc,
        "Fan等（2026）从热带太平洋内部补充了另一条机制：多年La Niña使西太平洋WWV异常增大，但也使后续增暖更向西延伸，削弱西太平洋负对流、东西向气压梯度和Southern Oscillation。Kim等（2025）和Lin等（2025）分别强调持续异常东风及冷NPMM—MJO/WWB不足。因而外海盆、长期趋势、多年La Niña遗产和NPMM并非互相排斥；它们是第三问的上游候选因素，而不是当前观测诊断已经唯一确定的答案。",
    )
    style.add_figure(
        doc,
        "literature_fan",
        "文献图L4  Fan等（2026）Figure 1：2023年具有接近1997年的前期WWV，但MJJAS Niño-3.4和低层风响应并不同步。原图来源：Fan et al., The Innovation Geoscience, 4, 100224, https://doi.org/10.59717/j.xinn-geo.2026.100224。此处按学术研究用途引用，版权归原作者/期刊。",
        width=6.0,
    )

    doc.add_heading("2.4 文献共识与真正的空白", level=2)
    style.add_table(
        doc,
        "表L2  文献共识、争议和本研究入口",
        ["层级", "相对稳健的共识", "仍有争议/缺失", "本研究如何推进"],
        [
            ["海洋记忆", "WWV/OHC可在暖事件前积累", "表面FEP暖是否代表同等次表层状态", "用FEP、0—300 m OHC、D20/SSH同时约束"],
            ["向东到达", "赤道暖水可经西—中太平洋向东输送", "多少真正跨140°W并转化为东部储热", "用固定断面Q20与东部H20联合判断"],
            ["成熟后放电", "5°N向极输送是重要出口", "南向、海表散热、储存及残差如何分配", "用三箱控制体闭合四类去向"],
            ["2023弱耦合", "强SST与弱风压/降水并存", "弱耦合如何落实成暖量出口差异", "分解5°N地转—Ekman，并做事件统计和换风反事实"],
        ],
        [1200, 2740, 2620, 2800],
        font_size=8.0,
    )

    style.add_section_heading(doc, "3 从文献缺口提出三个中心问题")
    style.add_para(
        doc,
        "三个问题不是并列罗列，而是一个因果顺序：先证明暖量确实进入东部，才能讨论其后去向；先量化去向差异，才能把事件差异收束到风驱和压力场竞争；最后再解释相似FEP先暖为何未产生相同海气接通。",
    )
    questions = [
        (
            "问题一",
            "暖量是否、以及通过什么路径到达东太平洋？",
            "140°W东向Q20、东部H20、OHC/D20/SSH",
            "排除‘2023只是东部表面局地变暖、没有暖量输入’的解释",
        ),
        (
            "问题二",
            "到达后的暖量在北向输出、南向输出、海表散热和海洋储存之间如何分配？",
            "三箱预算、5°N/10°S Q20、表面通量、储量与残差",
            "把‘暖水去哪了’从概念相图变成ZJ级去向",
        ),
        (
            "问题三",
            "为什么相似的FEP先暖和东部暖量输入，会通过不同的海气耦合产生完全不同的暖量命运？",
            "FEP—OHC/D20—OLR/u850—Ekman/地转链条",
            "解释1997强排放、2015部分排放、2023出口近闭合",
        ),
    ]
    style.add_table(
        doc,
        "表Q1  三个固定问题及其判据",
        ["编号", "中心问题", "核心判据", "科学价值"],
        [list(row) for row in questions],
        [900, 3420, 2500, 2540],
        font_size=8.1,
    )
    for label, question, _, _ in questions:
        style.add_callout(doc, label, question, fill=style.LIGHT_BLUE)
    style.add_figure(
        doc,
        "question_preview",
        "图Q1  本研究绘制的三事件‘到达—输出—散热’问题预览。它直接显示：2015和2023发展期输入近似，但次年输出和散热不同，因而必须把三个问题连接起来回答。",
        width=6.1,
    )

    style.add_section_heading(doc, "4 回答阶段的数据选择：为什么用、各自做什么", page_break=True)
    style.add_para(
        doc,
        "数据设计遵循‘海洋主产品负责动力闭合、独立观测/再分析负责外部约束’。SODA3是主体，因为同一套温盐流速、SSH、风应力和净加热场可以在一致网格上计算H20、Q20、地转和Ekman；OISST、ERA5、GPCP与CPC OLR则防止海气结论完全依赖同一海洋同化系统。",
    )
    style.add_table(
        doc,
        "表D1  核心数据、选择理由和在三问中的作用",
        ["数据", "为什么选择", "关键变量", "具体作用", "对应问题"],
        [
            ["SODA3.15.2", "海洋状态与动力变量同网格、可做完整断面和控制体积分；这是本研究相对多数仅用SST/指数文章的主要优势", "温度、盐度、u/v、SSH、D20、风应力、net_heating", "计算H20/Q20、三箱预算、0—300 m OHC、地转—Ekman分解及2024延伸", "Q1/Q2/Q3主体"],
            ["OISST v2.1", "卫星—现场融合的独立表面观测，不受SODA3海洋动力闭合约束", "SST、FEP、Niño-3.4", "定义表面触发、事件强度和residual FEP；检验表面暖是否先于地下调整", "Q1/Q3"],
            ["ERA5", "高时空一致性的独立大气再分析，可提供风、气压和完整表面热通量分量", "u850、SLP、风应力、短长波、潜热、感热", "判定大气接通、独立估计Ekman与海表散热、解释风驱出口", "Q2/Q3"],
            ["GPCP / CPC OLR", "降水和OLR直接约束深对流，避免只用风场代替大气加热", "降水、OLR", "判断对流中心是否东移、是否形成降水—SLP—西风链条", "Q3"],
            ["TAO / ORAS5（稳健性）", "与SODA3独立的现场/再分析约束", "OHC、D20、温度断面", "检查关键时序和符号；不承担当前主文的完整预算", "Q1/Q3验证"],
        ],
        [1150, 2300, 1840, 2670, 1400],
        font_size=7.5,
    )
    style.add_table(
        doc,
        "表D2  每个问题需要的数据组合与不可替代性",
        ["问题", "必需组合", "为什么不能只用一个指标", "最终输出"],
        [
            ["Q1 到达", "SODA3 Q20/H20/OHC/D20/SSH + OISST", "SST能显示暖，却不能证明暖热量跨断面输入；单独Q20又不能证明在东部储存", "140°W输入、东部储热和次表层路径"],
            ["Q2 去向", "SODA3边界输送/储量/net_heating + ERA5表面通量", "WWV下降不能区分北出、南出、散热和跨等温面过程", "四类去向及剩余项"],
            ["Q3 耦合", "OISST FEP + SODA3 OHC/D20/SSH + ERA5 u850/SLP/风应力 + GPCP/OLR", "FEP或Niño指数不能代表对流效率、风驱强度和海洋压力场", "FEP—海洋预置—大气接通—动力出口链条"],
        ],
        [1100, 2600, 3100, 2560],
        font_size=8.0,
    )
    style.add_figure(
        doc,
        "data_stack",
        "图D1  本研究绘制的多资料时序示例：东部H20和SST（SODA3/OISST）、降水（GPCP）、5°N风应力与Ekman/总Q20（SODA3/ERA5）按月对齐。该图说明每套数据承担不同环节，不能互相替代。",
        width=6.15,
    )
    style.add_callout(
        doc,
        "数据边界",
        "SODA3是动力闭合主线，OISST/ERA5/GPCP/OLR是独立海气约束；TAO和ORAS5只作稳健性。当前报告不引入TSI、NMME或预报技巧比较，以免把‘暖量去哪了’的过程问题改写成预报排名问题。",
        fill="FFF4E5",
        color=style.GOLD,
    )
    doc.add_page_break()


def add_limitations_and_references(doc: Document) -> None:
    doc.add_heading("10.3 总结：现有闭环与仍然不足", level=2)
    style.add_para(
        doc,
        "三问已经形成一条可写成论文的方法—结果链：第一问用跨140°W Q20、东部H20及OHC/D20/SSH证明暖量到达；第二问把到达后的暖量分到5°N北向输出、10°S南界、海表热损失、储量变化和剩余项；第三问再用FEP—次表层预置—对流/风场接通—Ekman/地转竞争解释三事件分叉，并以13事件统计和局地换风试验完成定量过程归因。",
    )
    style.add_para(
        doc,
        "最稳健的新认识是：2015与2023拥有近似的发展期东部暖量输入和近似的南向地转阻力，但2015风驱Ekman足以打开北向出口，2023则不足；2023的异常不是表面不暖，而是给定FEP暖度后的中太平洋OHC缺口最大，随后大气西风和出口均弱。",
    )
    style.add_table(
        doc,
        "表LIM1  当前不足、对结论的影响与下一步",
        ["不足", "当前影响", "不能写成什么", "优先补强"],
        [
            ["欧拉预算而非拉格朗日追踪", "证明暖量通道和控制体收支，但不能识别同一团水的连续轨迹", "‘同一水团从西太平洋一路被追踪到5°N’", "离线粒子/水团标签或伴随诊断"],
            ["剩余项未完全物理拆分", "混合、夹卷、跨20°C等温面、同化增量和离散误差仍合并", "‘每一焦耳已闭合到唯一过程’", "五日三维热收支、同化增量与垂向通量"],
            ["2024年5°N通量面对温度插值敏感", "2023净方向可变，但出口受抑/近闭合稳健", "‘2023一定净向南’", "更高频断面、TAO/独立再分析与多种通量面"],
            ["事件样本仅13个", "相关和bootstrap支持过程联系，但极端事件影响仍大", "‘跨所有气候态的普适定律’", "延长资料、分层事件与多产品重复"],
            ["换风为离线反事实", "能隔离局地风驱贡献，不能再现完整耦合反馈", "‘已经完成严格因果模式归因’", "耦合/AGCM海盆因子试验与风应力敏感性集合"],
            ["上游背景尚未唯一归因", "外海盆、趋势、La Niña遗产、冷NPMM和WWB不足均可能参与", "‘某一个因素是2023唯一原因’", "完成IA×FEP集合并做因子贡献分解"],
        ],
        [1250, 2800, 2300, 3010],
        font_size=7.6,
    )
    style.add_figure(
        doc,
        "limitations",
        "图LIM1  本研究绘制的预算剩余项诊断。显式加入经向辐合与表面加热后，1997和2023残差有所下降，但混合、垂向/跨等温面过程及离散误差仍未完全分离；这正是第二问尚不能宣称‘每一焦耳完全闭合’的原因。",
        width=6.15,
    )
    style.add_callout(
        doc,
        "最终判断",
        "当前实验足以支撑一篇以‘暖量到达—去向—耦合命运’为主线的完整诊断论文；最需要补的不是再堆更多相关图，而是两项边界工作：拆分预算剩余项，以及用受控模式集合检验第三问的上游因果来源。",
        fill="FFF4E5",
        color=style.GOLD,
    )

    style.add_section_heading(doc, "参考文献与外部图件来源", page_break=True)
    references = [
        ("Jin, F.-F. (1997). An equatorial ocean recharge paradigm for ENSO. Part I: Conceptual model. Journal of the Atmospheric Sciences, 54, 811–829.", "https://doi.org/10.1175/1520-0469(1997)054%3C0811:AEORPF%3E2.0.CO;2"),
        ("Meinen, C. S., & McPhaden, M. J. (2000). Observations of warm water volume changes in the equatorial Pacific and their relationship to El Niño and La Niña. Journal of Climate, 13, 3551–3559.", "https://doi.org/10.1175/1520-0442(2000)013%3C3551:OOWWVC%3E2.0.CO;2"),
        ("Sun, D.-Z. (2000). The heat sources and sinks of the 1986–87 El Niño. Journal of Climate, 13, 3533–3550.", "https://doi.org/10.1175/1520-0442(2000)013%3C3533:THSASO%3E2.0.CO;2"),
        ("Bosc, C., & Delcroix, T. (2008). Observed equatorial Rossby waves and ENSO-related warm water volume changes in the equatorial Pacific Ocean. Journal of Geophysical Research: Oceans, 113, C06003.", "https://doi.org/10.1029/2007JC004613"),
        ("Kug, J.-S., Jin, F.-F., & An, S.-I. (2009). Two types of El Niño events: Cold Tongue El Niño and Warm Pool El Niño. Journal of Climate, 22, 1499–1515.", "https://doi.org/10.1175/2008JCLI2624.1"),
        ("Capotondi, A., et al. (2015). Understanding ENSO diversity. Bulletin of the American Meteorological Society, 96, 921–938.", "https://doi.org/10.1175/BAMS-D-13-00117.1"),
        ("Carton, J. A., Chepurin, G. A., & Chen, L. (2018). SODA3: A new ocean climate reanalysis. Journal of Climate, 31, 6967–6983.", "https://doi.org/10.1175/JCLI-D-18-0149.1"),
        ("Huguenin, M. F., Holmes, R. M., & England, M. H. (2020). Key role of diabatic processes in regulating warm water volume variability over ENSO events. Journal of Climate, 33, 9783–9801.", "https://doi.org/10.1175/JCLI-D-20-0198.1"),
        ("Chakravorty, S., Perez, R. C., Gnanaseelan, C., & Anderson, B. T. (2021). Revisiting the recharge and discharge processes for different flavors of El Niño. Journal of Geophysical Research: Oceans, 126, e2020JC017075.", "https://doi.org/10.1029/2020JC017075"),
        ("Li, X., Hu, Z.-Z., Huang, B., & Jin, F.-F. (2022). Oceanic meridional transports and their roles in warm water volume variability and ENSO in the tropical Pacific. Climate Dynamics, 59, 245–261.", "https://doi.org/10.1007/s00382-021-06124-w"),
        ("Verdy, A., Mazloff, M. R., Cornuelle, B. D., & Kim, S. B. (2023). Balancing volume, temperature, and salinity budgets during 2014–2018 in the Tropical Pacific Ocean State Estimate. Journal of Geophysical Research: Oceans, 128, e2022JC019576.", "https://doi.org/10.1029/2022JC019576"),
        ("Hong, C.-C., Sullivan, A., & Chang, C.-C. (2024). Impact of North Atlantic Tripole and Extratropical North Pacific Extreme SSTs on the 2023/24 El Niño. Geophysical Research Letters, 51, e2024GL110639.", "https://doi.org/10.1029/2024GL110639"),
        ("Kim, J.-W., Tian, B., & Yu, J.-Y. (2025). Unusual and persistent easterlies restrained the 2023/24 El Niño development after a triple-dip La Niña. npj Climate and Atmospheric Science, 8, 34.", "https://doi.org/10.1038/s41612-024-00890-0"),
        ("Lin, Y.-F., et al. (2025). Diverse NPMM conditions deviate the 2023/24 El Niño from the 1997/1998 and 2015/2016 extreme El Niño events. npj Climate and Atmospheric Science, 8, 133.", "https://doi.org/10.1038/s41612-025-01013-z"),
        ("Peng, Q., et al. (2025). Strong 2023–2024 El Niño generated by ocean dynamics. Nature Geoscience, 18, 471–478.", "https://doi.org/10.1038/s41561-025-01700-9"),
        ("Zhang, L., et al. (2025). The 2023/24 El Niño event exhibited unusually weak extratropical teleconnections. Communications Earth & Environment, 6, 595.", "https://doi.org/10.1038/s43247-025-02584-8"),
        ("Fan, H., Lin, S., Zhang, L., et al. (2026). Multi-year La Niña dampens the Southern Oscillation during subsequent El Niño through tropical Pacific-rooted dynamics. The Innovation Geoscience, 4, 100224.", "https://doi.org/10.59717/j.xinn-geo.2026.100224"),
        ("Hersbach, H., et al. (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146, 1999–2049.", "https://doi.org/10.1002/qj.3803"),
        ("Huang, B., et al. (2021). Improvements of the Daily Optimum Interpolation Sea Surface Temperature (DOISST) version 2.1. Journal of Climate, 34, 2923–2939.", "https://doi.org/10.1175/JCLI-D-20-0166.1"),
        ("Adler, R. F., et al. (2018). The Global Precipitation Climatology Project monthly analysis: An update. Atmosphere, 9, 138.", "https://doi.org/10.3390/atmos9040138"),
    ]
    for number, (citation, url) in enumerate(references, start=1):
        style.add_reference(doc, number, citation, url)
    style.add_para(
        doc,
        "外部图件说明：文献图L1、L3和L4均在图题中标明原作者、年份、原图号与DOI；文献图L2为依据Peng等（2025）公开数据重绘。所有其余主图、补充图和统计图均由本项目脚本从本地数据产物绘制。",
        indent=False,
        italic=True,
        color=style.MUTED,
    )


def add_method_figure(doc: Document) -> None:
    style.add_figure(
        doc,
        "method_decomp",
        "图M1  本研究绘制的5°N动力分解方法示例。SODA3总经向速度与由SSH梯度和温盐密度梯度重建的地转速度在同一断面上比较，再以风应力计算Ekman Q20，并把未解释部分定义为残差；右侧同时给出CT/CP事件合成及断面东端敏感性。",
        width=6.15,
    )
    style.add_callout(
        doc,
        "方法与问题的对应",
        "Q1用固定经向断面和三箱储量判断到达；Q2用控制体预算判断去向；Q3用同一5°N断面的总量—地转—Ekman分解，把大气接通差异落实为可量化的动力出口。",
        fill=style.LIGHT_BLUE,
    )


def rename_cover_and_sections(doc: Document) -> None:
    cover_replacements = {
        "论文主体草稿  |  方法与结果": "完整研究报告  |  文献—问题—回答",
        "厄尔尼诺暖量的到达、去向与海气耦合命运": "厄尔尼诺暖量去哪了？",
        "围绕三个固定科学问题的完整方法与结果部分": "从Chakravorty等（2021）到三问闭环的发现、提出与回答",
        "SODA3.15.2 · OISST · ERA5 · GPCP/OLR": "文献证据 · SODA3暖量预算 · 独立海气约束 · 事件级过程归因",
        "分析状态：前两问主体闭合；第三问形成过程机制答案": "分析状态：三问形成完整诊断链；剩余项与上游因果仍待补强",
        "2026年8月7日 · 公式严谨性修订版": "2026年8月7日 · 完整报告公式修订版",
    }
    cover_fonts = {
        "完整研究报告  |  文献—问题—回答": (11, True, False, style.GOLD),
        "厄尔尼诺暖量去哪了？": (27, True, False, style.NAVY),
        "从Chakravorty等（2021）到三问闭环的发现、提出与回答": (15, True, False, style.DARK_BLUE),
        "文献证据 · SODA3暖量预算 · 独立海气约束 · 事件级过程归因": (11, False, True, style.MUTED),
        "分析状态：三问形成完整诊断链；剩余项与上游因果仍待补强": (11, True, False, style.GREEN),
        "2026年8月7日 · 完整报告公式修订版": (10.5, False, False, style.MUTED),
    }
    for old, new in cover_replacements.items():
        paragraph = find_paragraph(doc, old)
        paragraph.text = new
        size, bold, italic, color = cover_fonts[new]
        style.set_run_font(
            paragraph.runs[0],
            font=style.HEAD_FONT if bold else style.BODY_FONT,
            size=size,
            bold=bold,
            italic=italic,
            color=color,
        )

    heading_map = {
        "1 研究问题与总体设计": "5 研究问题与总体设计",
        "2 数据与方法": "6 数据与方法",
        "2.1 数据资料": "6.1 数据资料",
        "2.2 空间区域、阶段和符号": "6.2 空间区域、阶段和符号",
        "2.3 异常、趋势和统计推断": "6.3 异常、趋势和统计推断",
        "2.4 问题一：暖量到达的诊断": "6.4 问题一：暖量到达的诊断",
        "2.5 问题二：暖量去向的固定控制体预算": "6.5 问题二：暖量去向的固定控制体预算",
        "2.6 问题三：FEP—海洋结构—大气接通—出口分解": "6.6 问题三：FEP—海洋结构—大气接通—出口分解",
        "2.7 数值质量控制": "6.7 数值质量控制",
        "3 结果一：暖量是否以及如何到达东太平洋": "7 结果一：暖量是否以及如何到达东太平洋",
        "3.1 CT与CP都能把暖量送到东部，但CT更强": "7.1 CT与CP都能把暖量送到东部，但CT更强",
        "3.2 1997输入最强，2015与2023的发展期输入近似": "7.2 1997输入最强，2015与2023的发展期输入近似",
        "4 结果二：到达后的暖量去了哪里": "8 结果二：到达后的暖量去了哪里",
        "4.1 CT/CP合成显示关键分叉位于5°N北界": "8.1 CT/CP合成显示关键分叉位于5°N北界",
        "4.2 三事件形成强排放—部分排放—出口近闭合阶梯": "8.2 三事件形成强排放—部分排放—出口近闭合阶梯",
        "4.3 2023出口方向敏感，但“出口受抑”稳健": "8.3 2023出口方向敏感，但“出口受抑”稳健",
        "5 结果三：相似FEP先暖为何产生不同暖量命运": "9 结果三：相似FEP先暖为何产生不同暖量命运",
        "5.1 FEP先暖只描述表面起点，不代表相同海洋状态": "9.1 FEP先暖只描述表面起点，不代表相同海洋状态",
        "5.2 决定性分叉位于表面增暖之后、深对流东移之前": "9.2 决定性分叉位于表面增暖之后、深对流东移之前",
        "5.3 大气接通差异通过Ekman—地转竞争落实为暖量命运": "9.3 大气接通差异通过Ekman—地转竞争落实为暖量命运",
        "5.4 第三问的机制答案及因果边界": "9.4 第三问的机制答案及因果边界",
        "5.5 十三事件统计与风应力替换完成定量过程归因": "9.5 十三事件统计与风应力替换完成定量过程归因",
        "6 稳健性、完成度与结果边界": "10 稳健性、完成度与结果边界",
        "6.1 基准和状态方程敏感性": "10.1 基准和状态方程敏感性",
        "6.2 三个问题的完成度": "10.2 三个问题的完成度",
    }
    for old, new in heading_map.items():
        replace_paragraph_text(doc, old, new)


def build_report() -> Path:
    base_path = base.build_report()
    doc = Document(base_path)
    rename_cover_and_sections(doc)
    bullet_id = style.create_numbering(doc, numbered=False)
    number_id = style.create_numbering(doc, numbered=True)

    section_anchor = find_paragraph(doc, "5 研究问题与总体设计")._p
    marker = doc.add_paragraph("__FRONT_INSERT__")._p
    add_front_matter(doc, bullet_id, number_id)
    move_new_body_elements_before(doc, marker, section_anchor)

    results_anchor = find_paragraph(doc, "7 结果一：暖量是否以及如何到达东太平洋")._p
    marker = doc.add_paragraph("__METHOD_INSERT__")._p
    add_method_figure(doc)
    move_new_body_elements_before(doc, marker, results_anchor)

    appendix_anchor = find_paragraph(doc, "附录A 直接产物与复现入口")._p
    marker = doc.add_paragraph("__BACK_INSERT__")._p
    add_limitations_and_references(doc)
    move_new_body_elements_before(doc, marker, appendix_anchor)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip() == "ENSO暖量去向研究报告":
                paragraph.text = "ENSO暖量到达—去向—耦合命运完整报告"
                style.set_run_font(paragraph.runs[0], size=8.5, bold=True, color=style.MUTED)
        for table in section.footer.tables:
            for cell in table._cells:
                for paragraph in cell.paragraphs:
                    if "厄尔尼诺暖量去向与海气耦合闭环" in paragraph.text:
                        for run in paragraph.runs:
                            if "厄尔尼诺暖量去向与海气耦合闭环" in run.text:
                                run.text = run.text.replace(
                                    "厄尔尼诺暖量去向与海气耦合闭环",
                                    "暖量到达—去向—耦合命运",
                                )
                                style.set_run_font(run, size=8.5, color=style.MUTED)

    props = doc.core_properties
    props.title = "厄尔尼诺暖量去哪了：三个中心问题的发现、提出与回答"
    props.subject = "从权威文献到SODA3暖量预算、海气耦合诊断和事件级过程归因"
    props.keywords = "ENSO; warm heat fate; SODA3; FEP; H20; Qwarm,20; Ekman; geostrophic; air-sea coupling"
    props.comments = "Formula-rigour revision generated from verified project outputs on 2026-08-07."

    normalize_warm_transport_notation(doc)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
