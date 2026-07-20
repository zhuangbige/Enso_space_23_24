#!/usr/bin/env python3
"""Build the January-June DOCX report for the 2015/16, 2023/24, and 2026 comparison."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/research_2026/reports/enso_event_comparison_unified_1980_2014_jan_jun_2026.docx"

EVENTS = {
    "2015/16": {
        "npz": ROOT / "data/interim/research_2026/monthly_panel_inputs/monthly_2015_detrended_fields_1980_2014.npz",
        "label": "2015/16 development year",
        "ssta": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2015_revised_ssta_u850_detrended_1980_2014.png",
        "precip": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2015_paperstyle_era5_precip_u850_strip_detrended_1980_2014.png",
        "d20": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2015_revised_d20_u850_detrended_1980_2014.png",
    },
    "2023/24": {
        "npz": ROOT / "data/interim/research_2026/monthly_panel_inputs/monthly_2023_detrended_fields_1980_2014.npz",
        "label": "2023/24 development year",
        "ssta": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2023_revised_ssta_u850_detrended_1980_2014.png",
        "precip": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2023_paperstyle_era5_precip_u850_strip_detrended_1980_2014.png",
        "d20": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2023_revised_d20_u850_detrended_1980_2014.png",
    },
    "2026": {
        "npz": ROOT / "data/interim/research_2026/monthly_panel_inputs/monthly_2026_detrended_fields_1980_2014.npz",
        "label": "2026 January-June",
        "ssta": ROOT / "outputs/research_2026/revised_event_figures/fig_2026_revised_ssta_u850_detrended_1980_2014.png",
        "precip": ROOT
        / "outputs/research_2026/revised_event_figures/fig_2026_paperstyle_era5_precip_u850_strip_detrended_1980_2014.png",
        "d20": ROOT / "outputs/research_2026/revised_event_figures/fig_2026_revised_d20_u850_detrended_1980_2014.png",
    },
}

ORDINARY_DIR = ROOT / "outputs/research_2026/revised_event_figures"
ORDINARY_EXPECTED_CSV = ORDINARY_DIR / "ordinary_el_nino_expected_response.csv"
SIMILARITY_CSV = ORDINARY_DIR / "event_2026_taylor_similarity_to_2015_2023.csv"
ORDINARY_FIGURES = {
    "expected": ORDINARY_DIR / "fig_targets_observed_expected_atmospheric_response_1982_2014.png",
    "staged": ORDINARY_DIR / "fig_2026_staged_air_sea_coupling_jan_jun_1980_2014.png",
}

REGIONS = {
    "Nino3.4": (-5, 5, 190, 240),
    "CP": (-5, 5, 180, 220),
}
D20_REGIONS = {"West": (130, 180), "Central": (180, 220), "East": (220, 275)}


def set_run_font(run, size: float | None = None, bold: bool = False, italic: bool = False) -> None:
    has_cjk = any(0x2E80 <= ord(char) <= 0x9FFF for char in run.text)
    # STSong is exposed to both Word on macOS and the headless LibreOffice
    # renderer used for report QA.
    font_name = "STSong" if has_cjk else "Times New Roman"
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:cs"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str = "FFFFFF") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def mark_table_header(table) -> None:
    """Mark the first row as a repeating semantic table header."""

    properties = table.rows[0]._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        properties.append(repeat_header)


def area_mean(field: np.ndarray, lat: np.ndarray, lon: np.ndarray, box: tuple[float, float, float, float]) -> np.ndarray:
    lat0, lat1, lon0, lon1 = box
    lm = (lat >= lat0) & (lat <= lat1)
    om = (lon >= lon0) & (lon <= lon1)
    sub = field[:, lm, :][:, :, om]
    weights = np.cos(np.deg2rad(lat[lm]))
    return np.nansum(sub * weights[None, :, None], axis=(1, 2)) / (np.sum(weights) * om.sum())


def d20_mean(field: np.ndarray, lon: np.ndarray, bounds: tuple[float, float]) -> np.ndarray:
    m = (lon >= bounds[0]) & (lon <= bounds[1])
    return np.nanmean(field[:, m], axis=1)


def load_event_metrics() -> dict[str, dict[str, float]]:
    metrics: dict[str, dict[str, float]] = {}
    for event, info in EVENTS.items():
        z = np.load(info["npz"])
        lat = z["lat"].astype(float)
        lon = z["lon"].astype(float)
        slat = z["sst_lat"].astype(float)
        slon = z["sst_lon"].astype(float)
        dlon = z["d20_lon"].astype(float)
        out: dict[str, float] = {}
        for region, box in REGIONS.items():
            sst_values = area_mean(z["sst"], slat, slon, box)
            pr_values = area_mean(z["precip"], lat, lon, box)
            out[f"sst_{region}"] = float(np.nanmean(sst_values))
            out[f"pr_{region}"] = float(np.nanmean(pr_values))
            out[f"sst_{region}_june"] = float(sst_values[-1])
            out[f"pr_{region}_june"] = float(pr_values[-1])
        for region, bounds in D20_REGIONS.items():
            d20_values = d20_mean(z["d20"], dlon, bounds)
            out[f"d20_{region}"] = float(np.nanmean(d20_values))
            out[f"d20_{region}_june"] = float(d20_values[-1])
        out["sst_baseline_start"] = 1982.0
        metrics[event] = out

    expected = pd.read_csv(ORDINARY_EXPECTED_CSV)
    labels = {"2015/16": "2015/16", "2023/24": "2023/24", "2026": "2026 developing state"}
    for event, target_label in labels.items():
        wind = expected.loc[
            expected["target_label"].eq(target_label) & expected["response"].eq("cwp_u10")
        ].sort_values("month")
        if wind["month"].astype(int).tolist() != list(range(1, 7)):
            raise ValueError(f"Incomplete Peng CWP u10 response rows for {target_label}")
        metrics[event]["u_CWP"] = float(wind["observed"].mean())
        metrics[event]["u_CWP_june"] = float(wind.iloc[-1]["observed"])
    return metrics


def load_ordinary_analysis() -> dict[str, pd.DataFrame]:
    return {
        "expected": pd.read_csv(ORDINARY_EXPECTED_CSV),
        "similarity": pd.read_csv(SIMILARITY_CSV),
    }


def add_paragraph(doc: Document, text: str = "", style: str | None = None, size: float = 10.5) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size={1: 15, 2: 13, 3: 11.5}.get(level, 11.5), bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    set_run_font(run, size=8.7, italic=True)


def add_figure(doc: Document, image: Path, caption: str, width: float = 6.35) -> None:
    if not image.exists():
        raise FileNotFoundError(image)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(image), width=Inches(width))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def add_metric_table(doc: Document, metrics: dict[str, dict[str, float]]) -> None:
    add_heading(doc, "表1  三次事件的关键区域指标", level=2)
    add_paragraph(
        doc,
        "每格为1-6月平均 / 6月单月；Peng CWP u10为正表示近地面西风异常。",
        size=9.2,
    )
    rows = [
        ("Niño3.4 SST (°C)", "sst_Nino3.4", "{:+.2f}"),
        ("Peng CWP u10 (m s⁻¹)", "u_CWP", "{:+.2f}"),
        ("CP降水 (mm day⁻¹)", "pr_CP", "{:+.2f}"),
        ("西部D20 (m)", "d20_West", "{:+.1f}"),
        ("东部D20 (m)", "d20_East", "{:+.1f}"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    mark_table_header(table)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ["指标", "2015/16", "2023/24", "2026"]):
        cell.text = text
    for label, key, fmt in rows:
        cells = table.add_row().cells
        cells[0].text = label
        for i, event in enumerate(["2015/16", "2023/24", "2026"], start=1):
            cells[i].text = f"{fmt.format(metrics[event][key])} / {fmt.format(metrics[event][key + '_june'])}"
    set_table_width(table, [2.2, 1.43, 1.43, 1.43])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "FFFFFF")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    set_run_font(run, size=8.6, bold=(row is table.rows[0]))


def format_compact_table(table, widths: list[float]) -> None:
    mark_table_header(table)
    set_table_width(table, widths)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=65, start=90, bottom=65, end=90)
            set_cell_shading(cell, "FFFFFF")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    set_run_font(run, size=8.1, bold=(row is table.rows[0]))


def add_expected_response_table(doc: Document, ordinary: dict[str, pd.DataFrame]) -> None:
    expected = ordinary["expected"]
    means = (
        expected.groupby(["target_label", "response"], as_index=False)[["observed", "expected", "residual"]]
        .mean(numeric_only=True)
        .set_index(["target_label", "response"])
    )
    add_heading(doc, "表2  SST条件下的观测—历史预期响应", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["响应（1-6 月平均）", "2015/16", "2023/24", "2026 发展状态"]):
        cell.text = label
    rows = [
        ("Peng CWP u10 (m/s)", "cwp_u10"),
        ("CP precipitation (mm/day)", "cp_precip"),
    ]
    labels = {2015: "2015/16", 2023: "2023/24", 2026: "2026 developing state"}
    for metric_label, response in rows:
        cells = table.add_row().cells
        cells[0].text = metric_label
        for index, year in enumerate((2015, 2023, 2026), start=1):
            row = means.loc[(labels[year], response)]
            cells[index].text = f"{row.observed:+.2f} / {row.expected:+.2f} ({row.residual:+.2f})"
    format_compact_table(table, [2.05, 1.48, 1.48, 1.48])
    add_paragraph(
        doc,
        "每格为观测 / 历史预期（残差）。预期由1982-2014各日历月Niño3 SST回归得到；三个目标年份均排除在拟合之外。",
        size=9.2,
    )


def add_similarity_table(doc: Document, ordinary: dict[str, pd.DataFrame]) -> None:
    similarity = ordinary["similarity"]
    add_heading(doc, "表3  2026与2015、2023的分层相似性", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["物理层", "参照事件", "相关 r", "中心化 RMSD", "标准差比"]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    layer_labels = {
        "Tropical-Pacific SST": "SST",
        "850 hPa zonal wind": "u850",
        "Precipitation": "降水",
        "Equatorial D20": "D20",
    }
    for layer in ("Tropical-Pacific SST", "850 hPa zonal wind", "Precipitation", "Equatorial D20"):
        subset = similarity.loc[similarity["layer"].eq(layer)].sort_values("candidate_year")
        for row in subset.itertuples(index=False):
            cells = table.add_row().cells
            cells[0].text = layer_labels[layer]
            cells[1].text = str(row.candidate_label)
            cells[2].text = f"{row.pattern_correlation:.3f}"
            cells[3].text = f"{row.centered_rmsd:.2f} {row.unit}"
            cells[4].text = f"{row.standard_deviation_ratio:.3f}"
    format_compact_table(table, [1.05, 1.35, 1.0, 1.65, 1.2])
    add_paragraph(
        doc,
        "Taylor统计量分别约束形态相关、中心化误差和相对振幅；各变量独立评价，不合成总分。",
        size=9.2,
    )


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    for style_name, size, bold, before, after in [
        ("Normal", 10.5, False, 0, 5),
        ("Heading 1", 15, True, 14, 6),
        ("Heading 2", 13, True, 10, 5),
        ("Heading 3", 11.5, True, 8, 4),
        ("Caption", 8.7, False, 1, 6),
    ]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for key in ("ascii", "hAnsi", "cs"):
            rfonts.set(qn(f"w:{key}"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "Hiragino Sans GB")


def build_doc() -> None:
    metrics = load_event_metrics()
    analysis = load_ordinary_analysis()
    expected = analysis["expected"]
    wind_2026 = expected.loc[
        expected["target_year"].eq(2026) & expected["response"].eq("cwp_u10")
    ].sort_values("month")
    precipitation_2026 = expected.loc[
        expected["target_year"].eq(2026) & expected["response"].eq("cp_precip")
    ].sort_values("month")
    if wind_2026["month"].astype(int).tolist() != list(range(1, 7)):
        raise ValueError("Incomplete January-June 2026 Peng CWP u10 response")
    wind_observed = wind_2026["observed"].to_numpy(float)
    wind_expected = wind_2026["expected"].to_numpy(float)
    wind_residual = wind_2026["residual"].to_numpy(float)
    wind_mean_residual = float(np.mean(wind_residual))
    precip_mean_residual = float(precipitation_2026["residual"].mean())
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "2026 ENSO海气耦合诊断"
    doc.core_properties.subject = "2015/16、2023/24与2026年1-6月观测对比"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("2026 ENSO 海气耦合诊断")
    set_run_font(run, size=20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("2015/16与2023/24参照；观测截至2026年6月")
    set_run_font(run, size=10.5)

    add_heading(doc, "摘要", level=1)
    add_paragraph(
        doc,
        "2026年初的海温分布带有2023式东部先行特征，但至6月，赤道温跃层和低层纬向风已更接近2015式强耦合状态。"
        f"Niño3 SST条件回归显示，2026年1-6月Peng CWP u10相对历史预期的平均残差为{wind_mean_residual:+.2f} m s⁻¹，"
        f"而中太平洋降水残差为{precip_mean_residual:+.2f} mm day⁻¹。"
        "观测序列因此表现为温跃层调整和西风建立在前、深对流响应在5-6月追赶。",
    )

    add_heading(doc, "研究定位与方法依据", level=1)
    add_paragraph(
        doc,
        "本文不提出新的综合耦合指数，而是把已有的ENSO过程诊断用于2015/16、2023/24与2026年发展阶段的统一比较。"
        "Peng et al.（2025）以Niño3 SST为海洋端，通过历史回归估计中西太平洋纬向风的条件响应，并比较2023/24的实际值与历史预期；"
        "Jiang et al.（2026）同时比较东太平洋SST、中太平洋风和降水，并以历史事件和回归重建识别空间受限的海气耦合；"
        "Li et al.（2024）进一步把大气Bjerknes反馈分解为SST梯度—降水加热—气压梯度—纬向风应力。"
        "因此，本文的实际—预期响应以及SST、降水和风场分环分析均有直接的方法学依据；本文的增量是将其用于2026年逐月观测并与2015/16和2023/24分层比较。",
    )
    add_paragraph(
        doc,
        "本文严格采用Peng et al.（2025）的CWP定义：5°S–5°N、140°E–160°W区域平均10 m纬向风（u10），"
        "用于衡量Niño3 SST对应的近地面大气响应。CP降水采用5°S–5°N、180°–140°W，与Sun et al.（2016）的中太平洋诊断区一致。"
        "CWP u10和CP降水均是给定坐标的研究诊断量，不作为类似Niño3或Niño3.4的业务监测指数；三次事件使用相同坐标、异常基准和处理。"
        "空间图仍叠加850 hPa风，用于展示低层环流结构，不与Peng的CWP u10指标混用。",
    )

    add_heading(doc, "1  资料与方法", level=1)
    add_paragraph(
        doc,
        "SST取自NOAA OISST；10 m风、850 hPa风和降水取自ERA5；20°C等温面深度（D20）取自ORAS5。"
        "其中2026年6月u10由ARCO-ERA5T公共镜像的720个逐小时场聚合为月平均，其余u10使用ERA5月平均资料。"
        "ERA5 850 hPa风、降水及ORAS5异常基准为1980-2014，OISST和Peng CWP u10为1982-2014；各变量逐日历月去气候态并去线性趋势。"
        "三次事件均比较发展年1-6月。",
    )
    add_paragraph(
        doc,
        "大气条件响应采用1982-2014各日历月Niño3 SST回归，2015、2023和2026不进入拟合。"
        "历史相似性采用Taylor（2001）的面积加权相关、中心化RMSD和标准差比，分别评价SST、u850、降水和D20；不构造综合指数。"
        "图11的三段阴影仅按月份归纳观测时序，不是正式耦合指数、统计识别的变点或事件强度预报；本文据此只判断共同演变是否与反馈逐步建立一致，不作因果闭合判定。",
    )
    add_metric_table(doc, metrics)

    add_heading(doc, "2  三次事件的海气演变", level=1)
    add_heading(doc, "2.1  SST与850 hPa风", level=2)
    add_paragraph(
        doc,
        "2015年中东太平洋增暖与西风异常同步发展；2023年暖中心偏东，西风响应较弱。2026年1-3月同样由东部先行，"
        "4月后暖异常向中太平洋扩展，6月850 hPa西风异常显著增强。该图用于描述低层环流结构；CWP实际—预期计算采用10 m风。",
    )
    add_figure(doc, EVENTS["2015/16"]["ssta"], "图1  2015/16发展年OISST海温异常与ERA5 850 hPa风异常。")
    add_figure(doc, EVENTS["2023/24"]["ssta"], "图2  2023/24发展年OISST海温异常与ERA5 850 hPa风异常。")
    add_figure(doc, EVENTS["2026"]["ssta"], "图3  2026年1-6月OISST海温异常与ERA5 850 hPa风异常。")

    add_heading(doc, "2.2  降水与850 hPa风", level=2)
    add_paragraph(
        doc,
        "2015年中太平洋正降水异常与西风共同增强；2023年的中太平洋对流响应偏弱。2026年1-4月中太平洋降水仍受抑制，"
        "5月转正并在6月增强，但其振幅仍低于同期SST所对应的历史条件响应。",
    )
    add_figure(doc, EVENTS["2015/16"]["precip"], "图4  2015/16发展年ERA5降水与850 hPa风异常。")
    add_figure(doc, EVENTS["2023/24"]["precip"], "图5  2023/24发展年ERA5降水与850 hPa风异常。")
    add_figure(doc, EVENTS["2026"]["precip"], "图6  2026年1-6月ERA5降水与850 hPa风异常。")

    add_heading(doc, "2.3  D20与赤道u850", level=2)
    add_paragraph(
        doc,
        "2026年6月西部D20为-12.1 m、东部为+46.6 m，形成清楚的西浅东深结构；2023年同期西部D20仅-2.7 m。"
        "2026年的盆尺度温跃层调整因而明显更接近2015。",
    )
    add_figure(doc, EVENTS["2015/16"]["d20"], "图7  2015/16发展年ORAS5赤道D20异常与ERA5赤道u850。")
    add_figure(doc, EVENTS["2023/24"]["d20"], "图8  2023/24发展年ORAS5赤道D20异常与ERA5赤道u850。")
    add_figure(doc, EVENTS["2026"]["d20"], "图9  2026年1-6月ORAS5赤道D20异常与ERA5赤道u850。")

    add_heading(doc, "3  SST条件下的大气响应", level=1)
    add_expected_response_table(doc, analysis)
    add_figure(
        doc,
        ORDINARY_FIGURES["expected"],
        "图10  三次事件的观测大气响应及Niño3 SST条件历史预期；回归期为1982-2014。",
    )
    add_paragraph(
        doc,
        "2015年近地面西风明显强于SST条件预期；2023年风和降水均偏弱。2026年则呈分层差异："
        f"1-6月Peng CWP u10平均残差为{wind_mean_residual:+.2f} m s⁻¹，CP降水残差为{precip_mean_residual:+.2f} mm day⁻¹。"
        "该结果用于区分近地面风响应和深对流响应，不能由850 hPa风替代。",
    )

    add_heading(doc, "4  2026年耦合建立次序", level=1)
    add_figure(
        doc,
        ORDINARY_FIGURES["staged"],
        "图11  2026年1-6月D20、Niño3 SST、Peng CWP u10及CP降水的观测与历史条件响应。阴影仅为描述性月份分组，不代表正式阶段或统计变点。",
    )
    add_heading(doc, "4.1  温跃层调整（图11a）", level=2)
    add_paragraph(
        doc,
        "1月东部D20已加深至+21.0 m，而西部仅为-2.3 m，表明事件开始时东太平洋已有明显的次表层暖水条件，但盆地尺度倾斜尚未充分建立。"
        "2-4月东部D20维持+22.1至+30.7 m，西部D20由-4.6 m发展到-7.9 m。5月西部浅化至-18.3 m、东部加深至+48.5 m，"
        "东西D20差达到66.8 m；6月仍保持-12.1 m与+46.6 m。该演变显示暖水体向中东太平洋重新分配，温跃层西浅东深结构在5月完成盆地尺度强化。",
    )
    add_heading(doc, "4.2  东太平洋表层增暖（图11b）", level=2)
    add_paragraph(
        doc,
        "Niño3 SST异常由1月+0.05°C升至2月+0.63°C，3-4月暂时维持在+0.49至+0.61°C，随后于5月和6月快速增至+1.25°C和+1.66°C。"
        "表层增暖的第二次加速与5月温跃层倾斜显著增强同步。东部温跃层加深可减弱冷水卷入混合层，西部温跃层浅化则表明暖水体东移；"
        "因此D20与SST的共同演变符合温跃层反馈对东太平洋增暖的维持作用。",
    )
    add_heading(doc, "4.3  低层纬向风响应（图11c）", level=2)
    add_paragraph(
        doc,
        f"Peng CWP u10在1月为{wind_observed[0]:+.2f} m s⁻¹，SST条件预期为{wind_expected[0]:+.2f} m s⁻¹；"
        f"2月观测为{wind_observed[1]:+.2f} m s⁻¹，预期为{wind_expected[1]:+.2f} m s⁻¹。"
        f"3-5月观测值依次为{wind_observed[2]:+.2f}、{wind_observed[3]:+.2f}和{wind_observed[4]:+.2f} m s⁻¹；"
        f"6月为{wind_observed[5]:+.2f} m s⁻¹，相对预期的残差为{wind_residual[5]:+.2f} m s⁻¹。"
        "该序列直接表示海面附近西风响应；850 hPa风仅作为空间环流背景。近地面西风、温跃层调整和东部增暖的共同演变"
        "与Bjerknes反馈逐步建立一致，但历史回归本身不证明因果闭合。",
    )
    add_heading(doc, "4.4  深对流响应（图11d）", level=2)
    add_paragraph(
        doc,
        "CP降水在1-4月分别为-1.24、-1.46、-2.14和-2.04 mm day⁻¹，始终低于SST条件预期，4月残差达到-3.87 mm day⁻¹。"
        "5月降水转为+0.42 mm day⁻¹，但仍比预期低2.22 mm day⁻¹；6月增至+2.40 mm day⁻¹，与+3.02 mm day⁻¹的预期值已较接近，"
        "残差收窄至-0.61 mm day⁻¹。深对流响应在时间上晚于D20、SST和西风调整，5-6月才明显恢复。至6月，观测顺序呈现温跃层调整与东部增暖同步，"
        "增暖伴随中西太平洋西风增强，随后中太平洋对流恢复并与西风异常并存。这一共同演变与海气正反馈逐步闭合一致，但月平均观测和历史回归不能证明反馈环已完成因果闭合；对流强度也尚未完全达到历史条件响应。",
    )

    add_heading(doc, "5  历史相似性与演变判断", level=1)
    add_similarity_table(doc, analysis)
    add_paragraph(
        doc,
        "2026的SST形态更接近2023（r=0.683；中心化RMSD=0.51°C），u850和D20则更接近2015：两者与2015的相关分别为"
        "0.504和0.880，D20标准差比为0.991。降水未稳定归入任一事件。物理上，2026是东部表层先行、随后转入盆尺度动力耦合的过程，"
        "不是2015或2023的简单复刻。",
    )
    add_paragraph(
        doc,
        "截至6月的观测支持继续向强耦合、盆尺度El Niño发展，但类比统计不能约束峰值。独立预报中，NOAA CPC于2026年7月9日"
        "给出持续至2027年早春97%、2026年10-12月达到very strong级别81%的概率；WMO于7月3日判断7-9月将快速增强为强El Niño。",
    )

    add_heading(doc, "6  结论", level=1)
    add_paragraph(
        doc,
        "2026年1-6月的核心特征是D20调整、Peng CWP近地面风响应和850 hPa低层环流共同向耦合状态发展，而表层SST仍保留2023式东部先行结构。"
        "深对流在5-6月明显恢复，但相对于SST条件预期仍偏弱。后续需要重点检验中太平洋对流能否在夏秋持续维持，以及西风—海洋调整能否保持协同。",
    )

    add_heading(doc, "参考方法与资料", level=1)
    add_paragraph(
        doc,
        "Peng et al. (2025), Nature Geoscience, doi:10.1038/s41561-025-01700-9：CWP（5°S–5°N、140°E–160°W）u10、Niño3条件回归及实际—预期响应。\n"
        "Jiang et al. (2026), Geophysical Research Letters, doi:10.1029/2026GL123227：东太平洋先行增暖、中太平洋风—降水响应及跨事件回归比较。\n"
        "Li et al. (2024), Geophysical Research Letters, doi:10.1029/2024GL109874：SST—降水加热—气压—风应力的大气Bjerknes反馈分环诊断。\n"
        "Sun et al. (2016), Advances in Atmospheric Sciences, doi:10.1007/s00376-015-5018-6：Niño3 SST对应的中太平洋降水和风应力响应，以及5°S–5°N、180°–140°W诊断区。\n"
        "Kim, Tian & Yu (2025), npj Climate and Atmospheric Science, doi:10.1038/s41612-024-00890-0：D20、西太平洋暖水体和西—中太平洋纬向风的事件演变与跨事件对比。\n"
        "Taylor (2001), Journal of Geophysical Research, doi:10.1029/2000JD900719：相关、中心化RMSD和标准差比。\n"
        "Zuo et al. (2019), Ocean Science, doi:10.5194/os-15-779-2019：ORAS5系统、资料同化与海洋状态重建。\n"
        "NOAA CPC ENSO Diagnostic Discussion, 9 July 2026；WMO El Niño Update, 3 July 2026。",
        size=9.2,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
