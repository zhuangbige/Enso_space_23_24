#!/usr/bin/env python3
"""Translate the current Chinese ENSO event comparison DOCX into English."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs/research_2026/reports/enso_event_comparison_unified_1980_2014.docx"
OUT = ROOT / "outputs/research_2026/reports/enso_event_comparison_unified_1980_2014_en.docx"


TRANSLATIONS = {
    "2015/16、2023/24 与 2026 ENSO 发展期对比分析": "Comparison of the 2015/16, 2023/24, and 2026 ENSO Development Phases",
    "统一气候态和去趋势基准下的 1-5 月 SST-Wind、Precipitation-Wind、D20-Wind 与 PMM 诊断": (
        "January-May SST-Wind, Precipitation-Wind, D20-Wind, and PMM diagnostics using a unified climatology and detrending baseline"
    ),
    "分析口径": "Analysis Framework",
    "本报告只讨论最新重画的三组 1-5 月事件图。SST 使用 NOAA OISST，850 hPa 风和降水使用 ERA5，D20 使用 ORAS5。SST、风、降水和 D20 均按同月气候态计算异常，并移除线性趋势；ERA5/ORAS5 共同基准为 1980-2014，OISST 因资料起点限制使用 1982-2014。PMM 图为独立的 PMM-like EOF 诊断，EOF 训练期为 1982-2025，2026 年 1-5 月只投影到该 EOF 模态上，不参与 EOF 训练。": (
        "This report only discusses the three most recently redrawn January-May event figure sets. SST is from NOAA OISST, 850 hPa wind and precipitation are from ERA5, and D20 is from ORAS5. SST, wind, precipitation, and D20 anomalies are all calculated relative to the same-month climatology and then linearly detrended. The common ERA5/ORAS5 baseline is 1980-2014, while OISST uses 1982-2014 because of its data availability. The PMM figure is an independent PMM-like EOF diagnostic: the EOF training period is 1982-2025, and January-May 2026 is only projected onto that EOF mode rather than included in the EOF training."
    ),
    "空间上重点比较中太平洋、东太平洋和远东太平洋的 SST 异常；时间上重点看 1 月到 5 月暖中心、低层风、降水和 D20 是否同步推进。海气耦合的判据不是单一变量，而是 SST、u850、降水和 D20 是否在相近区域、相近月份共同增强。": (
        "Spatially, the focus is on comparing SST anomalies in the central Pacific, eastern Pacific, and far eastern Pacific. Temporally, the focus is on whether the warm center, lower-tropospheric winds, precipitation, and D20 progress together from January to May. Air-sea coupling is not judged from a single variable; it is evaluated by whether SST, u850, precipitation, and D20 strengthen together in similar regions and months."
    ),
    "统一气候态后的关键区域指标": "Key Regional Metrics after Applying the Unified Climatology",
    "下表均为 1-5 月平均。SST 单位为 degC，u850 单位为 m/s，降水单位为 mm/day，D20 单位为 m。u850 为正代表西风异常，为负代表东风异常。": (
        "The table reports January-May means. SST is in degC, u850 is in m/s, precipitation is in mm/day, and D20 is in m. Positive u850 indicates westerly anomalies; negative u850 indicates easterly anomalies."
    ),
    "综合判断": "Overall Assessment",
    "统一气候态后，2026 不再只是一个表层偏弱的东太平洋暖事件。它的 Niño3.4、CP 和 EP SST 均比原先不统一基准时更高，同时 CWP/CP 低层西风和东太平洋 D20 也更强。因此，2026 在动力结构上更接近 2015，而不是 2023。": (
        "After applying the unified climatology, 2026 no longer appears to be only a weak surface warming event in the eastern Pacific. Its Nino-3.4, CP, and EP SST anomalies are all higher than in the earlier non-unified baseline comparison, while CWP/CP lower-tropospheric westerlies and eastern Pacific D20 are also stronger. Dynamically, 2026 is therefore closer to 2015 than to 2023."
    ),
    "三者最主要的差异是海气耦合的成熟度。2015 的西风、D20、SST 和中太平洋降水较早协同增强；2023 则以远东太平洋暖异常先行，中太平洋和 PMM 支持都较弱；2026 的 D20 和低层风已经接近 2015 式结构，但降水仍滞后，说明深对流耦合尚未完全跟上。": (
        "The main difference among the three events is the maturity of air-sea coupling. In 2015, westerlies, D20, SST, and central Pacific precipitation strengthened together relatively early. In 2023, far eastern Pacific warming led the evolution, while central Pacific and PMM support remained weaker. In 2026, D20 and lower-tropospheric winds have already approached a 2015-like structure, but precipitation is still lagging, indicating that deep-convective coupling has not fully caught up."
    ),
    "SST 异常与 850 hPa 风：暖中心和东西风的空间传播": "SST Anomalies and 850 hPa Winds: Spatial Propagation of the Warm Center and Zonal Wind Anomalies",
    "2015 的早春信号更偏中太平洋到中东太平洋，并伴随持续西风异常。2023 的暖中心明显偏远东太平洋，随后逐步向中东太平洋扩展。2026 早期也有东部增强，但 4-5 月中太平洋 SST 与西风都开始接入，使它不像 2023 那样纯粹偏东。": (
        "The early-spring 2015 signal was more centered from the central Pacific to the central-eastern Pacific and was accompanied by persistent westerly anomalies. In 2023, the warm center was clearly displaced toward the far eastern Pacific and then gradually expanded toward the central-eastern Pacific. In 2026, eastern Pacific strengthening was also present early on, but by April-May both central Pacific SST and westerlies had begun to connect with the eastern signal, making it less purely eastern-Pacific-focused than 2023."
    ),
    "降水与 850 hPa 风：大气响应是否跟上海温和风场": "Precipitation and 850 hPa Winds: Whether the Atmospheric Response Keeps Pace with SST and Wind",
    "降水图显示 2015 的中太平洋对流响应最早、最一致；2023 在远东太平洋 SST 增暖期间，中太平洋降水仍偏弱，说明大气耦合偏晚。2026 的关键点是：西风异常已经比较像 2015，但中太平洋降水 1-4 月仍偏抑制，5 月才局部恢复。": (
        "The precipitation maps show that the central Pacific convective response in 2015 was the earliest and most coherent. During the 2023 far eastern Pacific SST warming, central Pacific precipitation remained relatively weak, indicating delayed atmospheric coupling. The key point for 2026 is that the westerly anomalies already resemble 2015, but central Pacific precipitation remained suppressed from January to April and only recovered locally in May."
    ),
    "D20 与赤道 u850：次表层动力结构": "D20 and Equatorial u850: Subsurface Dynamical Structure",
    "D20 是2015 和 2026 相像的核心证据。2015 和 2026 都表现出西太平洋相对浅化、中东太平洋加深的结构，对应暖水体向中东太平洋倾斜或输送。2023 虽然东太平洋 D20 也加深，但西太平洋并未呈现同样清楚的浅化，因此动力结构不如 2015/2026 相似。": (
        "D20 is the core evidence for the similarity between 2015 and 2026. Both 2015 and 2026 show relative shoaling in the western Pacific and deepening in the central-eastern Pacific, corresponding to a tilt or transport of warm water toward the central-eastern Pacific. Although eastern Pacific D20 also deepened in 2023, the western Pacific did not show the same clear shoaling; therefore, its dynamical structure is less similar to 2015/2026."
    ),
    "PMM：北太平洋副热带背景是否参与": "PMM: Whether the North Pacific Subtropical Background Is Involved",
    "PMM-like EOF 表示扣除 Niño3.4 线性影响以后，北太平洋副热带 SST 与近地面风场的主要耦合模态。正位相对应东北太平洋偏暖、信风减弱和较弱蒸发冷却的背景。2015 和 2026 的 joint SST+wind PMM 指数在 1-5 月均明显为正，而 2023 较弱甚至偏负，说明 2026 的副热带海气背景更像 2015。": (
        "The PMM-like EOF represents the leading coupled mode of North Pacific subtropical SST and near-surface winds after removing the linear influence of Nino-3.4. Its positive phase corresponds to warm northeastern Pacific SST, weakened trade winds, and reduced evaporative cooling. The joint SST+wind PMM index is clearly positive during January-May in both 2015 and 2026, whereas it is weak or even negative in 2023. This suggests that the subtropical air-sea background in 2026 is more similar to 2015."
    ),
    "三次事件的异同点": "Similarities and Differences among the Three Events",
    "2015/16 是最典型的强耦合发展型事件：中太平洋和中东太平洋 SST 较早为正，西风异常强，中太平洋降水也较早转正，D20 呈现清楚的西浅东深结构。它代表了 ocean dynamics、low-level winds 和 convection 的同步增强。": (
        "2015/16 is the most typical strongly coupled developing event: SST in the central and central-eastern Pacific became positive relatively early, westerly anomalies were strong, central Pacific precipitation also turned positive early, and D20 showed a clear west-shallow/east-deep structure. It represents the synchronized strengthening of ocean dynamics, low-level winds, and convection."
    ),
    "2023/24 的发展期更偏远东太平洋先行。FEP SST 是三者中相对最突出的，CP SST 和 PMM 支持较弱，降水和低层风对 SST 的响应也不如 2015 同步。因此它更像东部启动、随后向西连接的事件。": (
        "The 2023/24 development phase was more strongly led by the far eastern Pacific. FEP SST was the most prominent among the three events, while CP SST and PMM support were weaker, and the precipitation and low-level wind responses to SST were less synchronized than in 2015. It therefore looks more like an eastern-Pacific-starting event that later connected westward."
    ),
    "2026 的综合特征介于两者之间，但更靠近 2015 的动力框架。表层 SST 的早期暖中心带有东部先行特征，但统一基准后 CP/EP SST、西风异常、D20 西浅东深结构和 PMM 正位相都使 2026 更接近 2015。它目前与 2015 的主要差别是降水和深对流滞后，即海洋和低层风已经先行，大气深对流尚未完全闭合反馈。": (
        "The combined characteristics of 2026 lie between the other two events, but they are closer to the 2015 dynamical framework. The early surface SST warm center has an eastern-Pacific-leading component; however, after applying the unified baseline, CP/EP SST, westerly anomalies, the west-shallow/east-deep D20 structure, and the positive PMM phase all make 2026 more similar to 2015. The main difference from 2015 so far is the lag in precipitation and deep convection: the ocean and lower-tropospheric winds have led the evolution, but atmospheric deep convection has not yet fully closed the feedback."
    ),
    "结论": "Conclusion",
    "如果只看 1-5 月表层 SST 暖中心，2026 会显得有一部分 2023 式的东太平洋先行特征；但如果把 D20、u850 和 PMM 一起看，2026 与 2015 的相似性更强。更准确的表述是：2026 不是 2015 的简单复刻，而是一个东部启动明显、但动力背景逐渐转向 2015 式中东太平洋耦合框架的事件。后续判断是否真正接近 2015，还需要观察夏季以后中太平洋降水和 Walker circulation 是否持续响应。": (
        "If only the January-May surface SST warm center is considered, 2026 shows some 2023-like eastern-Pacific-leading characteristics. However, when D20, u850, and PMM are considered together, 2026 is more similar to 2015. A more accurate statement is that 2026 is not a simple repeat of 2015; rather, it is an event with a clear eastern Pacific start whose dynamical background is gradually shifting toward a 2015-like central-eastern Pacific coupled framework. To determine whether it truly approaches 2015, it will be necessary to monitor whether central Pacific precipitation and the Walker circulation continue to respond after summer."
    ),
}


def set_font(run) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")


def replace_paragraph_text(paragraph, text: str) -> None:
    size = None
    bold = False
    italic = False
    if paragraph.runs:
        size = paragraph.runs[0].font.size
        bold = bool(paragraph.runs[0].bold)
        italic = bool(paragraph.runs[0].italic)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    set_font(run)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = size


def translate() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)
    missing: list[str] = []
    for paragraph in doc.paragraphs:
        source = paragraph.text.strip()
        if not source:
            continue
        if source in TRANSLATIONS:
            replace_paragraph_text(paragraph, TRANSLATIONS[source])
        elif any("\u4e00" <= ch <= "\u9fff" for ch in source):
            missing.append(source)
        else:
            for run in paragraph.runs:
                set_font(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_font(run)
    if missing:
        raise RuntimeError("Missing translations for: " + " || ".join(missing))
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    translate()
