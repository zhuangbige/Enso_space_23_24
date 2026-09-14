#!/usr/bin/env python3
"""Replace Figure 7 in the latest English report with the D20 anomaly figure."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "outputs/research_2026/reports"
FIGURE_DIR = ROOT / "outputs/research_2026/report_visual_revision"
SOURCE = REPORT_DIR / "enso_event_comparison_unified_1980_2014_jan_jun_2026_en_visual_revision.docx"
OUTPUT = REPORT_DIR / "enso_event_comparison_unified_1980_2014_jan_jun_2026_en_d20_anomaly_revision.docx"
FIGURE = FIGURE_DIR / "fig07_three_event_d20_geographic_anomalies_clean.png"

CAPTION = (
    "Figure 7. Equatorial Pacific ORAS5 D20 anomalies and ERA5 u850 anomalies for the "
    "2015/16, 2023/24, and 2026 development periods. Positive D20 denotes thermocline "
    "deepening and negative D20 denotes shoaling; all panels show detrended calendar-month "
    "anomalies relative to 1980-2014. The locator map shows the equatorial section and the "
    "5°S-5°N wind-averaging band."
)
ALT_TEXT = (
    "Three-event comparison of January-June equatorial Pacific ORAS5 D20 anomalies and "
    "ERA5 850 hPa zonal-wind anomalies for 2015, 2023, and 2026."
)


def find_caption(document: Document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Figure 7."):
            return paragraph
    raise ValueError("Could not find the Figure 7 caption")


def set_drawing_size(image_element, image_path: Path, width_inches: float) -> None:
    with Image.open(image_path) as image:
        aspect = image.width / image.height
    width = int(width_inches * 914400)
    height = int(width / aspect)
    for extent in image_element.xpath(".//wp:extent"):
        extent.set("cx", str(width))
        extent.set("cy", str(height))
    for extent in image_element.xpath(".//a:xfrm/a:ext"):
        extent.set("cx", str(width))
        extent.set("cy", str(height))


def main() -> int:
    document = Document(SOURCE)
    caption = find_caption(document)
    image_element = caption._p.getprevious()
    if image_element is None or not image_element.xpath(".//w:drawing"):
        raise ValueError("The Figure 7 image was not found immediately before its caption")

    blips = image_element.xpath(".//a:blip")
    if len(blips) != 1:
        raise ValueError("Figure 7 should contain exactly one embedded image")
    relationship_id = blips[0].get(qn("r:embed"))
    document.part.related_parts[relationship_id]._blob = FIGURE.read_bytes()
    set_drawing_size(image_element, FIGURE, 6.80)

    for properties in image_element.xpath(".//wp:docPr"):
        properties.set("title", "Figure 7")
        properties.set("descr", ALT_TEXT)

    caption.text = ""
    caption_run = caption.add_run(CAPTION)
    caption_run.italic = True
    document.core_properties.subject = (
        "January-June comparison with 2015/16 and 2023/24; D20 anomaly revision"
    )
    document.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
