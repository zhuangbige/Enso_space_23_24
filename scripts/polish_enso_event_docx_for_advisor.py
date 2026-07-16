from pathlib import Path

from docx import Document


REPORT_DIR = Path("/Users/daiyue/Documents/Maryland/outputs/research_2026/reports")
SOURCE = REPORT_DIR / "enso_event_comparison_unified_1980_2014_en.docx"
OUTPUT = REPORT_DIR / "enso_event_comparison_unified_1980_2014_en_paper.docx"


REPLACEMENTS = {
    "January-May SST-Wind, Precipitation-Wind, D20-Wind, and PMM diagnostics using a unified climatology and detrending baseline":
        "January-May SST-wind, precipitation-wind, D20-wind, and PMM diagnostics under a unified climatology and detrending baseline",
    "Analysis Framework":
        "Data and Methods",
    "This report only discusses the three most recently redrawn January-May event figure sets. SST is from NOAA OISST, 850 hPa wind and precipitation are from ERA5, and D20 is from ORAS5. SST, wind, precipitation, and D20 anomalies are all calculated relative to the same-month climatology and then linearly detrended. The common ERA5/ORAS5 baseline is 1980-2014, while OISST uses 1982-2014 because of its data availability. The PMM figure is an independent PMM-like EOF diagnostic: the EOF training period is 1982-2025, and January-May 2026 is only projected onto that EOF mode rather than included in the EOF training.":
        "The analysis compares January-May diagnostics for three ENSO development phases. SST is from NOAA OISST, 850 hPa winds and precipitation are from ERA5, and D20 is from ORAS5. For SST, winds, precipitation, and D20, anomalies are calculated relative to the same-calendar-month climatology and then linearly detrended. The common baseline is 1980-2014 for ERA5/ORAS5 and 1982-2014 for OISST, reflecting OISST data availability. The PMM-like diagnostic is constructed independently using an EOF trained over 1982-2025; January-May 2026 is projected onto this fixed EOF mode and is not used in the EOF training.",
    "Spatially, the focus is on comparing SST anomalies in the central Pacific, eastern Pacific, and far eastern Pacific. Temporally, the focus is on whether the warm center, lower-tropospheric winds, precipitation, and D20 progress together from January to May. Air-sea coupling is not judged from a single variable; it is evaluated by whether SST, u850, precipitation, and D20 strengthen together in similar regions and months.":
        "The spatial comparison focuses on SST anomalies in the central Pacific, eastern Pacific, and far eastern Pacific. The temporal comparison examines whether the warm center, lower-tropospheric winds, precipitation, and D20 evolve in step from January to May. Air-sea coupling is assessed from the joint evolution of SST, u850, precipitation, and D20, rather than from any single field.",
    "The table reports January-May means. SST is in degC, u850 is in m/s, precipitation is in mm/day, and D20 is in m. Positive u850 indicates westerly anomalies; negative u850 indicates easterly anomalies.":
        "The table reports January-May means. SST is in \u00b0C, u850 is in m/s, precipitation is in mm/day, and D20 is in m. Positive u850 denotes westerly anomalies; negative u850 denotes easterly anomalies.",
    "Key Regional Metrics after Applying the Unified Climatology":
        "Regional Mean Diagnostics under the Unified Climatology",
    "Overall Assessment":
        "Results Overview",
    "After applying the unified climatology, 2026 no longer appears to be only a weak surface warming event in the eastern Pacific. Its Nino-3.4, CP, and EP SST anomalies are all higher than in the earlier non-unified baseline comparison, while CWP/CP lower-tropospheric westerlies and eastern Pacific D20 are also stronger. Dynamically, 2026 is therefore closer to 2015 than to 2023.":
        "Under the unified climatology, the 2026 case shows more than a weak eastern-Pacific surface warming signal. Its Ni\u00f1o-3.4, CP, and EP SST anomalies are higher than in the earlier non-unified comparison, and its CWP/CP lower-tropospheric westerlies and eastern Pacific D20 are stronger. In dynamical terms, 2026 is closer to 2015 than to 2023.",
    "The main difference among the three events is the maturity of air-sea coupling. In 2015, westerlies, D20, SST, and central Pacific precipitation strengthened together relatively early. In 2023, far eastern Pacific warming led the evolution, while central Pacific and PMM support remained weaker. In 2026, D20 and lower-tropospheric winds have already approached a 2015-like structure, but precipitation is still lagging, indicating that deep-convective coupling has not fully caught up.":
        "The three events differ mainly in the maturity of air-sea coupling. In 2015, westerlies, D20, SST, and central Pacific precipitation strengthened together early in the season. In 2023, far eastern Pacific warming led the development, with weaker support from the central Pacific and PMM. In 2026, D20 and lower-tropospheric winds have already approached a 2015-like structure, but precipitation still lags, indicating that deep-convective coupling is not yet fully established.",
    "SST Anomalies and 850 hPa Winds: Spatial Propagation of the Warm Center and Zonal Wind Anomalies":
        "SST Anomalies and 850 hPa Winds: Warm-Center Evolution and Zonal Wind Anomalies",
    "The early-spring 2015 signal was more centered from the central Pacific to the central-eastern Pacific and was accompanied by persistent westerly anomalies. In 2023, the warm center was clearly displaced toward the far eastern Pacific and then gradually expanded toward the central-eastern Pacific. In 2026, eastern Pacific strengthening was also present early on, but by April-May both central Pacific SST and westerlies had begun to connect with the eastern signal, making it less purely eastern-Pacific-focused than 2023.":
        "The early-spring 2015 signal was centered from the central to the central-eastern Pacific and was accompanied by persistent westerly anomalies. In 2023, the warm center was displaced toward the far eastern Pacific before expanding westward toward the central-eastern Pacific. In 2026, eastern Pacific warming was evident early, but by April-May both central Pacific SST and westerlies had begun to connect with the eastern signal. This makes 2026 less purely eastern-Pacific-focused than 2023.",
    "Precipitation and 850 hPa Winds: Whether the Atmospheric Response Keeps Pace with SST and Wind":
        "Precipitation and 850 hPa Winds: Atmospheric Response Relative to SST and Wind",
    "The precipitation maps show that the central Pacific convective response in 2015 was the earliest and most coherent. During the 2023 far eastern Pacific SST warming, central Pacific precipitation remained relatively weak, indicating delayed atmospheric coupling. The key point for 2026 is that the westerly anomalies already resemble 2015, but central Pacific precipitation remained suppressed from January to April and only recovered locally in May.":
        "The precipitation maps show that the central Pacific convective response was earliest and most coherent in 2015. During the 2023 far eastern Pacific warming, central Pacific precipitation remained weak, indicating delayed atmospheric coupling. In 2026, the westerly anomalies already resemble 2015, whereas central Pacific precipitation was still suppressed from January to April and only recovered locally in May.",
    "D20 is the core evidence for the similarity between 2015 and 2026. Both 2015 and 2026 show relative shoaling in the western Pacific and deepening in the central-eastern Pacific, corresponding to a tilt or transport of warm water toward the central-eastern Pacific. Although eastern Pacific D20 also deepened in 2023, the western Pacific did not show the same clear shoaling; therefore, its dynamical structure is less similar to 2015/2026.":
        "D20 provides the clearest evidence for the similarity between 2015 and 2026. Both events show relative shoaling in the western Pacific and deepening in the central-eastern Pacific, consistent with a tilt or transport of warm water toward the central-eastern Pacific. Although eastern Pacific D20 also deepened in 2023, the western Pacific did not show the same clear shoaling. Its subsurface dynamical structure is therefore less similar to 2015/2026.",
    "PMM: Whether the North Pacific Subtropical Background Is Involved":
        "PMM: North Pacific Subtropical Background",
    "The PMM-like EOF represents the leading coupled mode of North Pacific subtropical SST and near-surface winds after removing the linear influence of Nino-3.4. Its positive phase corresponds to warm northeastern Pacific SST, weakened trade winds, and reduced evaporative cooling. The joint SST+wind PMM index is clearly positive during January-May in both 2015 and 2026, whereas it is weak or even negative in 2023. This suggests that the subtropical air-sea background in 2026 is more similar to 2015.":
        "The PMM-like EOF represents the leading coupled mode of North Pacific subtropical SST and near-surface winds after the linear influence of Ni\u00f1o-3.4 is removed. Its positive phase corresponds to warm northeastern Pacific SST, weakened trade winds, and reduced evaporative cooling. The joint SST+wind PMM index is clearly positive during January-May in both 2015 and 2026, whereas it is weak or negative in 2023. This indicates that the subtropical coupled state in 2026 is closer to 2015 than to 2023.",
    "2015/16 is the most typical strongly coupled developing event: SST in the central and central-eastern Pacific became positive relatively early, westerly anomalies were strong, central Pacific precipitation also turned positive early, and D20 showed a clear west-shallow/east-deep structure. It represents the synchronized strengthening of ocean dynamics, low-level winds, and convection.":
        "2015/16 is the most strongly coupled developing event among the three. SST in the central and central-eastern Pacific became positive early, westerly anomalies were strong, central Pacific precipitation also turned positive early, and D20 showed a clear west-shallow/east-deep structure. This event represents a synchronized strengthening of ocean dynamics, low-level winds, and convection.",
    "The 2023/24 development phase was more strongly led by the far eastern Pacific. FEP SST was the most prominent among the three events, while CP SST and PMM support were weaker, and the precipitation and low-level wind responses to SST were less synchronized than in 2015. It therefore looks more like an eastern-Pacific-starting event that later connected westward.":
        "The 2023/24 development phase was led more strongly by the far eastern Pacific. FEP SST was the most prominent of the three cases, while CP SST and PMM support were weaker, and the precipitation and low-level wind responses to SST were less synchronized than in 2015. It therefore appears to be an eastern-Pacific-starting event that later connected westward.",
    "The combined characteristics of 2026 lie between the other two events, but they are closer to the 2015 dynamical framework. The early surface SST warm center has an eastern-Pacific-leading component; however, after applying the unified baseline, CP/EP SST, westerly anomalies, the west-shallow/east-deep D20 structure, and the positive PMM phase all make 2026 more similar to 2015. The main difference from 2015 so far is the lag in precipitation and deep convection: the ocean and lower-tropospheric winds have led the evolution, but atmospheric deep convection has not yet fully closed the feedback.":
        "The 2026 event lies between the other two, but its dynamical framework is closer to 2015. The early surface SST warm center has an eastern-Pacific component; however, under the unified baseline, CP/EP SST, westerly anomalies, the west-shallow/east-deep D20 structure, and the positive PMM phase all point toward a stronger resemblance to 2015. The main difference from 2015 so far is that precipitation and deep convection lag behind: the ocean and lower-tropospheric winds have led the development, while atmospheric deep convection has not yet closed the feedback.",
    "Conclusion":
        "Summary",
    "If only the January-May surface SST warm center is considered, 2026 shows some 2023-like eastern-Pacific-leading characteristics. However, when D20, u850, and PMM are considered together, 2026 is more similar to 2015. A more accurate statement is that 2026 is not a simple repeat of 2015; rather, it is an event with a clear eastern Pacific start whose dynamical background is gradually shifting toward a 2015-like central-eastern Pacific coupled framework. To determine whether it truly approaches 2015, it will be necessary to monitor whether central Pacific precipitation and the Walker circulation continue to respond after summer.":
        "Based only on the January-May surface SST warm center, 2026 retains some 2023-like east-Pacific-led characteristics. When D20, u850, and PMM are considered together, however, 2026 is more similar to 2015. These diagnostics suggest that 2026 is not a simple recurrence of the 2015 development, but an event with a clear eastern Pacific start whose dynamical background is shifting toward a 2015-like central-eastern Pacific coupled framework. Its subsequent evolution will depend on whether central Pacific precipitation and the Walker circulation continue to respond after summer.",
    "Ni\u00f1o3.4 SST":
        "Ni\u00f1o-3.4 SST",
}


def replace_paragraph_text(paragraph, new_text):
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    first = paragraph.runs[0]
    first.text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_table_text(cell, old, new):
    if old not in cell.text:
        return False
    for paragraph in cell.paragraphs:
        if old in paragraph.text:
            replace_paragraph_text(paragraph, paragraph.text.replace(old, new))
            return True
    return False


def main():
    doc = Document(SOURCE)
    missing = []
    changed = 0

    for old, new in REPLACEMENTS.items():
        found = False
        for paragraph in doc.paragraphs:
            if paragraph.text.strip() == old:
                replace_paragraph_text(paragraph, new)
                found = True
                changed += 1
                break
        if not found:
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if replace_table_text(cell, old, new):
                            found = True
                            changed += 1
                            break
                    if found:
                        break
                if found:
                    break
        if not found:
            missing.append(old[:90])

    doc.save(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"changed={changed}")
    if missing:
        print("missing replacements:")
        for item in missing:
            print(f"- {item}")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
