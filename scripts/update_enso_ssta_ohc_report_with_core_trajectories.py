#!/usr/bin/env python3
"""Add the 13-event warm-core movement analysis to the advisor report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "outputs/research_2026/reports"
INPUT = REPORTS / "ENSO发展期_SSTA-OHC不同步_阶段性汇报_修订版_20260815.docx"
OUTPUT = REPORTS / "ENSO发展期_SSTA-OHC不同步_阶段性汇报_暖核移动补充版_20260816.docx"

TRAJECTORIES = (
    ROOT
    / "outputs/research_2026/soda3_all_warm_event_core_trajectories"
    / "fig_soda3_13_warm_events_jan_jun_core_trajectory_small_multiples.png"
)
OVERVIEW = (
    ROOT
    / "outputs/research_2026/soda3_all_warm_event_core_trajectories"
    / "fig_soda3_13_warm_events_core_trajectory_overview.png"
)

FONT = "Noto Sans CJK SC"
FIGURE_WIDTH = Inches(6.45)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def insert_text(anchor, text: str = "", style: str | None = None):
    paragraph = anchor.insert_paragraph_before(style=style)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def insert_labeled_paragraph(anchor, label: str, text: str):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, bold=True)
    body_run = paragraph.add_run(text)
    set_run_font(body_run)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def insert_picture(anchor, path: Path, *, page_break_before: bool = False):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run()
    set_run_font(run)
    run.add_picture(str(path), width=FIGURE_WIDTH)
    return paragraph


def insert_caption(anchor, text: str):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    set_run_font(run, size=8)
    return paragraph


def paragraph_starting(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def add_method_line(document: Document) -> None:
    anchor = paragraph_starting(document, "权重：")
    paragraph = insert_text(
        anchor,
        "暖核质心：先对5°S—5°N的SODA温度距平作余弦纬度加权，得到"
        "q′=ρ₀CpT′；再在120°E—80°W、0—300 m内，以"
        "max(q′,0)乘原生层厚和纬向格距为权重，计算经度—深度质心。"
        "该计算不使用经度或深度分箱。",
        "Normal",
    )
    paragraph.paragraph_format.space_after = Pt(5)


def add_analysis(document: Document) -> None:
    appendix = paragraph_starting(document, "附录")
    appendix.paragraph_format.page_break_before = True

    heading = insert_text(
        appendix,
        "暖核移动：从三事件剖面到13次暖事件",
        "Heading 1",
    )
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True

    intro = insert_text(
        appendix,
        "为把剖面中的“向东上翻”转化为可比较量，本节以正热含量密度距平"
        "的加权质心定义暖核位置。每个点给出当月0—300 m正暖异常在经度—深度"
        "空间中的整体中心，箭头表示1—6月的顺序变化。该指标刻画暖异常的整体"
        "再分配，不追踪单个水团。",
        "Normal",
    )
    intro.paragraph_format.space_after = Pt(6)

    insert_picture(appendix, TRAJECTORIES)
    insert_caption(
        appendix,
        "图5  13次暖事件发展年1—6月赤道太平洋暖核质心轨迹。"
        "数字1—6表示月份，箭头表示时间方向；暖核为5°S—5°N、120°E—80°W、"
        "0—300 m内正热含量密度距平的加权质心。所有事件采用1980—2024年同月"
        "线性去趋势基准，计算不使用经度或深度分箱。",
    )

    first_finding = insert_text(
        appendix,
        "逐事件轨迹显示，向东、向浅层移动是发展期常见但并非普遍的几何变化。"
        "1982、1997、2009和2023表现出较清楚的东移上翻；部分中太平洋型事件"
        "移动较弱，甚至出现西移或下沉。这说明暖核移动能够区分事件的发展路径，"
        "但不能单独作为厄尔尼诺类型判据。",
        "Normal",
    )
    first_finding.paragraph_format.space_after = Pt(6)

    net_heading = insert_text(appendix, "跨事件净位移", "Heading 2")
    net_heading.paragraph_format.page_break_before = True
    net_heading.paragraph_format.keep_with_next = True
    insert_picture(appendix, OVERVIEW)
    insert_caption(
        appendix,
        "图6  13次暖事件暖核轨迹汇总及1—6月净位移。"
        "（a）全部事件轨迹，其中突出1997、2015和2023年；（b）1月至6月的"
        "向东与向上净位移，正值分别表示东移和上移。净位移代表正暖异常的整体"
        "重分配，不等同于单个水体的拉格朗日路径。",
    )

    result_heading = insert_text(appendix, "结果解释", "Heading 2")
    result_heading.paragraph_format.keep_with_next = True

    insert_labeled_paragraph(
        appendix,
        "1997。",
        "暖核由1月约185°E、160 m移动到6月约237°E、97 m，"
        "累计东移约52.1°、上移约63.3 m。其深层暖核在春季持续东移上翻，"
        "与OHC先行、随后东部海表快速增暖的演变一致。",
    )
    insert_labeled_paragraph(
        appendix,
        "2015。",
        "1月暖核已经位于约212°E、91 m，6月位于约233°E、90 m；"
        "东移约21.1°，垂向位置几乎不变。该事件的关键不是暖核大幅上翻，而是"
        "浅层偏东暖异常持续存在，因此2—3月海表短暂回冷并未消除次表层热储备。",
    )
    insert_labeled_paragraph(
        appendix,
        "2023。",
        "暖核由1月约160°E、165 m移动到6月约222°E、94 m，"
        "累计东移约62.6°、上移约71.2 m，在13次事件中均为最大。"
        "这表明2023年初并非完全缺少次表层正暖异常，而是暖量更偏西、偏深，随后发生"
        "快速的东移上翻；在暖核尚未充分进入远东整层水柱时，远东表层已可率先转暖。",
    )
    insert_labeled_paragraph(
        appendix,
        "与SSTA—OHC不同步的联系。",
        "三次事件对应三种不同几何状态：1997是深层暖核向东上翻并带动"
        "整层—表层共同增强；2015是偏东浅层暖核维持而海表短时波动；2023则是"
        "偏西深层暖核快速迁移，远东表层信号短暂领先于当地0—300 m整层异常。"
        "因此，发展期不同步不仅是异常幅度或过零月份不同，也反映暖量在"
        "经度—深度空间中的配置和迁移速度不同。",
    )
    insert_labeled_paragraph(
        appendix,
        "解释边界。",
        "质心位移会同时受到平流、赤道波动、温跃层起伏和混合过程影响，"
        "且只描述正异常的整体位置；它不能证明某一水团从西太平洋移动到东太平洋，"
        "也不能仅凭位移大小判断暖量是否增加。机制归因仍需结合热收支、输送分解"
        "或拉格朗日诊断。",
    )


def main() -> int:
    for path in (INPUT, TRAJECTORIES, OVERVIEW):
        if not path.exists():
            raise FileNotFoundError(path)
    document = Document(INPUT)
    add_method_line(document)
    add_analysis(document)
    document.core_properties.title = "ENSO发展期SSTA—OHC不同步：暖核移动补充分析"
    document.core_properties.subject = "导师阶段性汇报"
    document.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
